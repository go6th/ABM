import openpyxl, sys, glob, re
for fn in sorted(glob.glob(sys.argv[1]+'/*.xlsx')):
    wb=openpyxl.load_workbook(fn)
    tot=bad=0
    for ws in wb:
        hdr={}
        for row in ws.iter_rows():
            for c in row:
                if c.hyperlink:
                    tot+=1
                    t=c.hyperlink.target or ''
                    if not re.match(r'^https?://',t): bad+=1; print('  BAD',fn,ws.title,c.coordinate,t)
        # header named Link-ish columns whose cell has value but no hyperlink
        for row in ws.iter_rows(min_row=1,max_row=12):
            for c in row:
                if isinstance(c.value,str) and c.value.strip() in ('Link','Archived link','LinkedIn','Second profile'):
                    col=c.column; 
                    for r in range(c.row+1, ws.max_row+1):
                        cell=ws.cell(r,col)
                        if cell.value and not cell.hyperlink and not str(cell.value).startswith('None'):
                            print('  NOLINK',fn.split('/')[-1],ws.title,cell.coordinate,str(cell.value)[:50]); 
    print(fn.split('/')[-1],'hyperlinks',tot,'bad',bad)
