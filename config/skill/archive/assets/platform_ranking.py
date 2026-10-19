#!/usr/bin/env python3
"""Add a Platform Ranking sheet (A, B, C) to {Account}_Tech_Stack_and_Teams.xlsx.

Ranks the platforms named in the account's job postings from the last 12 months by how many
distinct postings name each one. Client facing: no tool names, no dashes as punctuation.

Usage:
    python platform_ranking.py ACCOUNT jobs.json Tech_Stack_and_Teams.xlsx \
        [--stack confirmed_stack_names.json] [--drop false_positives.json] [--checked checked.json]

Inputs
    jobs.json             the 12 month postings pull, keys t (title), d (description), f (first seen),
                          l (last seen), status ("closed" or null), url. Same file the jobs workbook uses.
    --stack               optional JSON list of product names from the confirmed stack pull.
    --drop                optional JSON list of [platform, url] pairs read and found to be false
                          positives (for example "salesforces" meaning sales teams). They are removed.
    --checked             optional JSON list of platform names whose matches were confirmed by reading
                          at least three postings. Every A rank must be in this list (gate check).

Rank rule (base = postings that name at least one platform in the dictionary):
    A  named in 15 percent or more of the base AND in at least 5 postings
    B  named in 5 percent or more of the base AND in at least 3 postings
    C  named in at least 2 postings
    Platforms named in only one posting are listed under "Seen once" and not ranked.
The rank cell is a formula on the count and share cells, so corrections flow through.
"""
import argparse, json, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Platform, category, lane (blank when the platform is outside the six lanes), pattern.
# Out of scope on every account and never listed: Databricks, Siebel, ServiceMax.
PLATFORMS = [
    ("Salesforce", "CRM", "Salesforce", r"salesforce|\bsfdc\b|sales cloud|service cloud|marketing cloud|health cloud|life sciences cloud|agentforce|\bcpq\b"),
    ("MuleSoft", "Integration", "Salesforce", r"mulesoft|anypoint"),
    ("Tableau", "Analytics and BI", "Tableau and analytics", r"tableau"),
    ("SAP", "ERP", "SAP", r"\bSAP\b|s/4 ?hana|s4hana|\bariba\b|successfactors|\bconcur\b"),
    ("ServiceNow", "IT service management", "ServiceNow", r"servicenow"),
    ("Veeva", "Life sciences CRM and content", "Veeva", r"\bveeva\b|vault crm|promomats|\bmedcomms\b|vault clinical|vault rim|vault quality|vault safety|\bcrossix\b"),
    ("Workday", "HR and finance", "", r"\bworkday\b"),
    ("Oracle ERP and EBS", "ERP", "", r"oracle (?:ebs|e-business|fusion|cloud erp|erp)|\bnetsuite\b|peoplesoft"),
    ("Microsoft Dynamics 365", "CRM and ERP", "", r"dynamics 365|\bd365\b|dynamics crm"),
    ("Snowflake", "Data platform", "Data and AI", r"snowflake"),
    ("Amazon Web Services", "Cloud", "Data and AI", r"\baws\b|amazon web services|redshift|sagemaker|bedrock"),
    ("Microsoft Azure", "Cloud", "Data and AI", r"\bazure\b|synapse|microsoft fabric"),
    ("Google Cloud", "Cloud", "Data and AI", r"google cloud|\bgcp\b|bigquery|vertex ai"),
    ("Power BI", "Analytics and BI", "Data and AI", r"power ?bi"),
    ("Qlik", "Analytics and BI", "Data and AI", r"\bqlik"),
    ("Looker", "Analytics and BI", "Data and AI", r"\blooker\b"),
    ("Informatica", "Data integration", "Data and AI", r"informatica"),
    ("dbt", "Data transformation", "Data and AI", r"\bdbt\b"),
    ("Kafka", "Data streaming", "Data and AI", r"\bkafka\b|confluent"),
    ("Alteryx", "Analytics", "Data and AI", r"alteryx"),
    ("SAS", "Analytics", "Data and AI", r"\bSAS\b"),
    ("OpenAI and Azure OpenAI", "Generative AI", "Data and AI", r"openai|chatgpt|\bgpt-?4"),
    ("Boomi", "Integration", "", r"\bboomi\b"),
    ("Epic", "Electronic health record", "", r"\bepic (?:systems|ehr|emr|certified|certification)|\bepic\b(?= (?:clarity|caboodle|tapestry|beaker|cadence))"),
    ("Oracle Health (Cerner)", "Electronic health record", "", r"cerner|oracle health"),
    ("Adobe Experience Cloud and Marketo", "Marketing", "", r"adobe experience|\bmarketo\b|adobe campaign"),
    ("HubSpot", "Marketing and CRM", "", r"hubspot"),
    ("Genesys and Five9", "Contact center", "", r"genesys|five9"),
    ("Zendesk", "Service desk", "", r"zendesk"),
    ("UiPath", "Automation", "", r"uipath"),
    ("Jira and Confluence", "Delivery tooling", "", r"\bjira\b|confluence(?! kafka)"),
]
OUT_OF_SCOPE = re.compile(r"databricks|siebel|servicemax", re.I)
assert not any(OUT_OF_SCOPE.search(p[0]) for p in PLATFORMS)


def clean(x):
    return re.sub(r"\s*[\u2013\u2014]\s*", ", ", x or "")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("account"); ap.add_argument("jobs"); ap.add_argument("workbook")
    ap.add_argument("--stack"); ap.add_argument("--drop"); ap.add_argument("--checked")
    a = ap.parse_args()

    jobs = json.load(open(a.jobs))
    stack = " | ".join(json.load(open(a.stack))) if a.stack else ""
    drops = {(p, u) for p, u in json.load(open(a.drop))} if a.drop else set()
    checked = set(json.load(open(a.checked))) if a.checked else set()

    hits = {p[0]: [] for p in PLATFORMS}
    base = set()
    for j in jobs:
        txt = (j.get("t") or "") + " " + (j.get("d") or "")
        for name, _, _, pat in PLATFORMS:
            if re.search(pat, txt, re.I) and (name, j.get("url")) not in drops:
                hits[name].append(j); base.add(j.get("url"))

    rows, once = [], []
    for name, cat, lane, pat in PLATFORMS:
        h = hits[name]
        if not h:
            continue
        open_now = sum(1 for j in h if j.get("status") != "closed")
        last = max((j.get("l") or j.get("f") or "") for j in h)
        in_stack = "Yes" if stack and re.search(pat, stack, re.I) else ("No" if stack else "Not pulled")
        rec = [name, cat, lane or "Outside the lanes", len(h), open_now, last, in_stack,
               "Yes" if name in checked else "No"]
        (rows if len(h) >= 2 else once).append(rec)
    rows.sort(key=lambda r: r[3], reverse=True)

    try:
        wb = load_workbook(a.workbook)
    except FileNotFoundError:
        wb = Workbook(); wb.remove(wb.active)
    if "Platform Ranking" in wb.sheetnames:
        del wb["Platform Ranking"]
    ws = wb.create_sheet("Platform Ranking", 0)

    F = "Arial"; navy = PatternFill("solid", fgColor="1F3A5F"); thin = Side(style="thin", color="BFC5CC")
    fills = {"A": "C6EFCE", "B": "FFEB9C", "C": "E7E6E6"}
    ws["A1"] = f"{a.account}: platforms seen in job postings, ranked A, B, C"; ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = "Postings in the last 12 months that name at least one platform (the base):"; ws["A2"].font = Font(name=F)
    ws["F2"] = len(base); ws["F2"].font = Font(name=F, bold=True)
    notes = [
        "How the rank is set. A: named in 15 percent or more of the base and in at least 5 postings. B: named in 5 percent or more of the base and in at least 3 postings. C: named in at least 2 postings. Platforms named in one posting are listed under Seen once.",
        "Each posting counts once per platform, however many times it names it. Open and closed postings both count. Matches found to be false on reading were removed.",
        "A platform named in hiring is a signal of skills the account wants, not proof the platform runs in production. The column In confirmed stack shows whether the separate stack pull also shows it.",
        "Research by 6th agents. Ranks are formulas on the counts, so a corrected count updates the rank.",
    ]
    for i, n in enumerate(notes):
        c = ws.cell(row=3 + i, column=1, value=n); c.font = Font(name=F, italic=True, size=9)
        ws.merge_cells(start_row=3 + i, start_column=1, end_row=3 + i, end_column=10)
        c.alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[3 + i].height = 28

    hdr = ["Rank", "Platform", "Category", "Lane", "Postings naming it", "Share of base", "Open now",
           "Last seen", "In confirmed stack", "Matches checked by reading"]
    H = 8
    for col, h in enumerate(hdr, 1):
        c = ws.cell(row=H, column=col, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = navy
        c.alignment = Alignment(wrap_text=True, vertical="center")
    r = H
    for rec in rows:
        r += 1
        ws.cell(row=r, column=2, value=clean(rec[0])); ws.cell(row=r, column=3, value=rec[1])
        ws.cell(row=r, column=4, value=rec[2]); ws.cell(row=r, column=5, value=rec[3])
        ws.cell(row=r, column=6, value=f"=IF($F$2=0,0,E{r}/$F$2)").number_format = "0%"
        ws.cell(row=r, column=7, value=rec[4]); ws.cell(row=r, column=8, value=rec[5])
        ws.cell(row=r, column=9, value=rec[6]); ws.cell(row=r, column=10, value=rec[7])
        ws.cell(row=r, column=1, value=f'=IF(AND(F{r}>=0.15,E{r}>=5),"A",IF(AND(F{r}>=0.05,E{r}>=3),"B",IF(E{r}>=2,"C","")))')
        for col in range(1, 11):
            c = ws.cell(row=r, column=col); c.font = Font(name=F, bold=(col == 1))
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    last_row = r
    if rows:
        from openpyxl.formatting.rule import FormulaRule
        for g, col in fills.items():
            ws.conditional_formatting.add(f"A{H+1}:A{last_row}", FormulaRule(formula=[f'$A{H+1}="{g}"'], fill=PatternFill("solid", fgColor=col)))

    r = last_row + 2
    ws.cell(row=r, column=1, value="Counts").font = Font(name=F, bold=True)
    for g in "ABC":
        r += 1
        ws.cell(row=r, column=1, value=f"Rank {g}").font = Font(name=F)
        ws.cell(row=r, column=2, value=f'=COUNTIF($A${H+1}:$A${max(last_row, H+1)},"{g}")').font = Font(name=F)
    r += 2
    ws.cell(row=r, column=1, value="Seen once (not ranked)").font = Font(name=F, bold=True)
    for rec in once:
        r += 1
        ws.cell(row=r, column=2, value=clean(rec[0])).font = Font(name=F)
        ws.cell(row=r, column=3, value=rec[1]).font = Font(name=F)
        ws.cell(row=r, column=4, value=rec[2]).font = Font(name=F)
    if not once:
        ws.cell(row=r + 1, column=2, value="None").font = Font(name=F)

    for col, w in zip("ABCDEFGHIJ", [8, 30, 26, 22, 12, 10, 10, 12, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{H+1}"
    ws.auto_filter.ref = f"A{H}:J{max(last_row, H)}"
    wb.save(a.workbook)
    n = len(base) or 1
    rank = lambda c: "A" if c / n >= 0.15 and c >= 5 else ("B" if c / n >= 0.05 and c >= 3 else "C")
    print(json.dumps({"base": len(base), "ranked": len(rows), "seen_once": len(once),
                      "by_rank": {g: [x[0] for x in rows if rank(x[3]) == g] for g in "ABC"},
                      "rank_A_not_checked_by_reading": [x[0] for x in rows if rank(x[3]) == "A" and x[7] == "No"]}))


if __name__ == "__main__":
    main()
