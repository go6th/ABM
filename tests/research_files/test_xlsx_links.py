"""The hyperlinks of an uploaded workbook (``tabular.sheet_links``, ``tabular.linked``): a link
cell showing a label reads as its address, safely. The package is untrusted, so the tests build
hostile ones by rewriting the zip parts directly: entity bombs, inflating parts, targets that
leave the package, ranges over the whole sheet, script and file addresses. SYNTHETIC data only
(no database)."""

from __future__ import annotations

import io
import time
import zipfile
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.hyperlink import Hyperlink

from research_files.config import load_config
from research_files.parsers import tabular
from research_files.parsers.base import ParseContext, Unit
from research_files.parsers.tabular import XlsxParser, linked, safe_url, sheet_links

CFG = load_config()
PROFILE = "https://www.linkedin.com/in/synthetic-person-{n}"
SHEET = "xl/worksheets/sheet1.xml"
SHEET_RELS = "xl/worksheets/_rels/sheet1.xml.rels"


def people_book(rows: int = 3, *, leading_empty: int = 0, links: bool = True) -> bytes:
    """A People sheet whose LinkedIn column shows "Open profile" and links to the profile."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "People"
    top = leading_empty + 1
    for j, name in enumerate(["Name", "Title as reported", "LinkedIn", "Note"], start=1):
        ws.cell(row=top, column=j, value=name)
    for i in range(1, rows + 1):
        r = top + i
        ws.cell(row=r, column=1, value=f"Synthetic Person {i}")
        ws.cell(row=r, column=2, value="Director")
        ws.cell(row=r, column=3, value="Open profile")
        ws.cell(row=r, column=4, value="note")
        if links:
            ws.cell(row=r, column=3).hyperlink = PROFILE.format(n=i)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def rewrite(data: bytes, parts: dict[str, bytes | None]) -> bytes:
    """The package with parts replaced, added (a new name) or removed (None)."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            if info.filename in parts:
                continue
            dst.writestr(info, src.read(info.filename), compress_type=zipfile.ZIP_DEFLATED)
        for name, body in parts.items():
            if body is not None:
                dst.writestr(name, body, compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


def part(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(name)


def parse(data: Any, cursor: dict[str, Any] | None = None) -> list[Unit]:
    stream = data if not isinstance(data, bytes) else io.BytesIO(data)
    return list(XlsxParser().units(stream, cursor or {}, ParseContext(CFG, "people.xlsx")))


def rows_of(units: list[Unit], sheet: str = "People") -> dict[int, list[Any]]:
    return {n: cells for u in units if u.table == sheet for n, cells in (u.rows or [])}


# ---- the address of a labelled link cell ---------------------------------------------------


def test_a_people_list_link_column_reads_as_profile_urls() -> None:
    rows = rows_of(parse(people_book()))
    assert [rows[r][2] for r in (2, 3, 4)] == [PROFILE.format(n=i) for i in (1, 2, 3)]
    assert rows[2][3] == "note"  # a cell without a link keeps its text
    # without hyperlinks (a CSV round trip) the label stays: the import then keeps no link
    assert rows_of(parse(people_book(links=False)))[2][2] == "Open profile"


def test_only_http_links_become_cell_values() -> None:
    hostile = [
        "javascript:alert(1)",
        "file:///etc/passwd",
        "mailto:someone@example.com",
        "../../secret.xml",
        "https://www.linkedin.com@attacker.example/in/x",
        "https://www.linkedin.com/in/x‮gnp.exe",
        "https://exa mple.com/",
        "https://" + "a" * 3000 + ".example/",
    ]
    for url in hostile:
        assert safe_url(url) is None, url
    assert safe_url(" https://www.linkedin.com/in/x ") == "https://www.linkedin.com/in/x"
    assert safe_url("HTTP://example.com/a?b=c#d") == "HTTP://example.com/a?b=c#d"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "People"
    ws.append(["Name", "LinkedIn"])
    for i, url in enumerate(hostile[:5], start=2):
        ws.append([f"Synthetic Person {i}", "Open profile"])
        ws.cell(row=i, column=2).hyperlink = url
    buf = io.BytesIO()
    wb.save(buf)
    rows = rows_of(parse(buf.getvalue()))
    assert all(cells[1] == "Open profile" for cells in rows.values())


def test_internal_links_and_merged_link_cells() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "People"
    ws.append(["Name", "LinkedIn", "Spare", "Jump"])
    ws.append(["Synthetic Person", "Open profile", None, "Open sheet"])
    ws.merge_cells("B2:C2")
    ws["B2"].hyperlink = PROFILE.format(n=1)
    ws["D2"].hyperlink = Hyperlink(ref="D2", location="'People'!A1")  # internal: no address
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    sheet = part(data, SHEET).replace(b'ref="B2"', b'ref="B2:C2"')  # Excel's merged-cell link
    rows = rows_of(parse(rewrite(data, {SHEET: sheet})))
    assert rows[2][1] == PROFILE.format(n=1)
    assert len(rows[2]) < 3 or rows[2][2] is None  # the merged range's other cell stays empty
    assert rows[2][3] == "Open sheet"


# ---- row numbers ---------------------------------------------------------------------------


def test_links_follow_the_row_numbers_when_a_sheet_resumes_from_a_cursor() -> None:
    data = people_book(rows=40)
    whole = rows_of(parse(data))
    header = next(u.header for u in parse(data) if u.table == "People" and u.header)
    resumed = rows_of(parse(data, {"sheet": 1 - 1, "row": 25, "header": header}))
    assert min(resumed) == 25
    assert all(resumed[r] == whole[r] for r in resumed)
    assert resumed[25][2] == PROFILE.format(n=24)  # row 25 is the 24th person


def test_a_sheet_that_opens_with_empty_rows_reads_its_header_from_the_first_row_with_values() -> (
    None
):
    units = parse(people_book(rows=2, leading_empty=2))
    first = next(u for u in units if u.table == "People" and u.header)
    assert first.header == ["Name", "Title as reported", "LinkedIn", "Note"]
    assert first.header_row == 3
    rows = rows_of(units)
    assert sorted(rows) == [4, 5]
    assert rows[4][2] == PROFILE.format(n=1)


def test_two_sheets_with_one_name_keep_their_own_links() -> None:
    wb = Workbook()
    first = wb.active
    assert first is not None
    first.title = "People"
    second = wb.create_sheet("Other")
    for ws, n in ((first, 1), (second, 2)):
        ws.append(["Name", "LinkedIn"])
        ws.append([f"Synthetic Person {n}", "Open profile"])
        ws["B2"].hyperlink = PROFILE.format(n=n)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    workbook = part(data, "xl/workbook.xml").replace(b'name="Other"', b'name="people"')
    data = rewrite(data, {"xl/workbook.xml": workbook})
    units = parse(data)
    got = sorted(cells[1] for u in units for _, cells in (u.rows or []))
    assert got == [PROFILE.format(n=1), PROFILE.format(n=2)]
    named = sheet_links(io.BytesIO(data))
    assert len(named) == 2 and {v for links in named.values() for v in links.values()} == set(got)


# ---- hostile packages ----------------------------------------------------------------------


def test_a_huge_or_malformed_hyperlink_range_is_read_in_bounded_time() -> None:
    data = people_book(rows=3)
    sheet = part(data, SHEET)
    sheet = sheet.replace(b'ref="C2"', b'ref="A1:XFD1048576"', 1)  # every cell of the sheet
    sheet = sheet.replace(b'ref="C3"', b'ref="ZZZZ9"', 1)  # malformed: skipped alone
    started = time.monotonic()
    units = parse(rewrite(data, {SHEET: sheet}))
    assert time.monotonic() - started < 5
    rows = rows_of(units)
    assert rows[2][2] == "Open profile"  # the giant range keeps only its first cell (A1)
    assert rows[3][2] == "Open profile"
    assert rows[4][2] == PROFILE.format(n=3)  # one bad ref never drops the others
    links = sheet_links(io.BytesIO(rewrite(data, {SHEET: sheet})))["People"]
    assert len(links) == 2 and (1, 1) in links.anchors


def test_the_linked_cells_of_a_workbook_are_capped(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(tabular, "MAX_LINKS", 2)
    rows = rows_of(parse(people_book(rows=5)))
    assert [rows[r][2] for r in range(2, 7)].count("Open profile") == 3


def test_a_part_that_inflates_past_its_cap_gives_no_links(monkeypatch: pytest.MonkeyPatch) -> None:
    data = people_book()
    padded = part(data, SHEET_RELS).replace(
        b"<Relationship ", b" " * 2_000_000 + b"<Relationship ", 1
    )
    bomb = rewrite(data, {SHEET_RELS: padded})
    assert len(bomb) < len(data) + 50_000  # two megabytes of spaces compress to almost nothing
    monkeypatch.setattr(tabular, "MAX_PART_BYTES", 64 * 1024)
    rows = rows_of(parse(bomb))
    assert rows[2][2] == "Open profile"  # the values are still read, without the links
    monkeypatch.setattr(tabular, "MAX_PART_BYTES", 64 * 1024 * 1024)
    monkeypatch.setattr(tabular, "MAX_SHEET_BYTES", 1024)
    assert rows_of(parse(data))[2][2] == "Open profile"  # the sheet scan is capped too


def test_a_part_that_declares_a_dtd_gives_no_links() -> None:
    data = people_book()
    rels = part(data, SHEET_RELS)
    laughs = (
        b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY a "aaaaaaaaaa">'
        + b"".join(
            b'<!ENTITY a%d "%s">' % (i, (b"&a%d;" % (i - 1) if i else b"&a;") * 10)
            for i in range(9)
        )
        + b"]>"
    )
    bomb = laughs + rels[rels.index(b"<Relationships") :].replace(
        b'Target="https', b'Target="&a8;https', 1
    )
    xxe = (
        b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/hostname">]>'
        + rels[rels.index(b"<Relationships") :]
    )
    for hostile in (bomb, xxe):
        started = time.monotonic()
        assert sheet_links(io.BytesIO(rewrite(data, {SHEET_RELS: hostile}))) in ({}, {"People": {}})
        units = parse(rewrite(data, {SHEET_RELS: hostile}))
        assert time.monotonic() - started < 5
        assert all("synthetic-person" not in str(u.rows) for u in units)


def test_sheet_targets_that_leave_the_package_are_not_read() -> None:
    assert tabular._part_name("xl/workbook.xml", "../../x.xml") is None
    assert tabular._part_name("xl/workbook.xml", "worksheets/../../../x.xml") is None
    assert tabular._part_name("xl/workbook.xml", "https://example.com/x.xml") is None
    assert tabular._part_name("xl/workbook.xml", "/xl/worksheets/sheet1.xml") == SHEET
    assert tabular._part_name("xl/workbook.xml", "worksheets/sheet1.xml") == SHEET
    data = people_book()
    rels = part(data, "xl/_rels/workbook.xml.rels")
    absolute = rels.replace(
        b'Target="/xl/worksheets/sheet1.xml"', b'Target="worksheets/sheet1.xml"'
    )
    absolute = absolute.replace(
        b'Target="worksheets/sheet1.xml"', b'Target="/xl/worksheets/sheet1.xml"'
    )
    assert sheet_links(io.BytesIO(rewrite(data, {"xl/_rels/workbook.xml.rels": absolute})))[
        "People"
    ]
    climbing = absolute.replace(b'Target="/xl/worksheets/sheet1.xml"', b'Target="../../evil.xml"')
    planted = rewrite(
        data, {"xl/_rels/workbook.xml.rels": climbing, "../evil.xml": part(data, SHEET)}
    )
    assert sheet_links(io.BytesIO(planted)) == {}


def test_a_non_seekable_or_broken_stream_is_reported_not_crashed() -> None:
    class OneWay(io.RawIOBase):
        def __init__(self, data: bytes) -> None:
            self.inner = io.BytesIO(data)

        def readable(self) -> bool:
            return True

        def seekable(self) -> bool:
            return False

        def readinto(self, buffer: Any) -> int:
            got = self.inner.read(len(buffer))
            buffer[: len(got)] = got
            return len(got)

    units = parse(OneWay(people_book()))
    assert units[-1].fatal and units[-1].findings[0].check == "corrupt_file"
    assert sheet_links(OneWay(people_book())) == {}  # type: ignore[arg-type]


def test_a_sheet_part_with_broken_xml_is_a_finding_not_an_exception() -> None:
    data = people_book(rows=3)
    sheet = part(data, SHEET).replace(b"</sheetData>", b"<row r=")  # cut inside the rows
    units = parse(rewrite(data, {SHEET: sheet}))
    assert units[-1].fatal
    assert units[-1].findings[0].check == "truncated_or_corrupt"


def test_a_plain_dict_of_links_still_fills_an_empty_cell() -> None:
    # callers that pass a plain mapping (no anchors) keep the earlier behaviour
    assert linked([None, "Open posting", "Senior"], 2, {(2, 1): "https://a.example/1"}) == [
        "https://a.example/1",
        "Open posting",
        "Senior",
    ]
