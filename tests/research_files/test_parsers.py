"""Every parser on a small fixture and on a corrupt copy; resume gives the same units; archive
guards; detection; sensitive data and research checks (no database)."""

from __future__ import annotations

import io
from typing import Any

import pytest

from research_files.checks import MetricLedger, mask, sensitive_kinds, unsourced_claims
from research_files.config import load_config
from research_files.detect import detect
from research_files.parsers.base import ParseContext, Unit
from research_files.parsers.registry import parser_for
from tests.research_files import fixtures

CFG = load_config()


def run(kind: str, data: bytes, name: str, cursor: dict[str, Any] | None = None) -> list[Unit]:
    parser = parser_for(kind)
    assert parser is not None, kind
    return list(parser.units(io.BytesIO(data), cursor or {}, ParseContext(CFG, name)))


@pytest.mark.parametrize("kind", sorted(fixtures.valid()))
def test_every_supported_type_is_detected_and_parsed(kind: str) -> None:
    name, data = fixtures.valid()[kind]
    detected = detect(io.BytesIO(data), name)
    expected = {"gz": "gz"}.get(kind, kind)
    assert detected == expected
    units = run(detected, data, name)
    assert units and units[-1].cursor.get("done") or any(u.cursor.get("done") for u in units)
    assert not any(u.fatal for u in units), [f.message for u in units for f in u.findings]
    if kind in ("csv", "tsv", "xlsx", "parquet", "json", "jsonl"):
        rows = [r for u in units if u.rows for r in u.rows]
        assert len(rows) >= 3


@pytest.mark.parametrize(
    "kind",
    [
        "csv",
        "tsv",
        "json",
        "jsonl",
        "pdf",
        "docx",
        "pptx",
        "xlsx",
        "parquet",
        "png",
        "wav",
        "zip",
        "gz",
        "html",
        "md",
    ],
)
def test_a_corrupt_file_is_reported_not_crashed(kind: str) -> None:
    name, data = fixtures.valid()[kind]
    broken = fixtures.corrupt(kind, data)
    detected = detect(io.BytesIO(broken), name)
    parser = parser_for(detected)
    if parser is None:  # truncated beyond recognition: stored, unsupported
        return
    units = list(parser.units(io.BytesIO(broken), {}, ParseContext(CFG, name)))
    findings = [f for u in units for f in u.findings]
    assert findings, f"{kind}: a corrupt file must produce a finding"
    assert any(f.severity in ("error", "warn") for f in findings), kind


def test_csv_resume_from_a_byte_offset_reads_the_same_rows() -> None:
    rows = [["k", "v"]] + [[str(i), f"value {i}"] for i in range(12000)]
    data = fixtures.csv_bytes(rows)
    full = run("csv", data, "big.csv")
    all_rows = [r for u in full if u.rows for r in u.rows]
    assert len(all_rows) == 12000 and all_rows[0][0] == 2
    first = full[0]
    resumed = run("csv", data, "big.csv", first.cursor)
    rest = [r for u in resumed if u.rows for r in u.rows]
    assert [r for u in full[1:] if u.rows for r in u.rows] == rest


@pytest.mark.parametrize(
    "kind", ["pdf", "pptx", "docx", "md", "html", "jsonl", "json", "xlsx", "parquet"]
)
def test_every_parser_resumes_to_the_same_units(kind: str) -> None:
    name, data = fixtures.valid()[kind]
    full = run(kind, data, name)
    content = [(u.kind, u.text, u.location) for u in full]
    for i, unit in enumerate(full[:-1]):
        resumed = run(kind, data, name, unit.cursor)
        assert [(u.kind, u.text, u.location) for u in resumed] == content[i + 1 :], (kind, i)


def test_zip_bomb_and_path_traversal_are_blocked() -> None:
    bomb = run("zip", fixtures.zip_bomb(), "bomb.zip")
    checks = {f.check for u in bomb for f in u.findings}
    assert "zip_bomb_blocked" in checks
    assert not any(u.rows for u in bomb)
    trav = run("zip", fixtures.traversal_zip(), "trav.zip")
    blocked = [f for u in trav for f in u.findings if f.check == "path_traversal_blocked"]
    assert {f.location["member"] for f in blocked} == {"../../etc/evil.csv", "/abs/evil.csv"}
    parsed_members = {u.location.get("member") for u in trav if u.rows}
    assert parsed_members == {"ok/data.csv"}


def test_a_gzip_bomb_is_stopped_by_the_member_cap() -> None:
    import gzip

    bomb = gzip.compress(b"0" * (300 << 20), mtime=0)
    units = run("gz", bomb, "bomb.txt.gz")
    assert any(f.check == "zip_bomb_blocked" for u in units for f in u.findings)


def test_sensitive_data_is_flagged_without_its_value() -> None:
    text = "mail ann@example.com card 4111 1111 1111 1111 key sk-abcdefghijklmnopqrstuv"
    kinds = sensitive_kinds(text)
    assert {"email", "card", "provider_api_key"} <= set(kinds)
    assert "ann@example.com" not in mask(text) and "sk-abcdefghijkl" not in mask(text)
    assert sensitive_kinds("order 4111 1111 1111 1112") == []  # fails the Luhn check


def test_research_checks_find_unsourced_and_contradicting_numbers() -> None:
    claims = unsourced_claims(fixtures.DOC_TEXT)
    assert any("25%" in c for c in claims)
    assert not any("40%" in c for c in claims)  # cited [1]
    ledger = MetricLedger()
    found = ledger.check("Revenue of $5 million in 2024.", {"page": 1})
    found += ledger.check("Revenue of $7 million in 2024.", {"page": 3})
    assert len(found) == 1 and found[0].check == "contradicting_numbers"


def test_detection_prefers_content_over_the_name() -> None:
    assert detect(io.BytesIO(fixtures.pdf_bytes()), "looks.csv") == "pdf"
    assert detect(io.BytesIO(fixtures.png_bytes()), "photo.txt") == "png"
    assert detect(io.BytesIO(b"\x00\x01\x02binary"), "mystery.bin") == "unknown"


@pytest.mark.asyncio
async def test_local_blob_store_streams_and_composes_without_holding_the_file(tmp_path) -> None:  # type: ignore[no-untyped-def]
    import hashlib

    from research_files.blobstore import LocalBlobStore, PartTooLarge

    store = LocalBlobStore(tmp_path)
    parts = [b"a" * 3000, b"b" * 3000, b"c" * 10]

    async def stream(data: bytes):  # type: ignore[no-untyped-def]
        for i in range(0, len(data), 1000):
            yield data[i : i + 1000]

    for n, data in enumerate(parts):
        size, digest = await store.write_stream(f"t/p{n}", stream(data), max_bytes=4000)
        assert (size, digest) == (len(data), hashlib.sha256(data).hexdigest())
    with pytest.raises(PartTooLarge):
        await store.write_stream("t/big", stream(b"x" * 5000), max_bytes=4000)
    assert not (tmp_path / "t/big").exists()  # a refused part leaves nothing behind
    size, digest = await store.compose("t/all", ["t/p0", "t/p1", "t/p2"])
    whole = b"".join(parts)
    assert (size, digest) == (len(whole), hashlib.sha256(whole).hexdigest())
    await store.move("t/all", "t/final")
    with store.open("t/final") as handle:
        assert handle.read() == whole
    with pytest.raises(ValueError):
        store.path("../escape")
    assert await store.delete_prefix("t/") >= 3


def _deck_with_notes_in_reverse_order() -> bytes:
    """Two slides whose deck order (presentation.xml) is the reverse of their file names, the
    first with speaker notes."""
    import zipfile

    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    out = io.BytesIO()

    def slide(text: str) -> str:
        return (
            f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><p:cSld><p:spTree><p:sp><p:txBody>'
            f"<a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
        )

    with zipfile.ZipFile(out, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("ppt/slides/slide1.xml", slide("Appendix"))
        z.writestr("ppt/slides/slide2.xml", slide("Opening"))
        z.writestr(
            "ppt/presentation.xml",
            f'<p:presentation xmlns:p="{p}" xmlns:r="{r}"><p:sldIdLst>'
            f'<p:sldId id="256" r:id="rId2"/><p:sldId id="257" r:id="rId1"/></p:sldIdLst>'
            f"</p:presentation>",
        )
        z.writestr(
            "ppt/_rels/presentation.xml.rels",
            f'<Relationships xmlns="{rel}"><Relationship Id="rId1" Target="slides/slide1.xml"/>'
            f'<Relationship Id="rId2" Target="slides/slide2.xml"/></Relationships>',
        )
        z.writestr(
            "ppt/slides/_rels/slide2.xml.rels",
            f'<Relationships xmlns="{rel}"><Relationship Id="rId9"'
            f' Target="../notesSlides/notesSlide1.xml"/></Relationships>',
        )
        z.writestr(
            "ppt/notesSlides/notesSlide1.xml",
            f'<p:notes xmlns:p="{p}" xmlns:a="{a}"><p:cSld><p:spTree>'
            f'<p:sp><p:nvSpPr><p:nvPr><p:ph type="sldImg"/></p:nvPr></p:nvSpPr></p:sp>'
            f'<p:sp><p:nvSpPr><p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr><p:txBody>'
            f"<a:p><a:r><a:t>Mention the budget cycle in March.</a:t></a:r></a:p>"
            f"</p:txBody></p:sp></p:spTree></p:cSld></p:notes>",
        )
    return out.getvalue()


def test_pptx_follows_the_deck_order_and_reads_speaker_notes() -> None:
    units = [u for u in run("pptx", _deck_with_notes_in_reverse_order(), "deck.pptx") if u.text]
    assert [u.heading for u in units] == ["Opening", "Appendix"]
    assert "Speaker notes: Mention the budget cycle in March." in units[0].text
    assert "Speaker notes" not in units[1].text


def test_docx_reads_headers_footnotes_and_comments() -> None:
    import zipfile

    from docx import Document

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Confidential: Northwind account plan"
    doc.add_paragraph("Body text about the renewal.")
    raw = io.BytesIO()
    doc.save(raw)
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    patched = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as src, zipfile.ZipFile(patched, "w") as z:
        for item in src.infolist():
            z.writestr(item, src.read(item.filename))
        z.writestr(
            "word/comments.xml",
            f'<w:comments xmlns:w="{w}"><w:comment w:id="0"><w:p><w:r>'
            f"<w:t>Ask about the March budget freeze.</w:t></w:r></w:p></w:comment></w:comments>",
        )
    units = run("docx", patched.getvalue(), "plan.docx")
    by_heading = {u.heading: u.text for u in units if u.heading}
    assert "Confidential: Northwind account plan" in by_heading["Headers and footers"]
    assert "Ask about the March budget freeze." in by_heading["Comments"]
    assert any("renewal" in u.text for u in units)
