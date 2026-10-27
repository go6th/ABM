"""Every sheet of the three workbooks rendered from the synthetic bundle and spot-checked:
headers, formulas, freeze panes, filters, merges, widths, fills, number formats, conditional
formatting; and, with LibreOffice present, the recalculated values."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from parity.bundle import AccountBundle
from parity.excel.people import security_line, summary_notes
from parity.recalc import check_workbook, soffice_binary
from parity.render import render_workbooks, workbook_names

office = pytest.mark.skipif(soffice_binary() is None, reason="LibreOffice is not installed")


@pytest.fixture(scope="module")
def rendered(synthetic_bundle: AccountBundle, tmp_path_factory: pytest.TempPathFactory) -> Path:
    out = tmp_path_factory.mktemp("render")
    render_workbooks(synthetic_bundle, out, recalc=False)
    return out


def test_file_names(synthetic_bundle: AccountBundle, rendered: Path) -> None:
    names = workbook_names(synthetic_bundle)
    assert names["people"] == "Acme_People_List.xlsx"
    assert sorted(p.name for p in rendered.iterdir()) == sorted(names.values())


def test_people_workbook(synthetic_bundle: AccountBundle, rendered: Path) -> None:
    wb = load_workbook(rendered / "Acme_People_List.xlsx")
    assert wb.sheetnames == ["Summary", "People"]
    ws = wb["People"]
    assert [c.value for c in ws[1]][:3] == ["No.", "Name", "Title as reported"]
    assert ws.freeze_panes == "C2"
    assert ws.auto_filter.ref == "A1:M8"
    assert ws.row_dimensions[1].height == 32
    assert ws.column_dimensions["C"].width == 56
    assert ws["A2"].value == 1 and ws["D2"].value == "C-level, EVP and SVP"
    assert ws["G2"].fill.fgColor.rgb.endswith("DDEFE0")
    assert ws["J2"].hyperlink.target == "https://profiles.example/in/pat-example"
    assert ws["J2"].font.color.rgb.endswith("1155CC")
    assert ws["B8"].value == "Jo Specimen" and ws["E8"].value == "No"
    s = wb["Summary"]
    assert s["A1"].value == "Acme people list, deduplicated"
    assert s["A1"].font.size == 14 and s["A1"].font.bold
    notes = summary_notes(synthetic_bundle)
    assert len(notes) == 6 and notes[2] == security_line(synthetic_bundle)
    assert notes[2].startswith("Security, IT compliance and IT audit people are excluded")
    assert "(2 rows): Sam Secure (Director, IT Security); Sam Secure (IT Audit Lead)." in notes[2]
    assert sorted(str(r) for r in s.merged_cells.ranges)[:2] == ["A3:F3", "A4:F4"]
    assert s["A11"].value == "Level" and s["A11"].fill.fgColor.rgb.endswith("1F3A5F")
    assert s["B12"].value == "=COUNTIF(People!$D$2:$D$8,A12)"
    assert s["C12"].value == '=COUNTIFS(People!$D$2:$D$8,A12,People!$G$2:$G$8,"A")'
    assert s["A21"].value == "Total" and s["B21"].value == "=SUM(B12:B20)"
    assert s["B22"].value == '=COUNTIF(People!$E$2:$E$8,"Yes")'
    assert s["A24"].value == "Lane, senior people only"
    assert s["B25"].value == (
        '=COUNTIFS(People!$F$2:$F$8,$A25,People!$G$2:$G$8,"A",People!$E$2:$E$8,"Yes")'
    )
    assert s["E33"].value == "=SUM(E25:E32)" and s["A33"].font.bold
    assert s.column_dimensions["A"].width == 44


def test_jobs_workbook(rendered: Path) -> None:
    wb = load_workbook(rendered / "Acme_Job_Postings.xlsx")
    assert wb.sheetnames == ["Summary", "Jobs"]
    ws = wb["Jobs"]
    assert [c.value for c in ws[1]][-2:] == ["Full description", "Link"]
    assert ws.freeze_panes == "C2" and ws.auto_filter.ref == "A1:N7"
    assert ws["B2"].value == "Clinical Systems Lead" and ws["F2"].value == "2026-02-28"
    assert ws["B3"].value == "SAP Architect" and ws["A3"].value == "R-1000"
    assert ws["B5"].value == "CRM Manager, Commercial" and ws["C5"].value == "Salesforce, SAP"
    assert ws["A6"].value is None and ws["I7"].value is None
    assert ws["N2"].hyperlink.target == "https://jobs.acme.example/job/6"
    assert ws["N2"].value == "Open posting"
    assert ws["B2"].alignment.wrap_text and not ws["C2"].alignment.wrap_text
    assert ws.column_dimensions["L"].width == 70
    s = wb["Summary"]
    assert s["A1"].value == (
        "Acme platform job postings, last 12 months (first seen 1 March 2025 to 1 March 2026)"
    )
    assert sorted(str(r) for r in s.merged_cells.ranges) == ["A3:H3", "A4:H4", "A5:H5"]
    assert s["A8"].value == "Platform (named anywhere), Acme positions only"
    assert s["A9"].value == "Salesforce"
    assert s["B9"].value == '=COUNTIFS(Jobs!$C$2:$C$7,"*Salesforce*",Jobs!$J$2:$J$7,"Acme*")'
    assert s["E9"].value == (
        '=COUNTIFS(Jobs!$C$2:$C$7,"*Salesforce*",Jobs!$J$2:$J$7,"Acme*",'
        'Jobs!$D$2:$D$7,"Build role")'
    )
    assert s["A17"].value == "Rows on the Jobs sheet (all entities)"
    assert s["B17"].value == "=COUNTA(Jobs!$B$2:$B$7)"
    assert s["A18"].value == "Rows that are Acme positions"
    assert s["A19"].value == "Rows that are Widgetco positions"
    assert s["B19"].value == (
        '=COUNTIF(Jobs!$J$2:$J$7,"Acme*")*0+COUNTA(Jobs!$B$2:$B$7)-COUNTIF(Jobs!$J$2:$J$7,"Acme*")'
    )
    assert s["A8"].fill.fgColor.rgb.endswith("1F3A5F") and s["B8"].font.color.rgb.endswith("FFFFFF")
    assert s.column_dimensions["A"].width == 52


def test_stack_workbook(rendered: Path) -> None:
    wb = load_workbook(rendered / "Acme_Tech_Stack_and_Teams.xlsx")
    assert wb.sheetnames == ["Platform Ranking", "Confirmed stack", "Platform teams"]
    r = wb["Platform Ranking"]
    assert r["A1"].value == "Acme: platforms seen in job postings, ranked A, B, C"
    assert r["F2"].value == 8 and r["F2"].font.bold
    assert sorted(str(m) for m in r.merged_cells.ranges) == ["A3:J3", "A4:J4", "A5:J5", "A6:J6"]
    assert r["A8"].value == "Rank" and r["J8"].value == "Matches checked by reading"
    assert r["B9"].value == "SAP" and r["E9"].value == 3
    assert r["A9"].value == (
        '=IF(AND(F9>=0.15,E9>=5),"A",IF(AND(F9>=0.05,E9>=3),"B",IF(E9>=2,"C","")))'
    )
    assert r["F9"].value == "=IF($F$2=0,0,E9/$F$2)" and r["F9"].number_format == "0%"
    assert r["I9"].value == "Yes" and r["J9"].value == "Yes" and r["J10"].value == "No"
    assert r["A13"].value == "Counts" and r["B14"].value == '=COUNTIF($A$9:$A$11,"A")'
    assert r["A18"].value == "Seen once (not ranked)" and r["B19"].value == "ServiceNow"
    assert r["D21"].value == "Outside the lanes"
    assert r.freeze_panes == "A9" and r.auto_filter.ref == "A8:J11"
    assert [str(c.sqref) for c in r.conditional_formatting] == ["A9:A11"]
    rule_fills = sorted(
        rule.dxf.fill.fgColor.rgb[-6:] for cf in r.conditional_formatting for rule in cf.rules
    )
    assert rule_fills == ["C6EFCE", "E7E6E6", "FFEB9C"]
    assert r.column_dimensions["B"].width == 30
    c = wb["Confirmed stack"]
    assert c["A1"].value.startswith("Acme: technologies seen in use")
    assert [str(m) for m in c.merged_cells.ranges] == ["A2:E2"]
    assert [x.value for x in c[4]] == [
        "Business function",
        "Category",
        "Technology",
        "Postings naming it",
        "Lane",
    ]
    assert c["C7"].value == "C#" and c["D5"].value == 12
    assert c.freeze_panes == "A5" and c.auto_filter.ref == "A4:E7"
    t = wb["Platform teams"]
    assert [str(m) for m in t.merged_cells.ranges] == ["A2:H2"]
    assert t["A5"].value == "Commercial Operations" and t["C5"].value == 40
    assert t["D5"].value is None and t["E5"].value == "2026-02-10"
    assert t.freeze_panes == "A5" and t.auto_filter.ref is None
    assert t.column_dimensions["F"].width == 70


@office
def test_recalculated_values_match_the_counts(
    synthetic_bundle: AccountBundle, tmp_path: Path
) -> None:
    written = render_workbooks(synthetic_bundle, tmp_path, recalc=True)
    for path in written:
        assert check_workbook(path) == []
    people = load_workbook(tmp_path / "Acme_People_List.xlsx", data_only=True)["Summary"]
    c = synthetic_bundle.computed
    assert people["B21"].value == c.people_total
    assert people["B22"].value == c.senior
    assert people["C21"].value == sum(1 for p in synthetic_bundle.people_rows if p.priority == "A")
    jobs = load_workbook(tmp_path / "Acme_Job_Postings.xlsx", data_only=True)["Summary"]
    assert jobs["B9"].value == c.jobs_counts["Salesforce"].total
    assert jobs["B12"].value == c.jobs_counts["SAP"].total
    assert jobs["B17"].value == c.lane_rows_all and jobs["B18"].value == c.lane_rows_account
    assert jobs["B19"].value == c.lane_rows_all - c.lane_rows_account
    rank = load_workbook(tmp_path / "Acme_Tech_Stack_and_Teams.xlsx", data_only=True)
    r = rank["Platform Ranking"]
    assert (r["A9"].value, r["A10"].value, r["A11"].value) == ("B", "C", "C")
    assert r["F9"].value == pytest.approx(3 / 8)
    assert (r["B14"].value, r["B15"].value, r["B16"].value) == (0, 1, 2)
    # formulas survive the re-save
    again = load_workbook(tmp_path / "Acme_People_List.xlsx")["Summary"]
    assert again["B21"].value == "=SUM(B12:B20)"
