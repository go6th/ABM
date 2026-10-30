"""Citations, dashes, stale phrases and excluded names on the rendered text."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook

from parity.gate.doc_checks import (
    dash_hits,
    excluded_names,
    orphan_citations,
    run_doc_checks,
)
from parity.gate.report import FAIL, PASS
from tests.parity.gate.conftest import EXCLUDED_NAME, PREFIX, make_docx, office_only


def test_orphans_are_cited_ids_outside_the_maps() -> None:
    assert orphan_citations("Seen (S1, J1) and (S9, J2) and U3.", {"S1", "J1"}) == [
        "J2",
        "S9",
        "U3",
    ]


def test_dashes_counted_except_inside_verbatim_quotes() -> None:
    text = "Plain text – with an en dash — and an em dash - and a spaced hyphen."
    assert dash_hits(text) == 3
    assert dash_hits("The quote says 5+ years – Salesforce.", ["5+ years – Salesforce"]) == 0


def test_excluded_names_strip_the_title(bundle: dict[str, Any]) -> None:
    assert excluded_names(bundle) == [EXCLUDED_NAME]


def test_clean_document_passes_without_a_converter(bundle: dict[str, Any], tmp_path: Path) -> None:
    docx = make_docx(tmp_path / f"{PREFIX}_Account_Summary.docx", ["Confirmed (S1, J1)."])
    result = run_doc_checks([docx], bundle, office=None)
    assert result.status == PASS
    assert result.details["text_source"][docx.name] == "docx"


def test_every_defect_is_a_finding(bundle: dict[str, Any], tmp_path: Path) -> None:
    docx = make_docx(
        tmp_path / f"{PREFIX}_Account_Summary.docx",
        ["Cited (S7) – with a dash.", f"Met {EXCLUDED_NAME} at the site.", "The old claim stands."],
    )
    result = run_doc_checks([docx], bundle, office=None, stale=["old claim"])
    assert result.status == FAIL
    messages = [f.message for f in result.findings]
    assert any(m.startswith("orphan citations (1): S7") for m in messages)
    assert any(m == "banned dashes: 1" for m in messages)
    assert any(m.startswith("stale phrase still present (1)") for m in messages)
    assert any(m == "excluded names present: 1" for m in messages)
    assert all(EXCLUDED_NAME not in m for m in messages)


def test_excluded_name_allowed_only_in_the_people_summary_note(
    bundle: dict[str, Any], tmp_path: Path
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A5"] = f"Security people are excluded (1 row): {EXCLUDED_NAME} (Director)"
    people = tmp_path / f"{PREFIX}_People_List.xlsx"
    wb.save(str(people))
    assert run_doc_checks([people], bundle, office=None).status == PASS
    wb["Summary"]["A6"] = f"Also met {EXCLUDED_NAME}."
    wb.save(str(people))
    assert run_doc_checks([people], bundle, office=None).status == FAIL
    wb2 = Workbook()
    wb2.active.title = "Jobs"
    wb2.active["A1"] = f"Reports to {EXCLUDED_NAME}, excluded."
    jobs = tmp_path / f"{PREFIX}_Job_Postings.xlsx"
    wb2.save(str(jobs))
    assert run_doc_checks([jobs], bundle, office=None).status == FAIL


@office_only
def test_pdf_text_is_used_when_a_converter_exists(
    bundle: dict[str, Any], tmp_path: Path, office: Any
) -> None:
    docx = make_docx(tmp_path / f"{PREFIX}_Account_Summary.docx", ["Confirmed (S1) and (S8)."])
    result = run_doc_checks([docx], bundle, office=office)
    assert result.details["text_source"][docx.name] == "pdf"
    assert result.status == FAIL
    assert result.findings[0].message == "orphan citations (1): S8"
