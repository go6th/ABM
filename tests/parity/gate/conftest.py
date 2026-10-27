"""Synthetic fixtures for the gate tests: an invented account, invented people, invented
numbers. Nothing here comes from a real workspace."""

from __future__ import annotations

import copy
import shutil
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook

from parity import counts

PREFIX = "Zorvex"
LANES = ["Salesforce", "SAP"]
PEOPLE_LANES = ["Salesforce and commercial", "SAP and ERP", "Business leadership and other"]
OTHER_ENTITY = "Quill (divested business)"
CLIENT_INTERNAL = "Marlowe Systems"
CLIENT_ALIASES = ["Marlowe", "marlowe-systems.example"]
EXCLUDED_NAME = "Nils Varga"
HAVE_OFFICE = shutil.which("soffice") is not None and shutil.which("pdftotext") is not None
office_only = pytest.mark.skipif(not HAVE_OFFICE, reason="LibreOffice or pdftotext missing")


def _job(
    title: str,
    platforms: list[str],
    cls: str,
    status: str,
    entity: str = PREFIX,
    requisition: str | None = "R-100",
) -> dict[str, Any]:
    return {
        "requisition": requisition,
        "title": title,
        "platforms": platforms,
        "cls": cls,
        "status": status,
        "first_seen": "2026-03-01",
        "last_seen": "2026-09-01",
        "location": "Oslo, Norway",
        "seniority": "mid_senior",
        "entity": entity,
        "what": "",
        "platform_sentences": f"{platforms[0]}: works with {platforms[0]} daily.",
        "description": "A posting.",
        "url": "https://careers.zorvex.example/job/1",
    }


def _person(
    name: str, title: str, level: int, lane: str, priority: str, url: str | None
) -> dict[str, Any]:
    return {
        "name": name,
        "title": title,
        "level": level,
        "senior": level <= 7,
        "lane": lane,
        "priority": priority,
        "former_si": "",
        "found_by": "title search",
        "profile_url": url,
        "second_profile_url": None,
        "location": "Oslo, Norway",
        "note": "",
        "employer_label": PREFIX,
    }


def make_bundle() -> dict[str, Any]:
    """A small schema-shaped bundle for an invented account."""
    jobs_rows = [
        _job("Salesforce Developer", ["Salesforce"], "Build role", "Open", requisition="R-101"),
        _job(
            "CRM Administrator",
            ["Salesforce"],
            "Admin and owner role",
            "Closed",
            requisition="R-102",
        ),
        _job(
            "Sales Representative", ["Salesforce", "SAP"], "User role", "Open", requisition="R-103"
        ),
        _job("SAP Architect", ["SAP"], "Build role", "Closed", requisition="R-104"),
        _job("Finance Analyst", ["SAP"], "Working tool", "Open", requisition=None),
        _job("Plant Controller", ["SAP"], "Working tool", "Open", entity=OTHER_ENTITY),
    ]
    jobs_rows[0]["what"] = "Builds the CRM integrations."
    people_rows = [
        _person(
            "Ada Quorum",
            "Chief Information Officer",
            1,
            PEOPLE_LANES[2],
            "A",
            "https://profiles.example/in/ada-quorum",
        ),
        _person(
            "Bram Tessel",
            "VP Commercial Systems",
            2,
            PEOPLE_LANES[0],
            "A",
            "https://profiles.example/in/bram-tessel",
        ),
        _person(
            "Cyra Lindqvist",
            "Director ERP",
            4,
            PEOPLE_LANES[1],
            "B",
            "https://profiles.example/in/cyra-lindqvist",
        ),
        _person("Dov Marchetti", "Senior Manager CRM", 7, PEOPLE_LANES[0], "C", None),
        _person(
            "Elke Brandvold",
            "Salesforce Developer",
            9,
            PEOPLE_LANES[0],
            "C",
            "https://profiles.example/in/elke-brandvold",
        ),
    ]
    people_rows[0]["note"] = "New owner of IT."
    people_rows[1]["former_si"] = "Halvorsen Partners (2015 to 2019)"
    ranking_rows = [
        {
            "platform": "SAP",
            "category": "ERP",
            "lane": "SAP",
            "count": 40,
            "open": 12,
            "last_seen": "2026-09-01",
            "in_stack": "Yes",
            "checked": "Yes",
        },
        {
            "platform": "Salesforce",
            "category": "CRM",
            "lane": "Salesforce",
            "count": 20,
            "open": 5,
            "last_seen": "2026-09-01",
            "in_stack": "Yes",
            "checked": "Yes",
        },
        {
            "platform": "Workday",
            "category": "HR",
            "lane": "Outside the lanes",
            "count": 6,
            "open": 1,
            "last_seen": "2026-08-01",
            "in_stack": "No",
            "checked": "No",
        },
        {
            "platform": "Qlik",
            "category": "BI",
            "lane": "Data and AI",
            "count": 2,
            "open": 0,
            "last_seen": "2026-07-01",
            "in_stack": "No",
            "checked": "No",
        },
        {
            "platform": "Boomi",
            "category": "Integration",
            "lane": "Outside the lanes",
            "count": 1,
            "open": 0,
            "last_seen": "2026-07-01",
            "in_stack": "No",
            "checked": "No",
        },
    ]
    bundle: dict[str, Any] = {
        "schema_version": "1.0.0",
        "meta": {
            "account_short": PREFIX,
            "account_legal": "Zorvex Holdings",
            "domain": "zorvex.example",
            "careers_site": "careers.zorvex.example",
            "run_date": "20 October 2026",
            "pull_date": "19 October 2026",
            "pull_date_short": "19 Oct 2026",
            "window": {"from": "2025-10-19", "to": "2026-10-19"},
            "window_display": {"from": "19 October 2025", "to": "19 October 2026"},
            "pursuing_client_display": "Client",
            "pursuing_client_internal": CLIENT_INTERNAL,
            "pursuing_client_aliases": CLIENT_ALIASES,
            "byline": "6th agents",
            "lanes": LANES,
            "people_lanes": PEOPLE_LANES,
            "entities": [
                {"label": PREFIX, "counted_as_account": True},
                {"label": OTHER_ENTITY, "counted_as_account": False, "note": "Divested in 2026."},
            ],
            "out_of_scope_platforms": [],
            "live_capture_count": 7,
            "feed_unique_links": 120,
            "raw_board_counts": {"Salesforce": 9},
        },
        "sources": {
            "S1": {
                "name": "Zorvex newsroom",
                "url": "https://zorvex.example/news",
                "type": "Company page",
                "date": "2026-09-01",
                "how": "Read",
                "supports": "The picture",
            },
            "S2": {
                "name": "Zorvex careers",
                "url": "https://careers.zorvex.example",
                "type": "Careers site",
                "date": "2026-10-19",
                "how": "Captured",
                "supports": "Postings",
            },
        },
        "jobs_read": {
            "J1": {
                "name": "Salesforce Developer",
                "url": "https://careers.zorvex.example/job/1",
                "plats": "Salesforce",
                "cls": "Build role",
                "what": "Builds",
                "status": "Open",
                "f": "2026-03-01",
                "ent": PREFIX,
            },
        },
        "jobs_rows": jobs_rows,
        "people_rows": people_rows,
        "people_excluded": {
            "domain": ["Someone Else (Director)"],
            "security": [f"{EXCLUDED_NAME} (Director, Security Operations)"],
            "unresolved": [],
            "below_rule_no_stack": [],
        },
        "overrides": [],
        "stack_rows": [
            {
                "business_function": "Sales",
                "category": "CRM",
                "technology": "Salesforce",
                "postings": 20,
                "lane": "Salesforce",
            },
        ],
        "team_rows": [
            {
                "team": "Commercial Systems",
                "lane_technologies": "Salesforce",
                "postings": 20,
                "first_activity": "2026-01-01",
                "last_activity": "2026-09-01",
                "technologies": "Salesforce (20)",
                "locations": "Oslo, Norway (20)",
                "job_functions": "IT",
            },
        ],
        "ranking": {
            "base": 60,
            "rows": ranking_rows,
            "seen_once": [
                {"platform": "Boomi", "category": "Integration", "lane": "Outside the lanes"}
            ],
        },
        "computed": {
            "jobs_counts": {},
            "ent_counts": {PREFIX: 100, OTHER_ENTITY: 569},
            "lane_rows_account": 5,
            "lane_rows_all": 6,
            "people_total": 5,
            "senior": 4,
            "prio": {"A": 2, "B": 1, "C": 1},
            "excluded": {"domain": 1, "security": 1, "unresolved": 0, "below_rule_no_stack": 0},
            "pf": {"senior_with_profiles": 3},
            "tracker_tally": {"done": 3, "open": 1},
        },
        "prose": {
            "glossary": "CRM is the customer system.",
            "narrative": {
                "picture": [f"{PREFIX} runs {LANES[0]} and {LANES[1]} (S1)."],
                "stands_out": ["The ERP consolidation is staffed (J1)."],
                "lanes": [["Salesforce", "Our interpretation: the CRM is being re-platformed."]],
                "wrong": ["We excluded 569 Quill postings by their own wording."],
                "sharpen": ["Ask who the partner is."],
            },
            "news": [["2026-09-01", "Zorvex opened a plant", "Read", "S1", "Capacity"]],
            "news_not_covered": "Product launches are not covered.",
            "baseline": ["Baseline: 100 Zorvex postings after removing 569 Quill positions."],
            "signals": {
                "Tier 1": [["New CIO", "Confirmed", "S1", "Owner of IT", "Ask"]],
                "Tier 2": [],
                "Tier 3": [],
                "Context and history": [],
            },
            "platform": [
                ["Salesforce", "Confirmed", "In use as the CRM (S1, J1).", "Ask about CPQ."]
            ],
            "hiring_interpretation": "Hiring says skills wanted, not production use.",
            "people_text": {
                "built": "Two title searches on the domain.",
                "profiles": "4 of 5 people meet the rule.",
                "exclusions": "Security people are excluded (1 row).",
            },
            "key_people": [
                [
                    "Ada Quorum",
                    "Chief Information Officer",
                    PEOPLE_LANES[2],
                    "A",
                    "https://profiles.example/in/ada-quorum",
                    "New owner of IT.",
                ]
            ],
            "profile_findings": [["Tenure", "Confirmed", "Long tenures", "Stable team"]],
            "status_changes": [["Ada Quorum", "Joined", "2026-09-01", "S1"]],
            "people_interpretation": ["The CRM team is small."],
            "hypotheses": [["The CRM will be re-platformed.", "Ask who decides."]],
            "cautions": ["The security exclusion is at the owner's instruction."],
            "first_call": ["Who owns the CRM roadmap?"],
            "waiting": ["Decision on the ERP lane."],
            "tracker_p1": [
                ["Careers site", "Done", "High", "Nothing", "Captured 7 live postings", "Nothing"]
            ],
            "tracker_p2": [["Profiles", "Done", "High", "Nothing", "Read 5 profiles", "Nothing"]],
            "evidence": {
                "extracts": [
                    {
                        "id": "J1",
                        "title": "Salesforce Developer",
                        "url": "https://careers.zorvex.example/job/1",
                        "quotes": ["5+ years – Salesforce"],
                    }
                ],
                "live_rows": [
                    [
                        "Salesforce Developer",
                        "https://careers.zorvex.example/job/1",
                        "Salesforce",
                        PREFIX,
                        "Yes",
                        "Build role",
                    ]
                ],
                "live_lane_count": 3,
                "leads": [],
                "detections": [],
                "news_events": [],
            },
            "workbook_notes": {
                "people": ["This is the single people list for the account."],
                "jobs": ["Every posting that names a platform."],
                "stack": {
                    "confirmed_title": "Technologies seen in use",
                    "confirmed_note": "Pulled from postings.",
                    "teams_title": "Teams naming a lane platform",
                    "teams_note": "Team names as recorded.",
                    "ranking_notes": [
                        "A: named in 15 percent or more of the base and in at least 5 postings."
                    ],
                },
            },
            "lane_highlights": None,
            "internal_only": None,
        },
    }
    bundle["computed"]["jobs_counts"] = {
        lane: {k: v for k, v in vars(jc).items() if k != "closed"}
        for lane, jc in counts.jobs_summary(bundle).items()
    }
    return bundle


@pytest.fixture
def bundle() -> dict[str, Any]:
    """A fresh copy of the synthetic bundle."""
    return copy.deepcopy(make_bundle())


@pytest.fixture
def office() -> Any:
    """A LibreOffice converter, or a skip."""
    if not HAVE_OFFICE:
        pytest.skip("LibreOffice or pdftotext missing")
    from exports.verify import LibreOffice

    return LibreOffice(timeout_s=300.0)


def make_docx(
    path: Path, paragraphs: Sequence[str], tables: Sequence[Sequence[Sequence[str]]] = ()
) -> Path:
    """A Word file with the paragraphs and tables given."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    for rows in tables:
        table = doc.add_table(rows=0, cols=max(len(r) for r in rows))
        for row in rows:
            cells = table.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = text
    doc.save(str(path))
    return path


def _header(ws: Any, row: int, labels: Sequence[str], start_col: int = 1) -> None:
    for i, label in enumerate(labels):
        ws.cell(row=row, column=start_col + i, value=label)


def write_jobs_workbook(path: Path, bundle: dict[str, Any]) -> Path:
    """The Job Postings workbook in the reference layout, formulas included."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    _header(
        ws,
        1,
        [
            "Requisition",
            "Title",
            "Platforms named",
            "Class",
            "Status",
            "First seen",
            "Last seen",
            "Location",
            "Seniority",
            "Entity",
            "What the posting says (read)",
            "Platform sentences",
            "Full description",
            "Link",
        ],
    )
    for r in bundle["jobs_rows"]:
        ws.append(
            [
                r["requisition"],
                r["title"],
                counts.platforms_named(r),
                r["cls"],
                r["status"],
                r["first_seen"],
                r["last_seen"],
                r["location"],
                r["seniority"],
                r["entity"],
                r["what"],
                r["platform_sentences"],
                r["description"],
                "Open posting",
            ]
        )
    last = len(bundle["jobs_rows"]) + 1
    s = wb.create_sheet("Summary", 0)
    s["A1"] = f"{PREFIX} platform job postings, last 12 months"
    for k, note in enumerate(bundle["prose"]["workbook_notes"]["jobs"]):
        s.cell(row=3 + k, column=1, value=note)
    r0 = 10
    s.cell(row=r0, column=1, value=f"Platform (named anywhere), {PREFIX} positions only")
    _header(
        s,
        r0,
        [
            "Total",
            "Closed",
            "Open",
            "Build role",
            "Admin and owner role",
            "Working tool",
            "User role",
        ],
        2,
    )
    prefix = bundle["meta"]["account_short"]
    for k, lane in enumerate(bundle["meta"]["lanes"]):
        r = r0 + 1 + k
        s.cell(row=r, column=1, value=lane)
        base = f'Jobs!$C$2:$C${last},"*{lane}*",Jobs!$J$2:$J${last},"{prefix}*"'
        s.cell(row=r, column=2, value=f"=COUNTIFS({base})")
        s.cell(row=r, column=3, value=f'=COUNTIFS({base},Jobs!$E$2:$E${last},"Closed")')
        s.cell(row=r, column=4, value=f'=COUNTIFS({base},Jobs!$E$2:$E${last},"Open")')
        for j, cls in enumerate(
            ["Build role", "Admin and owner role", "Working tool", "User role"]
        ):
            s.cell(row=r, column=5 + j, value=f'=COUNTIFS({base},Jobs!$D$2:$D${last},"{cls}")')
    rt = r0 + 2 + len(bundle["meta"]["lanes"])
    s.cell(row=rt, column=1, value="Rows on the Jobs sheet (all entities)")
    s.cell(row=rt, column=2, value=f"=COUNTA(Jobs!$B$2:$B${last})")
    s.cell(row=rt + 1, column=1, value=f"Rows that are {prefix} positions")
    s.cell(row=rt + 1, column=2, value=f'=COUNTIF(Jobs!$J$2:$J${last},"{prefix}*")')
    s.cell(row=rt + 2, column=1, value="Rows that are other entities")
    s.cell(
        row=rt + 2,
        column=2,
        value=f'=COUNTA(Jobs!$B$2:$B${last})-COUNTIF(Jobs!$J$2:$J${last},"{prefix}*")',
    )
    wb.save(str(path))
    return path


def write_people_workbook(path: Path, bundle: dict[str, Any]) -> Path:
    """The People List workbook in the reference layout, formulas included."""
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    _header(
        ws,
        1,
        [
            "No.",
            "Name",
            "Title as reported",
            "Level",
            "Meets seniority rule",
            "Lane",
            "Priority",
            "Former SI and consulting employers (firm and years)",
            "Found by",
            "LinkedIn",
            "Second profile",
            "Location",
            "Note",
        ],
    )
    for i, p in enumerate(bundle["people_rows"], 1):
        ws.append(
            [
                i,
                p["name"],
                p["title"],
                counts.level_label(p["level"]),
                "Yes" if p["senior"] else "No",
                p["lane"],
                p["priority"],
                p["former_si"],
                p["found_by"],
                "Open profile" if p["profile_url"] else None,
                p["second_profile_url"],
                p["location"],
                p["note"],
            ]
        )
    last = len(bundle["people_rows"]) + 1
    s = wb.create_sheet("Summary", 0)
    s["A1"] = f"{PREFIX} people list, deduplicated"
    for k, note in enumerate(bundle["prose"]["workbook_notes"]["people"]):
        s.cell(row=3 + k, column=1, value=note)
    s.cell(
        row=5,
        column=1,
        value="Security, IT compliance and IT audit people are excluded (1 row): "
        + "; ".join(bundle["people_excluded"]["security"]),
    )
    r0 = 11
    _header(s, r0, ["Level", "People", "Priority A"])
    for k, label in enumerate(counts.LEVEL_LABELS.values()):
        r = r0 + 1 + k
        s.cell(row=r, column=1, value=label)
        s.cell(row=r, column=2, value=f"=COUNTIF(People!$D$2:$D${last},A{r})")
        s.cell(
            row=r,
            column=3,
            value=f'=COUNTIFS(People!$D$2:$D${last},A{r},People!$G$2:$G${last},"A")',
        )
    rt = r0 + 1 + len(counts.LEVEL_LABELS)
    s.cell(row=rt, column=1, value="Total")
    s.cell(row=rt, column=2, value=f"=SUM(B{r0 + 1}:B{rt - 1})")
    s.cell(row=rt, column=3, value=f"=SUM(C{r0 + 1}:C{rt - 1})")
    s.cell(row=rt + 1, column=1, value="Meets seniority rule")
    s.cell(row=rt + 1, column=2, value=f'=COUNTIF(People!$E$2:$E${last},"Yes")')
    r2 = rt + 3
    _header(s, r2, ["Lane, senior people only", "A", "B", "C", "Total"])
    lanes = bundle["meta"]["people_lanes"]
    for k, lane in enumerate(lanes):
        r = r2 + 1 + k
        s.cell(row=r, column=1, value=lane)
        for j, pv in enumerate(["A", "B", "C"]):
            s.cell(
                row=r,
                column=2 + j,
                value=f'=COUNTIFS(People!$F$2:$F${last},$A{r},People!$G$2:$G${last},"{pv}",People!$E$2:$E${last},"Yes")',
            )
        s.cell(row=r, column=5, value=f"=SUM(B{r}:D{r})")
    r3 = r2 + 1 + len(lanes)
    s.cell(row=r3, column=1, value="Total")
    for j, col in enumerate("BCDE"):
        s.cell(row=r3, column=2 + j, value=f"=SUM({col}{r2 + 1}:{col}{r3 - 1})")
    wb.save(str(path))
    return path


def write_tech_workbook(path: Path, bundle: dict[str, Any]) -> Path:
    """The Tech Stack and Teams workbook in the reference layout, formulas included."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Platform Ranking"
    ws["A1"] = f"{PREFIX}: platforms seen in job postings, ranked A, B, C"
    ws["A2"] = "Postings in the last 12 months that name at least one platform (the base):"
    ws["F2"] = bundle["ranking"]["base"]
    for k, note in enumerate(bundle["prose"]["workbook_notes"]["stack"]["ranking_notes"]):
        ws.cell(row=3 + k, column=1, value=note)
    h = 8
    _header(
        ws,
        h,
        [
            "Rank",
            "Platform",
            "Category",
            "Lane",
            "Postings naming it",
            "Share of base",
            "Open now",
            "Last seen",
            "In confirmed stack",
            "Matches checked by reading",
        ],
    )
    for k, row in enumerate(bundle["ranking"]["rows"]):
        r = h + 1 + k
        ws.cell(
            row=r,
            column=1,
            value=f'=IF(AND(F{r}>=0.15,E{r}>=5),"A",IF(AND(F{r}>=0.05,E{r}>=3),"B",IF(E{r}>=2,"C","")))',
        )
        ws.cell(row=r, column=2, value=row["platform"])
        ws.cell(row=r, column=3, value=row["category"])
        ws.cell(row=r, column=4, value=row["lane"])
        ws.cell(row=r, column=5, value=row["count"])
        ws.cell(row=r, column=6, value=f"=IF($F$2=0,0,E{r}/$F$2)")
        ws.cell(row=r, column=7, value=row["open"])
        ws.cell(row=r, column=8, value=row["last_seen"])
        ws.cell(row=r, column=9, value=row["in_stack"])
        ws.cell(row=r, column=10, value=row["checked"])
    last = h + len(bundle["ranking"]["rows"])
    r = last + 2
    ws.cell(row=r, column=1, value="Counts")
    for j, g in enumerate("ABC"):
        ws.cell(row=r + 1 + j, column=1, value=f"Rank {g}")
        ws.cell(row=r + 1 + j, column=2, value=f'=COUNTIF($A${h + 1}:$A${last},"{g}")')
    stack = wb.create_sheet("Confirmed stack")
    _header(stack, 1, ["Business function", "Category", "Technology", "Postings naming it", "Lane"])
    for row in bundle["stack_rows"]:
        stack.append(
            [
                row["business_function"],
                row["category"],
                row["technology"],
                row["postings"],
                row["lane"],
            ]
        )
    teams = wb.create_sheet("Platform teams")
    _header(
        teams,
        1,
        [
            "Team",
            "Lane technology",
            "Jobs count",
            "First activity",
            "Last activity",
            "Team technologies",
            "Team locations",
            "Job functions",
        ],
    )
    for row in bundle["team_rows"]:
        teams.append(
            [
                row["team"],
                row["lane_technologies"],
                row["postings"],
                row["first_activity"],
                row["last_activity"],
                row["technologies"],
                row["locations"],
                row["job_functions"],
            ]
        )
    wb.save(str(path))
    return path


def write_workbooks(out_dir: Path, bundle: dict[str, Any]) -> list[Path]:
    """The three workbooks of the standing set."""
    prefix = bundle["meta"]["account_short"]
    return [
        write_jobs_workbook(out_dir / f"{prefix}_Job_Postings.xlsx", bundle),
        write_people_workbook(out_dir / f"{prefix}_People_List.xlsx", bundle),
        write_tech_workbook(out_dir / f"{prefix}_Tech_Stack_and_Teams.xlsx", bundle),
    ]


def write_documents(out_dir: Path, bundle: dict[str, Any]) -> list[Path]:
    """A clean Account Summary and Sources and Evidence for the bundle."""
    prefix = bundle["meta"]["account_short"]
    summary = make_docx(
        out_dir / f"{prefix}_Account_Summary.docx",
        [
            f"{prefix} account summary. Research date 20 October 2026. Pursuing client: Client.",
            "Zorvex runs Salesforce and SAP (S1). The CRM developer role is open (J1).",
            "We excluded 569 Quill postings by their own wording.",
            "4 of 5 people meet the seniority rule; 2 priority A senior people.",
        ],
        [
            [
                ["Name", "Title"],
                ["Ada Quorum", "Chief Information Officer"],
                ["Bram Tessel", "VP Commercial Systems"],
            ]
        ],
    )
    evidence = make_docx(
        out_dir / f"{prefix}_Sources_and_Evidence.docx",
        ["Part A: source register.", "S1 Zorvex newsroom, read 1 September 2026."],
        [[["ID", "Source"], ["S1", "Zorvex newsroom"], ["S2", "Zorvex careers"]]],
    )
    return [summary, evidence]


def write_outputs(out_dir: Path, bundle: dict[str, Any]) -> list[Path]:
    """The full standing set for the bundle."""
    return write_documents(out_dir, bundle) + write_workbooks(out_dir, bundle)


def make_workspace(root: Path) -> Path:
    """A workspace folder with notes, data and build, the client name planted in the notes."""
    (root / "notes").mkdir(parents=True)
    (root / "data").mkdir()
    (root / "build").mkdir()
    (root / "notes" / "pull_log_phase1.md").write_text(
        "# Zorvex phase 1 pull log\n\n"
        f"Run date: 2026-10-19. Pursuing client: {CLIENT_INTERNAL}.\n\n"
        "| Step | Tool | Credits used |\n|---|---|---|\n| 2.1 | web search | 12 credits |\n\n"
        "Credits remaining: 4,500\n"
        "Contact: ops@marlowe-systems.example for the account team.\n",
        encoding="utf-8",
    )
    (root / "data" / "pull.json").write_text(
        '{"contact": "recruiter@zorvex.example", "client": "marlowe"}\n', encoding="utf-8"
    )
    (root / "build" / "build_people.py").write_text("print('build')\n", encoding="utf-8")
    (root / "build" / f"build_{PREFIX}_Internal_Only.js").write_text(
        "// internal\n", encoding="utf-8"
    )
    (root / "data" / "blob.bin").write_bytes(b"\x00\x01\x02binary")
    return root
