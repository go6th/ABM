"""Source links as the client sees them (owner request 2026-10-21): "6th / <site>", opening the
original page; a data connector's name or host never shows, in the tables, the Excel exports or
an Ask answer's sources. Without a database."""

from __future__ import annotations

import io
from collections.abc import AsyncIterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from account_intelligence.replica import sheets as SH
from account_intelligence.replica import table_export as TE
from account_intelligence.replica.answers import Acct, Row, source_out
from account_intelligence.replica.source_label import NO_SITE, SourceLink, source_link
from documents.gate import provider_hosts


def _provider_url() -> str:
    """A page on a data provider's own host, read from the provider config (never in code)."""
    return f"https://api.{provider_hosts()[0]}/v1/records/42"


def _vendor_terms() -> tuple[str, ...]:
    from agents.naming import provider_terms

    return tuple(t.lower() for t in provider_terms() if len(t) > 3)


@pytest.mark.parametrize(
    ("url", "label"),
    [
        ("https://www.linkedin.com/jobs/view/123", "6th / LinkedIn"),
        ("https://jobs.lever.co/acme/abc", "6th / Lever"),
        ("https://careers.quidelortho.com/global/en/job/R1", "6th / quidelortho.com"),
        ("https://finance.yahoo.com/news/x.html", "6th / Yahoo Finance"),
        ("https://acme.wd5.myworkdayjobs.com/External/job/1", "6th / Workday"),
        ("https://news.example.co.uk/story", "6th / example.co.uk"),
    ],
)
def test_a_source_link_names_6th_and_the_site_never_the_connector(url: str, label: str) -> None:
    assert source_link(url) == SourceLink(label, url)
    found = source_link(_provider_url())
    assert found == SourceLink(NO_SITE, None)
    assert not any(term in found.label.lower() for term in _vendor_terms())


@pytest.mark.parametrize("url", [None, "", "mailto:a@example.com", "javascript:alert(1)", "x"])
def test_no_web_address_is_no_link(url: object) -> None:
    assert source_link(url) is None


def test_link_fields_drop_a_provider_page_and_keep_the_original() -> None:
    assert SH.link_fields("https://www.linkedin.com/jobs/view/1") == {
        "link": "https://www.linkedin.com/jobs/view/1",
        "link_label": "6th / LinkedIn",
    }
    assert SH.link_fields(_provider_url()) == {"link": None, "link_label": NO_SITE}
    assert SH.link_fields(None) == {"link": None, "link_label": None}


def _rows() -> list[dict[str, Any]]:
    return [
        {
            "n": 1,
            "cells": {"title": "Salesforce Architect"},
            **SH.link_fields("https://jobs.lever.co/acme/1"),
        },
        {"n": 2, "cells": {"title": "Data Engineer"}, **SH.link_fields(_provider_url())},
    ]


def test_the_workbook_link_reads_6th_and_opens_the_original_page() -> None:
    sheet = {
        "title": "Job postings",
        "private": False,
        "columns": [
            {"key": "title", "label": "Role", "kind": "strong"},
            {"key": "link", "label": "Posting", "kind": "link"},
        ],
        "filters": [{"key": "all", "label": "All", "count": 2}],
        "filter": "all",
        "total": 2,
        "rows": _rows(),
    }
    data = SH.workbook(sheet, "Acme")
    ws = load_workbook(io.BytesIO(data))["Sheet 1"]
    assert ws["C2"].value == "6th / Lever"
    assert ws["C2"].hyperlink.target == "https://jobs.lever.co/acme/1"
    assert ws["C3"].value == NO_SITE and ws["C3"].hyperlink is None
    assert provider_hosts()[0].encode() not in data


@pytest.mark.asyncio
async def test_the_streamed_table_export_links_the_same_way(tmp_path: Path) -> None:
    q = SH.parse("jobs", None, None, None)

    async def source() -> AsyncIterator[dict[str, Any]]:
        for row in _rows():
            yield {
                "n": row["n"],
                "cells": {c.key: row["cells"].get(c.key, "") for c in q.sheet.cols},
                "link": row["link"],
                "link_label": row["link_label"],
            }

    path = tmp_path / "jobs.xlsx"
    assert await TE.write_xlsx(path, q, "Acme", source()) == 2
    ws = load_workbook(path)["Sheet 1"]
    col = 2 + [c.kind for c in q.sheet.cols].index("link")
    first, second = ws.cell(row=2, column=col), ws.cell(row=3, column=col)
    assert first.value == "6th / Lever"
    assert first.hyperlink.target == "https://jobs.lever.co/acme/1"
    assert second.value == NO_SITE and second.hyperlink is None
    # a CSV cannot link: it keeps the original page's address, never a provider's
    lines = [
        b
        async for b in TE.csv_chunks(
            q,
            (r async for r in source()),
        )
    ]
    text = b"".join(lines).decode()
    assert "https://jobs.lever.co/acme/1" in text and provider_hosts()[0] not in text


def test_an_ask_source_carries_the_label_and_the_original_page() -> None:
    acct = Acct(uuid4(), "acme.com", "Acme", [], None, None, 0)

    def row(url: str | None) -> Row:
        return Row("j1", "job", acct, "t", "Data Engineer", "m", datetime.now(UTC), url)

    out = source_out(1, row("https://www.linkedin.com/jobs/view/9"))
    assert (out["label"], out["url"]) == ("6th / LinkedIn", "https://www.linkedin.com/jobs/view/9")
    hidden = source_out(2, row(_provider_url()))
    assert (hidden["label"], hidden["url"]) == (NO_SITE, None)
    assert source_out(3, row(None))["label"] is None


def test_a_saved_answer_is_relabelled_when_reopened() -> None:
    from account_intelligence.replica.source_label import relabel_answer

    saved = {
        "answered": True,
        "sources": [
            {"n": 1, "title": "Role", "url": "https://jobs.lever.co/acme/1"},
            {"n": 2, "title": "Feed", "url": _provider_url()},
            {"n": 3, "title": "Your file", "url": None},
        ],
    }
    out = relabel_answer(saved)
    assert [(s["label"], s["url"]) for s in out["sources"]] == [
        ("6th / Lever", "https://jobs.lever.co/acme/1"),
        (NO_SITE, None),
        (None, None),
    ]
    assert relabel_answer(None) is None
