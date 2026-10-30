"""Cell rules of the deep-dive table export, without a database: dashes, formula text, data tool
names and credit counts; and the whole-file check that refuses anything left over."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from account_intelligence.replica import sheets as SH


def _vendor() -> str:
    """A data tool's display name, read from the provider config (never written in code)."""
    from account_intelligence.connectors.registry_internal import vendor_names
    from account_intelligence.connectors.settings import load_providers_config

    return sorted(n for n in vendor_names(load_providers_config()).values() if n != "Parallel")[0]


def _vendors() -> tuple[str, ...]:
    from account_intelligence.connectors.registry_internal import vendor_names
    from account_intelligence.connectors.settings import load_providers_config

    return tuple(
        sorted(n for n in vendor_names(load_providers_config()).values() if n != "Parallel")
    )


@pytest.mark.parametrize(
    ("raw", "safe"),
    [
        ("=SUM(A1:A3)", "'=SUM(A1:A3)"),
        ("+44 20 7946 0000", "'+44 20 7946 0000"),
        ("-1 point", "'-1 point"),
        ("@cmd", "'@cmd"),
        ("Director – Sales", "Director, Sales"),
        ("2019–2021", "2019-2021"),
        ("Cloud - data", "Cloud, data"),
        ("Coca-Cola", "Coca-Cola"),
        # a dash at the edge of a cell: the finished-file check joins cells and refused the
        # whole Job postings workbook for one such title (staging export 2026-10-30)
        ("Senior Analyst -", "Senior Analyst"),
        ("Plan --  ", "Plan"),
        ("- Lead role", "Lead role"),
        (12, 12),
        (None, None),
    ],
)
def test_cells_lose_formulas_and_punctuation_dashes(raw: object, safe: object) -> None:
    assert SH.safe_cell(raw, shareable=True) == (safe, False)


@pytest.mark.parametrize(
    "raw", ["Found by {vendor}", "120 credits used", "call ref " + "a" * 64, "p_search hit"]
)
def test_tool_names_credit_counts_and_references_are_withheld(raw: str) -> None:
    raw = raw.format(vendor=_vendor())
    expected = (
        ("Found by 6th agents", False) if raw.startswith("Found by ") else (SH.WITHHELD, True)
    )
    assert SH.safe_cell(raw, shareable=True) == expected
    assert SH.safe_cell(raw, shareable=False) == (raw, False)


def _sheet(title: str, private: bool = False) -> dict[str, object]:
    return {
        "title": "Signals",
        "private": private,
        "columns": [{"key": "title", "label": "Signal", "kind": "strong"}],
        "filters": [{"key": "all", "label": "All", "count": 1}],
        "filter": "all",
        "total": 1,
        "rows": [{"n": 1, "cells": {"title": title}, "link": None}],
    }


def test_a_planted_tool_name_never_reaches_the_file() -> None:
    data = SH.workbook(_sheet(f"{_vendor()} shows growth"), "Acme")
    assert _vendor().lower().encode() not in data.lower()
    ws = load_workbook(io.BytesIO(data))["Sheet 1"]
    assert ws["B2"].value == "6th agents shows growth"


def test_people_workbook_scrubs_vendor_names_from_cells_and_summary() -> None:
    """The People workbook uses the same guarded cell writer as other client tables."""
    names = _vendors()[:2]
    sheet = {
        **_sheet("ordinary title"),
        "title": "People",
        "columns": [
            {"key": "person", "label": "Person", "kind": "strong"},
            {"key": "source", "label": "Found by", "kind": "plain"},
        ],
        "rows": [
            {
                "n": n,
                "cells": {"person": f"Person {n}", "source": name},
                "link": None,
            }
            for n, name in enumerate(names, start=1)
        ],
    }
    data = SH.workbook(sheet, "Acme")

    assert all(name.lower().encode() not in data.lower() for name in names)
    ws = load_workbook(io.BytesIO(data))["Sheet 1"]
    assert [ws.cell(row, 3).value for row in (2, 3)] == ["6th agents", "6th agents"]


def test_model_text_with_tool_name_and_provider_domain_is_withheld_from_client_file() -> None:
    from documents.gate import provider_hosts

    host = provider_hosts()[0]
    source = f"Found by {_vendor()} at https://{host}/search"
    data = SH.workbook(_sheet(source), "Acme")
    ws = load_workbook(io.BytesIO(data))["Sheet 1"]
    assert ws["B2"].value == SH.WITHHELD
    assert _vendor().lower().encode() not in data.lower()
    assert host.encode() not in data.lower()


def test_provider_domain_in_shareable_cell_is_withheld() -> None:
    from documents.gate import provider_hosts

    assert SH.safe_cell(f"Found at https://{provider_hosts()[0]}/search", shareable=True) == (
        SH.WITHHELD,
        True,
    )


@pytest.mark.asyncio
async def test_provider_domain_and_tool_name_are_blocked_in_streamed_csv() -> None:
    from account_intelligence.replica import table_export as TX
    from documents.gate import provider_hosts

    q = SH.parse("signals", None, None, None, 0, 10)
    vendor = _vendor()
    host = provider_hosts()[0]

    async def rows():
        yield {
            "n": 1,
            "cells": {
                col.key: f"{vendor} at https://{host}/search" if col.key == "title" else "ok"
                for col in q.sheet.cols
            },
            "link": None,
        }

    body = b"".join([chunk async for chunk in TX.csv_chunks(q, rows())])
    assert vendor.lower().encode() not in body.lower()
    assert host.encode() not in body.lower()
    assert SH.WITHHELD.encode() in body


def test_the_about_sheet_is_checked_too() -> None:
    data = SH.workbook(_sheet("fine"), _vendor())
    about = dict(load_workbook(io.BytesIO(data))["About"].iter_rows(values_only=True))
    assert about["Account"] == "6th agents"


def test_anything_left_in_the_finished_file_blocks_the_export(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(SH, "safe_cell", lambda value, *, shareable: (value, False))
    with pytest.raises(SH.ExportBlocked) as err:
        SH.workbook(_sheet(f"Seen by {_vendor()} \u2014 twice"), "Acme")
    assert "client_dashes" in err.value.checks
    assert any(c.startswith("client_terms:") for c in err.value.checks)


def test_a_private_sheet_is_not_scanned_as_shareable(monkeypatch: pytest.MonkeyPatch) -> None:
    title = f"Thread with {_vendor()}"
    data = SH.workbook(_sheet(title, private=True), "Acme")
    assert load_workbook(io.BytesIO(data))["Private to you"]["B2"].value == title


def test_the_gate_reads_dashes_that_a_workbook_stores_as_character_references() -> None:
    from openpyxl import Workbook

    from documents.gate import _DASHES, client_text
    from exports.excel.common import text_cell

    book = Workbook()
    assert book.active is not None
    text_cell(book.active["A1"], "Seen — twice")
    out = io.BytesIO()
    book.save(out)
    assert _DASHES.search(client_text(out.getvalue()))
