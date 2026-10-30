"""The delivered workbooks' sheets as tabs of the deep-dive tables (``replica.file_tabs``), the
formula values they show (``replica.workbook_formulas``), the file's own table standing in when
the stored table cannot be shown, and the LinkedIn links of the People tables (owner request
2026-10-27: "the table should be created in tabs the way the zip file has it; the LinkedIn
profile link is the main thing"). SYNTHETIC data only (the export test corpus and world)."""

from __future__ import annotations

import hashlib
import io
import json
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from account_intelligence.grading.config import load_grading_config
from account_intelligence.replica import file_tabs as FT
from account_intelligence.replica import sheets as SH
from account_intelligence.replica.workbook_formulas import Evaluator, matcher
from exports.builder import map_snapshot
from exports.config import load_export_config
from exports.excel.common import expected_of
from exports.excel.jobs import plan_jobs, render_jobs
from exports.excel.people import plan_people, render_people
from exports.excel.tech import plan_tech, render_tech

CFG, GRADING = load_export_config(), load_grading_config()
RENDER = {"people": render_people, "jobs": render_jobs, "tech": render_tech}


def _books() -> dict[str, tuple[bytes, Any]]:
    """The three workbooks of the rich synthetic corpus, as the builder renders them."""
    from tests.exports.corpus import rich

    mapped, _ = map_snapshot(rich(), CFG, GRADING)
    return {
        k: (RENDER[k](mapped[k].model, CFG, f"Synthetic {k}")[0], mapped[k].model) for k in RENDER
    }


def _grid(data: bytes) -> dict[tuple[str, int, int], Any]:
    wb = load_workbook(io.BytesIO(data))
    return {
        (ws.title, c.row, c.column): c.value
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    }


def test_the_web_tabs_compute_the_builders_formulas() -> None:
    """Every formula the builder writes evaluates to the value the builder computed in code (the
    file stores no cached value, so this is what the tab shows)."""
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    books = _books()
    plans = {
        "people": lambda m: plan_people(m, CFG)[1],
        "jobs": lambda m: plan_jobs(m, CFG),
        "tech": lambda m: plan_tech(m, CFG),
    }
    checked = 0
    for kind, (data, model) in books.items():
        ev = Evaluator(_grid(data))
        for (sheet, coord), want in expected_of(plans[kind](model)).items():
            col, row = coordinate_from_string(coord)
            got = ev.value(sheet, row, column_index_from_string(col))
            assert got == pytest.approx(want), (kind, sheet, coord)
            checked += 1
    assert checked > 50


def test_the_web_tabs_are_the_delivered_workbooks_sheets() -> None:
    books = _books()
    for kind, (data, model) in books.items():
        layout = getattr(CFG.rules.workbooks, kind)
        book = FT.read(data, layout.sheet)
        names = load_workbook(io.BytesIO(data)).sheetnames
        # every sheet, in the file's order, the configured one marked as the main table
        assert [t.name for t in book.tabs] == names
        assert [t.name for t in book.tabs if t.main] == [layout.sheet]
        assert len({t.key for t in book.tabs}) == len(names)
        main = book.tab(None)
        assert main is not None and main.table is not None
        assert list(main.table.header) == [c.header for c in layout.columns]
        # the Tech workbook's main sheet is the Confirmed stack (its detections are not printed)
        assert len(main.table.rows) == len(model.stack if kind == "tech" else model.rows)
    # the People List: each LinkedIn cell shows the file's text and opens the profile
    people = FT.read(books["people"][0], CFG.rules.workbooks.people.sheet)
    main = people.tab(None)
    assert main is not None and main.table is not None
    j = list(main.table.header).index("LinkedIn")
    model = books["people"][1]
    for cell, row in zip((r[j] for r in main.table.rows), model.rows, strict=True):
        assert cell.text == CFG.rules.sheet_texts.people.link_text
        assert cell.href == row.profile_url
    # the Summary sheet: its title line first, then its tables with computed counts
    summary = people.tab(FT.tab_key(CFG.rules.workbooks.people.summary_sheet))
    assert summary is not None and summary.blocks[0].kind == "text"
    tables = [b for b in summary.blocks if b.kind == "table"]
    assert tables and all(r[1].text != "" for t in tables for r in t.rows)
    # the Tech workbook opens on the Platform Ranking, whose rank cells are computed
    tech = FT.read(books["tech"][0], CFG.rules.workbooks.tech.sheet)
    ranking = tech.tabs[0]
    assert ranking.name == CFG.rules.workbooks.tech.ranking_sheet and not ranking.main
    table = ranking.table
    assert table is not None and table.rows
    assert {r[0].text for r in table.rows} <= {"A", "B", "C", ""}


def test_the_file_table_sorts_and_pages_like_a_sheet() -> None:
    data, model = _books()["people"]
    main = FT.read(data, CFG.rules.workbooks.people.sheet).tab(None)
    assert main is not None and main.table is not None
    live = SH.SHEETS["people"]
    q = FT.grid_query(main, live, None, None, 0, 5)
    page = FT.grid_page(main, q)
    assert page["from_file"] and page["total"] == len(model.rows) and len(page["rows"]) == 5
    assert [c["label"] for c in page["columns"]] == list(main.table.header)
    name = f"c{list(main.table.header).index('Name') + 1}"
    q = FT.grid_query(main, live, name, "desc", 0, 500)
    names = [r["cells"][name] for r in FT.grid_page(main, q)["rows"]]
    assert names == sorted(names, key=lambda t: t.casefold(), reverse=True)
    link = f"c{list(main.table.header).index('LinkedIn') + 1}"
    assert all(r["links"][link] for r in FT.grid_page(main, q)["rows"])
    with pytest.raises(SH.BadSheetQuery):
        FT.grid_query(main, live, "nope", None, 0, 5)


def test_criteria_follow_excel() -> None:
    assert matcher("*Sales*")("Salesforce, SAP") and not matcher("*Sales*")("SAP")
    assert matcher("open")("Open") and not matcher("open")("Opened")
    assert matcher("=~*x")("*x") and not matcher("=~*x")("ax")
    assert matcher(">=5")(7) and not matcher(">=5")(3) and not matcher(">=5")("text")
    assert matcher("<>Yes")("No") and not matcher("<>Yes")("yes")
    grid = {("S", 1, 1): 2, ("S", 2, 1): 4, ("S", 3, 1): "=SUM(A1:A2)/A1", ("S", 4, 1): "=A4+1"}
    ev = Evaluator(grid)
    assert ev.value("S", 3, 1) == 3
    assert ev.value("S", 4, 1) is None  # a cycle shows empty, never a wrong number
    assert Evaluator({("S", 1, 1): "=VLOOKUP(1,A2:B3,2)"}).value("S", 1, 1) is None


def test_an_uploaded_workbooks_link_cells_keep_their_address() -> None:
    """A People List uploaded as a file keeps each "Open profile" cell's LinkedIn address (the
    read-only reader drops hyperlinks; ``tabular.sheet_links`` reads them from the package)."""
    from research_files.parsers.tabular import linked, sheet_links

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "People"
    ws.append(["Name", "Title as reported", "LinkedIn", "Note"])
    ws.append(["Synthetic Person", "Director", "Open profile", "Open to work"])
    ws.cell(row=2, column=3).hyperlink = "https://www.linkedin.com/in/synthetic-person"
    ws.cell(row=2, column=4).hyperlink = "https://example.com/not-a-label"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    links = sheet_links(buf)["People"]
    cells = linked(["Synthetic Person", "Director", "Open profile", "Open to work"], 2, links)
    assert cells[2] == "https://www.linkedin.com/in/synthetic-person"
    assert cells[3] == "https://example.com/not-a-label"  # "Open to work" reads as a link label
    assert sheet_links(io.BytesIO(b"not a zip")) == {}


# ---- the API ---------------------------------------------------------------------------------


async def _world() -> Any:
    import psycopg

    from tests.conftest import ADMIN_DSN
    from tests.exports.test_db import builder, seed_world

    w, payloads = await seed_world()
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        c.execute(
            "UPDATE account_runs SET status = 'COMPLETED', finished_at = now() WHERE id = %s",
            (w.run,),
        )
    return w, payloads, builder(w, payloads)


async def _deliver_files(w: Any, payloads: Any, build: Any) -> Any:
    """The run's real workbooks stored as the builder stores them (manifest row counts
    included)."""
    import psycopg

    from tests.conftest import ADMIN_DSN

    content = await build.reader.read(str(w.tenant), str(w.run))
    mapped, _ = map_snapshot(content, CFG, GRADING)
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        for kind, render in RENDER.items():
            body, _ = render(mapped[kind].model, CFG, f"Synthetic {kind}")
            sha = hashlib.sha256(body).hexdigest()
            key = f"{w.tenant}/{w.account}/{w.run}/deliverables/{sha}.xlsx"
            payloads.objects[key] = body
            n = len(mapped[kind].model.rows)
            c.execute(
                "INSERT INTO deliverables (tenant_id, account_id, run_id, set_hash, kind, version,"
                " filename, sha256, size_bytes, object_key, manifest_json, content_hash,"
                " snapshot_hash, theme_hash, rules_hash, built_at, owner_user_id) VALUES"
                " (%s, %s, %s, %s, %s, 1, %s, %s, %s, %s, %s, %s, %s, %s, %s, now(),"
                " coalesce((SELECT owner_user_id FROM accounts WHERE id = %s), 'test-owner'))"
                " ON CONFLICT DO NOTHING",
                (
                    w.tenant, w.account, w.run, sha, kind, f"{kind}.xlsx", sha,
                    len(body), key, json.dumps({"rows": n, "content_hash": sha}), sha, sha, sha,
                    sha, w.account,
                ),
            )  # fmt: skip
    return mapped


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_the_tables_show_every_sheet_of_the_delivered_workbooks_as_tabs() -> None:
    from tests.replica.test_file_sheets import _app, _get

    w, payloads, build = await _world()
    plane = _app(w, payloads, build)
    base = f"/v1/accounts/{w.domain}/sheets"
    # before a delivery: no tabs, and the live People table has its LinkedIn column
    live = (await _get(plane, f"{base}/people")).json()
    assert live["tabs"] == [] and live["from_file"] is False
    assert "LinkedIn" in [c["label"] for c in live["columns"]]
    await _deliver_files(w, payloads, build)
    for key, kind in FT.KINDS.items():
        r = await _get(plane, f"{base}/{key}")
        assert r.status_code == 200, (key, r.text)
        sheet = r.json()
        assert sheet["from_file"] is True and sheet["file_id"]
        labels = [t["label"] for t in sheet["tabs"]]
        layout = getattr(CFG.rules.workbooks, kind)
        assert layout.sheet in labels and [t["main"] for t in sheet["tabs"]].count(True) == 1
        for t in sheet["tabs"]:
            page = await _get(plane, f"{base}/{key}/tabs/{t['key']}")
            assert page.status_code == 200, (key, t, page.text)
            body = page.json()
            assert body["tab_label"] == t["label"] and body["tabs"] == sheet["tabs"]
    assert (await _get(plane, f"{base}/people/tabs/no-such-sheet")).status_code == 404
    assert (await _get(plane, f"{base}/signals/tabs/summary")).status_code == 404


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_a_stale_stored_table_shows_the_files_own_table() -> None:
    """When the delivered run's stored table cannot be shown (here: it was stored under an
    earlier column set, as runs delivered before a column change are), the People sheet is the
    delivered file's table, links included, never the live research table with other columns;
    the counts and exports are the same rows."""
    import psycopg

    from tests.conftest import ADMIN_DSN
    from tests.replica.test_file_sheets import _app, _get

    w, payloads, build = await _world()
    plane = _app(w, payloads, build)
    mapped = await _deliver_files(w, payloads, build)
    base = f"/v1/accounts/{w.domain}/sheets/people"
    assert (await _get(plane, base)).json()["from_file"] is True  # stored table in use
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        c.execute(
            "UPDATE view_rows SET cells = cells - 'note'"
            " WHERE run_id = %s AND table_kind = 'people'",
            (w.run,),
        )
    sheet = (await _get(plane, base)).json()
    headers = [c.header for c in CFG.rules.workbooks.people.columns]
    assert sheet["from_file"] is True and sheet["tabs"]
    assert [c["label"] for c in sheet["columns"]] == headers
    assert sheet["total"] == len(mapped["people"].model.rows)
    link = f"c{headers.index('LinkedIn') + 1}"
    for got, row in zip(sheet["rows"], mapped["people"].model.rows, strict=False):
        assert got["links"].get(link) == row.profile_url
        assert got["cells"][link] == CFG.rules.sheet_texts.people.link_text
    ws = (await _get(plane, f"/v1/accounts/{w.domain}/workspace")).json()
    assert ws["counts"]["people"] == sheet["total"]
    csv_body = (await _get(plane, f"{base}/csv")).content.decode("utf-8-sig")
    assert csv_body.splitlines()[0].split(",")[1:3] == ["No.", "Name"]
    xlsx = await _get(plane, f"{base}/xlsx")
    assert xlsx.status_code == 200, xlsx.text


@pytest.mark.requires_postgres
def test_a_people_fact_that_is_only_a_title_is_not_a_person() -> None:
    """The People table leaves out stored people facts whose value is only a job title (the
    owner's exported People list showed "Vice President of Technology Product Management" as a
    name); a title followed by a name, or a name, stays."""
    import psycopg

    from account_intelligence.replica.workspace import title_only
    from tests.conftest import ADMIN_DSN

    values = {
        "Vice President of Technology Product Management": True,
        "Senior VP of Product and Strategy": True,
        "Principal Product Manager, Synthetic Unit by Synthetic Co": True,
        "Synthetic Person, President and Chief Executive Officer": False,
        "Chief Information Officer, Synthetic Person": False,
        "Head of Data, Synthetic Q. Person": False,
        "Synthetic Person": False,
    }
    with psycopg.connect(ADMIN_DSN) as c:
        for value, title in values.items():
            sql = f"SELECT {title_only('x.v')} FROM (SELECT CAST(%s AS text) AS v) x"
            got = c.execute(sql, (value,)).fetchone()
            assert got is not None and got[0] is title, value


# ---- the audit of 2026-10-29 (F1, F2, F4, F6, F7, F9, F10) -----------------------------------


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_a_files_header_may_gain_or_lack_columns_but_not_reorder_them() -> None:
    """``headers_agree``: a delivered file with columns added (the Archived link), or one
    delivered before the configuration gained a column, is still the builder's table; another
    order or two lists that each lack the other's names is another table (F4)."""
    want = ["Requisition", "Title", "Link"]
    assert FT.headers_agree(want, want)
    assert FT.headers_agree([*want, "Reviewer link", "Reviewer note"], want)
    assert FT.headers_agree(["Requisition", "Title"], want)  # the configuration gained Link
    assert FT.headers_agree(["", "Requisition", "Title", "Link", ""], want)  # blank headings
    assert not FT.headers_agree(["Title", "Requisition", "Link"], want)
    assert not FT.headers_agree(["Requisition", "Other"], ["Title", "Link"])
    assert not FT.headers_agree([], want) and not FT.headers_agree(["Name"], want)
    # builder_main takes the file's own table with two extra columns, rows counted by the manifest
    data, model = _books()["jobs"]
    wb = load_workbook(io.BytesIO(data))
    ws = wb[CFG.rules.workbooks.jobs.sheet]
    width = ws.max_column
    ws.cell(row=1, column=width + 1, value="Reviewer link")
    ws.cell(row=1, column=width + 2, value="Reviewer note")
    for r in range(2, ws.max_row + 1):
        ws.cell(
            row=r, column=width + 1, value="Open review"
        ).hyperlink = (
            f"https://web.example/v1/accounts/x/archive?url=https%3A%2F%2Fjobs.example%2F{r}"
        )
        ws.cell(row=r, column=width + 2, value="Yes" if r % 2 else "")
    book = FT.read(_bytes(wb), CFG.rules.workbooks.jobs.sheet)
    found = FT.Delivered("f", "jobs.xlsx", "0" * 64, "k", rows=len(model.rows))
    main = FT.builder_main(book, found, "jobs")
    assert main is not None and main.table is not None
    assert list(main.table.header)[-2:] == ["Reviewer link", "Reviewer note"]
    # both link columns of a row work on the file-grid path: each under its own key (F2, F8)
    q = FT.grid_query(main, SH.SHEETS["jobs"], None, None, 0, 3)
    page = FT.grid_page(main, q)
    kinds = {c["label"]: c["kind"] for c in page["columns"]}
    assert kinds["Reviewer link"] == "url" and kinds["Reviewer note"] == "text"
    reviewed = f"c{width + 1}"
    for row in page["rows"]:
        assert row["links"][reviewed].startswith("https://web.example/v1/accounts/x/archive?url=")
        assert row["cells"][reviewed] == "Open review"
    # two hyperlink columns in one row, each under its own key (F8)
    ws.cell(row=2, column=width + 2, value="Open posting").hyperlink = "https://jobs.example/2"
    two = FT.read(_bytes(wb), CFG.rules.workbooks.jobs.sheet).tab(None)
    assert two is not None
    row = FT.grid_page(two, FT.grid_query(two, SH.SHEETS["jobs"], None, None, 0, 1))["rows"][0]
    assert row["links"][reviewed] and row["links"][f"c{width + 2}"] == "https://jobs.example/2"
    assert row["cells"][f"c{width + 2}"] == "Open posting"
    # the file's own table stands even when the manifest count differs (never the live sheet)
    one_row = FT.Delivered("f", "jobs.xlsx", "0" * 64, "k", rows=1)
    assert FT.builder_main(book, one_row, "jobs") is None
    assert FT.main_table(book) is main
    # the filter is validated on the file-grid path (F10): All and the stored table's chips
    assert FT.grid_query(main, SH.SHEETS["jobs"], None, None, 0, 3, "all").filter.key == "all"
    with pytest.raises(SH.BadSheetQuery):
        FT.grid_query(main, SH.SHEETS["jobs"], None, None, 0, 3, "status_no_such_status")
    # one whole cell by row number and column key
    cell = FT.grid_cell(main, q, 2, reviewed)
    assert cell is not None and cell["text"] == "Open review" and cell["href"]
    assert FT.grid_cell(main, q, 10**6, reviewed) is None
    assert FT.grid_cell(main, q, 1, "c999") is None and FT.grid_cell(main, q, 1, "zz") is None


def test_the_stored_table_judge_lets_a_stored_row_hold_extra_keys() -> None:
    """``_judge`` (F4): a stored first row with every configured key stands, extra keys or not;
    one missing a configured key is a table stored under an earlier column set."""
    from uuid import uuid4

    from views.tables import specs

    keys = [c.key for c in specs()["people"].columns]
    run = uuid4()

    def pre(first: dict[str, Any]) -> SH._FilePrefetch:
        manifests = {"people": {"rows": 2}}
        return SH._FilePrefetch(run, True, manifests, {"people": first}, {"people": 2})

    assert SH._judge("people", pre(dict.fromkeys(keys, "x"))) == (run, 2)
    assert SH._judge("people", pre(dict.fromkeys([*keys, "archived_url"], "x"))) == (run, 2)
    assert SH._judge("people", pre(dict.fromkeys(keys[:-1], "x"))) is None


def test_fills_are_theme_tokens_on_both_screen_paths(monkeypatch: pytest.MonkeyPatch) -> None:
    """A cell's solid fill becomes the theme's token (F9): the hex value comes from the theme,
    never from code; an unnamed fill is dropped; the token travels in the tab's cells and in the
    grid row's ``fills``."""
    from openpyxl.styles import PatternFill

    monkeypatch.setattr(FT, "fill_tokens", lambda: {"DDEFE0": "priority_a", "C6EFCE": "rank_a"})
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "People"
    ws.append(["Name", "Priority", "Rank", "Note"])
    ws.append(["Synthetic One", "A", "A"])
    ws.append(["Synthetic Two", "B", ""])
    ws["D2"].fill = PatternFill("solid", start_color="FFDDEFE0")  # an empty cell of a filled row
    ws["B2"].fill = PatternFill("solid", start_color="FFDDEFE0")
    ws["C2"].fill = PatternFill("solid", start_color="C6EFCE")
    ws["B3"].fill = PatternFill("solid", start_color="FFFFF1CC")  # not in the theme map
    book = FT.read(_bytes(wb), "People")
    main = book.tab(None)
    assert main is not None and main.table is not None
    assert [c.fill for c in main.table.rows[0]] == [None, "priority_a", "rank_a", "priority_a"]
    assert [c.fill for c in main.table.rows[1]] == [None, None, None, None]
    out = FT.blocks_out(main)
    assert out[0]["rows"][0][1] == {"text": "A", "href": None, "fill": "priority_a"}
    assert "fill" not in out[0]["rows"][1][1]
    q = FT.grid_query(main, SH.SHEETS["people"], None, None, 0, 5)
    rows = FT.grid_page(main, q)["rows"]
    assert rows[0]["fills"] == {"c2": "priority_a", "c3": "rank_a", "c4": "priority_a"}
    assert "fills" not in rows[1]
    # the real theme map: every named value is six hex digits keyed by the brief's token names
    monkeypatch.undo()
    FT.fill_tokens.cache_clear()
    for rgb, token in FT.fill_tokens().items():
        assert len(rgb) == 6 and token in FT.FILL_TOKENS


def test_an_unsupported_formula_shows_the_files_cached_value() -> None:
    """F6: a re-saved file carries cached values; a function the evaluator lacks shows that
    value instead of a blank. An error value is not a value. The builder's own files, with no
    cached values, still compute."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    ws.append(["Rows", 1])
    ws.append(["More", 3])
    ws.append(["Biggest", "=MAX(B1:B2)"])
    ws.append(["Broken", "=B1/0"])
    ws.append(["Counted", "=COUNTA(A1:A4)"])
    raw = _bytes(wb)

    def cells(book: FT.Book) -> dict[str, str]:
        table = book.tabs[0].table
        assert table is not None  # two-cell rows without a styled header: the first is the header
        return {table.header[0]: table.header[1], **{r[0].text: r[1].text for r in table.rows}}

    # plant cached values as a re-save would: 3 for MAX, an error for the division
    import re
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = body.decode()
                text = re.sub(
                    r"<f>MAX\(B1:B2\)</f><v\s*/?>(</v>)?", "<f>MAX(B1:B2)</f><v>3</v>", text
                )
                text = re.sub(
                    r'<c r="B4"><f>B1/0</f><v\s*/?>(</v>)?</c>',
                    '<c r="B4" t="e"><f>B1/0</f><v>#DIV/0!</v></c>',
                    text,
                )
                body = text.encode()
            dst.writestr(item, body)
    got = cells(FT.read(out.getvalue(), "S"))
    assert got == {"Rows": "1", "More": "3", "Biggest": "3", "Broken": "", "Counted": "4"}
    # the builder's file: nothing cached, the evaluator's value shows and MAX stays blank
    assert cells(FT.read(raw, "S")) == {
        "Rows": "1", "More": "3", "Biggest": "", "Broken": "", "Counted": "4",
    }  # fmt: skip


def test_a_gap_closes_a_table_before_the_next_header_is_judged() -> None:
    """F7: in a sheet without the builder's header style, a header row right after a blank row
    starts its own table (the reference People Summary's lane table lost its header)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Summary"
    ws.append(["Synthetic people list"])
    ws.append([])
    ws.append(["Level", "People", "Priority A"])
    ws.append(["Director", 4, 2])
    ws.append(["Manager", 6, 1])
    ws.append([])
    ws.append(["Lane, senior people only", "A", "B"])
    ws.append(["Analytics", 2, 3])
    ws.append(["Commercial", 1, 0])
    tabs = FT.read(_bytes(wb), "Summary").tabs
    tables = [b for b in tabs[0].blocks if b.kind == "table"]
    assert [t.header for t in tables] == [
        ("Level", "People", "Priority A"),
        ("Lane, senior people only", "A", "B"),
    ]
    assert [len(t.rows) for t in tables] == [2, 2]
    assert tabs[0].blocks[0] == FT.Block("text", text="Synthetic people list")
    # the configured main sheet missing: the sheet with the largest table is the main one
    wb.create_sheet("Data").append(["A", "B"])
    for i in range(5):
        wb["Data"].append([i, i])
    book = FT.read(_bytes(wb), "Nope")
    assert [t.name for t in book.tabs if t.main] == ["Data"]


def test_the_page_cuts_long_cells_and_the_guard_withholds_a_cell_never_the_page() -> None:
    """F1 and F3 on a page of rows: a cell past ``CELL_CAP`` is cut and recorded, "in parallel"
    in a description is shown, a provider code withholds that cell only, and a link to a
    provider's host is dropped."""
    from agents.naming import provider_terms

    code = next(t for t in provider_terms() if t.startswith("p_"))
    long_text = "word " * 1000
    rows: list[dict[str, Any]] = [
        {
            "n": 1,
            "cells": {
                "a": long_text,
                "b": "The teams work in parallel on the rollout.",
                "c": f"from {code}",
            },
            "links": {"c": "https://example.com/x"},
            "link": None,
        },
        {"n": 2, "cells": {"a": "short", "b": "", "c": "x"}, "links": {}, "link": None},
    ]
    SH.cap_rows(rows)
    assert rows[0]["cut"] == {"a": len(long_text)} and rows[1]["cut"] == {}
    assert len(rows[0]["cells"]["a"]) == SH.CELL_CAP and rows[0]["cells"]["a"].endswith("…")
    held = SH.scrub_rows(rows)
    assert held == 1
    assert rows[0]["cells"]["b"] == "The teams work in parallel on the rollout."
    assert rows[0]["cells"]["c"] == SH.WITHHELD and rows[0]["links"] == {
        "c": "https://example.com/x"
    }
    assert rows[1]["cells"] == {"a": "short", "b": "", "c": "x"}
    # a tab's blocks: the same, cell by cell and line by line
    blocks: list[dict[str, Any]] = [
        {"kind": "text", "text": f"Pulled with {code}"},
        {"kind": "text", "text": "Counts by lane"},
        {
            "kind": "table",
            "header": ["A", code],
            "rows": [[{"text": "in parallel", "href": None}, {"text": code, "href": None}]],
        },
    ]
    assert FT.scrub_blocks(blocks) == 3
    assert blocks[0]["text"] == SH.WITHHELD and blocks[1]["text"] == "Counts by lane"
    assert blocks[2]["header"] == ["A", SH.WITHHELD]
    assert [c["text"] for c in blocks[2]["rows"][0]] == ["in parallel", SH.WITHHELD]


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_a_workbook_with_extra_columns_and_a_guarded_cell_shows_as_the_file() -> None:
    """The API end to end (F1, F2, F3, F4, F8): the delivered Jobs workbook gains two columns
    (a reviewer-link column and a reviewer-note column) and a description with a provider code; the
    stored table is passed over; the page is the file's own table with every column, both link
    columns open, the guarded cell alone is withheld, long cells are cut and whole from the cell
    route, a filter the file does not have is refused, and the exports carry the same table."""
    import psycopg

    from tests.conftest import ADMIN_DSN
    from tests.replica.test_file_sheets import _app, _get

    w, payloads, build = await _world()
    plane = _app(w, payloads, build)
    mapped = await _deliver_files(w, payloads, build)
    base = f"/v1/accounts/{w.domain}/sheets/jobs"
    assert (await _get(plane, base)).json()["from_file"] is True  # the stored table in use
    # rewrite the delivered Jobs workbook with two extra columns and planted cells
    from agents.naming import provider_terms

    code = next(t for t in provider_terms() if t.startswith("p_"))
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        row = c.execute(
            "SELECT id, object_key FROM deliverables WHERE run_id = %s AND kind = 'jobs'", (w.run,)
        ).fetchone()
        assert row is not None
        wb = load_workbook(io.BytesIO(payloads.objects[row[1]]))
        ws = wb[CFG.rules.workbooks.jobs.sheet]
        headers = [c.value for c in ws[1]]
        width = len(headers)
        desc = headers.index("Full description") + 1
        ws.cell(row=1, column=width + 1, value="Reviewer link")
        ws.cell(row=1, column=width + 2, value="Reviewer note")
        for r in range(2, ws.max_row + 1):
            ws.cell(
                row=r, column=width + 1, value="Open review"
            ).hyperlink = f"https://web.example/v1/accounts/{w.domain}/archive?url=https%3A%2F%2Fjobs.example%2F{r}"
            ws.cell(row=r, column=width + 2, value="Yes")
        ws.cell(row=2, column=desc, value="Teams work in parallel. " + "Detail. " * 400)
        ws.cell(row=3, column=desc, value=f"Pulled from {code} feed")
        body = _bytes(wb)
        sha = hashlib.sha256(body).hexdigest()
        payloads.objects[row[1]] = body
        c.execute("UPDATE deliverables SET sha256 = %s WHERE id = %s", (sha, row[0]))
        c.execute(
            "UPDATE view_rows SET cells = cells - 'title'"
            " WHERE run_id = %s AND table_kind = 'jobs'",
            (w.run,),
        )
    sheet = (await _get(plane, base)).json()
    assert sheet["from_file"] is True and sheet["tabs"] and sheet["columns"][0]["key"] == "c1"
    labels = [c["label"] for c in sheet["columns"]]
    assert labels == [*headers, "Reviewer link", "Reviewer note"]
    assert sheet["total"] == len(mapped["jobs"].model.rows)
    reviewed, link = f"c{width + 1}", f"c{headers.index('Link') + 1}"
    first, second = sheet["rows"][0], sheet["rows"][1]
    assert first["cells"][reviewed] == "Open review"
    assert first["links"][reviewed].startswith(
        f"https://web.example/v1/accounts/{w.domain}/archive"
    )
    assert link in first["links"] or first["cells"][link]  # the posting link column of the file
    d = f"c{desc}"
    assert first["cells"][d].startswith("Teams work in parallel.")
    assert first["cells"][d].endswith("\u2026")
    assert first["cut"] == {d: len("Teams work in parallel. " + "Detail. " * 400)}
    assert second["cells"][d] == SH.WITHHELD and second["cut"] == {}
    # the whole cell, under the same order
    whole = await _get(plane, f"{base}/cell", params={"n": 1, "col": d})
    assert whole.status_code == 200, whole.text
    assert whole.json()["text"] == "Teams work in parallel. " + "Detail. " * 400
    held = (await _get(plane, f"{base}/cell", params={"n": 2, "col": d})).json()
    assert held["text"] == SH.WITHHELD and held["withheld"] is True
    assert (await _get(plane, f"{base}/cell", params={"n": 10**6, "col": d})).status_code == 404
    assert (await _get(plane, f"{base}/cell", params={"n": 1, "col": "zz"})).status_code == 404
    # a filter the file's table does not have is refused like an unknown sort; All is fine
    assert (await _get(plane, f"{base}?filter=status_no_such_status")).status_code == 422
    # a link column is not sortable
    assert (await _get(plane, f"{base}?filter=all&sort={reviewed}")).status_code == 422
    assert (await _get(plane, f"{base}?filter=all&sort=c2&dir=desc")).status_code == 200
    # the exports write the same table, extra columns included
    csv_body = (await _get(plane, f"{base}/csv")).content.decode("utf-8-sig")
    assert csv_body.splitlines()[0].split(",")[-2:] == ["Reviewer link", "Reviewer note"]
    xlsx = await _get(plane, f"{base}/xlsx")
    assert xlsx.status_code == 200, xlsx.text
    ws = (await _get(plane, f"/v1/accounts/{w.domain}/workspace")).json()
    assert ws["counts"]["jobs"] == sheet["total"]


# ---- the web-table finish of 2026-10-29: chips on the file grid, link texts, fills, the guard ---


def test_the_file_grid_has_the_stored_tables_chips_and_they_filter() -> None:
    """The buttons work on file-shaped tabs: the chips the stored table has (Priority A, the
    posting statuses) are matched to the file's columns by header, counted over the whole table,
    and filter the page, the whole-cell route and the exports alike."""
    from views.tables import specs

    books = _books()
    for key, kind, chip_key, column in (
        ("people", "people", "priority_a", "priority"),
        ("jobs", "jobs", f"status_{next(iter(CFG.rules.labels.posting_status))}", "status"),
    ):
        data, model = books[key]
        main = FT.read(data, getattr(CFG.rules.workbooks, kind).sheet).tab(None)
        assert main is not None and main.table is not None
        header = {c.key: c.header for c in specs()[kind].exported_columns}
        j = list(main.table.header).index(header[column])
        live = SH.SHEETS[key]
        chip = next(f for f in SH.file_sheet(key).filters if f.key == chip_key)
        want = [r for r in main.table.rows if r[j].text.strip() == _chip_value(kind, chip_key)]
        assert want, key  # the synthetic corpus holds rows for the chip
        q = FT.grid_query(main, live, None, None, 0, 500, chip_key)
        page = FT.grid_page(main, q)
        assert page["filter"] == chip_key and page["total"] == len(want)
        assert len(page["rows"]) == len(want)
        counts = {f["key"]: f["count"] for f in page["filters"]}
        assert counts["all"] == len(main.table.rows) and counts[chip_key] == len(want)
        assert chip.label in [f["label"] for f in page["filters"]]
        # a sort inside the chip, the whole cell under the chip, and the export see the same rows
        sorted_q = FT.grid_query(main, live, "c1", "desc", 0, 500, chip_key)
        assert FT.grid_page(main, sorted_q)["total"] == len(want)
        cell = FT.grid_cell(main, q, len(want), "c1")
        assert cell is not None and FT.grid_cell(main, q, len(want) + 1, "c1") is None
        assert len(FT.grid_rows(main, q)) == len(want)
    with pytest.raises(SH.BadSheetQuery):
        FT.grid_query(main, SH.SHEETS["jobs"], None, None, 0, 5, "no_such_chip")


def _chip_value(kind: str, chip_key: str) -> str:
    """The cell text a chip of the stored table matches (``sheets.chip_specs``)."""
    from views.tables import specs

    for key, _label, _column, value in SH.chip_specs(
        kind, specs()[kind].exported_columns, CFG.rules
    ):
        if key == chip_key:
            return value
    raise AssertionError(chip_key)


def test_a_url_column_shows_its_configured_link_text_on_the_stored_table_path() -> None:
    """``sheets._link_text``: a url column's label is the workbook column's own ``link_text``
    ("Open posting", "Open archived copy"), else the sheet text ("Open profile"), else the
    address; every url column of a table gets one."""
    rules = CFG.rules
    jobs = rules.workbooks.jobs
    columns = tuple(
        c.model_copy(update={"link_text": "Open posting"}) if c.key == "url" else c
        for c in jobs.columns
    )
    configured = rules.model_copy(
        update={
            "workbooks": rules.workbooks.model_copy(
                update={"jobs": jobs.model_copy(update={"columns": columns})}
            )
        }
    )
    assert SH._link_text(configured, "jobs", "url") == "Open posting"
    assert SH._link_text(configured, "jobs", "no_such_column") is None
    # the people sheet text covers both profile columns
    for key in ("profile_url", "second_profile_url"):
        assert SH._link_text(rules, "people", key) == rules.sheet_texts.people.link_text
    # a rendered row: the cell reads the label and the address opens
    col = SH.Col("url", "Link", None, "url", "Open posting")
    sheet = SH.Sheet("jobs", "Jobs", "f.xlsx", "", "", (col,), (), "x", table="jobs")
    q = SH.Query(sheet, SH.Filter("all", "All", "true"), col, False, 0, 10)
    row = SH.row_out(q, 1, {"url": "https://jobs.example/p/1", "link": None}, None)  # type: ignore[arg-type]
    assert row["cells"] == {"url": "Open posting"} and row["links"] == {
        "url": "https://jobs.example/p/1"
    }


def test_the_stylesheet_fills_are_the_themes_fills() -> None:
    """Colours live in the theme (``excel.fills``) and the stylesheet that mirrors them for the
    screen (CSP forbids inline styles): a class per token with the theme's hex value."""
    import re as _re
    from pathlib import Path

    import yaml

    from account_intelligence.config import CONFIG_DIR

    theme = yaml.safe_load((CONFIG_DIR / "export_theme.yaml").read_text())["excel"]["fills"]
    css = (Path(__file__).resolve().parents[2] / "web" / "ui" / "components.css").read_text()
    for token in FT.FILL_TOKENS:
        name = token.replace("_", "-")
        found = _re.search(
            rf"table\.xl td\.fill-{name}\b[^{{]*\{{\s*background:\s*#([0-9A-Fa-f]{{6}})", css
        )
        assert found, token
        assert found.group(1).upper() == str(theme[token]).upper()[-6:], token
    assert set(FT.fill_tokens().values()) == set(FT.FILL_TOKENS)


def test_ordinary_english_is_not_a_leak_in_the_gate_and_the_page() -> None:
    """F1: "work in parallel" and "parallel processing" in a posting's description pass the
    exports' leak check and the page's guard; the vendor's own name is still caught."""
    from agents.naming import assert_client_clean
    from exports.policy import leak_hits

    description = (
        "You will work in parallel with the platform team. Experience with parallel processing "
        "of loads is a plus. Salesforce, MuleSoft and SAP experience required."
    )
    assert leak_hits(description) == ()
    assert_client_clean([description])
    rows: list[dict[str, Any]] = [
        {"n": 1, "cells": {"c1": description}, "links": {}, "link": None},
    ]
    SH.cap_rows(rows)
    assert SH.scrub_rows(rows) == 0 and rows[0]["cells"]["c1"] == description
    vendor = "Findings come from the parallel search API."
    assert leak_hits(vendor)
    with pytest.raises(Exception, match="G9|client"):
        assert_client_clean([vendor])
    # deny by default (security review 2026-10-29, F3): a capitalised "Parallel processing" at the
    # start of a sentence is withheld, one cell, never a leak
    sentence_start = "Parallel processing of loads is a plus."
    rows = [{"n": 1, "cells": {"c1": sentence_start}, "links": {}, "link": None}]
    assert SH.scrub_rows(rows) == 1


def test_a_long_text_is_scanned_not_remembered() -> None:
    """The memo of the guard keeps cut page cells; a whole 32,000-character description (the cell
    route) is scanned each time and never held."""
    from agents import naming

    naming._scrubbed_memo.cache_clear()
    big = "word " * 7000
    assert len(big) > naming.MEMO_MAX_CHARS
    assert naming.scrub_client_text(big) == (big, False)
    assert naming._scrubbed_memo.cache_info().currsize == 0
    small = "A short, clean sentence."
    naming.scrub_client_text(small)
    naming.scrub_client_text(small)
    info = naming._scrubbed_memo.cache_info()
    assert info.currsize == 1 and info.hits >= 1
