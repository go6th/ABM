"""Accounts, the Account workspace, deliverable downloads, /v1/me and Ask through the real
application on PostgreSQL, over the demo world the dev seed builds (SYNTHETIC: .example domains,
"[Name]" people). Two clients hold a Cigna account each; neither ever sees the other's."""

from __future__ import annotations

import io
import json
import zipfile
from typing import Any
from uuid import uuid4

import psycopg
import pytest

pytestmark = [pytest.mark.requires_postgres, pytest.mark.asyncio]

from tests.replica.conftest import _call, _plane  # noqa: E402


def _frames(body: str) -> list[tuple[str, dict[str, Any]]]:
    out = []
    for raw in body.strip().split("\n\n"):
        lines = raw.split("\n")
        event = lines[0].removeprefix("event: ")
        out.append((event, json.loads(lines[1].removeprefix("data: "))))
    return out


async def test_accounts_rows_carry_the_designs_sub_line_states(world: dict[str, Any]) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    r = await _call(plane, "GET", "/v1/accounts?limit=200")
    assert r.status_code == 200, r.text
    rows = {a["name"]: a for a in r.json()["accounts"]}
    assert (rows["Highmark"]["state"], rows["Highmark"]["new_signals"]) == ("new_signals", 2)
    assert (rows["Cambia Health"]["state"], rows["Cambia Health"]["new_signals"]) == (
        "new_signals",
        1,
    )
    assert rows["Cigna"]["state"] == "first_scan"
    assert rows["Medica"]["state"] == "no_change"
    assert rows["Highmark"]["lanes"] == ["Salesforce", "Data and AI"]
    assert rows["Highmark"]["cadence_label"] == "Weekly"
    assert rows["Highmark"]["last_researched_label"] == "2 days ago"


async def test_workspace_after_a_refresh_shows_what_changed_and_live_counts(
    world: dict[str, Any],
) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    r = await _call(plane, "GET", "/v1/accounts/highmark.example/workspace")
    assert r.status_code == 200, r.text
    ws = r.json()
    assert ws["first_scan"] is False and ws["changes"] is not None
    assert ws["subline"] == "Refreshed 2 days ago · Weekly schedule · Web: not used for answers"
    assert {row["grade"] for row in ws["changes"]["rows"]} <= {"confirmed", "probable"}
    assert ws["counts"] == {
        "signals": 9,
        "people": 41,
        "jobs": 27,
        "tech": 14,
        "news": 12,
        "mail_drive": 3,
        "sources": 86,
        "coverage_done": 24,
        "coverage_total": 33,
    }
    assert [f["name"] for f in ws["files"]] == [
        "Account summary",
        "People list",
        "Job postings",
        "Tech stack and teams",
        "Evidence",
    ]
    assert all(s["source_url"] for s in ws["narrative"]["sentences"])


async def test_workspace_on_a_first_scan_has_findings_and_no_changes(world: dict[str, Any]) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    ws = (await _call(plane, "GET", "/v1/accounts/cigna.example/workspace")).json()
    assert ws["first_scan"] is True and ws["changes"] is None
    assert [f["grade"] for f in ws["findings"]] == ["confirmed", "probable", "hypothesis"]
    assert ws["subline"].startswith("First scan · ")


async def test_a_colleagues_account_is_not_found(world: dict[str, Any]) -> None:
    # private workspaces (migration 0063): another user of the same client cannot open it
    plane = _plane(world["a"], "someone-else", world["payloads"])
    got = await _call(plane, "GET", "/v1/accounts/highmark.example/workspace")
    assert got.status_code == 404


async def test_one_file_downloads_as_itself_and_several_as_a_zip(world: dict[str, Any]) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    files = (await _call(plane, "GET", "/v1/accounts/highmark.example/deliverables")).json()[
        "files"
    ]
    one = await _call(
        plane, "GET", f"/v1/accounts/highmark.example/deliverables/download?ids={files[0]['id']}"
    )
    assert one.status_code == 200
    assert one.headers["content-disposition"] == f'attachment; filename="{files[0]["filename"]}"'
    assert one.content[:2] == b"PK"
    ids = ",".join(f["id"] for f in files[:3])
    several = await _call(
        plane, "GET", f"/v1/accounts/highmark.example/deliverables/download?ids={ids}"
    )
    names = zipfile.ZipFile(io.BytesIO(several.content)).namelist()
    assert sorted(names) == sorted(f["filename"] for f in files[:3])
    everything = await _call(
        plane, "GET", "/v1/accounts/highmark.example/deliverables/download?ids=all"
    )
    assert len(zipfile.ZipFile(io.BytesIO(everything.content)).namelist()) == 5


async def test_another_clients_account_and_files_are_not_found(world: dict[str, Any]) -> None:
    a = _plane(world["a"], "priya-test", world["payloads"])
    b = _plane(world["b"], "other-test", world["payloads"])
    files = (await _call(a, "GET", "/v1/accounts/highmark.example/deliverables")).json()["files"]
    r = await _call(b, "GET", "/v1/accounts/highmark.example/workspace")
    assert r.status_code == 404
    r = await _call(
        b, "GET", f"/v1/accounts/cigna.example/deliverables/download?ids={files[0]['id']}"
    )
    assert r.status_code == 404
    own = (await _call(b, "GET", "/v1/accounts/cigna.example/workspace")).json()
    theirs = (await _call(a, "GET", "/v1/accounts/cigna.example/workspace")).json()
    assert own["id"] != theirs["id"]


async def test_me_counts_the_users_own_accounts_against_the_cap(world: dict[str, Any]) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    me = (await _call(plane, "GET", "/v1/me")).json()
    assert (me["accounts_used"], me["account_cap"], me["accounts_total"]) == (4, 34, 4)
    other = _plane(world["a"], "nobody-test", world["payloads"])
    assert (await _call(other, "GET", "/v1/me")).json()["accounts_used"] == 0


async def test_ask_streams_a_cited_answer_and_keeps_it(world: dict[str, Any]) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    request_id = str(uuid4())
    body = {"request_id": request_id, "question": "Who owns ServiceNow at Cigna?"}
    r = await _call(plane, "POST", "/v1/chat/ask", json=body)
    frames = dict(_frames(r.text))
    assert "error" not in frames, frames
    answer, thread = frames["answer"], frames["thread"]
    assert answer["answered"] and answer["sources"] and answer["used"]["web"] == "not_used"
    assert answer["offer"]["account_domain"] == "cigna.example"
    got = (await _call(plane, "GET", f"/v1/chat/threads/{thread['thread_id']}")).json()
    assert got["turns"][0]["answer"] == answer and got["title"] == "Who owns ServiceNow at Cigna?"
    again = dict(
        _frames(
            (
                await _call(
                    plane, "POST", "/v1/chat/ask", json={**body, "thread_id": thread["thread_id"]}
                )
            ).text
        )
    )
    assert again["done"]["replayed"] is True and again["answer"] == answer
    from tests.conftest import ADMIN_DSN

    with psycopg.connect(ADMIN_DSN) as c:
        n = c.execute(
            "SELECT sum(questions) FROM question_categories WHERE tenant_id = %s AND category ="
            " 'people'",
            (world["a"],),
        ).fetchone()
    assert n is not None and n[0] >= 1


async def test_ask_refuses_another_clients_account_and_thread(world: dict[str, Any]) -> None:
    a = _plane(world["a"], "priya-test", world["payloads"])
    b = _plane(world["b"], "other-test", world["payloads"])
    first = dict(
        _frames(
            (
                await _call(
                    a,
                    "POST",
                    "/v1/chat/ask",
                    json={"request_id": str(uuid4()), "question": "What changed this week?"},
                )
            ).text
        )
    )
    r = await _call(
        b,
        "POST",
        "/v1/chat/ask",
        json={
            "request_id": str(uuid4()),
            "question": "x",
            "thread_id": first["thread"]["thread_id"],
        },
    )
    assert r.status_code == 404
    ws_a = (await _call(a, "GET", "/v1/accounts/highmark.example/workspace")).json()
    r = await _call(
        b,
        "POST",
        "/v1/chat/ask",
        json={"request_id": str(uuid4()), "question": "x", "account_id": ws_a["id"]},
    )
    assert r.status_code == 404


async def test_people_sheet_counts_filters_and_sorts_on_the_server(world: dict[str, Any]) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    base = "/v1/accounts/highmark.example/sheets/people"
    sheet = (await _call(plane, "GET", base)).json()
    assert {f["label"]: f["count"] for f in sheet["filters"]} == {
        "All": 41,
        "Priority A": 8,
        "Technology": 17,
        "Commercial": 11,
        "Changed this month": 2,
    }
    assert [r["n"] for r in sheet["rows"]] == list(range(1, 42))
    page2 = (await _call(plane, "GET", f"{base}?limit=10&offset=10")).json()
    assert [r["n"] for r in page2["rows"]] == list(range(11, 21))
    a = (await _call(plane, "GET", f"{base}?filter=priority_a&sort=title&dir=asc")).json()
    titles = [r["cells"]["title"] for r in a["rows"]]
    assert a["total"] == 8 and titles == sorted(titles)
    assert {r["cells"]["priority"] for r in a["rows"]} == {"A"}
    desc = (await _call(plane, "GET", f"{base}?filter=priority_a&sort=title&dir=desc")).json()
    assert [r["cells"]["title"] for r in desc["rows"]] == titles[::-1]


@pytest.mark.parametrize("query", ["sort=title;drop", "filter=nope", "dir=up", "sort=link"])
async def test_sheet_refuses_anything_outside_its_allow_lists(
    world: dict[str, Any], query: str
) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    r = await _call(plane, "GET", f"/v1/accounts/highmark.example/sheets/people?{query}")
    assert r.status_code == 422
    r = await _call(plane, "GET", "/v1/accounts/highmark.example/sheets/payroll")
    assert r.status_code == 404


async def test_sheet_export_is_a_real_workbook_of_the_filtered_rows(world: dict[str, Any]) -> None:
    from openpyxl import load_workbook

    plane = _plane(world["a"], "priya-test", world["payloads"])
    r = await _call(
        plane, "GET", "/v1/accounts/highmark.example/sheets/people/xlsx?filter=technology"
    )
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="Highmark-People-list.xlsx"'
    ws = load_workbook(io.BytesIO(r.content))["Sheet 1"]
    assert ws.max_row - 1 == 17
    assert [c.value for c in ws[1]][:3] == ["#", "Name", "Role"]


async def test_mail_and_drive_sheet_holds_only_the_callers_items(world: dict[str, Any]) -> None:
    owner = _plane(world["a"], "priya-test", world["payloads"])
    colleague = _plane(world["a"], "colleague-test", world["payloads"])
    mine = (await _call(owner, "GET", "/v1/accounts/highmark.example/sheets/mail-and-drive")).json()
    theirs = await _call(colleague, "GET", "/v1/accounts/highmark.example/sheets/mail-and-drive")
    assert mine["total"] == 3 and mine["private"] is True
    assert theirs.status_code == 404  # the account itself is the owner's (migration 0063)


async def test_another_clients_sheets_are_not_found(world: dict[str, Any]) -> None:
    b = _plane(world["b"], "other-test", world["payloads"])
    assert (await _call(b, "GET", "/v1/accounts/highmark.example/sheets/people")).status_code == 404
    assert (
        await _call(b, "GET", "/v1/accounts/highmark.example/sheets/people/xlsx")
    ).status_code == 404
    own = (await _call(b, "GET", "/v1/accounts/cigna.example/sheets/people")).json()
    assert own["total"] == 3


async def test_coverage_shows_every_tracker_item_with_its_status_and_reason(
    world: dict[str, Any],
) -> None:
    plane = _plane(world["a"], "priya-test", world["payloads"])
    ws = (await _call(plane, "GET", "/v1/accounts/highmark.example/workspace")).json()
    assert (ws["counts"]["coverage_done"], ws["counts"]["coverage_total"]) == (24, 33)
    sheet = (await _call(plane, "GET", "/v1/accounts/highmark.example/sheets/coverage")).json()
    assert sheet["total"] == 33
    assert {f["key"]: f["count"] for f in sheet["filters"]} == {
        "all": 33,
        "done": 24,
        "partial": 2,
        "open": 7,
        "high": 19,
    }
    by_status = {r["cells"]["status"]: r["cells"] for r in sheet["rows"]}
    assert by_status["Blocked"]["reason"] == "Source unavailable"  # never silently missing
    assert by_status["Not done"]["reason"] and by_status["Done"]["reason"] == ""
    assert all(r["cells"]["item"] and "_" not in r["cells"]["item"] for r in sheet["rows"])
    worst = await _call(
        plane, "GET", "/v1/accounts/highmark.example/sheets/coverage?sort=status&filter=open"
    )
    assert worst.json()["rows"][0]["cells"]["status"] == "Blocked"
    xlsx = await _call(plane, "GET", "/v1/accounts/highmark.example/sheets/coverage/xlsx")
    assert xlsx.status_code == 200  # the tracker export passes the client-file checks
    other = _plane(world["b"], "other-test", world["payloads"])
    gone = await _call(other, "GET", "/v1/accounts/highmark.example/sheets/coverage")
    assert gone.status_code == 404
