"""Deep-dive table exports pass the same checks as every other client file: no data tool name,
credit count or internal reference, no dash used as punctuation, no cell a spreadsheet would run
as a formula, and no private row in a workbook that could be shared (SYNTHETIC data only).

GUARANTEE(test_a_shareable_export_withholds_tool_names_and_neutralises_formulas)
GUARANTEE(test_private_rows_never_enter_a_shareable_workbook)
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any
from uuid import uuid4

import psycopg
import pytest
from openpyxl import load_workbook

from account_intelligence.replica import sheets as SH
from tests.replica.conftest import _call, _plane
from tests.replica.test_sheet_cells import _vendor

pytestmark = [pytest.mark.requires_postgres, pytest.mark.asyncio]


def _plant_signal(tenant: Any, domain: str, title: str, *, private: bool = False) -> None:
    """Add one signal to the account; ``private`` points it at a private copy of a finding (an
    internal fact from an uploaded file)."""
    import hashlib

    from tests.conftest import ADMIN_DSN

    with psycopg.connect(ADMIN_DSN) as c:
        acct, fact = c.execute(
            "SELECT a.id, (SELECT s.fact_id FROM signal_events s WHERE s.account_id = a.id LIMIT 1)"
            " FROM accounts a WHERE a.tenant_id = %s AND a.domain = %s",
            (tenant, domain),
        ).fetchone()  # type: ignore[misc]
        if private:
            copy = uuid4()
            c.execute(
                "INSERT INTO facts SELECT (jsonb_populate_record(NULL::facts, to_jsonb(f)"
                " || jsonb_build_object('id', %s::text, 'visibility', 'private',"
                " 'fingerprint', %s::text))).* FROM facts f WHERE f.id = %s",
                (copy, hashlib.sha256(copy.bytes).hexdigest(), fact),
            )
            c.execute(
                "INSERT INTO fact_evidence SELECT (jsonb_populate_record(NULL::fact_evidence,"
                " to_jsonb(e) || jsonb_build_object('id', gen_random_uuid()::text,"
                " 'fact_id', %s::text))).* FROM fact_evidence e WHERE e.fact_id = %s",
                (copy, fact),
            )
            fact = copy
        c.execute(
            "INSERT INTO signal_events (tenant_id, account_id, fingerprint, signal_type,"
            " event_type, entity, title, bucket, grade, lead, sensitive, fact_id, score,"
            " components, config_version, detected_at) VALUES (%s, %s, %s, 'news', 'new', 'x',"
            " %s, 1, 'probable', false, false, %s, 1, '[]', 'test', %s)",
            (
                tenant,
                acct,
                hashlib.sha256(title.encode()).hexdigest(),
                title,
                fact,
                datetime.now(UTC),
            ),
        )
        c.commit()


def _cells(content: bytes, sheet: str = "Sheet 1") -> list[Any]:
    ws = load_workbook(io.BytesIO(content))[sheet]
    return [c.value for row in ws.iter_rows(min_row=2) for c in row]


async def test_a_shareable_export_withholds_tool_names_and_neutralises_formulas(
    world: dict[str, Any],
) -> None:
    _plant_signal(world["a"], "highmark.example", f"Hiring surge seen by {_vendor()} crawl")
    _plant_signal(world["a"], "highmark.example", '=HYPERLINK("http://evil.example","open")')
    _plant_signal(world["a"], "highmark.example", "Data team grows — fast")
    plane = _plane(world["a"], "priya-test", world["payloads"])
    r = await _call(plane, "GET", "/v1/accounts/highmark.example/sheets/signals/xlsx")
    assert r.status_code == 200, r.text
    values = [v for v in _cells(r.content) if isinstance(v, str)]
    assert not any(_vendor().lower() in v.lower() for v in values)
    assert SH.WITHHELD in values
    assert '\'=HYPERLINK("http://evil.example","open")' in values
    assert "Data team grows, fast" in values
    book = load_workbook(io.BytesIO(r.content))
    assert all(c.data_type != "f" for ws in book for row in ws.iter_rows() for c in row)
    about = dict(book["About"].iter_rows(values_only=True))  # type: ignore[arg-type]
    assert about["Withheld"] == "1 cells held back for review"


async def test_private_rows_never_enter_a_shareable_workbook(world: dict[str, Any]) -> None:
    _plant_signal(world["a"], "cambiahealth.example", "From the call notes", private=True)
    plane = _plane(world["a"], "priya-test", world["payloads"])
    shown = (await _call(plane, "GET", "/v1/accounts/cambiahealth.example/sheets/signals")).json()
    titles = [r["cells"]["title"] for r in shown["rows"]]
    assert "From the call notes" in titles  # the user still sees it in the app
    r = await _call(plane, "GET", "/v1/accounts/cambiahealth.example/sheets/signals/xlsx")
    assert r.status_code == 200
    exported = _cells(r.content)
    assert "From the call notes" not in exported
    assert len([v for v in exported if v in titles]) == len(titles) - 1


async def test_the_mail_and_drive_export_is_the_owners_and_labelled_private(
    world: dict[str, Any],
) -> None:
    owner = _plane(world["a"], "priya-test", world["payloads"])
    colleague = _plane(world["a"], "colleague-test", world["payloads"])
    mine = await _call(owner, "GET", "/v1/accounts/highmark.example/sheets/mail-and-drive/xlsx")
    theirs = await _call(
        colleague, "GET", "/v1/accounts/highmark.example/sheets/mail-and-drive/xlsx"
    )
    book = load_workbook(io.BytesIO(mine.content))
    assert book.sheetnames == ["Private to you", "About"]
    assert book.properties.title == "Private to you"
    assert ("Note", "Private to you. Do not share.") in list(
        book["About"].iter_rows(values_only=True)
    )
    assert len(_cells(mine.content, "Private to you")) > 0
    assert theirs.status_code == 404  # the account itself is the owner's (migration 0063)
