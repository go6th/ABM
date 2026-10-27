"""The edges of the workbook tabs (``replica.file_tabs``) and of the formula values they show
(``replica.workbook_formulas``): Excel's rules where the subset meets them, bounded work on long
chains and huge ranges, a book cache that is the tenant's and reads a book once, oversized files
refused before they are opened, the file table's sort and paging, the tab response cap, and the
exports of the file table passing the client-file checks. SYNTHETIC data only."""

from __future__ import annotations

import asyncio
import hashlib
import io
import sys
import threading
import time
from collections import OrderedDict
from types import SimpleNamespace
from typing import Any
from uuid import UUID, uuid4

import pytest
from openpyxl import Workbook, load_workbook

from account_intelligence.replica import file_tabs as FT
from account_intelligence.replica import sheets as SH
from account_intelligence.replica.workbook_formulas import Evaluator, matcher


def _one(formula: str, **cells: Any) -> Any:
    """The value of ``formula`` in S!Z1 over the cells given as ``A1=...``."""
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    grid: dict[tuple[str, int, int], Any] = {("S", 1, 26): formula}
    for coord, value in cells.items():
        col, row = coordinate_from_string(coord)
        grid[("S", row, column_index_from_string(col))] = value
    return Evaluator(grid).value("S", 1, 26)


# ---- the evaluator against Excel -------------------------------------------------------------


def test_formulas_follow_excel_on_the_edges() -> None:
    # IF computes only the branch it picks: the builder's share formula with a zero base is 0
    assert _one("=IF(B1=0,0,A1/B1)", A1=3, B1=0) == 0
    assert _one("=IF(B1=0,0,A1/B1)", A1=2, B1=4) == 0.5
    assert _one("=A1/B1", A1=3, B1=0) is None  # #DIV/0! shows empty
    assert _one('=IF(A1>=2,"C","")', A1=1) == ""
    assert _one('=IF(TRUE,"y","n")') == "y" and _one("=IF(FALSE,1)") is False
    assert _one('=IF("maybe",1,2)') is None  # #VALUE!
    # COUNTA counts formula cells whatever they return, and every value argument
    assert _one("=COUNTA(A1:A3)", A1='=IF(1=1,"","x")', A2="x") == 2
    assert _one('=COUNTA("",1)') == 2
    # a cell that cannot be computed makes its dependents unknown, never a count of blanks
    assert _one('=COUNTIF(A1:A3,"")', A1="=VLOOKUP(1,B1:C2,2)", A2="x") is None
    assert _one("=A1+1", A1="=1/0") is None
    # comparisons: a number never equals text, text sorts after numbers, a blank is 0 or ""
    assert _one('=5="5"') is False and _one('="a">5') is True
    assert _one("=A1=0", A2=1) is True and _one('=A1=""', A2=1) is True
    assert _one('="abc"="ABC"') is True and _one("=TRUE>1") is True
    # arithmetic and joins
    assert _one("=A1", A2=1) == 0  # a reference to a blank cell shows 0
    assert _one('=(0.1+0.2)&""') == "0.3"
    assert _one('="3"+1') == 4 and _one('=""+1') is None and _one('="x"*2') is None
    assert _one("=1e400") is None  # #NUM!
    # SUM adds a range's numbers only; value arguments are converted
    assert _one("=SUM(A1:A3)", A1=True, A2=2, A3="3") == 2
    assert _one('=SUM(TRUE,"3")') == 4
    # AND and OR pass over a range's text and blanks
    assert _one("=AND(A1:A3)", A1=True, A2="text") is True
    assert _one("=OR(A1:A2)", A1="text") is None
    # a sheet the workbook does not have is Excel's #REF!
    assert _one("=COUNTA('Nope'!A1:A3)") is None
    assert Evaluator({("S", 1, 1): "=COUNTA('Empty'!A1:A3)"}, ["S", "Empty"]).value("S", 1, 1) == 0


def test_criteria_follow_excel_on_the_edges() -> None:
    # a number criterion never matches TRUE; TRUE matches only TRUE
    assert not matcher(1)(True) and matcher(1)(1.0) and matcher(1)("1")
    assert matcher(True)(True) and not matcher(True)(1)
    # number text is what Excel reads as a number, never "nan" or "1_000"
    assert matcher("nan")("NaN") and not matcher("nan")(float("nan"))
    assert not matcher("1_000")(1000) and matcher("1e3")(1000)
    # blank criteria
    assert matcher("")(None) and matcher("")("")
    assert matcher("=")(None) and not matcher("=")("")
    assert matcher("<>")("") and not matcher("<>")(None)
    # text comparisons compare text, case-insensitively, and never numbers
    assert matcher(">b")("C") and not matcher(">b")("a") and not matcher(">b")(5)
    # an empty criterion cell is not supported
    assert _one("=COUNTIF(A1:A2,B1)", A1="x") is None
    # an exact text criterion counted from the tally is the same count as the scan
    grid: dict[tuple[str, int, int], Any] = {("S", r, 1): w for r, w in enumerate("aAbB", start=1)}
    grid[("S", 1, 2)] = '=COUNTIF(A1:A4,"a")'
    grid[("S", 2, 2)] = '=COUNTIF(A1:A4,"=a*")'
    ev = Evaluator(grid)
    assert ev.value("S", 1, 2) == 2 and ev.value("S", 2, 2) == 2


def test_a_long_formula_chain_and_huge_ranges_stay_bounded() -> None:
    limit = sys.getrecursionlimit()
    # a chain of 5,000 formulas, each the next one plus 1, computed from its far end
    n = 5000
    grid: dict[tuple[str, int, int], Any] = {("S", r, 1): f"=A{r + 1}+1" for r in range(1, n)}
    grid[("S", n, 1)] = "=1"
    started = time.monotonic()
    ev = Evaluator(grid)
    assert ev.value("S", 1, 1) == n
    assert ev.value("S", 2500, 1) == n - 2499
    # a cycle longer than one step shows empty on every cell, never a number
    cycle = {("S", r, 1): f"=A{r % 300 + 1}+1" for r in range(1, 301)}
    ev = Evaluator(cycle)
    assert ev.value("S", 1, 1) is None and ev.value("S", 150, 1) is None
    # deep nesting and long sign runs are refused or read, never a RecursionError
    assert _one("=" + "(" * 300 + "1" + ")" * 300) is None
    assert _one("=" + "-" * 5000 + "1") == 1
    assert _one("=" + "+".join(["1"] * 500)) is None
    assert _one("=" + "1" * 9000) is None  # longer than Excel allows
    # a range over a whole sheet is read only as far as the sheet has cells
    assert _one('=COUNTIF(A1:A1048576,"")', A1="x", A3="y") == 1_048_576 - 2
    assert _one('=COUNTIFS(A1:A1048576,"",B1:B1048576,"")', A1="x", B2="y") == 1_048_574
    assert _one("=SUM(A1:Y1048576)", A1=2, C5=3) == 5
    assert _one("=COUNTA(A1:Y1048576)", A1=2, C5="=1") == 2
    assert _one("=SUM(A1:XFD1048576)", A1=2) is None  # the range holds the formula: a cycle
    # many counts over one 20,000-cell range read its values once
    rows = 20_000
    big: dict[tuple[str, int, int], Any] = {
        ("D", r, 1): f'=IF({r}>10000,"open","closed")' for r in range(1, rows + 1)
    }
    for k in range(200):
        big[("C", k + 1, 1)] = f'=COUNTIF(D!$A$1:$A${rows},"{"open" if k % 2 else "closed"}")'
        big[("C", k + 1, 2)] = f'=COUNTIF(D!$A$1:$A${rows},"*e*")+{k}'
    ev = Evaluator(big)
    assert all(ev.value("C", k + 1, 1) == 10_000 for k in range(200))
    assert all(ev.value("C", k + 1, 2) == rows + k for k in range(200))
    assert len(ev._ranges) == 1
    assert time.monotonic() - started < 20
    assert sys.getrecursionlimit() == limit


# ---- reading and the book cache --------------------------------------------------------------


def _workbook(rows: int = 3, title: str = "People") -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title
    ws.append(["Name", "Score"])
    for i in range(rows):
        ws.append([f"Synthetic {i}", i])
    wb.create_sheet("Summary").append(["Rows", f"=COUNTA('{title}'!A2:A{rows + 1})"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _Store:
    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}
        self.gets = 0

    async def get(self, key: str) -> bytes:
        self.gets += 1
        await asyncio.sleep(0)
        return self.objects[key]


def _delivered(tenant: UUID, body: bytes, sha: str | None = None, **kw: Any) -> FT.Delivered:
    digest = sha or hashlib.sha256(body).hexdigest()
    key = f"{tenant}/{uuid4()}/{uuid4()}/deliverables/{digest}.xlsx"
    return FT.Delivered(str(uuid4()), "people.xlsx", digest, key, **kw)


@pytest.fixture
def fresh_cache(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(FT, "_CACHE", OrderedDict())
    monkeypatch.setattr(FT, "_READING", {})


@pytest.mark.asyncio
@pytest.mark.usefixtures("fresh_cache")
async def test_the_book_cache_is_the_tenants_and_reads_a_book_once(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    body = _workbook()
    a, b = uuid4(), uuid4()
    store = _Store()
    mine = _delivered(a, body)
    store.objects[mine.object_key] = body
    book = await FT.book_of(store, a, mine, "People")
    assert book is not None and [t.name for t in book.tabs] == ["People", "Summary"]
    summary = book.tab("summary")
    assert summary is not None and summary.blocks[0].header == ("Rows", "3")
    # another tenant naming the same object and hash is refused before the cache is looked at
    stolen = FT.Delivered("x", "people.xlsx", mine.sha256, mine.object_key)
    assert await FT.book_of(store, b, stolen, "People") is None
    # the same file stored by another tenant is read for that tenant, under its own key
    theirs = _delivered(b, body)
    store.objects[theirs.object_key] = body
    gets = store.gets
    assert await FT.book_of(store, b, theirs, "People") is not None
    assert store.gets == gets + 1 and len(FT._CACHE) == 2
    assert {k[0] for k in FT._CACHE} == {str(a), str(b)}
    # a cached book is served without the store
    gets = store.gets
    assert await FT.book_of(store, a, mine, "People") is book and store.gets == gets
    # two requests for one uncached book read it once
    reads = 0
    real = FT.read
    gate = threading.Event()

    def counted(data: bytes, main: str | None) -> FT.Book:
        nonlocal reads
        reads += 1
        gate.wait(5)
        return real(data, main)

    monkeypatch.setattr(FT, "read", counted)
    other = _delivered(a, _workbook(5))
    store.objects[other.object_key] = _workbook(5)
    first = asyncio.ensure_future(FT.book_of(store, a, other, "People"))
    second = asyncio.ensure_future(FT.book_of(store, a, other, "People"))
    await asyncio.sleep(0.2)
    gate.set()
    got = await asyncio.gather(first, second)
    assert reads == 1 and got[0] is got[1] and got[0] is not None
    assert FT._READING == {}
    # bytes that do not match the recorded hash are shown but never cached
    wrong = _delivered(a, body, sha="0" * 64)
    store.objects[wrong.object_key] = body
    assert await FT.book_of(store, a, wrong, "People") is not None
    assert await FT.book_of(store, a, wrong, "People") is not None
    assert reads == 3
    # a file recorded as larger than the bound is not fetched
    gets = store.gets
    big = _delivered(a, body, size=FT.MAX_BYTES + 1)
    store.objects[big.object_key] = body
    assert await FT.book_of(store, a, big, "People") is None and store.gets == gets


@pytest.mark.asyncio
@pytest.mark.usefixtures("fresh_cache")
async def test_an_oversized_workbook_is_refused_before_it_is_opened(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from account_intelligence.api.routes import account_workspace as AW

    body = _workbook(50)
    opened = 0
    real = load_workbook

    def counting(*args: Any, **kw: Any) -> Any:
        nonlocal opened
        opened += 1
        return real(*args, **kw)

    monkeypatch.setattr("openpyxl.load_workbook", counting)
    monkeypatch.setattr(FT, "MAX_XML_BYTES", 1000)
    with pytest.raises(FT.TooLarge):
        FT.read(body, "People")
    assert opened == 0
    store = _Store()
    tenant = uuid4()
    found = _delivered(tenant, body)
    store.objects[found.object_key] = body
    assert await FT.book_of(store, tenant, found, "People") is None
    monkeypatch.setattr(FT, "MAX_XML_BYTES", 64 * 1024 * 1024)
    monkeypatch.setattr(FT, "MAX_CELLS", 20)
    with pytest.raises(FT.TooLarge):
        FT.read(body, "People")
    # bytes that are not a workbook give no tabs, never an error
    junk = _delivered(tenant, b"not a zip")
    store.objects[junk.object_key] = b"not a zip"
    assert await FT.book_of(store, tenant, junk, "People") is None
    # a file without the builder's row count is not read when only the file's table is wanted
    request = SimpleNamespace(
        app=SimpleNamespace(state=SimpleNamespace(container=SimpleNamespace(payloads=store)))
    )
    ident = SimpleNamespace(tenant_id=tenant)
    gets = store.gets
    unbuilt = _delivered(tenant, body)
    store.objects[unbuilt.object_key] = body
    got = await AW._read_book(request, ident, unbuilt, "people", builder_only=True)  # type: ignore[arg-type]
    assert got is None and store.gets == gets


def test_tab_keys_are_unique_and_never_end_in_a_hyphen() -> None:
    assert FT.tab_key("Platform Ranking") == "platform-ranking"
    assert not FT.tab_key("a" * 39 + " b").endswith("-")
    assert FT.tab_key("!!!") == "sheet"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Summary 2"
    wb.create_sheet("Summary")
    wb.create_sheet("Summary!")
    buf = io.BytesIO()
    wb.save(buf)
    keys = [t.key for t in FT.read(buf.getvalue(), None).tabs]
    assert keys == ["summary-2", "summary", "summary-3"]


# ---- the file table --------------------------------------------------------------------------


def _tab(cells: list[tuple[str, str, str | None]]) -> FT.Tab:
    rows = tuple(
        (FT.Cell(name), FT.Cell(score), FT.Cell("Open profile" if link else "", link))
        for name, score, link in cells
    )
    table = FT.Block("table", header=("Name", "Score", "LinkedIn"), rows=rows)
    return FT.Tab("people", "People", (table,), main=True)


ROWS = [
    ("Ann", "10", "https://www.linkedin.com/in/synthetic-ann"),
    ("bob", "", None),
    ("Cy", "2", None),
    ("Di", "10", None),
    ("Ed", "", None),
    ("Fay", "-3", None),
]


def test_the_file_table_sorts_blanks_last_and_pages_past_the_end() -> None:
    tab = _tab(ROWS)
    live = SH.SHEETS["people"]

    def names(sort: str | None, direction: str | None) -> list[str]:
        q = FT.grid_query(tab, live, sort, direction, 0, 500)
        return [r["cells"]["c1"] for r in FT.grid_page(tab, q)["rows"]]

    # numbers as numbers, ties in the file's order, blanks last in either direction
    assert names("c2", "asc") == ["Fay", "Cy", "Ann", "Di", "bob", "Ed"]
    assert names("c2", "desc") == ["Ann", "Di", "Cy", "Fay", "bob", "Ed"]
    assert names("c1", "asc") == ["Ann", "bob", "Cy", "Di", "Ed", "Fay"]
    assert names(None, "desc") == ["Fay", "Ed", "Di", "Cy", "bob", "Ann"]
    # a link column is not sortable, and an unknown sort or direction is refused
    q = FT.grid_query(tab, live, None, None, 0, 500)
    page = FT.grid_page(tab, q)
    assert [c["sortable"] for c in page["columns"]] == [True, True, False]
    assert page["rows"][0]["links"] == {"c3": "https://www.linkedin.com/in/synthetic-ann"}
    for sort, direction in (("c3", None), ("c9", None), ("_pos", None), (None, "up")):
        with pytest.raises(SH.BadSheetQuery):
            FT.grid_query(tab, live, sort, direction, 0, 5)
    # row numbers are absolute; an offset past the end is an empty page with the total
    q = FT.grid_query(tab, live, "c2", "asc", 2, 2)
    page = FT.grid_page(tab, q)
    assert [r["n"] for r in page["rows"]] == [3, 4] and page["total"] == 6
    q = FT.grid_query(tab, live, None, None, 100_000, 5)
    page = FT.grid_page(tab, q)
    assert page["rows"] == [] and page["total"] == 6 and page["offset"] == 100_000
    # the page is clamped like a stored table's
    assert FT.grid_query(tab, live, None, None, -4, 0).limit == 1
    assert FT.grid_query(tab, live, None, None, 0, 10**6).limit == SH.MAX_PAGE
    assert FT.grid_query(tab, live, None, None, -4, 5).offset == 0


@pytest.mark.asyncio
async def test_the_file_tables_exports_pass_the_client_file_checks(tmp_path: Any) -> None:
    from account_intelligence.replica import table_export as TX

    tab = _tab(
        [
            ("=HYPERLINK(1)", "1", None),
            ("Synthetic — Person", "2", "https://www.linkedin.com/in/synthetic"),
        ]
    )
    q = FT.grid_query(tab, SH.SHEETS["people"], None, None, 0, SH.MAX_PAGE)
    every = [r async for r in FT.iter_rows(tab, q)]
    assert [r["n"] for r in every] == [1, 2]
    body = b"".join([c async for c in TX.csv_chunks(q, FT.iter_rows(tab, q))]).decode("utf-8-sig")
    lines = body.splitlines()
    assert lines[1].startswith("1,'=HYPERLINK(1)")
    assert "Synthetic, Person" in lines[2] and "—" not in body
    assert "https://www.linkedin.com/in/synthetic" in lines[2]
    path = tmp_path / "out.xlsx"
    assert await TX.write_xlsx(path, q, "Synthetic", FT.iter_rows(tab, q)) == 2
    ws = load_workbook(path).worksheets[0]
    assert ws.cell(row=2, column=2).value == "'=HYPERLINK(1)"
    link = ws.cell(row=3, column=4).hyperlink
    assert link is not None and link.target == "https://www.linkedin.com/in/synthetic"


def test_a_tab_response_is_capped(monkeypatch: pytest.MonkeyPatch) -> None:
    rows = tuple((FT.Cell(f"r{i}"), FT.Cell("x")) for i in range(2500))
    first = FT.Block("table", header=("A", "B"), rows=rows, cut=7)
    second = FT.Block("table", header=("A", "B"), rows=rows[:10])
    lines = tuple(FT.Block("text", text=f"line {i}") for i in range(1500))
    tab = FT.Tab("t", "T", (first, *lines, second))
    out = FT.blocks_out(tab)
    tables = [b for b in out if b["kind"] == "table"]
    assert len(tables[0]["rows"]) == FT.SCREEN_ROWS and tables[0]["cut"] == 7 + 500
    assert len([b for b in out if b["kind"] == "text"]) == FT.SCREEN_LINES
    assert len(tables[1]["rows"]) == 10
    # past the tab's cell budget, the rest of every table is counted, not sent
    monkeypatch.setattr(FT, "SCREEN_CELLS", 100)
    tables = [b for b in FT.blocks_out(tab) if b["kind"] == "table"]
    assert len(tables[0]["rows"]) == 50 and tables[0]["cut"] == 7 + 2450
    assert tables[1]["rows"] == [] and tables[1]["cut"] == 10
    monkeypatch.setattr(FT, "SCREEN_CELLS", 10**9)
    monkeypatch.setattr(FT, "SCREEN_CHARS", 30)
    tables = [b for b in FT.blocks_out(tab) if b["kind"] == "table"]
    assert sum(len(t["rows"]) for t in tables) <= 10


# ---- the API ---------------------------------------------------------------------------------


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_the_file_table_route_refuses_bad_requests_and_pages_past_the_end() -> None:
    import psycopg

    from tests.conftest import ADMIN_DSN
    from tests.replica.test_file_sheets import _app, _get
    from tests.replica.test_file_tabs import _deliver_files, _world

    w, payloads, build = await _world()
    plane = _app(w, payloads, build)
    mapped = await _deliver_files(w, payloads, build)
    base = f"/v1/accounts/{w.domain}/sheets/people"
    assert (await _get(plane, base)).json()["from_file"] is True  # materialises the stored table
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        c.execute(
            "UPDATE view_rows SET cells = cells - 'note'"
            " WHERE run_id = %s AND table_kind = 'people'",
            (w.run,),
        )
    sheet = (await _get(plane, base)).json()
    assert sheet["from_file"] is True and sheet["columns"][0]["key"] == "c1"
    link = next(c["key"] for c in sheet["columns"] if c["kind"] == "url")
    for query in ("?sort=nope", f"?sort={link}", "?dir=up"):
        r = await _get(plane, base + query)
        assert r.status_code == 422, (query, r.text)
        assert "bad_sheet_query" in r.text
    page = (await _get(plane, f"{base}?offset=100000")).json()
    assert page["rows"] == [] and page["total"] == len(mapped["people"].model.rows)
    for tab in ("Summary", "a" * 49, "has_underscore"):
        assert (await _get(plane, f"{base}/tabs/{tab}")).status_code == 404
    ws = (await _get(plane, f"/v1/accounts/{w.domain}/workspace")).json()
    assert ws["counts"]["people"] == page["total"]
