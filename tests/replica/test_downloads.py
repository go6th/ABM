"""Everything collected is downloadable (Phase 4B section 6): a table exports whole, past a page
and past Excel's sheet limit, as CSV and as Excel; Download everything builds its archive in zip
parts with every table, the deliverables and a README, hands out short-lived links, serves byte
ranges, and is the asking user's alone."""

from __future__ import annotations

import csv
import io
import zipfile
from typing import Any
from uuid import uuid4

import psycopg
import pytest

from account_intelligence.db.session import tenant_txn
from account_intelligence.replica import downloads as D
from account_intelligence.replica import table_export as TX
from tests.conftest import ADMIN_DSN
from tests.replica.test_workspace_api import _call, _plane

pytestmark = [pytest.mark.requires_postgres, pytest.mark.asyncio]
PEOPLE = "/v1/accounts/highmark.example/sheets/people"


async def test_a_table_exports_whole_past_the_page_and_excel_limits(
    world: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl import load_workbook

    plane = _plane(world["a"], "priya-test", world["payloads"])
    total = (await _call(plane, "GET", f"{PEOPLE}?limit=1")).json()["total"]
    assert total > 12
    monkeypatch.setattr(TX, "PAGE", 5)  # many pages
    monkeypatch.setattr(TX, "EXCEL_ROWS", 6)  # a header and five rows per sheet
    got = await _call(plane, "GET", f"{PEOPLE}/csv")
    assert got.status_code == 200 and got.headers["content-type"].startswith("text/csv")
    rows = list(csv.reader(io.StringIO(got.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["#", "Name", "Title as reported"]
    assert [int(r[0]) for r in rows[1 : total + 1]] == list(range(1, total + 1))
    book = load_workbook(io.BytesIO((await _call(plane, "GET", f"{PEOPLE}/xlsx")).content))
    sheets = [ws for ws in book.worksheets if ws.title != "About"]
    assert len(sheets) == -(-total // 5)  # the table continues sheet after sheet
    assert sum(ws.max_row - 1 for ws in sheets) == total
    assert all([c.value for c in ws[1]][:2] == ["#", "Name"] for ws in sheets)
    about = dict(book["About"].iter_rows(values_only=True))
    assert about["Rows"] == total


async def test_download_everything_builds_parts_with_every_table_and_a_short_lived_link(
    world: dict[str, Any],
) -> None:
    from account_intelligence.api.routes.tasks import sources

    plane = _plane(world["a"], "priya-test", world["payloads"])
    started = await _call(
        plane,
        "POST",
        "/v1/accounts/highmark.example/download-everything",
        json={"request_id": str(uuid4())},
    )
    assert started.status_code == 202, started.text
    job = started.json()
    assert job["status"] == "queued" and job["parts"] == []
    queued = plane.app.state.container.scheduler.downloads
    assert (world["a"], D.UUID(job["job_id"])) in queued
    src = sources(plane.app.state.container)
    src.part_bytes = 20_000  # small parts: the archive splits
    from account_intelligence.db.session import acting_as

    with acting_as("priya-test"):
        assert await D.build(world["a"], D.UUID(job["job_id"]), src) == "ready"
    ready = (await _call(plane, "GET", f"/v1/downloads/{job['job_id']}")).json()
    assert ready["status"] == "ready" and len(ready["parts"]) > 1 and ready["links_expire_at"]
    names: list[str] = []
    for part in ready["parts"]:
        got = await _call(plane, "GET", part["url"])
        assert got.status_code == 200 and int(got.headers["content-length"]) == part["bytes"]
        names += zipfile.ZipFile(io.BytesIO(got.content)).namelist()
    assert "README.txt" in names
    assert {"Tables/People.csv", "Tables/People.xlsx", "Tables/Coverage.csv"} <= set(names)
    assert any(n.startswith("Deliverables/") for n in names)
    # a download that broke off continues with a byte range
    first = ready["parts"][0]
    tail = await _call(plane, "GET", first["url"], headers={"Range": "bytes=100-199"})
    assert tail.status_code == 206 and len(tail.content) == 100
    assert tail.headers["content-range"] == f"bytes 100-199/{first['bytes']}"
    # every status read renews the links: the old one stops working
    again = (await _call(plane, "GET", f"/v1/downloads/{job['job_id']}")).json()
    assert (await _call(plane, "GET", first["url"])).status_code == 404
    assert (await _call(plane, "GET", again["parts"][0]["url"])).status_code == 200
    # an expired link is refused
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute(
            "UPDATE download_jobs SET token_expires_at = now() - interval '1 second' WHERE id = %s",
            (job["job_id"],),
        )
    assert (await _call(plane, "GET", again["parts"][0]["url"])).status_code == 404
    # it is the user's own: a colleague does not see it
    colleague = _plane(world["a"], "colleague-test", world["payloads"])
    assert (await _call(colleague, "GET", f"/v1/downloads/{job['job_id']}")).status_code == 404
    async with tenant_txn(world["a"], user="colleague-test", system=False) as s:
        assert await D.view(s, D.UUID(job["job_id"]), D.datetime.now(D.UTC)) is None


async def test_a_build_that_died_is_taken_over_and_an_abandoned_job_ends_failed(
    world: dict[str, Any],
) -> None:
    """Staging 2026-10-20: the instance building an archive was killed, the retry found the job
    running and skipped it, and the page waited for ever. A running job with a fresh beat is busy
    (the queue retries), one whose beat stopped is taken over, and one nobody builds any more is
    ended failed when its status is read."""
    from account_intelligence.api.routes.tasks import sources
    from account_intelligence.db.session import acting_as

    plane = _plane(world["a"], "priya-test", world["payloads"])
    path = "/v1/accounts/highmark.example/download-everything"
    job = (await _call(plane, "POST", path, json={"request_id": str(uuid4())})).json()
    job_id = D.UUID(job["job_id"])
    src = sources(plane.app.state.container)
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute(
            "UPDATE download_jobs SET status = 'running', started_at = now() WHERE id = %s",
            (job["job_id"],),
        )
    with acting_as("priya-test"):
        assert await D.build(world["a"], job_id, src) == "busy"
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute(
            "UPDATE download_jobs SET started_at = now() - interval '5 minutes' WHERE id = %s",
            (job["job_id"],),
        )
    with acting_as("priya-test"):
        assert await D.build(world["a"], job_id, src) == "ready"
    assert (await _call(plane, "GET", f"/v1/downloads/{job_id}")).json()["status"] == "ready"

    left = (await _call(plane, "POST", path, json={"request_id": str(uuid4())})).json()
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute(
            "UPDATE download_jobs SET created_at = now() - interval '1 hour' WHERE id = %s",
            (left["job_id"],),
        )
    seen = (await _call(plane, "GET", f"/v1/downloads/{left['job_id']}")).json()
    assert seen["status"] == "failed" and seen["failure"] == D.STOPPED
    with acting_as("priya-test"):
        assert await D.build(world["a"], D.UUID(left["job_id"]), src) == "skipped"
