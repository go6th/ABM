"""Once an account's files are delivered, its People, Job postings, Tech stack, Platform picture
and Coverage sheets are the files' own tables: the same columns, headers and order as the
workbooks and the Account Summary (``views.tables``, built from the export configuration), and the
same cells, row for row, in the file's order (``views.rowsources``, materialised once per run).
Before a delivery the live sheets stand in. SYNTHETIC data only (the export test world)."""

from __future__ import annotations

import csv
import hashlib
import io
import json
from typing import Any

import psycopg
import pytest
import pytest_asyncio
from openpyxl import load_workbook

from account_intelligence.grading.config import load_grading_config
from account_intelligence.replica import sheets as SH
from exports.builder import map_snapshot
from exports.config import load_export_config
from views.rowsources import rows_of
from views.tables import specs

pytestmark = [pytest.mark.requires_postgres, pytest.mark.asyncio]
CFG, GRADING = load_export_config(), load_grading_config()


@pytest_asyncio.fixture(scope="module")
async def world() -> Any:
    from tests.conftest import ADMIN_DSN
    from tests.exports.test_db import builder, seed_world

    w, payloads = await seed_world()
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        c.execute(
            "UPDATE account_runs SET status = 'COMPLETED', finished_at = now() WHERE id = %s",
            (w.run,),
        )
        c.execute("UPDATE accounts SET name = 'Synthetic DB Ltd' WHERE id = %s", (w.account,))
    return w, payloads, builder(w, payloads)


def _deliver(w: Any, rows: dict[str, int]) -> None:
    """The run's delivered files, as the builder stores them: one row per file whose manifest
    counts the file's rows (the sheets key on the run and check those counts)."""
    from tests.conftest import ADMIN_DSN

    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        for kind, n in rows.items():
            h = hashlib.sha256(f"{w.run}{kind}".encode()).hexdigest()
            ext = "docx" if kind == "summary" else "xlsx"
            c.execute(
                "INSERT INTO deliverables (tenant_id, account_id, run_id, set_hash, kind, version,"
                " filename, sha256, size_bytes, object_key, manifest_json, content_hash,"
                " snapshot_hash, theme_hash, rules_hash, built_at, owner_user_id) VALUES"
                " (%s, %s, %s, %s, %s, 1, %s, %s, 10, %s, %s, %s, %s, %s, %s, now(),"
                " coalesce((SELECT owner_user_id FROM accounts WHERE id = %s), 'test-owner'))"
                " ON CONFLICT DO NOTHING",
                (
                    w.tenant,
                    w.account,
                    w.run,
                    h,
                    kind,
                    f"{kind}.{ext}",
                    h,
                    f"{w.tenant}/{w.account}/{w.run}/deliverables/{h}.{ext}",
                    json.dumps({"rows": n, "content_hash": h}),
                    h,
                    h,
                    h,
                    h,
                    w.account,
                ),
            )


async def _delivered(w: Any, build: Any) -> dict[str, Any]:
    """The run's client tables (the mapping the files print) after its files are delivered."""
    content = await build.reader.read(str(w.tenant), str(w.run))
    mapped, _ = map_snapshot(content, CFG, GRADING)
    expected = rows_of(mapped)
    files = {k: len(expected[k]) for k in ("people", "jobs", "tech")}
    _deliver(w, {**files, "summary": len(mapped["summary"].row_ids)})
    return expected  # type: ignore[return-value]


def _app(w: Any, payloads: Any, build: Any) -> Any:
    from tests.phase3.test_ws4_views import _app as ws4_app

    return ws4_app(w, payloads, build)


async def _get(plane: Any, path: str, **kw: Any) -> Any:
    from tests.phase3.test_ws4_views import _get as ws4_get

    return await ws4_get(plane, path, **kw)


def _text(value: Any) -> str:
    return "" if value is None else str(value)


async def test_sheets_show_the_delivered_files_tables(world: Any) -> None:
    w, payloads, build = world
    plane = _app(w, payloads, build)
    base = f"/v1/accounts/{w.domain}/sheets"

    # before a delivery: the live sheets, and an empty Platform picture with the file's columns
    live = await _get(plane, f"{base}/people")
    assert live.status_code == 200, live.text
    assert live.json()["from_file"] is False
    picture = (await _get(plane, f"{base}/facts")).json()
    assert picture["from_file"] is True and picture["total"] == 0
    assert [c["label"] for c in picture["columns"]] == list(
        CFG.rules.documents.tables["platform_picture"]
    )

    expected = await _delivered(w, build)
    assert all(expected[t] for t in ("people", "jobs", "tech", "facts", "coverage")), {
        t: len(v) for t, v in expected.items()
    }
    ws = (await _get(plane, f"/v1/accounts/{w.domain}/workspace")).json()
    for key, table in SH.FILE_TABLES.items():
        spec = specs()[table]  # type: ignore[index]
        r = await _get(plane, f"{base}/{key}", params={"limit": 500})
        assert r.status_code == 200, (key, r.text)
        sheet = r.json()
        assert sheet["from_file"] is True, key
        # the file's headers, in the file's order
        assert [c["label"] for c in sheet["columns"]] == [c.header for c in spec.exported_columns]
        assert [c["key"] for c in sheet["columns"]] == [c.key for c in spec.exported_columns]
        # the file's rows, in the file's order, cell for cell
        want = expected[table]
        assert sheet["total"] == len(want), key
        for got, row in zip(sheet["rows"], want, strict=True):
            for c in spec.exported_columns:
                value = _text(row.cells.get(c.key))
                if c.kind == "url" and value:
                    assert got["links"][c.key] == value, (key, c.key)
                    shown = SH._link_text(CFG.rules, table, c.key) or value
                    assert got["cells"][c.key] == shown, (key, c.key)
                elif len(value) > SH.CELL_CAP:
                    # a long cell is cut on the page and recorded; the whole cell is one
                    # request away (the cell route) and in the file
                    assert got["cells"][c.key] == value[: SH.CELL_CAP - 1].rstrip() + "\u2026"
                    assert got["cut"][c.key] == len(value), (key, c.key)
                else:
                    assert got["cells"][c.key] == value, (key, c.key)
        counted = "tech" if key == "technology" else key
        if counted in ws["counts"]:
            assert ws["counts"][counted] == len(want), key


async def test_file_sheets_sort_filter_and_export_like_the_file(world: Any) -> None:
    w, payloads, build = world
    plane = _app(w, payloads, build)
    base = f"/v1/accounts/{w.domain}/sheets/people"
    first = (await _get(plane, base)).json()
    if not first["from_file"]:
        await _delivered(w, build)
        first = (await _get(plane, base)).json()
    assert first["from_file"] and first["sort"] == SH.FILE_ORDER.key
    assert first["sheet_name"] == CFG.rules.workbooks.people.sheet
    names = [r["cells"]["name"] for r in first["rows"]]
    desc = (await _get(plane, base, params={"sort": "name", "dir": "desc"})).json()
    assert [r["cells"]["name"] for r in desc["rows"]] == sorted(names, reverse=True)
    a = (await _get(plane, base, params={"filter": "priority_a"})).json()
    assert {r["cells"]["priority"] for r in a["rows"]} <= {"A"}
    assert a["total"] == sum(1 for r in first["rows"] if r["cells"]["priority"] == "A")
    # a live sheet's sort or filter is refused on the file's table
    assert (await _get(plane, base, params={"filter": "technology"})).status_code == 422

    headers = [c.header for c in specs()["people"].exported_columns]
    xlsx = await _get(plane, f"{base}/xlsx")
    assert xlsx.status_code == 200, xlsx.text
    book = load_workbook(io.BytesIO(xlsx.content))
    sheet = book[book.sheetnames[0]]
    assert [c.value for c in sheet[1]] == ["#", *headers]
    linked = headers.index("LinkedIn") + 2
    if first["rows"] and first["rows"][0]["links"].get("profile_url"):
        cell = sheet.cell(row=2, column=linked)
        assert cell.value == CFG.rules.sheet_texts.people.link_text
        assert cell.hyperlink is not None
        assert cell.hyperlink.target == first["rows"][0]["links"]["profile_url"]
    body = (await _get(plane, f"{base}/csv")).content.decode("utf-8-sig")
    lines = list(csv.reader(io.StringIO(body)))
    assert lines[0] == ["#", *headers]
    if first["rows"] and first["rows"][0]["links"].get("profile_url"):
        assert lines[1][linked - 1] == first["rows"][0]["links"]["profile_url"]


async def test_a_file_that_does_not_hold_the_materialised_rows_keeps_the_live_sheet() -> None:
    """A delivered workbook whose stored row count is not the materialised table's (a set not
    made by this builder, such as the demo seed's stand-in files) is never shown as the file's
    table: the live sheet stands in."""
    from tests.conftest import ADMIN_DSN
    from tests.exports.test_db import builder, seed_world

    w, payloads = await seed_world()
    with psycopg.connect(ADMIN_DSN, autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        c.execute(
            "UPDATE account_runs SET status = 'COMPLETED', finished_at = now() WHERE id = %s",
            (w.run,),
        )
    build = builder(w, payloads)
    plane = _app(w, payloads, build)
    _deliver(w, {"people": 9999})
    sheet = (await _get(plane, f"/v1/accounts/{w.domain}/sheets/people")).json()
    assert sheet["from_file"] is False
    assert [c["label"] for c in sheet["columns"]] == [c.label for c in SH.SHEETS["people"].cols]
