"""WS4 views and exports (VW1 to VW6, TN5). SYNTHETIC data only.

Pure: query allow-lists, parity on the export corpus and on generated snapshots (a mutated file
is caught), render-time exclusions, branding inputs. On PostgreSQL through the real application:
the five tables against a hand-written oracle (filter, sort, search, keyset pagination), injection
through every parameter, saved views, and every export kind delivered only with parity.

Lying scenarios: an export whose cell differs from the table, a dropped or reordered row, a CSV
cell that would run a formula, a saved view naming a removed column, a logo that is an SVG or a
polyglot, an accent no one can read, a display name carrying a script, a run read after the
tenant excluded someone.
"""

from __future__ import annotations

import asyncio
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any
from uuid import uuid4

import psycopg
import pytest
import pytest_asyncio
from hypothesis import HealthCheck, given, settings
from openpyxl import load_workbook

from account_intelligence.grading.config import load_grading_config
from exports.builder import ExportBuilder, map_snapshot
from exports.config import load_export_config
from exports.readers.exclusions import apply_exclusions
from tenancy.settings import ExclusionFilter, Exclusions, PeopleScope, TenantSettings
from tests.exports.corpus import rich, sparse
from tests.fakes.exports import FakeReader, InProcessRenderer, MemoryStore
from views import branding, query
from views.packages import csv_bytes
from views.parity import check, check_package
from views.rowsources import rows_of
from views.tables import TABLES, specs

ROOT = Path(__file__).resolve().parents[2]
CFG, GRADING = load_export_config(), load_grading_config()
SCAN = ROOT / ".claude/skills/account-intelligence-platform/scripts/scan_client_output.py"


def _assembled(content: Any) -> tuple[dict[str, Any], dict[str, tuple[str, bytes]]]:
    builder = ExportBuilder(FakeReader(content), MemoryStore(), InProcessRenderer(), CFG, GRADING)
    out = builder.assemble(content)
    assert not out.errors, out.errors
    rows = {k: [r.cells for r in v] for k, v in rows_of(out.mapped).items()}
    files = {k: (name.rsplit(".", 1)[1], data) for k, (name, data) in out.files.items()}
    return rows, files


# ---------------------------------------------------------------- query allow-lists (VW3)


@pytest.mark.parametrize(
    ("raw", "code"),
    [
        (
            {"filters": [{"column": "name; DROP TABLE facts", "op": "eq", "value": "x"}]},
            "unknown_column",
        ),
        ({"filters": [{"column": "cells->>'x'", "op": "eq", "value": "x"}]}, "unknown_column"),
        (
            {"filters": [{"column": "profile_url", "op": "eq", "value": "x"}]},
            "operator_not_allowed",
        ),
        ({"filters": [{"column": "name", "op": "in", "value": "x"}]}, "operator_not_allowed"),
        ({"filters": [{"column": "level", "op": "in", "value": []}]}, "bad_in_values"),
        ({"filters": [{"column": "level", "op": "in", "value": ["x"] * 51}]}, "bad_in_values"),
        ({"filters": [{"column": "name", "op": "eq", "value": "x" * 201}]}, "bad_value"),
        ({"filters": [{"column": "name", "op": "is_empty", "value": "x"}]}, "value_not_allowed"),
        ({"filters": [{"column": "name", "op": "eq", "value": ""}]}, "bad_value"),
        ({"filters": [{"column": "name", "op": "eq", "value": "x"}] * 9}, "too_many_filters"),
        ({"sort": [{"column": "profile_url"}]}, "unsortable_column"),
        ({"sort": [{"column": "name"}, {"column": "name"}]}, "duplicate_sort"),
        (
            {"sort": [{"column": "name"}, {"column": "title"}, {"column": "level"}]},
            "too_many_sorts",
        ),
        ({"search": "x" * 201}, "search_too_long"),
        ({"page_size": 201}, "page_too_large"),
    ],
)
def test_query_validation_refuses_everything_outside_the_allow_lists(
    raw: dict[str, Any], code: str
) -> None:
    with pytest.raises(query.QueryError) as refused:
        query.validate(query.parse(raw), specs()["people"])
    assert refused.value.code == code


def test_malformed_specs_and_cursors_are_refused() -> None:
    for raw in (
        {"filters": "x"},
        {"sql": "SELECT 1"},
        {"filters": [{"column": "a"}]},
        {"filters": [{"column": "name", "op": "like", "value": "x"}]},
    ):
        with pytest.raises(query.QueryError):
            query.parse(raw)
    jobs = specs()["jobs"]
    with pytest.raises(query.QueryError):
        query.validate(
            query.parse({"filters": [{"column": "posted_at", "op": "gte", "value": "yesterday"}]}),
            jobs,
        )
    v = query.validate(query.parse({"sort": [{"column": "name"}]}), specs()["people"])
    for bad in (
        "!!!",
        query.encode_cursor([], 1),
        query.encode_cursor(["a"], -1),
        query.encode_cursor([True], 1),
        "e30",
    ):
        with pytest.raises(query.QueryError):
            query.compile_query(v, run_id=str(uuid4()), cursor=bad)
    sql, params = query.compile_query(
        query.validate(
            query.parse(
                {"filters": [{"column": "name", "op": "contains", "value": "50%_off'; DROP"}]}
            ),
            specs()["people"],
        ),
        run_id=str(uuid4()),
    )
    assert "DROP" not in sql and params["f0"] == "50\\%\\_off'; DROP"


# ---------------------------------------------------------------- parity (VW1)


@pytest.mark.parametrize("make", [rich, sparse])
def test_exported_rows_equal_api_rows_and_a_mutation_is_caught(make: Any) -> None:
    rows, files = _assembled(make())
    verdicts = check_package(dict(specs()), rows, files)
    assert all(v.ok for v in verdicts.values()), {k: v.problems for k, v in verdicts.items()}
    if not rows["people"]:
        return
    # a cell changed in the workbook
    ext, data = files["people"]
    wb = load_workbook(io.BytesIO(data))
    wb["People"]["C2"].value = "Synthetic forged title"
    buf = io.BytesIO()
    wb.save(buf)
    bad = check(specs()["people"], rows["people"], "xlsx", buf.getvalue())
    assert not bad.ok and "title differs" in bad.problems[0]
    # a dropped row, and reordered rows
    assert not check(specs()["people"], rows["people"][1:], "xlsx", data).ok
    swapped = [rows["people"][1], rows["people"][0], *rows["people"][2:]]
    assert not check(specs()["people"], swapped, "xlsx", data).ok
    # a changed word in the Account Summary's lane table
    ext, doc = files["summary"]
    finding = str(rows["facts"][-1]["finding"])
    forged = _replace_docx(doc, finding, finding + " forged")
    assert not check(specs()["facts"], rows["facts"], "docx", forged).ok
    # an unreadable file
    assert not check(specs()["people"], rows["people"], "xlsx", b"not a workbook").ok


def _replace_docx(data: bytes, old: str, new: str) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "word/document.xml":
                body = body.replace(old.encode(), new.encode(), 1)
            dst.writestr(item, body)
    return out.getvalue()


def test_csv_parity_and_formula_neutralisation() -> None:
    rows, _ = _assembled(rich())
    hostile = [r for r in rows["people"] if str(r["name"]).startswith("@")]
    assert hostile, "the rich corpus carries a formula-like name"
    body = csv_bytes(specs()["people"], rows["people"])
    assert b"\n'@SUM" in body or b",'@SUM" in body
    assert check(specs()["people"], rows["people"], "csv", body).ok
    mutated = body.replace(b"Synthetic Director 0", b"Synthetic Director X", 1)
    assert not check(specs()["people"], rows["people"], "csv", mutated).ok


try:
    from tests.exports.test_properties import fact_sets, postings
except ImportError:  # pragma: no cover - the strategies live with the export tests
    fact_sets = postings = None


@settings(max_examples=12, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(fact_sets())
def test_generated_accounts_export_exactly_their_rows(content: Any) -> None:
    builder = ExportBuilder(FakeReader(content), MemoryStore(), InProcessRenderer(), CFG, GRADING)
    out = builder.assemble(content)
    if out.errors:
        return  # the builder refused the snapshot: nothing is delivered, nothing to compare
    rows = {k: [r.cells for r in v] for k, v in rows_of(out.mapped).items()}
    files = {k: (n.rsplit(".", 1)[1], d) for k, (n, d) in out.files.items()}
    verdicts = check_package(dict(specs()), rows, files)
    assert all(v.ok for v in verdicts.values()), {k: v.problems for k, v in verdicts.items()}


@settings(max_examples=12, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(postings())
def test_generated_postings_export_exactly_their_rows(case: Any) -> None:
    content, _ = case
    builder = ExportBuilder(FakeReader(content), MemoryStore(), InProcessRenderer(), CFG, GRADING)
    out = builder.assemble(content)
    if out.errors:
        return
    rows = {k: [r.cells for r in v] for k, v in rows_of(out.mapped).items()}
    ext, data = out.files["jobs"][0].rsplit(".", 1)[1], out.files["jobs"][1]
    assert check(specs()["jobs"], rows["jobs"], ext, data).ok


def test_every_row_links_to_its_source_or_says_it_has_none() -> None:
    m, _ = map_snapshot(rich(), CFG, GRADING)
    rows = rows_of(m)
    assert all(r.source_state == "linked" for r in rows["facts"]), "a fact without a source"
    assert all(r.source_state == "no_source" for r in rows["coverage"])
    for table in ("people", "jobs", "tech", "facts"):
        for r in rows[table]:
            if r.source_state == "linked":
                assert r.source_code and r.source_url and r.source_url.startswith("http")


def test_exclusions_at_render_time_drop_people_and_forbid_companies() -> None:
    content = rich()
    target = content.people[0]
    rule = ExclusionFilter.of(
        TenantSettings(
            platform_lanes=("sap",),
            people_scope=PeopleScope(levels_in_scope=("VP",)),
            exclusions=Exclusions(people=(target.profile_url,), companies=("partner.example",)),
        )
    )
    trimmed = apply_exclusions(content, rule)
    assert all(p.id != target.id for p in trimmed.people)
    assert (target.name, False) in trimmed.people_names or not content.people_names
    assert "partner.example" in trimmed.forbidden_terms
    m, _ = map_snapshot(trimmed, CFG, GRADING)
    names = [r.cells["name"] for r in rows_of(m)["people"]]
    assert target.name not in names
    assert apply_exclusions(content, ExclusionFilter(frozenset(), frozenset())) is content


# ---------------------------------------------------------------- branding (VW6)


def _png(w: int = 40, h: int = 20, colour: str = "red") -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), colour).save(buf, format="PNG")
    return buf.getvalue()


def test_branding_refuses_every_hostile_input() -> None:
    from PIL import Image

    lim = branding.load_limits()
    svg = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
    gif = b"GIF89a" + b"\x00" * 40
    for bad, code in (
        (svg, "logo_type_refused"),
        (gif, "logo_type_refused"),
        (b"<?xml version='1.0'?><svg/>", "logo_type_refused"),
        (b"", "logo_too_large"),
        (b"\x89PNG\r\n\x1a\n" + b"\x00" * 50, "logo_unreadable"),
        (b"x" * (lim.max_logo_bytes + 1), "logo_too_large"),
    ):
        with pytest.raises(branding.BrandingRefused) as refused:
            branding.process_logo(bad)
        assert refused.value.code == code
    big = io.BytesIO()
    Image.new("RGB", (lim.max_logo_side + 1, 10), "red").save(big, format="PNG")
    with pytest.raises(branding.BrandingRefused) as wide:
        branding.process_logo(big.getvalue())
    assert wide.value.code == "logo_dimensions_refused"
    # a polyglot: a valid PNG with a script appended; the re-encoded logo carries none of it
    polyglot = _png() + b"<script>alert(1)</script><?php system('id'); ?>"
    logo = branding.process_logo(polyglot)
    assert b"<script>" not in logo.png and b"php" not in logo.png
    assert logo.png.startswith(branding.PNG_MAGIC) and (logo.width, logo.height) == (40, 20)
    jpeg = io.BytesIO()
    Image.new("RGB", (30, 30), "blue").save(jpeg, format="JPEG")
    assert branding.process_logo(jpeg.getvalue()).png.startswith(branding.PNG_MAGIC)
    # accent: a token with enough contrast
    assert branding.accent("#1F4E79") == "#1f4e79"
    for bad in ("#ffff00", "#eeeeee", "red", "#12345", "url(javascript:1)", "#1f4e79;}"):
        with pytest.raises(branding.BrandingRefused):
            branding.accent(bad)
    # display name: no markup, no forbidden terms, bounded
    assert branding.display_name("  Synthetic   Partners  ") == "Synthetic Partners"
    from account_intelligence.connectors.registry_internal import vendor_names
    from account_intelligence.connectors.settings import load_providers_config

    vendor = sorted(n for n in vendor_names(load_providers_config()).values() if n != "Parallel")[0]
    for bad in (
        "<script>alert(1)</script>",
        "x" * 61,
        "   ",
        f"Powered by {vendor}",  # a provider name (from config; never written in source)
        "1,000 credits remaining",
        "{{evil}}",
    ):
        with pytest.raises(branding.BrandingRefused):
            branding.display_name(bad)
    assert branding.tokens(None)["attribution"] == "6th agents"
    assert branding.contrast_with_white("#000000") > 20


def test_document_header_carries_the_brand_and_the_footer_keeps_the_attribution() -> None:
    content = rich().model_copy(update={"brand_name": "Synthetic Partners"})
    _, files = _assembled(content)
    with zipfile.ZipFile(io.BytesIO(files["summary"][1])) as z:
        headers = b"".join(z.read(n) for n in z.namelist() if n.startswith("word/header"))
        footers = b"".join(z.read(n) for n in z.namelist() if n.startswith("word/footer"))
    assert b"Synthetic Partners" in headers
    assert b"6th agents" in footers


# ---------------------------------------------------------------- on PostgreSQL

pytestmark_db = [pytest.mark.requires_postgres]


@pytest_asyncio.fixture(scope="module")
async def seeded() -> Any:
    from tests.exports.test_db import builder, seed_world

    w, payloads = await seed_world()
    with psycopg.connect(_admin(), autocommit=True) as c:
        c.execute("SET session_replication_role = replica")
        c.execute(
            "UPDATE account_runs SET status = 'COMPLETED', finished_at = now() WHERE id = %s",
            (w.run,),
        )
        c.execute("UPDATE accounts SET name = 'Synthetic DB Ltd' WHERE id = %s", (w.account,))
    return w, payloads, builder(w, payloads)


def _admin() -> str:
    from tests.conftest import ADMIN_DSN

    return ADMIN_DSN


def _app(w: Any, payloads: Any, build: Any, role: str = "admin", user: str = "test-owner") -> Any:
    from datetime import UTC, datetime

    from account_intelligence.api.auth import Identity, identity
    from exports.verify import LibreOffice
    from tests.phase3.conftest import add_member, make_app

    add_member(w.tenant, user, role)
    plane = make_app()
    plane.app.state.container.builder = build
    plane.app.state.container.payloads = payloads
    build.pdf = LibreOffice()
    now = datetime.now(UTC)
    plane.clock.current = now

    async def ident() -> Identity:
        return Identity(user, w.tenant, role, now)

    plane.app.dependency_overrides[identity] = ident
    return plane


async def _get(plane: Any, path: str, **kw: Any) -> Any:
    import httpx

    t = httpx.ASGITransport(app=plane.app, raise_app_exceptions=False)
    async with httpx.AsyncClient(transport=t, base_url="https://app.test") as c:
        return await c.get(path, **kw)


async def _post(plane: Any, method: str, path: str, **kw: Any) -> Any:
    import httpx

    t = httpx.ASGITransport(app=plane.app, raise_app_exceptions=False)
    async with httpx.AsyncClient(transport=t, base_url="https://app.test") as c:
        return await c.request(method, path, **kw)


async def _all_rows(
    plane: Any, domain: str, table: str, spec: dict[str, Any] | None = None, page: int = 2
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    cursor = None
    # LOOP(test): cap=500 pages; exit=no next cursor
    for _ in range(500):
        params = {"spec": json.dumps({**(spec or {}), "page_size": page})}
        if cursor:
            params["cursor"] = cursor
        r = await _get(plane, f"/v1/accounts/{domain}/tables/{table}", params=params)
        assert r.status_code == 200, r.text
        body = r.json()
        rows.extend(body["rows"])
        cursor = body["next_cursor"]
        if not cursor:
            break
    return rows


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_tables_match_the_mapping_and_paginate_without_gaps(seeded: Any) -> None:
    w, payloads, build = seeded
    plane = _app(w, payloads, build)
    content = await build.reader.read(str(w.tenant), str(w.run))
    mapped, _ = map_snapshot(content, CFG, GRADING)
    expected = {k: [r.cells for r in v] for k, v in rows_of(mapped).items()}
    assert any(expected[t] for t in TABLES)
    for table in TABLES:
        got = await _all_rows(plane, w.domain, table)
        assert [r["cells"] for r in got] == expected[table], table
        assert [r["position"] for r in got] == list(range(1, len(got) + 1))
        for r in got:
            assert r["source"]["state"] in ("linked", "no_source")


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_queries_match_a_hand_written_oracle(seeded: Any) -> None:
    w, payloads, build = seeded
    plane = _app(w, payloads, build)
    everything = [r["cells"] for r in await _all_rows(plane, w.domain, "coverage", page=200)]
    assert len(everything) >= 10
    statuses = sorted({str(r["status"]) for r in everything})
    target = statuses[0]
    cases = [
        (
            {"filters": [{"column": "status", "op": "eq", "value": target}]},
            [r for r in everything if r["status"] == target],
        ),
        (
            {"filters": [{"column": "status", "op": "ne", "value": target}]},
            [r for r in everything if r["status"] != target],
        ),
        (
            {"filters": [{"column": "status", "op": "in", "value": statuses[:2]}]},
            [r for r in everything if r["status"] in statuses[:2]],
        ),
        (
            {"filters": [{"column": "reason", "op": "is_empty"}]},
            [r for r in everything if r["reason"] == ""],
        ),
        (
            {"sort": [{"column": "status", "direction": "desc"}]},
            sorted(everything, key=lambda r: str(r["status"]), reverse=True),
        ),
        (
            {"sort": [{"column": "priority"}, {"column": "status", "direction": "desc"}]},
            _stable_sort(everything, ("priority", False), ("status", True)),
        ),
    ]
    for spec, want in cases:
        got = [r["cells"] for r in await _all_rows(plane, w.domain, "coverage", spec, page=3)]
        if "sort" in spec:
            keys = [c["column"] for c in spec["sort"]]
            assert [[r[k] for k in keys] for r in got] == [[r[k] for k in keys] for r in want]
            assert sorted(r["ref"] for r in got) == sorted(r["ref"] for r in want)
        else:
            assert [r["ref"] for r in got] == [r["ref"] for r in want], spec
    word = str(everything[0]["row"]).split()[0]
    hits = [r["cells"] for r in await _all_rows(plane, w.domain, "coverage", {"search": word})]
    assert hits and all(word.lower() in json.dumps(h).lower() for h in hits)


def _stable_sort(rows: list[dict[str, Any]], *keys: tuple[str, bool]) -> list[dict[str, Any]]:
    out = list(rows)
    for key, reverse in reversed(keys):
        out.sort(key=lambda r, k=key: str(r[k]), reverse=reverse)
    return out


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_injection_through_every_parameter_is_inert(seeded: Any) -> None:
    w, payloads, build = seeded
    plane = _app(w, payloads, build)
    base = f"/v1/accounts/{w.domain}/tables/coverage"
    probes = [
        {
            "spec": json.dumps(
                {"filters": [{"column": "status", "op": "eq", "value": "x' OR '1'='1"}]}
            )
        },
        {
            "spec": json.dumps(
                {"filters": [{"column": "status) OR (1=1", "op": "eq", "value": "x"}]}
            )
        },
        {"spec": json.dumps({"search": "'); DROP TABLE view_rows; --"})},
        {"spec": "{not json"},
        {"cursor": "W1siYScgT1IgJzEnPScxIl0sMV0"},
        {"view_id": "not-a-uuid"},
        {"run_id": "' OR 1=1 --"},
    ]
    for params in probes:
        r = await _get(plane, base, params=params)
        assert r.status_code in (200, 404, 409, 422), (params, r.status_code, r.text)
        if r.status_code == 200:
            assert r.json()["rows"] == [] or "search" not in params.get("spec", "")
    for path in (
        f"/v1/accounts/{w.domain}' OR '1'='1/tables/people",
        f"/v1/accounts/{w.domain}/tables/people;DROP",
        "/v1/accounts/..%2F..%2Fetc/tables/people",
    ):
        assert (await _get(plane, path)).status_code in (404, 422)
    with psycopg.connect(_admin()) as c:
        assert (
            c.execute("SELECT count(*) FROM view_rows WHERE run_id = %s", (w.run,)).fetchone()[0]
            > 0
        )


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_every_export_kind_is_delivered_only_with_parity(
    seeded: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    w, payloads, build = seeded
    plane = _app(w, payloads, build)
    for body in (
        {"domain": w.domain, "kind": "files"},
        {"domain": w.domain, "kind": "csv", "table": "coverage"},
        {"domain": w.domain, "kind": "pdf"},
    ):
        r = await _post(plane, "POST", "/v1/exports", json=body)
        assert r.status_code == 201, r.text
        pkg = r.json()
        if body["kind"] == "pdf" and pkg["reason"] == "pdf_not_configured":
            continue  # no converter on this host: blocked, never delivered unchecked
        assert pkg["status"] == "delivered", pkg
        assert all(v["ok"] for v in pkg["parity"].values()), pkg["parity"]
        for f in pkg["files"]:
            got = await _get(plane, f"/v1/exports/{pkg['package_id']}/files/{f['index']}")
            assert got.status_code == 200
            import hashlib

            assert hashlib.sha256(got.content).hexdigest() == f["sha256"]
            out = tmp_path / f["name"]
            out.write_bytes(got.content)
            if not f["name"].endswith(".pdf"):
                scan = subprocess.run(
                    [sys.executable, str(SCAN), str(out)], capture_output=True, text=True
                )
                assert scan.returncode == 0, (f["name"], scan.stdout)
    # a mutated workbook from the store is caught and the package is blocked
    real_get = payloads.get

    async def tampered(key: str) -> bytes:
        body = await real_get(key)
        if key.endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(body))
            if "People" in wb.sheetnames and wb["People"].max_row >= 2:
                wb["People"]["B2"].value = "Synthetic Forged Name"
                buf = io.BytesIO()
                wb.save(buf)
                return buf.getvalue()
        return body

    monkeypatch.setattr(payloads, "get", tampered)
    r = await _post(plane, "POST", "/v1/exports", json={"domain": w.domain, "kind": "files"})
    pkg = r.json()
    people_rows = await _all_rows(plane, w.domain, "people", page=200)
    if people_rows:
        assert pkg["status"] == "blocked" and pkg["reason"] == "parity_mismatch"
        dl = await _get(plane, f"/v1/exports/{pkg['package_id']}/files/0")
        assert dl.status_code == 409


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_saved_views_validate_and_fail_safe_when_a_column_is_removed(seeded: Any) -> None:
    w, payloads, build = seeded
    admin = _app(w, payloads, build)
    spec = {
        "filters": [{"column": "status", "op": "ne", "value": "Done"}],
        "sort": [{"column": "priority"}],
    }
    r = await _post(
        admin,
        "POST",
        "/v1/views",
        json={"table": "coverage", "name": "Open rows", "spec": spec, "shared": True},
    )
    assert r.status_code == 201, r.text
    view = r.json()
    t = await _get(
        admin, f"/v1/accounts/{w.domain}/tables/coverage", params={"view_id": view["id"]}
    )
    assert t.status_code == 200 and all(x["cells"]["status"] != "Done" for x in t.json()["rows"])
    conflict = await _post(
        admin, "PUT", f"/v1/views/{view['id']}", json={"expected_version": 9, "name": "x"}
    )
    assert conflict.status_code == 409
    seller = _app(w, payloads, build, role="seller", user="idp:ws4-seller")
    shared = await _post(
        seller,
        "POST",
        "/v1/views",
        json={"table": "people", "name": "Mine", "spec": {}, "shared": True},
    )
    assert shared.status_code == 403
    viewer = _app(w, payloads, build, role="viewer", user="idp:ws4-viewer")
    assert (
        await _post(viewer, "POST", "/v1/views", json={"table": "people", "name": "V", "spec": {}})
    ).status_code == 403
    with psycopg.connect(_admin(), autocommit=True) as c:  # the table lost a column since
        c.execute(
            "UPDATE saved_views SET spec = jsonb_set(spec, '{filters,0,column}', '\"retired\"')"
            " WHERE id = %s",
            (view["id"],),
        )
    stale = await _get(
        admin, f"/v1/accounts/{w.domain}/tables/coverage", params={"view_id": view["id"]}
    )
    assert stale.status_code == 409 and "retired" in stale.text
    archived = await _post(
        admin, "POST", f"/v1/views/{view['id']}/archive", json={"expected_version": 1}
    )
    assert archived.status_code == 200
    assert (
        await _get(
            admin, f"/v1/accounts/{w.domain}/tables/coverage", params={"view_id": view["id"]}
        )
    ).status_code == 404


@pytest.mark.requires_postgres
@pytest.mark.asyncio
async def test_materialisation_is_once_per_run_and_matches_the_mapping(seeded: Any) -> None:
    from datetime import UTC, datetime

    from account_intelligence.db.session import tenant_txn
    from tests.tracker.world import make_world
    from views import materialize as mat

    w, payloads, build = seeded
    now = datetime.now(UTC)
    with psycopg.connect(_admin(), autocommit=True) as c:  # start from a clean slate for this run
        c.execute("SET session_replication_role = replica")
        c.execute("DELETE FROM view_rows WHERE run_id = %s", (w.run,))
        c.execute("DELETE FROM view_materializations WHERE run_id = %s", (w.run,))

    async def once() -> mat.Materialized:
        return await mat.materialize(tenant_txn, build.reader, CFG, GRADING, w.tenant, w.run, now)

    first, second = await asyncio.gather(once(), once())
    assert first == second and first.status == "ready"
    content = await build.reader.read(str(w.tenant), str(w.run))
    mapped, _ = map_snapshot(content, CFG, GRADING)
    expected = {k: len(v) for k, v in rows_of(mapped).items()}
    assert first.counts == expected
    with psycopg.connect(_admin()) as c:
        assert (
            c.execute(
                "SELECT count(*) FROM view_materializations WHERE run_id = %s", (w.run,)
            ).fetchone()[0]
            == 1
        )
        assert c.execute("SELECT count(*) FROM view_rows WHERE run_id = %s", (w.run,)).fetchone()[
            0
        ] == sum(expected.values())
    running = make_world(tenant=w.tenant)
    with pytest.raises(mat.NotReady):
        await mat.materialize(tenant_txn, build.reader, CFG, GRADING, w.tenant, running.run, now)
    with pytest.raises(mat.UnknownRun):
        await mat.materialize(tenant_txn, build.reader, CFG, GRADING, w.tenant, uuid4(), now)
    other = make_world()  # another tenant's run id is unknown in this tenant
    with pytest.raises(mat.UnknownRun):
        await mat.materialize(tenant_txn, build.reader, CFG, GRADING, w.tenant, other.run, now)


def test_a_refused_snapshot_is_withheld_not_shown() -> None:
    from exports.builder import errors_of

    content = rich()
    forbidden = content.model_copy(
        update={"forbidden_terms": (*content.forbidden_terms, content.account_name)}
    )
    _, ctx = map_snapshot(forbidden, CFG, GRADING)
    assert errors_of(ctx.violations), "naming the account itself as forbidden must be refused"
