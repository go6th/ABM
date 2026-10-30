"""The pre-delivery gate on Postgres for both input modes: a clean run passes; an unset tracker
row, a dash used as punctuation, a broker name and a credit count, and a client file carrying the
account's private uploaded material each fail it, with the verdict and every leak check stored.
SYNTHETIC data only."""

from __future__ import annotations

import hashlib
import io
import json
from uuid import UUID, uuid4

import pytest

from account_intelligence.db.session import tenant_txn
from account_intelligence.platform.hermetic import MemoryPayloadStore
from documents.config import load_config
from documents.embed import HashingEmbedder
from documents.gate import PreDeliveryGate
from tests.conftest import _make_account
from tests.documents import fixtures
from tests.documents.world import admin

pytestmark = [pytest.mark.requires_postgres, pytest.mark.asyncio]
CLEAN = (
    "Northwind Traders opened a new distribution centre in Ohio, according to its press release."
)


def docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    for para in text.split("\n\n"):
        doc.add_paragraph(para)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _run(tenant: UUID, account: UUID, mode: str = "company") -> UUID:
    (row,) = admin(
        "INSERT INTO account_runs (tenant_id, account_id, graph_version, budget_json,"
        " reserved_credits, idempotency_key, input_mode) VALUES (%s, %s, 'test', '{}', 1, %s, %s)"
        " RETURNING id",
        tenant,
        account,
        f"k-{uuid4().hex[:10]}",
        mode,
    )
    admin(
        "INSERT INTO runs (id, tenant_id, account_id, run_type, status, config_version,"
        " credit_cap) VALUES (%s, %s, %s, 'full', 'queued', 'test', 1)",
        row[0],
        tenant,
        account,
    )
    return UUID(str(row[0]))


def _rows(tenant: UUID, run: UUID, unset: bool = False) -> None:
    proof = json.dumps({"call_refs": [], "calls": [], "notes": []})
    admin(
        "INSERT INTO coverage_rows (run_id, row_key, tenant_id, priority, status, reason_code,"
        " proof, proof_hash) VALUES (%s, 'company_news', %s, 'high', 'done', NULL, %s, 'h'),"
        " (%s, 'podcasts', %s, 'medium', 'not_done', 'not_searched_file_only', %s, 'h')",
        run,
        tenant,
        proof,
        run,
        tenant,
        proof,
    )
    if unset:
        admin(
            "INSERT INTO coverage_rows (run_id, row_key, tenant_id, priority) VALUES"
            " (%s, 'youtube', %s, 'medium')",
            run,
            tenant,
        )


async def _deliver(
    store: MemoryPayloadStore,
    tenant: UUID,
    account: UUID,
    run: UUID,
    data: bytes,
    version: int,
    kind: str = "summary",
    manifest: dict[str, object] | None = None,
) -> None:
    sha = hashlib.sha256(data).hexdigest()
    key = f"{tenant}/{account}/{run}/deliverables/{sha}.docx"
    await store.put(key, data)
    h = "a" * 64
    supersedes = None
    if version > 1:  # a later version names the one it replaces
        (previous,) = admin(
            "SELECT id FROM deliverables WHERE account_id = %s AND kind = %s AND version = %s",
            account,
            kind,
            version - 1,
        )
        supersedes = previous[0]
    admin(
        "INSERT INTO deliverables (tenant_id, account_id, run_id, set_hash, kind, version,"
        " supersedes, filename, sha256, size_bytes, object_key, manifest_json, content_hash,"
        " snapshot_hash, theme_hash, rules_hash, built_at) VALUES (%s, %s, %s, %s, %s, %s, %s,"
        " %s, %s, %s, %s, %s, %s, %s, %s, %s, now())",
        tenant,
        account,
        run,
        h,
        kind,
        version,
        supersedes,
        f"{kind}.docx",
        sha,
        len(data),
        key,
        json.dumps(manifest or {}),
        h,
        h,
        h,
        h,
    )


async def _deliver_set(
    store: MemoryPayloadStore,
    tenant: UUID,
    account: UUID,
    run: UUID,
    data: bytes,
    version: int = 1,
) -> None:
    """The five client files (the same clean body each; enough for the gate's checks)."""
    for kind in ("summary", "evidence", "people", "jobs", "tech"):
        await _deliver(store, tenant, account, run, data, version, kind)


def _private_passage(
    tenant: UUID,
    account: UUID,
    text: str,
    name: str = "notes.docx",
    *,
    member: str | None = None,
    document: UUID | None = None,
    seq: int = 1,
) -> UUID:
    """A private document (or, with ``document``, one more passage of it) holding ``text``;
    ``member`` is the file inside a zip the passage came from."""
    if document is None:
        (doc,) = admin(
            "INSERT INTO documents (tenant_id, account_id, name, declared_bytes, declared_type,"
            " quarantine_key, uploaded_by, upload_expires_at) VALUES (%s, %s, %s, 10,"
            " 'docx', 'q', 'u', now()) RETURNING id",
            tenant,
            account,
            name,
        )
        document = UUID(str(doc[0]))
    sha = hashlib.sha256(text.encode()).hexdigest()
    location = json.dumps({"member": member} if member else {})
    admin(
        "INSERT INTO document_passages (tenant_id, account_id, document_id, seq, kind, location,"
        " content, content_sha256, token_count)"
        " VALUES (%s, %s, %s, %s, 'section', %s::jsonb, %s, %s, 10)",
        tenant,
        account,
        document,
        seq,
        location,
        text,
        sha,
    )
    vec = HashingEmbedder().vector(text)
    admin(
        "INSERT INTO chunks (tenant_id, account_id, doc_id, content, embedding, passage_seq,"
        " embedding_model, location) VALUES (%s, %s, %s, %s, %s::halfvec, %s, 'local-hash-v1',"
        " %s::jsonb)",
        tenant,
        account,
        document,
        text,
        "[" + ",".join(f"{v:.6f}" for v in vec) + "]",
        seq,
        location,
    )
    return document


def _delivered(tenant: UUID, account: UUID, kind: str, filename: str) -> None:
    """A file this product delivered for the account (``deliverables``)."""
    digest = hashlib.sha256(filename.encode()).hexdigest()
    run = _run(tenant, account)
    ext = filename.rsplit(".", 1)[-1]
    admin(
        "INSERT INTO deliverables (tenant_id, account_id, run_id, set_hash, kind, version,"
        " filename, sha256, size_bytes, object_key, manifest_json, content_hash, snapshot_hash,"
        " theme_hash, rules_hash, built_at) VALUES (%s, %s, %s, %s, %s, 1, %s, %s, 1, %s,"
        " '{}', %s, %s, %s, %s, now())",
        tenant,
        account,
        run,
        digest,
        kind,
        filename,
        digest,
        f"{tenant}/{account}/{run}/deliverables/{digest}.{ext}",
        digest,
        digest,
        digest,
        digest,
    )


def _gate(store: MemoryPayloadStore) -> PreDeliveryGate:
    return PreDeliveryGate(tenant_txn, store, HashingEmbedder(), load_config())


async def test_gate_blocks_a_run_with_a_defect_and_passes_a_clean_one(tenant_a: UUID) -> None:
    account = _make_account(tenant_a, f"northwind-{uuid4().hex[:6]}.example")
    store = MemoryPayloadStore()
    _private_passage(
        tenant_a, account, f"Call notes. {fixtures.CANARY}. The budget renews in April."
    )
    clean = _run(tenant_a, account)
    _rows(tenant_a, clean)
    await _deliver_set(store, tenant_a, account, clean, docx(CLEAN))
    passed = await _gate(store).check(str(tenant_a), str(clean))
    assert not any(passed.values()), passed
    assert admin("SELECT verdict FROM run_gate_results WHERE run_id = %s", clean) == [("pass",)]
    leaks = admin(
        "SELECT verdict, passages_checked, hits FROM leak_checks WHERE run_id = %s", clean
    )
    assert len(leaks) == 5 and all(v == "pass" and n == 1 and h == [] for v, n, h in leaks)

    other = _make_account(tenant_a, f"northwind-{uuid4().hex[:6]}.example")
    _private_passage(tenant_a, other, f"Call notes. {fixtures.CANARY}. The budget renews in April.")
    bad = _run(tenant_a, other)
    _rows(tenant_a, bad, unset=True)
    from account_intelligence.connectors.registry_internal import vendor_names
    from account_intelligence.connectors.settings import load_providers_config

    # a broker name from config (no vendor is named in code); the denylist must catch it
    broker = sorted(n for n in vendor_names(load_providers_config()).values() if n != "Parallel")[0]
    text = (
        f"{CLEAN}\n\nThe pilot \u2014 if approved \u2014 starts soon. The search used {broker}"
        " and 40 credits.\n\n" + fixtures.CANARY + "."
    )
    await _deliver(store, tenant_a, other, bad, docx(text), 1)
    # a file whose bytes changed after the builder stored it is not the builder's file: the gate
    # scans it whole (the builder vouches for the bytes it stored, see documents.gate)
    admin("UPDATE deliverables SET sha256 = repeat('b', 64) WHERE run_id = %s", bad)
    failing = await _gate(store).check(str(tenant_a), str(bad))
    assert failing["tracker_rows"] == ["youtube:unset"]
    assert failing["deliverable_set"] == ["evidence", "jobs", "people", "tech"]
    assert failing["client_dashes"] and failing["client_terms"] and failing["leak_check"]
    (verdict, defects) = admin(
        "SELECT verdict, defects FROM run_gate_results WHERE run_id = %s", bad
    )[0]
    assert verdict == "fail" and {d["check"] for d in defects} >= {
        "tracker_rows",
        "client_dashes",
        "client_terms",
        "leak_check",
    }
    # the stored records name codes and counts, never the private text
    stored = json.dumps(defects) + json.dumps(
        admin("SELECT hits FROM leak_checks WHERE run_id = %s", bad)
    )
    assert "Zebrafinch" not in stored and broker.lower() not in stored.lower()


async def test_an_uploaded_earlier_deliverable_is_not_a_leak_but_a_memo_is(tenant_a: UUID) -> None:
    """The owner refreshes a brief by uploading the previous workspace; the new files share
    titles and headings with it by design (staging 2026-10-21: every delivery held back). A
    file named as one of the account's own deliveries (also under a browser's ``(1)`` suffix,
    also as a member inside the workspace zip), or a workbook the import recognised, is
    skipped; a private memo with the same words still blocks, and the hit names that memo.
    A client's own file merely named like a deliverable is never skipped."""
    store = MemoryPayloadStore()
    shared = "Senior Salesforce developer in Ohio, full time, posted in April, source company site."

    def account_with_uploads(memo: bool, *, lookalike: bool = False) -> UUID:
        """An account that was delivered a workspace and uploaded it back: two files by
        name, one the import recognised, and the zip with a deliverable member inside; when
        asked, a private memo (or a client file merely named like a deliverable) too."""
        account = _make_account(tenant_a, f"northwind-{uuid4().hex[:6]}.example")
        for kind, name in (
            ("summary", "Northwind_Account_Summary.docx"),
            ("evidence", "Northwind_Sources_and_Evidence.docx"),
            ("people", "Northwind_People_List.xlsx"),
            ("jobs", "Northwind_Job_Postings.xlsx"),
            ("tech", "Northwind_Tech_Stack_and_Teams.xlsx"),
            ("workspace", "Northwind_workspace.zip"),
        ):
            _delivered(tenant_a, account, kind, name)
        _private_passage(tenant_a, account, shared, name="Northwind_Job_Postings.xlsx")
        _private_passage(tenant_a, account, shared, name="Northwind_Account_Summary (1).docx")
        recognised = _private_passage(tenant_a, account, shared, name="renamed.xlsx")
        admin(
            "INSERT INTO account_runs (tenant_id, account_id, status, graph_version, budget_json,"
            " reserved_credits, idempotency_key, finished_reason, finished_at, input_mode) VALUES"
            " (%s, %s, 'COMPLETED', 'file-import', '{}', 0, %s, 'file_import', now(), 'file')",
            tenant_a,
            account,
            f"file-import-{recognised}",
        )
        _private_passage(
            tenant_a,
            account,
            shared,
            name="Northwind_workspace (2).zip",
            member="Northwind_Account_Summary.docx",
        )
        if memo:
            _private_passage(tenant_a, account, f"Memo. {shared}", name="notes.docx")
        if lookalike:
            _private_passage(tenant_a, account, f"Memo. {shared}", name="People List.xlsx")
        return account

    async def gate(account: UUID) -> tuple[UUID, dict[str, list[str]]]:
        run = _run(tenant_a, account)
        _rows(tenant_a, run)
        # version 2: the account's version 1 is the earlier delivery it uploaded back
        await _deliver_set(store, tenant_a, account, run, docx(f"{CLEAN}\n\n{shared}"), 2)
        return run, await _gate(store).check(str(tenant_a), str(run))

    run, passed = await gate(account_with_uploads(memo=False))
    assert passed["leak_check"] == [], passed
    assert all(
        n == 0 for (n,) in admin("SELECT passages_checked FROM leak_checks WHERE run_id = %s", run)
    )

    with_memo = account_with_uploads(memo=True)
    (memo,) = admin(
        "SELECT id FROM documents WHERE account_id = %s AND name = 'notes.docx'", with_memo
    )
    run, failing = await gate(with_memo)
    assert len(failing["leak_check"]) == 5 and all(
        item.endswith(",documents=1") for item in failing["leak_check"]
    )
    hits = admin("SELECT hits FROM leak_checks WHERE run_id = %s", run)
    assert hits and all(h["document"] == str(memo[0]) for (stored,) in hits for h in stored)
    assert "Salesforce" not in json.dumps([h for (h,) in hits])

    # a client file that only carries a deliverable-like name was never delivered: it blocks
    lookalike = account_with_uploads(memo=False, lookalike=True)
    run, failing = await gate(lookalike)
    assert failing["leak_check"] and all(
        item.endswith(",documents=1") for item in failing["leak_check"]
    )

    # a memo inside the workspace zip blocks too, next to the delivered member that does not
    zipped = account_with_uploads(memo=False)
    (zip_doc,) = admin(
        "SELECT id FROM documents WHERE account_id = %s AND name = 'Northwind_workspace (2).zip'",
        zipped,
    )
    _private_passage(
        tenant_a,
        zipped,
        f"Memo. {shared}",
        document=UUID(str(zip_doc[0])),
        member="notes/board memo.docx",
        seq=2,
    )
    run, failing = await gate(zipped)
    assert failing["leak_check"] and all(
        item.endswith(",documents=1") for item in failing["leak_check"]
    )
    hits = admin("SELECT hits FROM leak_checks WHERE run_id = %s", run)
    assert hits and all(h["document"] == str(zip_doc[0]) for (stored,) in hits for h in stored)

    # the same words as a public fact of the account are public wherever they appear
    public = account_with_uploads(memo=True)
    admin(
        "WITH f AS (INSERT INTO facts (tenant_id, account_id, attribute_key, value_json,"
        " fingerprint, product_grade, visibility) VALUES (%s, %s, 'jobs.1.title', %s, 'fp-1',"
        " 'probable', 'public') RETURNING id) INSERT INTO fact_evidence (tenant_id, fact_id,"
        " evidence_kind, url, supporting_sentence, span_sha256, source_tier, evidence_date)"
        " SELECT %s, id, 'source', 'https://northwind.example/jobs', 'A public posting.',"
        " repeat('a', 64), 'T1', now() FROM f",
        tenant_a,
        public,
        json.dumps({"title": shared}),
        tenant_a,
    )
    _, passed = await gate(public)
    assert passed["leak_check"] == [], passed


async def test_file_mode_gate_requires_the_labelled_internal_file(tenant_a: UUID) -> None:
    account = _make_account(tenant_a, f"northwind-{uuid4().hex[:6]}.example")
    store = MemoryPayloadStore()
    run = _run(tenant_a, account, "file")
    _rows(tenant_a, run)
    missing = await _gate(store).check(str(tenant_a), str(run))
    assert missing["internal_label"] == ["missing"]
    admin(
        "INSERT INTO internal_deliverables (tenant_id, account_id, run_id, label, object_key,"
        " sha256, size_bytes) VALUES (%s, %s, %s, 'Internal: from your uploaded material', %s,"
        " %s, 1)",
        tenant_a,
        account,
        run,
        f"tenants/{tenant_a}/internal/{run}/{'b' * 64}.docx",
        "b" * 64,
    )
    assert not any((await _gate(store).check(str(tenant_a), str(run))).values())


async def test_a_dash_in_a_workbook_cell_is_a_source_title_not_a_dash_error(
    tenant_a: UUID,
) -> None:
    """Staging 2026-10-21: a 896 row job postings run failed the gate on `client_dashes`; the
    workbooks print stored job titles ("Engineer \u2013 Data") as they are, and the rule's scope
    is authored prose. A dash in a written document still fails, and names its file kind."""
    from openpyxl import Workbook

    def xlsx(cell: str) -> bytes:
        wb = Workbook()
        wb.active.append(["Title"])
        wb.active.append([cell])
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    account = _make_account(tenant_a, f"northwind-{uuid4().hex[:6]}.example")
    store = MemoryPayloadStore()
    run = _run(tenant_a, account)
    _rows(tenant_a, run)
    titled = xlsx("Senior Salesforce Engineer \u2013 Data Cloud")
    for kind in ("people", "jobs", "tech"):
        await _deliver(store, tenant_a, account, run, titled, 1, kind)
    for kind in ("summary", "evidence"):
        await _deliver(store, tenant_a, account, run, docx(CLEAN), 1, kind)
    assert (await _gate(store).check(str(tenant_a), str(run)))["client_dashes"] == []

    other = _make_account(tenant_a, f"northwind-{uuid4().hex[:6]}.example")
    written = _run(tenant_a, other)
    _rows(tenant_a, written)
    await _deliver(
        store, tenant_a, other, written, docx(f"{CLEAN}\n\nA pilot \u2014 if approved."), 1
    )
    # a file whose bytes changed after the builder stored it is not the builder's file: the gate
    # scans it whole (the builder vouches for the bytes it stored, see documents.gate)
    admin("UPDATE deliverables SET sha256 = repeat('b', 64) WHERE run_id = %s", written)
    failing = await _gate(store).check(str(tenant_a), str(written))
    assert failing["client_dashes"] == ["summary:1"]


async def test_gate_finds_a_lost_row_and_a_provider_host(tenant_a: UUID) -> None:
    from documents.gate import provider_hosts

    account = _make_account(tenant_a, f"rows-{uuid4().hex[:6]}.example")
    store = MemoryPayloadStore()
    run = _run(tenant_a, account)
    _rows(tenant_a, run)
    # the builder recorded a row the Word file does not hold, and a workbook names a host
    await _deliver(
        store, tenant_a, account, run, docx(CLEAN), 1, "summary", {"refs": {"summary": ["F-0001"]}}
    )
    for kind in ("evidence", "people", "jobs"):
        await _deliver(store, tenant_a, account, run, docx(CLEAN), 1, kind)
    await _deliver(
        store, tenant_a, account, run, docx(f"{CLEAN} See {provider_hosts()[0]}."), 1, "tech"
    )
    failing = await _gate(store).check(str(tenant_a), str(run))
    assert failing["rows_dropped"] == ["summary:1"]
    assert failing["third_party_domains"] == ["tech:1"]


def _with_hyperlink(data: bytes, target: str) -> bytes:
    """The Word file with one external hyperlink relationship (a target the visible text never
    shows): the member the text-only host check did not read."""
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                rel = (
                    '<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships/hyperlink" '
                    f'Target="{target}" TargetMode="External"/></Relationships>'
                )
                body = body.decode().replace("</Relationships>", rel).encode()
            z.writestr(item, body)
    return out.getvalue()


async def test_gate_reads_every_member_for_a_provider_host_in_a_file_it_did_not_make(
    tenant_a: UUID,
) -> None:
    """F11: a hyperlink target is not visible text. For a file the builder did not make (its
    bytes no longer match the stored hash) the host check covers every member, as the builder's
    own scan does."""
    from documents.gate import provider_hosts

    account = _make_account(tenant_a, f"hyper-{uuid4().hex[:6]}.example")
    store = MemoryPayloadStore()
    run = _run(tenant_a, account)
    _rows(tenant_a, run)
    sneaky = _with_hyperlink(docx(CLEAN), f"https://{provider_hosts()[0]}/profile")
    for kind in ("summary", "evidence", "people", "jobs", "tech"):
        await _deliver(
            store, tenant_a, account, run, sneaky if kind == "tech" else docx(CLEAN), 1, kind
        )
    # the builder vouches for the bytes it stored: its own scan covered them
    assert (await _gate(store).check(str(tenant_a), str(run)))["third_party_domains"] == []
    # a file whose bytes changed after the builder stored it is scanned whole, every member
    admin(
        "UPDATE deliverables SET sha256 = repeat('b', 64) WHERE run_id = %s AND kind = 'tech'", run
    )
    failing = await _gate(store).check(str(tenant_a), str(run))
    assert failing["third_party_domains"] == ["tech:1"]
