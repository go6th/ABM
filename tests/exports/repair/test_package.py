"""CSV and XLSX output, the dash fix, README generation and the deterministic zip."""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date
from typing import Any

import pytest
from openpyxl import load_workbook

from account_intelligence.replica import sheets as SH
from exports.repair import package as P
from exports.repair.core import (
    Exclusion,
    RepairedTable,
    Sheet,
    WorkspaceProfile,
)

RUN = date(2026, 10, 30)
BOM = b"\xef\xbb\xbf"


def profile(name: str = "Nimbus Orchards Ltd") -> WorkspaceProfile:
    return WorkspaceProfile(
        account_name=name,
        account_short="Nimbus",
        domain="nimbus.example",
        run_date=RUN,
        window_from=date(2025, 10, 30),
        window_to=RUN,
    )


def sheet(name: str, columns: tuple[str, ...], rows: list[dict[str, str]], **kw: Any) -> Sheet:
    return Sheet(name=name, columns=columns, rows=rows, **kw)


def table(key: str, title: str, sheets: list[Sheet], **kw: Any) -> RepairedTable:
    return RepairedTable(key=key, title=title, sheets=sheets, **kw)


def read_csv(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig"), newline="")))


def open_book(data: bytes) -> Any:
    return load_workbook(io.BytesIO(data))


# ---- CSV --------------------------------------------------------------------------------------


def test_csv_has_bom_quotes_every_field_crlf_and_one_header() -> None:
    s = sheet("Main", ("Name", "Note"), [{"Name": "Ada Byrne", "Note": 'says "hi", ok'}])
    data = P.csv_bytes(s)
    assert data.startswith(BOM)
    text = data.decode("utf-8-sig")
    assert text.endswith("\r\n")
    lines = text.split("\r\n")[:-1]
    assert lines == ['"Name","Note"', '"Ada Byrne","says ""hi"", ok"']
    assert read_csv(data) == [["Name", "Note"], ["Ada Byrne", 'says "hi", ok']]


@pytest.mark.parametrize("lead", ["=", "+", "-", "@", "\t", "\r"])
def test_csv_prefixes_apostrophe_on_formula_like_cells(lead: str) -> None:
    s = sheet("Main", ("A",), [{"A": f"{lead}1+1"}])
    cell = read_csv(P.csv_bytes(s))[1][0]
    assert cell == f"'{lead}1+1"


def test_csv_private_sheet_gets_no_extra_rows() -> None:
    s = sheet(
        "Mail", ("Title", "From"), [{"Title": "Re: plan", "From": "a@b.example"}], private=True
    )
    rows = read_csv(P.csv_bytes(s))
    assert len(rows) == 2 and rows[0] == ["Title", "From"]


def test_csv_empty_sheet_is_header_only() -> None:
    assert read_csv(P.csv_bytes(sheet("Main", ("A", "B"), []))) == [["A", "B"]]


def test_csv_missing_key_is_an_empty_cell() -> None:
    assert read_csv(P.csv_bytes(sheet("Main", ("A", "B"), [{"A": "x"}])))[1] == ["x", ""]


# ---- clean_cell -------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "want"),
    [
        ("Senior Analyst -", "Senior Analyst"),
        ("Senior Analyst --", "Senior Analyst"),
        ("Senior Analyst ---", "Senior Analyst"),
        ("Senior Analyst –", "Senior Analyst"),
        ("- Senior Analyst", "Senior Analyst"),
        ("Lead - Data", "Lead, Data"),
        ("plain", "plain"),
        ("", ""),
    ],
)
def test_clean_cell_removes_dash_punctuation(raw: str, want: str) -> None:
    assert P.clean_cell(raw, shareable=True) == want


def test_clean_cell_escapes_formula_start_and_keeps_numbers_as_text() -> None:
    assert P.clean_cell("=SUM(A1)", shareable=True) == "'=SUM(A1)"
    assert P.clean_cell("-5", shareable=True) == "'-5"
    assert P.clean_cell(None, shareable=True) == ""
    assert P.clean_cell(7, shareable=True) == "7"


def test_clean_cell_never_leaves_a_gate_dash_between_cells() -> None:
    a = P.clean_cell("Senior Analyst -", shareable=True)
    b = P.clean_cell("Finance", shareable=True)
    assert SH._GATE_DASHES.search(f"{a} {b}") is None


# ---- XLSX -------------------------------------------------------------------------------------


def make_xlsx(sheets: list[Sheet], *, private: bool = False, title: str = "Demo") -> bytes:
    return P.xlsx_bytes(sheets, "Nimbus Orchards Ltd", private=private, title=title, run_date=RUN)


def test_xlsx_sheets_header_widths_and_about() -> None:
    s1 = sheet("Main", ("Name", "Note"), [{"Name": "Ada Byrne", "Note": "x" * 200}])
    s2 = sheet("Second [sheet]: a/b?", ("K",), [{"K": "v"}])
    book = open_book(make_xlsx([s1, s2]))
    assert book.sheetnames == ["Main", "Second sheet a b", "About"]
    ws = book["Main"]
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold and ws["A1"].value == "Name"
    assert ws.column_dimensions["B"].width == 60
    about = {r[0].value: r[1].value for r in book["About"].iter_rows()}
    assert about["Account"] == "Nimbus Orchards Ltd"
    assert about["Table"] == "Demo"
    assert about["Exported"] == "2026-10-30"
    assert about["Rows per sheet"] == "Main: 1; Second [sheet]: a/b?: 1"
    assert "Note" not in about


def test_xlsx_sheet_titles_are_unique_and_at_most_31_chars() -> None:
    long = "L" * 50
    book = open_book(
        make_xlsx([sheet(long, ("A",), []), sheet(long, ("A",), []), sheet("About", ("A",), [])])
    )
    names = book.sheetnames
    assert all(len(n) <= 31 for n in names)
    assert len({n.casefold() for n in names}) == len(names)


def test_xlsx_escapes_formula_like_text_and_stores_strings() -> None:
    s = sheet("Main", ("A",), [{"A": "=1+1"}, {"A": "@cmd"}, {"A": "+3"}, {"A": "plain"}])
    ws = open_book(make_xlsx([s]))["Main"]
    assert [ws.cell(row=r, column=1).value for r in (2, 3, 4, 5)] == [
        "'=1+1",
        "'@cmd",
        "'+3",
        "plain",
    ]
    assert all(ws.cell(row=r, column=1).data_type == "s" for r in (2, 3, 4, 5))


def test_xlsx_formula_cells_are_real_formulas() -> None:
    cols = ("Platform", "Postings naming it", "Share of base")
    rows = [
        {
            "Platform": "Alpha",
            "Postings naming it": "6",
            "Share of base": "60%",
            "__formula__:Share of base": "=B2/$B$1",
        },
        {
            "Platform": "Beta",
            "Postings naming it": "4",
            "Share of base": "40%",
            "__formula__:Share of base": "=IF(B3>=5,B3/10,0)",
        },
    ]
    ws = open_book(make_xlsx([sheet("Platform Ranking", cols, rows)]))["Platform Ranking"]
    assert ws["C2"].data_type == "f" and ws["C2"].value == "=B2/$B$1"
    assert ws["C3"].data_type == "f" and ws["C3"].value == "=IF(B3>=5,B3/10,0)"


@pytest.mark.parametrize(
    "formula",
    [
        '=HYPERLINK("http://x.example","y")',
        "=cmd|' /C calc'!A0",
        '=B2 - B3 + INDIRECT("A1")',
        "=SUM(B2",
        "B2/3",
        "=B2/3)",
        '=1+"unterminated',
        "=" + "1+" * 600 + "1",
    ],
)
def test_xlsx_unsafe_formula_is_written_as_its_value(formula: str) -> None:
    s = sheet("Main", ("A", "B"), [{"A": "7", "B": "42", "__formula__:B": formula}])
    cell = open_book(make_xlsx([s]))["Main"]["B2"]
    assert cell.data_type == "s" and cell.value == "42"


def test_valid_formula_removes_spaces_so_no_spaced_dash_reaches_the_file() -> None:
    assert P.valid_formula("=B2 - B3") == "=B2-B3"
    assert P.valid_formula("=SUM(B2:B9) / $B$1") == "=SUM(B2:B9)/$B$1"
    assert P.valid_formula("='Platform Ranking'!B2*2") == "='Platform Ranking'!B2*2"
    assert P.valid_formula('=IF(B2>1,"A","B")') == '=IF(B2>1,"A","B")'
    assert P.valid_formula("=") is None


def test_xlsx_dash_in_formula_still_passes_the_whole_workbook_check() -> None:
    cols = ("A", "B", "C")
    rows = [{"A": "3", "B": "1", "C": "2", "__formula__:C": "=A2 - B2"}]
    data = make_xlsx([sheet("Main", cols, rows)])
    SH.check_finished(data)


def test_xlsx_hyperlinks_only_for_safe_urls_in_link_columns() -> None:
    cols = ("Source URL", "Link", "LinkedIn", "Note")
    rows = [
        {
            "Source URL": "https://news.example/a",
            "Link": "javascript:alert(1)",
            "LinkedIn": "https://social.example/in/x",
            "Note": "https://notacolumn.example/z",
        },
        {
            "Source URL": "http://user:pw@x.example/",
            "Link": "Not found",
            "LinkedIn": "Not found",
            "Note": "n",
        },
    ]
    ws = open_book(make_xlsx([sheet("Main", cols, rows)]))["Main"]
    assert ws["A2"].hyperlink.target == "https://news.example/a"
    assert ws["B2"].hyperlink is None
    assert ws["C2"].hyperlink.target == "https://social.example/in/x"
    assert ws["D2"].hyperlink is None
    assert ws["A3"].hyperlink is None


def test_xlsx_long_text_is_cut_at_a_word_boundary_with_a_visible_note() -> None:
    words = " ".join(["lorem"] * 20000)
    s = sheet("Main", ("Text",), [{"Text": words}])
    data, dropped = P.write_xlsx_checked(
        [s], "Nimbus Orchards Ltd", private=False, title="Demo", run_date=RUN
    )
    value = open_book(data)["Main"]["A2"].value
    assert len(value) <= P.CELL_MAX
    assert value.endswith(P.CUT_NOTE)
    assert value[: -len(P.CUT_NOTE)].endswith("lorem")
    assert [d.reason for d in dropped] == ["cell cut at 32,000 characters"]
    assert dropped[0].table == "Demo" and dropped[0].key == "Main!A2"


def test_xlsx_private_table_is_marked_and_not_scanned_as_shareable() -> None:
    s = sheet("Mail", ("Title",), [{"Title": "Re: Q3 numbers"}], private=True)
    book = open_book(make_xlsx([s], private=True, title="From your mail and drive"))
    about = {r[0].value: r[1].value for r in book["About"].iter_rows()}
    assert about["Note"] == "Internal: do not share"


def test_xlsx_is_deterministic() -> None:
    s = sheet("Main", ("A", "B"), [{"A": "1", "B": "https://x.example/p"}])
    assert make_xlsx([s]) == make_xlsx([s])


# ---- the dash regression ---------------------------------------------------------------------


def dash_sheet() -> Sheet:
    cols = ("Title", "Department", "Note")
    rows = [
        {"Title": "Senior Analyst -", "Department": "Finance", "Note": "x"},
        {"Title": "Platform Lead --", "Department": "Data", "Note": "Ops - Core"},
        {"Title": "Engineer —", "Department": "- Legal", "Note": "line one\n- bullet two"},
    ]
    return sheet("Job postings", cols, rows)


def test_dash_regression_workbook_ships_and_passes_the_finished_file_check() -> None:
    data, dropped = P.write_xlsx_checked(
        [dash_sheet()], "Nimbus Orchards Ltd", private=False, title="Job postings", run_date=RUN
    )
    SH.check_finished(data)  # does not raise ExportBlocked
    ws = open_book(data)["Job postings"]
    assert ws["A2"].value == "Senior Analyst"
    assert ws["B4"].value == "Legal"
    assert dropped == []


def test_dash_regression_loop_repairs_what_the_cell_function_missed(monkeypatch: Any) -> None:
    """Cells go in raw (clean_cell switched off): the workbook check trips, the offending cells
    are found alone and in pairs, repaired and the workbook is checked again."""
    monkeypatch.setattr(P, "clean_cell", lambda v, *, shareable: "" if v is None else str(v))
    s = sheet(
        "Main", ("A", "B"), [{"A": "Senior Analyst -", "B": "Next"}, {"A": "ok", "B": "fine"}]
    )
    raw = P._render(
        P._prepare([s], "Main", True)[0],
        [["Account", "x"]],
        title="t",
        run_date=RUN,
        shareable=True,
    )
    with pytest.raises(SH.ExportBlocked):
        SH.check_finished(raw)
    data, dropped = P.write_xlsx_checked([s], "Nimbus", private=False, title="Main", run_date=RUN)
    SH.check_finished(data)
    assert open_book(data)["Main"]["A2"].value == "Senior Analyst"
    assert dropped == []


def test_cells_that_cannot_be_repaired_are_withheld_and_logged(monkeypatch: Any) -> None:
    monkeypatch.setattr(P, "clean_cell", lambda v, *, shareable: "" if v is None else str(v))
    monkeypatch.setattr(P, "_harden", lambda text: text)
    s = sheet(
        "Main", ("A", "B"), [{"A": "Senior Analyst -", "B": "Next"}, {"A": "ok", "B": "fine"}]
    )
    data, dropped = P.write_xlsx_checked([s], "Nimbus", private=False, title="Main", run_date=RUN)
    SH.check_finished(data)  # still ships, and passes
    ws = open_book(data)["Main"]
    assert SH.WITHHELD in {ws["A2"].value, ws["B2"].value}
    assert dropped and all(d.reason == "cell withheld for review" for d in dropped)
    assert {d.key for d in dropped} <= {"Main!A2", "Main!B2"}


def test_shareable_cells_with_a_provider_domain_are_withheld_not_shipped() -> None:
    from documents.gate import provider_hosts

    hosts = provider_hosts()
    if not hosts:
        pytest.skip("no provider host in this configuration")
    s = sheet("Main", ("A",), [{"A": f"see https://api.{hosts[0]}/x for more"}])
    data, dropped = P.write_xlsx_checked([s], "Nimbus", private=False, title="T", run_date=RUN)
    SH.check_finished(data)
    assert open_book(data)["Main"]["A2"].value == SH.WITHHELD
    assert dropped[0].reason == "cell withheld for review"
    assert dropped[0].key == "Main!A2"


# ---- README -----------------------------------------------------------------------------------

MEMBERS = [
    "README.txt",
    "Deliverables/Account Summary.docx",
    "Tables/Job postings.csv",
    "Tables/Job postings.xlsx",
    "Internal - your files/From your mail and drive.csv",
    "Logs/repair_report.md",
    "Logs/exclusions.csv",
]


def test_readme_is_generated_from_members_and_lists_every_path() -> None:
    text = P.render_readme(profile(), MEMBERS, ["Web searches: capped by the execution limit."])
    assert "Nimbus Orchards Ltd" in text and "2026-10-30" in text
    for heading in ("Deliverables/", "Tables/", "Internal - your files/", "Logs/"):
        assert heading in text
    assert "Internal: do not share" in text
    assert "Not done / Partial:" in text and "capped by the execution limit" in text
    assert "Contents:" in text
    assert P.readme_paths(text) == [MEMBERS[0], *sorted(MEMBERS[1:])]
    assert P.validate_readme(text, MEMBERS) == []


def test_readme_without_notes_says_nothing_is_open() -> None:
    text = P.render_readme(profile(), MEMBERS, [])
    assert "Nothing is marked Not done or Partial." in text


def test_readme_marks_empty_folders() -> None:
    text = P.render_readme(profile(), ["README.txt"], [])
    assert text.count("(nothing in this folder)") == 4


def test_validate_readme_reports_both_directions() -> None:
    text = P.render_readme(profile(), MEMBERS, [])
    fewer = [m for m in MEMBERS if m != "Tables/Job postings.xlsx"]
    assert P.validate_readme(text, fewer) == ["missing from archive: Tables/Job postings.xlsx"]
    more = [*MEMBERS, "Tables/Extra.csv"]
    assert P.validate_readme(text, more) == ["not in README: Tables/Extra.csv"]


# ---- the zip ----------------------------------------------------------------------------------


def world() -> dict[str, RepairedTable]:
    jobs = table(
        "jobs",
        "Job postings",
        [
            sheet(
                "Job postings",
                ("Title", "Source URL"),
                [{"Title": "Senior Analyst -", "Source URL": "https://careers.example/1"}],
            ),
            sheet("All postings (unfiltered)", ("Title",), [{"Title": "Cook"}]),
        ],
        exclusions=[Exclusion("Job postings", "J9", "duplicate posting", "merged into J1")],
        notes=["Live board not reached."],
    )
    mail = table(
        "mail",
        "From your mail and drive",
        [
            sheet(
                "Mail",
                ("Title",),
                [{"Title": "Not collected: mail and drive not connected"}],
                private=True,
            )
        ],
        internal=True,
        notes=["Internal: do not share"],
    )
    people = table("people", "People", [sheet("People", ("Name",), [{"Name": "Ada Byrne"}])])
    return {"jobs": jobs, "mail": mail, "people": people}


def build(**kw: Any) -> bytes:
    args: dict[str, Any] = {
        "deliverables": {"Account Summary.docx": b"docx-bytes"},
        "internal_files": {"My notes.docx": b"private"},
        "report_md": "# Repair report\n",
        "notes": ["Searches: Partial, cap reached."],
    }
    args.update(kw)
    return P.build_zip(profile(), world(), **args)


def test_zip_layout_order_and_readme_agreement() -> None:
    data = build()
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        readme = z.read("README.txt").decode()
    assert names[0] == "README.txt"
    assert names[1:] == sorted(names[1:])
    assert "Deliverables/Account Summary.docx" in names
    assert {
        "Tables/Job postings.csv",
        "Tables/Job postings.xlsx",
        "Tables/Job postings - All postings (unfiltered).csv",
        "Tables/People.csv",
        "Tables/People.xlsx",
        "Internal - your files/My notes.docx",
        "Internal - your files/From your mail and drive (internal).csv",
        "Internal - your files/From your mail and drive (internal).xlsx",
        "Logs/repair_report.md",
        "Logs/exclusions.csv",
    } <= set(names)
    assert not any(n.startswith("Tables/From your mail") for n in names)
    assert P.validate_readme(readme, names) == []
    assert "Searches: Partial, cap reached." in readme
    assert "Job postings: Live board not reached." in readme


def test_zip_xlsx_members_open_and_pass_the_finished_check() -> None:
    with zipfile.ZipFile(io.BytesIO(build())) as z:
        data = z.read("Tables/Job postings.xlsx")
    SH.check_finished(data)
    book = open_book(data)
    assert book["Job postings"]["A2"].value == "Senior Analyst"
    assert "All postings (unfiltered)" in book.sheetnames


def test_zip_logs_hold_the_report_and_exclusions() -> None:
    with zipfile.ZipFile(io.BytesIO(build())) as z:
        assert z.read("Logs/repair_report.md").decode() == "# Repair report\n"
        rows = read_csv(z.read("Logs/exclusions.csv"))
    assert rows[0] == ["Table", "Key", "Reason", "Detail"]
    assert ["Job postings", "J9", "duplicate posting", "merged into J1"] in rows


def test_zip_is_deterministic_with_fixed_timestamps() -> None:
    first, second = build(), build()
    assert first == second
    with zipfile.ZipFile(io.BytesIO(first)) as z:
        assert {i.date_time for i in z.infolist()} == {(2026, 10, 30, 0, 0, 0)}


def test_zip_member_names_are_safe_unique_and_short() -> None:
    t1 = table("a", "Same / title: x", [sheet("S", ("A",), [{"A": "1"}])])
    t2 = table("b", "Same  title  x", [sheet("S", ("A",), [{"A": "1"}])])
    t3 = table("c", "T" * 300, [sheet("S", ("A",), [{"A": "1"}])])
    data = P.build_zip(profile(), {"a": t1, "b": t2, "c": t3}, deliverables={})
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
    assert len(names) == len(set(n.casefold() for n in names))
    assert all(
        len(n) <= 150 and ".." not in n and "\\" not in n and not n.startswith("/") for n in names
    )
    assert "Tables/Same title x.csv" in names and "Tables/Same title x (2).csv" in names


def test_duplicate_deliverable_names_after_cutting_get_a_counter() -> None:
    long_a = "A" * 200 + ".docx"
    long_b = "A" * 199 + "B.docx"
    data = P.build_zip(profile(), {}, deliverables={long_a: b"1", long_b: b"2"})
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        docs = [n for n in z.namelist() if n.startswith("Deliverables/")]
    assert len(docs) == 2 and all(len(n) <= 150 for n in docs)


@pytest.mark.parametrize("name", ["../evil.docx", "a/b.docx", "a\\b.docx", "", "..", "x..y.docx"])
def test_zip_refuses_unsafe_deliverable_names(name: str) -> None:
    with pytest.raises(P.PackageInvalid):
        P.build_zip(profile(), {}, deliverables={name: b"x"})


@pytest.mark.parametrize(
    "name",
    [
        "Summary internal only.docx",
        "OPERATOR_ONLY notes.docx",
        "x (Internal-Only).xlsx",
        "operator.only.txt",
    ],
)
def test_internal_only_names_are_rejected_in_deliverables_but_fine_as_internal_files(
    name: str,
) -> None:
    with pytest.raises(P.PackageInvalid, match="internal-only"):
        P.build_zip(profile(), {}, deliverables={name: b"x"})
    data = P.build_zip(profile(), {}, deliverables={}, internal_files={name: b"x"})
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert f"Internal - your files/{name}" in z.namelist()


def test_zip_raises_package_invalid_when_readme_and_members_disagree(monkeypatch: Any) -> None:
    real = P.render_readme

    def lossy(prof: WorkspaceProfile, members: Any, notes: Any) -> str:
        return real(prof, [m for m in members if m != "Tables/People.xlsx"], notes)

    monkeypatch.setattr(P, "render_readme", lossy)
    with pytest.raises(P.PackageInvalid) as err:
        build()
    assert err.value.missing == ["not in README: Tables/People.xlsx"]


def test_zip_raises_when_readme_lists_a_path_that_does_not_exist(monkeypatch: Any) -> None:
    real = P.render_readme

    def greedy(prof: WorkspaceProfile, members: Any, notes: Any) -> str:
        return real(prof, [*members, "Tables/Ghost.xlsx"], notes)

    monkeypatch.setattr(P, "render_readme", greedy)
    with pytest.raises(P.PackageInvalid) as err:
        build()
    assert "missing from archive: Tables/Ghost.xlsx" in err.value.missing


def test_withheld_cells_land_in_the_exclusions_csv(monkeypatch: Any) -> None:
    ex = Exclusion("People", "cell", "cell withheld for review", "column Name")
    data = P.build_zip(profile(), {}, deliverables={}, extra_exclusions=[ex])
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert ["People", "cell", "cell withheld for review", "column Name"] in read_csv(
            z.read("Logs/exclusions.csv")
        )


def test_the_fallback_guard_alone_escapes_and_the_live_guard_agrees() -> None:
    """Without an installed guard the package still escapes formulas and fixes edge dashes; the
    live guard (``replica.repair_guard``) is a superset with the same finished-file dash rule."""
    from account_intelligence.replica import repair_guard
    from account_intelligence.replica import sheets as SH

    P.install_guard(P.FALLBACK_GUARD)
    try:
        assert P.clean_cell("=SUM(A1)", shareable=True) == "'=SUM(A1)"
        assert P.clean_cell("Senior Analyst -", shareable=True) == "Senior Analyst"
    finally:
        P.install_guard(repair_guard.build())
    assert P.FALLBACK_GUARD.gate_dashes.pattern == SH._GATE_DASHES.pattern
    assert P.clean_cell("=SUM(A1)", shareable=True) == "'=SUM(A1)"


def test_free_text_members_are_scanned_line_by_line() -> None:
    """The report, README and exclusions are written as raw bytes: a line that names something
    the guard refuses is replaced, the clean lines (a spaced hyphen included) stay."""
    from dataclasses import replace

    P.install_guard(replace(P.FALLBACK_GUARD, text_trips=lambda t: "vendorx" in t))
    try:
        out = P.clean_text("# Report\nInternal - your files/\nseen on vendorx here").split("\n")
    finally:
        P.install_guard(P.FALLBACK_GUARD)
    assert out == ["# Report", "Internal - your files/", "(withheld)"]
