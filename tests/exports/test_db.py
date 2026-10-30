"""Export builders on real Postgres (spec 9.8): the full build through ``PgContentReader`` and
``PgDeliverableStore`` over rows written by the real pipelines; one consistent snapshot under a
concurrent writer; the database's two-output policy; append-only deliverables; a repeat build
writes nothing; a redelivery adds a superseding version; a kill leaves nothing complete; two
tenants never cross. All data synthetic."""

from __future__ import annotations

import contextlib
import hashlib
import io
import zipfile
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from typing import Any
from unittest.mock import AsyncMock, patch
from uuid import UUID, uuid4

import psycopg
import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from account_intelligence.db.enums import SourceKind
from account_intelligence.db.session import tenant_txn
from account_intelligence.db.source_ids import get_or_create_source
from account_intelligence.grading.config import load_grading_config
from account_intelligence.platform.hermetic import MemoryPayloadStore
from account_intelligence.urls import normalize_url, url_hash
from exports.builder import ExportBuilder
from exports.config import load_export_config
from exports.errors import MissingContent
from exports.readers.pg import PgContentReader
from exports.store import PgDeliverableStore
from exports.verify import docx_text
from exports.word.renderer import SubprocessWordRenderer, python_renderer
from jobs.repo import PgPostingRepo
from jobs.stage import run_jobs
from people.repo import PgPeopleRepo
from people.stage import run_people
from tests.fakes.exports import FixedClock
from tests.fakes.jobs import Corpus as JobsCorpus
from tests.fakes.jobs import FakeRunInput as JobsRunInput
from tests.fakes.jobs import Recorder
from tests.fakes.people import Corpus as PeopleCorpus
from tests.fakes.people import FakeRunInput as PeopleRunInput
from tests.jobs.conftest import build as build_jobs
from tests.people.kit import Kit
from tests.tracker.world import World, admin_dsn, fetch, make_world, txn

pytestmark = [pytest.mark.requires_postgres, pytest.mark.asyncio]
CFG = load_export_config()
GRADING = load_grading_config()
PRIVATE = "Synthetic private fact never shown"
SUPERSEDED = "Synthetic superseded fact never shown"


class Payloads:
    """The platform payload store the reader sees: people payloads and job texts together."""

    def __init__(self, people: MemoryPayloadStore, jobs: Recorder) -> None:
        self.people, self.jobs = people, jobs
        self.objects: dict[str, bytes] = {}
        self.fail_on_put: int | None = None
        self.puts = 0

    async def put(self, key: str, body: bytes) -> str:
        self.puts += 1
        if self.fail_on_put is not None and self.puts >= self.fail_on_put:
            raise RuntimeError("killed mid-build")
        self.objects[key] = body
        return key

    async def get(self, key: str) -> bytes:
        if key in self.objects:
            return self.objects[key]
        if key in self.jobs.texts:
            return self.jobs.texts[key].encode()
        return await self.people.get(key)


class DbRegistry:
    """The people stage's source registry over the real registry writer (as production)."""

    def __init__(self, w: World) -> None:
        self.w = w

    async def code_for(self, url: str) -> str:
        return await register(self.w, url)


async def register(w: World, url: str) -> str:
    async with tenant_txn(w.tenant) as s:
        source = await s.run_sync(
            lambda sync: get_or_create_source(
                sync, account_id=w.account, run_id=w.run, url=url, kind=SourceKind.SECONDARY
            )
        )
        return str(source.source_code)


def _source(
    c: psycopg.Connection[Any], w: World, n: int, path: str, kind: str = "primary"
) -> tuple[UUID, str, str]:
    url = f"https://{w.domain}/{path}"
    code = f"S-{n:04d}"
    sid = c.execute(
        "INSERT INTO sources (tenant_id, account_id, source_code, url, url_normalized, url_hash,"
        " kind, first_seen_run_id, last_seen_run_id, published_at) VALUES"
        " (%s, %s, %s, %s, %s, %s, %s, %s, %s, '2026-09-01') RETURNING id",
        (
            w.tenant,
            w.account,
            code,
            url,
            normalize_url(url),
            url_hash(normalize_url(url)),
            kind,
            w.run,
            w.run,
        ),
    ).fetchone()[0]
    return sid, code, url


def _fact(
    c: psycopg.Connection[Any],
    w: World,
    source: tuple[UUID, str, str],
    value: str,
    *,
    grade: str = "confirmed",
    lane: str | None = "salesforce",
    visibility: str = "public",
    superseded: bool = False,
) -> UUID:
    fid = uuid4()
    c.execute(
        "INSERT INTO facts (id, tenant_id, account_id, attribute_key, value_json, fingerprint,"
        " engine_grade, product_grade, read_in_primary_this_run, may_state_to_prospect, status,"
        " visibility, run_id, claim_kind, value_text, confirming_question, lead_status, lane,"
        " superseded_at) VALUES (%s, %s, %s, 'tech.synthetic', to_jsonb(%s::text), %s, 'A', %s,"
        " %s, %s, 'current', %s, %s, 'sourced', %s, %s, %s, %s, %s)",
        (
            fid,
            w.tenant,
            w.account,
            value,
            uuid4().hex,
            grade,
            grade == "confirmed",
            grade == "confirmed",
            visibility,
            w.run,
            value,
            None if grade == "confirmed" else "Synthetic question?",
            "complete" if grade == "lead" else None,
            lane,
            "2026-09-01" if superseded else None,
        ),
    )
    c.execute(
        "INSERT INTO fact_evidence (tenant_id, fact_id, evidence_kind, url, supporting_sentence,"
        " span_sha256, source_id, source_tier, evidence_date, origin, provider_code) VALUES"
        " (%s, %s, 'source', %s, %s, %s, %s, 'T1', now(), 'full_page_read', 'p_search')",
        (
            w.tenant,
            fid,
            source[2],
            f"Synthetic sentence for {value}.",
            hashlib.sha256(value.encode()).hexdigest(),
            source[0],
        ),
    )
    return fid


def seed_rows(w: World) -> dict[str, Any]:
    """Facts, evidence, sources, report sections and detections for the world's run."""
    with psycopg.connect(admin_dsn()) as c, c.transaction():
        c.execute("UPDATE accounts SET name = 'Synthetic DB Ltd' WHERE id = %s", (w.account,))
        s1 = _source(c, w, 1, "about")
        s2 = _source(c, w, 2, "news/a", "secondary")
        s3 = _source(c, w, 3, "tech/d", "aggregator")
        c.execute(
            "INSERT INTO account_source_counters (tenant_id, account_id, next_value)"
            " VALUES (%s, %s, 4)",
            (w.tenant, w.account),
        )
        f1 = _fact(c, w, s1, "Synthetic confirmed platform")
        _fact(c, w, s2, "Synthetic probable owner", grade="probable", lane=None)
        _fact(c, w, s2, PRIVATE, visibility="private")
        _fact(c, w, s1, SUPERSEDED, superseded=True)
        for version, text_ in (
            (1, "Old synthetic narrative."),
            (2, "Synthetic DB Ltd is a synthetic account for testing."),
        ):
            c.execute(
                "INSERT INTO report_sections (tenant_id, account_id, run_id, version, kind,"
                " position, text, cited_source_codes, fact_ids) VALUES (%s, %s, %s, %s,"
                " 'narrative', 0, %s, %s, %s)",
                (w.tenant, w.account, w.run, version, text_, [s1[1]], [f1]),
            )
        c.execute(
            "INSERT INTO tech_detections (tenant_id, account_id, run_id, technology,"
            " first_seen_at, last_seen_at, source_ref) VALUES (%s, %s, %s, 'Salesforce',"
            " '2025-01', '2026-08', %s)",
            (w.tenant, w.account, w.run, s3[2]),
        )
        # the organisation's tech stack, in the lookup's own order (position)
        for position, (function, category, technology, postings) in enumerate(
            (
                ("Sales", "Customer Relationship Manager (CRM)", "Salesforce", 12),
                ("Finance", "ERP and Business Management Software", "SAP", 5),
            )
        ):
            c.execute(
                "INSERT INTO tech_stack_items (tenant_id, owner_user_id, account_id, run_id,"
                " business_function, category, technology, postings, position)"
                " VALUES (%s, 'test-owner', %s, %s, %s, %s, %s, %s, %s)",
                (w.tenant, w.account, w.run, function, category, technology, postings, position),
            )
    return {"codes": (s1[1], s2[1], s3[1])}


async def seed_world(
    *, register_postings: bool = True, legacy_unregistered: bool = False
) -> tuple[World, Payloads]:
    """A run with every input the builders read, written by the real pipelines and stores.
    ``legacy_unregistered`` saves the postings the way a research did before the job repository
    registered their links (no source code for a posting no fact cites)."""
    w = make_world()
    seed_rows(w)
    async with txn(w) as tracker:
        await tracker.seed(w.run, w.tenant)
    corpus = JobsCorpus()
    pipeline, *_ = build_jobs(corpus)
    rec = Recorder()
    with contextlib.ExitStack() as stack:
        if legacy_unregistered:
            stack.enter_context(
                patch("jobs.repo.register_posting_sources", AsyncMock(return_value=0))
            )
        jobs = await run_jobs(
            str(w.run),
            pipeline,
            PgPostingRepo(tenant_txn, w.tenant, w.account),
            rec,
            rec,
            rec,
            JobsRunInput(corpus.lanes_off, corpus.as_of),
        )
    people_store = MemoryPayloadStore()
    people_repo = PgPeopleRepo(
        tenant_txn, w.tenant, w.account, people_store, lambda kind, name: f"t/{w.run}/{kind}/{name}"
    )
    pc = PeopleCorpus()
    kit = Kit(pc, repo=people_repo)  # type: ignore[arg-type]
    people = await run_people(
        str(w.run),
        pc.account,
        kit.pipeline,
        people_repo,
        DbRegistry(w),
        kit.rec,
        kit.rec,
        PeopleRunInput(pc),
    )
    for f in people.findings:
        await register(w, f.url)  # follow-up: the people stage does not register post links yet
    if register_postings:
        for p in jobs.postings:
            await register(w, p.url)
    return w, Payloads(people_store, rec)


def builder(w: World, payloads: Payloads, **reader: Any) -> ExportBuilder:
    store = PgDeliverableStore(tenant_txn, payloads, FixedClock(), CFG.rules_hash)
    return ExportBuilder(
        PgContentReader(payloads, batch=2, **reader),
        store,  # batch 2: pagination exercised
        SubprocessWordRenderer(python_renderer(), CFG.rules.renderer),
        CFG,
        GRADING,
    )


def _as_app_rw(tenant: UUID, sql: str, *params: object) -> list[Any]:
    from tests.conftest import APP_DSN

    with psycopg.connect(APP_DSN) as c, c.transaction():
        c.execute(
            "SELECT set_config('app.tenant_id', %s, true), set_config('app.scope', 'system', true)",
            (str(tenant),),
        )
        cur = c.execute(sql, params)
        return cur.fetchall() if cur.description else []


async def test_end_to_end_build_from_postgres() -> None:
    w, payloads = await seed_world()
    b = builder(w, payloads)
    result = await b.build_all(str(w.tenant), str(w.run))
    assert result.status == "complete", result.errors
    m = result.manifest
    assert m is not None
    assert m.snapshot_counts["facts"] == 2  # private and superseded facts are not selected
    assert m.snapshot_counts["sources"] >= 3 and m.rows_per_file["tech"] == 2
    # the Confirmed stack rows come from tech_stack_items, counted in the same snapshot
    assert m.snapshot_counts["tech_stack_items"] == 2 and m.snapshot_counts["tech_detections"] == 1
    assert m.refs["tech"] == (
        "Sales | Customer Relationship Manager (CRM) | Salesforce",
        "Finance | ERP and Business Management Software | SAP",
    )
    assert m.rows_per_file["people"] > 0 and m.rows_per_file["jobs"] > 0
    rows = fetch(
        "SELECT kind, version, supersedes, object_key FROM deliverables WHERE run_id = %s"
        " ORDER BY kind",
        w.run,
    )
    assert [r[0] for r in rows] == ["evidence", "jobs", "people", "summary", "tech", "workspace"]
    assert all(r[1] == 1 and r[2] is None for r in rows)
    for _, _, _, key in rows:
        assert key.startswith(f"{w.tenant}/{w.account}/{w.run}/deliverables/")
        assert key in payloads.objects
    summary = next(d for d in result.deliverables if d.kind == "summary")
    text_ = docx_text(payloads.objects[summary.object_key])
    assert "Synthetic DB Ltd is a synthetic account" in text_  # the latest section version
    assert "Old synthetic narrative" not in text_
    for d in result.deliverables:
        body = payloads.objects[d.object_key]
        if d.kind == "workspace":
            with zipfile.ZipFile(io.BytesIO(body)) as z:
                body = b"".join(z.read(n) for n in z.namelist())
        assert PRIVATE.encode() not in body and SUPERSEDED.encode() not in body
        assert b"p_search" not in body  # the provider code stays in the internal table


async def test_every_listed_posting_gets_a_source_code_though_no_fact_cites_it() -> None:
    """E2E 2026-10-29, D-1: the job repository registers the link of every posting the file
    lists, so a posting no graded fact cited is listed, not silently dropped. Nothing was
    registered by hand here (the sources table holds only what the research wrote)."""
    w, payloads = await seed_world(register_postings=False)
    result = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert result.status == "complete", result.errors
    assert not [b for b in result.manifest.withheld_blocks if b.startswith("jobs:")]
    listed = result.manifest.row_ids["jobs"]
    assert len(listed) > 0
    held = {r[0] for r in fetch("SELECT url FROM sources WHERE account_id = %s", w.account)}
    exported_urls = {
        r[0]
        for r in fetch(
            "SELECT url FROM job_postings WHERE run_id = %s AND full_text_read AND NOT hostile"
            " AND cardinality(platforms) > 0 AND window_bucket <> 'out_of_window'",
            w.run,
        )
    }
    assert len(exported_urls) == len(listed) and exported_urls <= held
    # the codes are the registry's own and stay put: a second registration changes nothing
    before = fetch(
        "SELECT url, source_code FROM sources WHERE account_id = %s ORDER BY url", w.account
    )
    again = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert again.status == "complete"
    assert before == fetch(
        "SELECT url, source_code FROM sources WHERE account_id = %s ORDER BY url", w.account
    )


async def test_the_counts_say_no_posting_was_left_out() -> None:
    """The Counts line "Postings not listed" is 0 when every read posting was listed."""
    w, payloads = await seed_world(register_postings=False)
    result = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert result.status == "complete", result.errors
    jobs = next(d for d in result.deliverables if d.kind == "jobs")
    book = load_workbook(io.BytesIO(payloads.objects[jobs.object_key]))
    counts = book[CFG.rules.workbooks.jobs.counts_sheet]
    lines = {str(r[0].value): r[1].value for r in counts.iter_rows() if r[0].value}
    not_listed = [v for k, v in lines.items() if k.startswith("Postings not listed")]
    assert not_listed == [0]
    ref_col = CFG.rules.workbooks.ref_column(CFG.rules.workbooks.jobs)
    sheet = book[CFG.rules.workbooks.jobs.sheet]
    refs = [c.value for c in sheet[ref_col][1:] if c.value]
    assert refs == [f"J-{n:04d}" for n in range(1, len(refs) + 1)]


async def test_a_run_saved_before_the_links_were_registered_is_repaired_on_rebuild() -> None:
    """A research saved by the old repository left uncited posting links without a code. The
    reader alone only reads (those rows are withheld and listed); with the registering
    transaction wired (as in production) a rebuild registers the missing links and lists them."""
    w, payloads = await seed_world(register_postings=False, legacy_unregistered=True)
    plain = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert plain.status == "complete", plain.errors
    withheld = [b for b in plain.manifest.withheld_blocks if b.startswith("jobs:")]
    assert withheld and all(b.endswith(":missing_source_code") for b in withheld), withheld
    assert plain.manifest.row_ids["jobs"] == ()
    fixed = await builder(w, payloads, register=tenant_txn).build_all(str(w.tenant), str(w.run))
    assert fixed.status == "complete", fixed.errors
    assert len(fixed.manifest.row_ids["jobs"]) == len(withheld)
    assert not [b for b in fixed.manifest.withheld_blocks if b.startswith("jobs:")]
    assert fetch("SELECT count(*) FROM deliverables WHERE run_id = %s", w.run)[0][0] >= 6


async def test_one_snapshot_under_a_concurrent_writer() -> None:
    w, payloads = await seed_world()
    wrote: list[UUID] = []

    @asynccontextmanager
    async def snapshot_then_write(tid: UUID) -> AsyncIterator[AsyncSession]:
        async with tenant_txn(tid, snapshot=True) as s:
            await s.execute(text("SELECT count(*) FROM facts"))  # the snapshot is taken here
            with psycopg.connect(admin_dsn()) as c, c.transaction():  # a writer commits now
                src = c.execute(
                    "SELECT id, source_code, url FROM sources WHERE account_id = %s"
                    " ORDER BY source_code LIMIT 1",
                    (w.account,),
                ).fetchone()
                wrote.append(_fact(c, w, src, "Synthetic fact written during the build"))
            yield s

    during = await PgContentReader(payloads, batch=2, snapshot=snapshot_then_write).read(
        str(w.tenant), str(w.run)
    )
    assert wrote and str(wrote[0]) not in {f.id for f in during.facts}
    assert during.snapshot_counts["facts"] == len(during.facts) == 2
    after = await PgContentReader(payloads, batch=2).read(str(w.tenant), str(w.run))
    assert str(wrote[0]) in {f.id for f in after.facts}


async def test_snapshot_transaction_is_read_only() -> None:
    w = make_world()
    with pytest.raises(Exception, match="read-only"):
        async with tenant_txn(w.tenant, snapshot=True) as s:
            await s.execute(
                text(
                    "INSERT INTO report_sections (account_id, run_id, version, kind,"
                    " position, text) VALUES (:a, :r, 1, 'caveat', 0, 'x')"
                ),
                {"a": w.account, "r": w.run},
            )


async def test_deliverables_and_sections_refuse_update_and_delete() -> None:
    w, payloads = await seed_world()
    result = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert result.status == "complete"
    for sql in (
        "UPDATE deliverables SET filename = 'x.docx' WHERE run_id = %s",
        "DELETE FROM deliverables WHERE run_id = %s",
        "UPDATE report_sections SET text = 'x' WHERE run_id = %s",
        "DELETE FROM report_sections WHERE run_id = %s",
    ):
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            _as_app_rw(w.tenant, sql, w.run)


async def test_repeat_build_writes_nothing_and_redelivery_supersedes() -> None:
    w, payloads = await seed_world()
    b = builder(w, payloads)
    first = await b.build_all(str(w.tenant), str(w.run))
    puts = payloads.puts
    again = await b.build_all(str(w.tenant), str(w.run))
    assert again.reused and payloads.puts == puts
    assert fetch("SELECT count(*) FROM deliverables WHERE run_id = %s", w.run)[0][0] == 6
    # a content change on the same run (a gate fix) is a redelivery: version 2 supersedes 1
    with psycopg.connect(admin_dsn()) as c, c.transaction():
        src = c.execute(
            "SELECT id, source_code, url FROM sources WHERE account_id = %s"
            " ORDER BY source_code LIMIT 1",
            (w.account,),
        ).fetchone()
        _fact(c, w, src, "Synthetic fact added by a gate fix")
    second = await b.build_all(str(w.tenant), str(w.run))
    assert second.status == "complete" and not second.reused
    old = {d.kind: d.id for d in first.deliverables}
    for d in second.deliverables:
        assert d.version == 2 and d.supersedes == old[d.kind]
    # the change note compares the run with the PREVIOUS research of the account, never with its
    # own earlier set (leakage review 2026-10-30, F7): there is none here
    assert second.change_note == first.change_note
    assert "First delivery" in (second.change_note or "")
    assert fetch("SELECT count(*) FROM deliverables WHERE run_id = %s", w.run)[0][0] == 12


async def test_a_kill_mid_build_leaves_nothing_complete() -> None:
    w, payloads = await seed_world()
    payloads.fail_on_put = payloads.puts + 3  # die on the third object put
    with pytest.raises(RuntimeError, match="killed"):
        await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert fetch("SELECT count(*) FROM deliverables WHERE run_id = %s", w.run)[0][0] == 0
    payloads.fail_on_put = None
    result = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert result.status == "complete"
    assert fetch("SELECT count(*) FROM deliverables WHERE run_id = %s", w.run)[0][0] == 6


async def test_two_tenants_never_cross() -> None:
    a, payloads = await seed_world()
    result = await builder(a, payloads).build_all(str(a.tenant), str(a.run))
    assert result.status == "complete"
    b = make_world()  # another tenant
    with pytest.raises(MissingContent):
        await PgContentReader(payloads, batch=2).read(str(b.tenant), str(a.run))
    for table in ("deliverables", "report_sections"):
        assert _as_app_rw(b.tenant, f"SELECT count(*) FROM {table}")[0][0] == 0
    same_domain = make_world(tenant=b.tenant)  # tenant B researching its own account
    with psycopg.connect(admin_dsn(), autocommit=True) as c:
        c.execute(
            "UPDATE accounts SET domain = %s WHERE id = %s",
            (f"x{uuid4().hex[:6]}.{a.domain}", same_domain.account),
        )
    built = await builder(same_domain, payloads).build_all(str(b.tenant), str(same_domain.run))
    # no narrative of its own: nothing of tenant A was used; the Summary alone is withheld and
    # the other files are delivered (a file's defect no longer fails the whole build)
    assert any(w.startswith("summary:") for w in built.withheld_files)
    assert all(d.kind != "summary" for d in built.deliverables)


async def test_reader_fills_the_posting_fields_and_the_archived_copy_links() -> None:
    """The reader fills last seen, seniority, the (platform, sentence) pairs, the new-row flag
    and the archived-copy address of each posting and source (absolute on the public origin;
    blank without one or without a copy)."""
    from datetime import UTC, date, datetime
    from urllib.parse import quote

    from account_intelligence import archive as LA

    w, payloads = await seed_world()
    # one source has an archived copy (Phase 1 would have stored it), the others none
    async with tenant_txn(w.tenant) as s:
        await LA.record(
            s,
            w.tenant,
            w.account,
            w.run,
            [
                LA.ArchiveEntry(
                    url=f"https://{w.domain}/about",
                    kind="page",
                    payload_ref=f"t/{w.tenant}/archive/about",
                    archived_at=datetime(2026, 10, 1, tzinfo=UTC),
                )
            ],
        )
    base = "https://app.example"
    content = await PgContentReader(payloads, batch=2, archive_base=base + "/").read(
        str(w.tenant), str(w.run)
    )
    by_key = {p.posting_key: p for p in content.postings}
    j18 = by_key["id:j18"]
    assert (j18.last_seen, j18.seniority) == (date(2026, 9, 15), "mid_senior")
    assert j18.platform_sentence_pairs == (
        ("ServiceNow", "The ServiceNow developer builds & tests flows."),
    )
    assert j18.new_since_previous is None  # a first research
    assert by_key["id:j13"].seniority == "manager" and by_key["id:j13"].last_seen is None
    # every exported posting has a stored text, so every one has an archived copy
    for p in content.postings:
        assert p.archive_url == (
            f"{base}/accounts/{quote(w.domain, safe='')}/archive?url={quote(p.url, safe='')}"
        ), p.posting_key
    sources = {x.url: x for x in content.sources}
    assert sources[f"https://{w.domain}/about"].archive_url == (
        f"{base}/accounts/{quote(w.domain, safe='')}/archive"
        f"?url={quote(f'https://{w.domain}/about', safe='')}"
    )
    assert sources[f"https://{w.domain}/news/a"].archive_url == ""
    # without a public origin no link is made
    plain = await PgContentReader(payloads, batch=2).read(str(w.tenant), str(w.run))
    assert all(p.archive_url == "" for p in plain.postings)
    assert all(x.archive_url == "" for x in plain.sources)
    assert {p.posting_key: p.last_seen for p in plain.postings}["id:j18"] == date(2026, 9, 15)


async def test_the_workspace_zip_carries_the_accounts_archived_copies() -> None:
    """Fix wave C (item 8): the reader selects the account's archived copies (client kinds,
    newest structured record per address, newest first) and the zip holds one text file each,
    never a profile copy; a rebuild changes nothing."""
    from datetime import UTC, datetime

    from account_intelligence import archive as LA

    w, payloads = await seed_world()

    class Clock:
        def now(self) -> datetime:
            return datetime(2026, 9, 20, 9, tzinfo=UTC)

    archive = LA.LinkArchive(tenant_txn, w.tenant, w.account, w.run, payloads, Clock())  # type: ignore[arg-type]
    await archive.put_text(
        "https://synthetic.example/news/a", "page", "Synthetic news text.", title="Synthetic news"
    )
    await archive.put_text(
        "https://www.linkedin.com/in/synthetic-person", "profile", "Profile text.", title="Person"
    )
    result = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert result.status == "complete", result.errors
    workspace = next(d for d in result.deliverables if d.kind == "workspace")
    with zipfile.ZipFile(io.BytesIO(payloads.objects[workspace.object_key])) as z:
        releases = {n: z.read(n).decode() for n in z.namelist() if n.startswith("data/releases/")}
    texts = "\n".join(releases.values())
    assert "title: Synthetic news" in texts and "Synthetic news text." in texts
    assert "Profile text." not in texts and "linkedin.com/in/synthetic-person" not in texts
    assert len(releases) >= 2  # the news page and the read postings the jobs repository archived
    assert all(t.splitlines()[3] == "source: 6th agents" for t in releases.values())
    again = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert again.reused and {d.sha256 for d in again.deliverables} == {
        d.sha256 for d in result.deliverables
    }
