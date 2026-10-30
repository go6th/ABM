"""Property tests (spec 9.3, hypothesis): for any generated row set every row appears exactly once
and in order in the produced file, no generated string becomes an executable formula, no cell
exceeds the limit (long text continues, counted, and joins back exactly), building twice gives the
same bytes, the mapping is total and pure, and the zip never accepts an unsafe member name."""

from __future__ import annotations

import io
import re
from datetime import date

from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook

from account_intelligence.grading.config import load_grading_config
from exports.config import load_export_config
from exports.content import (
    EvidenceRow,
    FactRow,
    PersonRow,
    PostingRow,
    SnapshotContent,
    SourceRow,
)
from exports.errors import UnsafePackage
from exports.excel.jobs import render_jobs
from exports.excel.people import render_people
from exports.policy import leak_hits, sanitise
from exports.readers.common import MapContext
from exports.readers.jobs import map_jobs
from exports.readers.people import map_people
from exports.readers.summary import map_summary
from exports.verify import xlsx_row_keys
from exports.workspace import safe_name
from jobs.normalise import strip_pii
from tests.exports.corpus import account

CFG = load_export_config()
GRADING = load_grading_config()
SETTINGS = settings(max_examples=40, deadline=None, suppress_health_check=[HealthCheck.too_slow])
# any text, including control characters, surrogate-free emoji, right-to-left scripts and
# formula prefixes; text the leak scanner refuses on its own (a provider-code shape such as
# "p_a", a denylisted term) is a correct violation, not a builder defect, so it is not generated
TEXT = st.text(alphabet=st.characters(blacklist_categories=("Cs",)), max_size=60).filter(
    lambda t: not leak_hits(t)
)
FORMULA = st.sampled_from(["=", "+", "-", "@", "\t", "\r", ""])


def _small_limit_cfg(limit: int):  # type: ignore[no-untyped-def]
    limits = CFG.rules.limits.model_copy(update={"excel_cell_chars": limit})
    return CFG.model_copy(update={"rules": CFG.rules.model_copy(update={"limits": limits})})


@st.composite
def postings(draw: st.DrawFn) -> tuple[SnapshotContent, list[str]]:
    n = draw(st.integers(min_value=0, max_value=25))
    rows, sources, descriptions = [], [], []
    for i in range(n):
        url = f"https://gen.example/p/{i}"
        sources.append(SourceRow(code=f"S-{i + 1:04d}", url=url, published=None, kind="primary"))
        description = draw(FORMULA) + draw(TEXT) * draw(st.integers(min_value=1, max_value=12))
        descriptions.append(description)
        rows.append(
            PostingRow(
                posting_key=f"k-{i:04d}",
                title=draw(FORMULA) + "Synthetic role " + draw(TEXT),
                url=url,
                status=draw(st.sampled_from(["open", "closed", "unknown"])),
                posting_class=draw(st.sampled_from(["build", "user", None])),
                platforms=("Salesforce",),
                platform_sentences=(draw(TEXT),),
                description=description,
                posted_at=date(2026, 1, 1),
                full_text_read=True,
            )
        )
    base = account(n_facts=0, n_people=0, n_postings=0, n_detections=0)
    content = base.model_copy(
        update={"postings": tuple(rows), "sources": tuple(sources), "jobs_counts": None}
    )
    return content, descriptions


@SETTINGS
@given(postings())
def test_jobs_rows_conserved_safe_and_deterministic(
    case: tuple[SnapshotContent, list[str]],
) -> None:
    content, descriptions = case
    cfg = _small_limit_cfg(400)  # small limit: continuation happens often
    ctx = MapContext(content, cfg, GRADING)
    mapped = map_jobs(ctx)
    assert mapped.model is not None and not ctx.violations
    body, parts = render_jobs(mapped.model, cfg, "Synthetic Jobs")
    assert body == render_jobs(mapped.model, cfg, "Synthetic Jobs")[0]
    keys = xlsx_row_keys(body, "Jobs", cfg.rules.workbooks.ref_column(cfg.rules.workbooks.jobs))
    assert keys == mapped.refs and len(keys) == len(content.postings)
    wb = load_workbook(io.BytesIO(body))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if ws.title in ("Counts", "Summary"):
                    continue  # the formula totals, the platform matrix (B12) and fixed texts
                assert cell.data_type != "f", (ws.title, cell.coordinate)
                if isinstance(cell.value, str):
                    assert len(cell.value) <= 400
    long_rows = list(wb["Long text"].iter_rows(min_row=2, values_only=True))
    assert len(long_rows) == parts
    header = "Full description"
    jobs_headers = [c.value for c in wb["Jobs"][1]]
    col = jobs_headers.index(header) + 1
    ref_col = jobs_headers.index("Ref") + 1
    for i, text in enumerate(descriptions, start=2):
        expected = sanitise(strip_pii(text))
        cell = wb["Jobs"].cell(row=i, column=col).value or ""
        ref = wb["Jobs"].cell(row=i, column=ref_col).value
        rest = "".join(t for r, c, _, _, t in long_rows if r == ref and c == header)
        head = re.sub(r" \[continued: part 1 of \d+; .*\]$", "", cell, flags=re.S)
        assert head + rest == expected


@SETTINGS
@given(st.lists(st.tuples(FORMULA, TEXT, st.sampled_from(["A", "B", "C"])), max_size=20))
def test_people_rows_conserved_in_order(people: list[tuple[str, str, str]]) -> None:
    base = account(n_people=0, n_facts=0, n_postings=0, n_detections=0)
    sources = list(base.sources)
    rows = []
    for i, (prefix, name, priority) in enumerate(people):
        url = f"https://gen.example/in/{i}"
        sources.append(
            SourceRow(code=f"S-{len(sources) + 1:04d}", url=url, published=None, kind="secondary")
        )
        rows.append(
            PersonRow(
                id=f"pk-{i}",
                name=(prefix + name).strip() or f"Synthetic {i}",
                title="",
                level="director",
                lanes="",
                priority=priority,
                profile_url=url,
                source_code=None,
            )
        )
    content = base.model_copy(update={"people": tuple(rows), "sources": tuple(sources)})
    ctx = MapContext(content, CFG, GRADING)
    mapped = map_people(ctx)
    # a generated name may contain a denylisted pattern (a leak) or nothing printable once
    # sanitised (name_missing). A name with nothing printable is withheld and listed
    # (``name_missing``); a name the leak scan refuses is replaced by the withheld note: never an
    # exception, never a violation that refuses the set
    assert not ctx.violations
    withheld = {row for _, row, rule in ctx.withheld_blocks if rule == "name_missing"}
    rows = [p for p in rows if p.id not in withheld]
    assert mapped.model is not None
    body, _ = render_people(mapped.model, CFG, "Synthetic People")
    ref_col = CFG.rules.workbooks.ref_column(CFG.rules.workbooks.people)
    assert xlsx_row_keys(body, "People", ref_col) == mapped.refs
    # one level and no lane: the list orders by priority, then name, then link
    ordered = sorted(rows, key=lambda p: (p.priority, p.name.casefold(), p.profile_url))
    assert [r.row_id for r in mapped.model.rows] == [p.id for p in ordered]


@st.composite
def fact_sets(draw: st.DrawFn) -> SnapshotContent:
    base = account(n_facts=0)
    n = draw(st.integers(min_value=0, max_value=12))
    facts, evidence, sources = [], [], list(base.sources)
    for i in range(n):
        grade = draw(st.sampled_from(["confirmed", "probable", "hypothesis", "lead", "certain"]))
        url = f"https://gen.example/f/{i}"
        sources.append(
            SourceRow(code=f"S-{len(sources) + 1:04d}", url=url, published=None, kind="primary")
        )
        fid = f"00000000-0000-4000-8000-{i:012d}"
        facts.append(
            FactRow(
                id=fid,
                attribute_key=f"tech.g{i}",
                value_text=draw(TEXT) or "v",
                lane=draw(st.sampled_from(["salesforce", "sap", None, "unknown_lane"])),
                claim_kind="sourced",
                product_grade=grade,
                engine_grade="B",
                read_in_primary_this_run=draw(st.booleans()),
                may_state_to_prospect=draw(st.booleans()),
                confirming_question=draw(st.sampled_from([None, "Why?"])),
                lead_status="complete",
            )
        )
        evidence.append(
            EvidenceRow(
                id=f"e{i}",
                fact_id=fid,
                evidence_kind="source",
                url=url,
                supporting_sentence="s",
                source_code=None,
                is_aggregator=draw(st.booleans()),
                origin=None,
                evidence_date=None,
            )
        )
    return base.model_copy(
        update={"facts": tuple(facts), "evidence": tuple(evidence), "sources": tuple(sources)}
    )


@SETTINGS
@given(fact_sets())
def test_summary_mapping_is_total_pure_and_conserves_rows(content: SnapshotContent) -> None:
    def run() -> tuple[object, list[tuple[str, str, str]]]:
        ctx = MapContext(content, CFG, GRADING)
        mapped = map_summary(ctx, map_people(ctx).model)
        listed = sorted((v.code, v.where, v.item) for v in ctx.violations)
        # a row the mapping withholds (and lists) is accounted for like a refused one
        held = sorted(("withheld", where, row) for where, row, _ in ctx.withheld_blocks)
        return mapped, listed + held

    first, violations = run()
    second, again = run()
    assert first == second and violations == again  # pure: same input, same output
    model = first.model  # type: ignore[attr-defined]
    assert model is not None
    rendered = [r.row_id for t in model.lanes for r in t.rows]
    assert len(rendered) == len(set(rendered))  # never twice
    refused = {item for _, _, item in violations}
    for f in content.facts:
        # each fact is rendered once or refused with a violation naming it
        assert f.id in rendered or f.id in refused or f.product_grade == "certain"


NAME = st.text(alphabet=st.characters(min_codepoint=32, max_codepoint=0x2FF), max_size=40)


@settings(max_examples=300, deadline=None)
@given(NAME)
def test_an_accepted_member_name_is_always_safe(name: str) -> None:
    try:
        safe_name(name, 120)
    except UnsafePackage:
        return
    assert not name.startswith("/") and "\\" not in name and ":" not in name
    assert ".." not in name.split("/") and len(name) <= 120
    assert re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._/-]*", name)
