"""The Tech Stack and Teams workbook against the owner's reference file set (2026-10-29): the
three sheets, their names, order, header rows, headers, freeze panes, filters, merges and column
widths equal the reference manifest; the Platform Ranking follows the reference's build script
(patterns, lane words, Open now, Last seen, rank fills); the Confirmed stack is the organisation's
tech stack with the lane the platform words give; the Platform teams sheet prints rows when teams
were stored and the "None recorded" line otherwise. All data synthetic."""

from __future__ import annotations

import io
import json
import re
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from account_intelligence.grading.config import GradingConfig
from exports.builder import ExportBuilder
from exports.config import ExportConfig
from exports.content import PostingRow, StackRow, TeamRow
from exports.excel.tech import LIST_HEADER, expected_tech, render_tech
from exports.models import TechModel
from exports.readers.common import MapContext
from exports.readers.tech import map_tech, stack_lane
from exports.verify import LibreOffice, recalc_check, xlsx_stack_keys
from tests.exports.corpus import account
from tests.exports.test_integration_checks import office
from tests.fakes.exports import FakeReader, InProcessRenderer, MemoryStore

MANIFEST = Path(__file__).resolve().parents[2] / "docs/parity/reference_manifest.json"
REFERENCE = json.loads(MANIFEST.read_text())["Tech_Stack_and_Teams.xlsx"]


def _posting(key: str, text: str, **over: Any) -> PostingRow:
    values: dict[str, Any] = {
        "posting_key": key,
        "title": f"Synthetic role {key}",
        "url": f"https://alpha.example/jobs/{key}",
        "status": "open",
        "posting_class": "build",
        "platforms": (),
        "platform_sentences": (),
        "description": text,
        "posted_at": date(2026, 6, 1),
        "first_seen": date(2026, 6, 2),
        "last_seen": date(2026, 9, 1),
        "full_text_read": True,
    }
    values.update(over)
    return PostingRow(**values)


def _content(postings: tuple[PostingRow, ...] = (), **update: Any) -> Any:
    return account(n_postings=0, n_detections=0).model_copy(update={"postings": postings, **update})


def _model(cfg: ExportConfig, grading: GradingConfig, content: Any) -> TechModel:
    ctx = MapContext(content, cfg, grading)
    model = map_tech(ctx).model
    assert model is not None and not ctx.violations, ctx.violations
    return model


def _workbook(cfg: ExportConfig, model: TechModel) -> tuple[Any, bytes]:
    body, _ = render_tech(model, cfg, "Synthetic tech")
    return load_workbook(io.BytesIO(body)), body


def _widths(ws: Worksheet) -> dict[str, float]:
    """Width per column letter, a grouped dimension (min to max) expanded."""
    out: dict[str, float] = {}
    for dim in ws.column_dimensions.values():
        for n in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
            out[get_column_letter(n)] = dim.width
    return out


@pytest.fixture(scope="module")
def rich_book(cfg: ExportConfig, grading: GradingConfig) -> tuple[Any, bytes, TechModel]:
    """Postings for a ranked and a seen-once platform, a stack with lane words, and two teams."""
    postings = (
        *(
            _posting(f"sf{i}", "Salesforce admin.", status="closed" if i < 2 else "open")
            for i in range(6)
        ),
        *(_posting(f"sn{i}", "ServiceNow work.") for i in range(3)),
        _posting("sap1", "SAP S/4HANA work."),
        _posting("sap2", "SAP work.", last_seen=None),
        _posting("wd", "Workday payroll."),
    )
    stack = (
        StackRow(
            id="a", business_function="Sales", category="CRM", technology="Salesforce", postings=30
        ),
        StackRow(id="b", business_function="Finance", category="ERP", technology="SAP", postings=9),
    )
    teams = (
        TeamRow(
            id="t1",
            name="Synthetic Commercial Systems",
            lane_technologies="Salesforce, SAP",
            postings=40,
            first_activity="2025-01-02",
            last_activity="2026-08-30",
            technologies="Salesforce (30); Microsoft Excel (12)",
            locations="United States (40); Ohio (12)",
            job_functions="Sales (30)",
        ),
        TeamRow(id="t2", name="Synthetic Finance", postings=3),
    )
    content = _content(
        postings,
        stack=stack,
        teams=teams,
        teams_matching=878,
        teams_since=date(2025, 4, 19),
    )
    model = _model(cfg, grading, content)
    wb, body = _workbook(cfg, model)
    return wb, body, model


# ---- the reference manifest ---------------------------------------------------------------------


def test_sheets_match_the_reference_manifest(rich_book: Any) -> None:
    wb, _body, _model = rich_book
    assert wb.sheetnames == [s["name"] for s in sorted(REFERENCE, key=lambda s: s["order"])]
    for ref in REFERENCE:
        ws = wb[ref["name"]]
        n = len(ref["headers"])
        assert [c.value for c in ws[ref["header_row"]]][:n] == ref["headers"], ref["name"]
        assert ws.max_column == n, ref["name"]
        assert ws.freeze_panes == ref["freeze"], ref["name"]
        widths = _widths(ws)
        for letter, width in ref["widths"].items():
            assert widths[letter] == width, (ref["name"], letter)
        have = {str(r) for r in ws.merged_cells.ranges}
        assert set(ref["merged"]) <= have, (ref["name"], set(ref["merged"]) - have)
        if ref["filter"] is None:
            assert ws.auto_filter.ref is None, ref["name"]
        else:
            # the filter spans the header and the data rows: same corner and same last column
            c0, r0, c1, _r1 = range_boundaries(ws.auto_filter.ref)
            ref0, ref_r0, ref1, _ = range_boundaries(ref["filter"])
            assert (c0, r0, c1) == (ref0, ref_r0, ref1), ref["name"]


def test_title_rows_and_notes_are_the_reference_wording(cfg: ExportConfig, rich_book: Any) -> None:
    wb, _body, model = rich_book
    ranking, stack, teams = (wb[n] for n in wb.sheetnames)
    assert (
        ranking["A1"].value == "Synthetic Alpha Ltd: platforms seen in job postings, ranked A, B, C"
    )
    assert ranking["A2"].value == (
        "Postings in the last 12 months that name at least one platform (the base):"
    )
    assert "In confirmed stack" in ranking["A5"].value
    assert stack["A1"].value == (
        "Synthetic Alpha Ltd: technologies seen in use, by business function and category"
    )
    assert stack["A2"].value.startswith("Pulled by 6th agents on 30 September 2026 from public")
    assert "masked" in stack["A2"].value and stack.row_dimensions[2].height == 40
    assert teams["A1"].value == (
        "Synthetic Alpha Ltd: teams whose public records name a lane platform "
        "(2 largest of 878 matching teams, by postings since 19 April 2025)"
    )
    assert teams["A2"].value == cfg.rules.sheet_texts.tech.teams_note
    # the Evidence file's line counts the teams on the sheet
    assert model.teams_note.startswith("2 teams")


# ---- the Platform Ranking -----------------------------------------------------------------------


def test_ranking_columns_lanes_open_now_last_seen_and_stack(rich_book: Any) -> None:
    wb, _body, model = rich_book
    ws = wb["Platform Ranking"]
    rows = {r[1]: r for r in ws.iter_rows(min_row=9, max_row=11, values_only=True)}
    sf, sn, sap = rows["Salesforce"], rows["ServiceNow"], rows["SAP"]
    # columns: rank, platform, category, lane, count, share, open now, last seen, in stack, read
    assert sf[3] == "Salesforce" and sn[3] == "ServiceNow" and sap[3] == "SAP"
    assert sf[4] == 6 and sf[6] == 4  # open now: postings not closed
    assert sf[7] == "2026-09-01"  # the feed's last seen, not the posted date
    assert sf[8] == "Yes" and sn[8] == "No" and sap[8] == "Yes"
    assert [r.platform for r in model.seen_once] == ["Workday"]
    # a posting with no last seen falls back to its first seen, the latest date still wins
    assert sap[7] == "2026-09-01"
    # the seen-once lane of a platform outside the lanes
    assert ws["B19"].value == "Workday" and ws["D19"].value == "Outside the lanes"


def test_last_seen_falls_back_to_first_seen(cfg: ExportConfig, grading: GradingConfig) -> None:
    postings = tuple(_posting(f"p{i}", "ServiceNow.", last_seen=None) for i in range(2))
    model = _model(cfg, grading, _content(postings))
    assert model.ranking[0].last_seen == "2026-06-02"


def test_rank_cells_are_bold_on_their_theme_fill(rich_book: Any, cfg: ExportConfig) -> None:
    wb, _body, _model = rich_book
    ws = wb["Platform Ranking"]
    fills = cfg.theme.excel.fills
    assert all(str(ws[f"A{r}"].value).startswith("=IF(") for r in (9, 10, 11))  # formulas
    letters = {9: "rank_a", 10: "rank_b", 11: "rank_c"}
    for row, name in letters.items():
        cell = ws[f"A{row}"]
        assert cell.font.b is True
        assert cell.fill.fgColor.rgb[-6:] == fills[name]
    # the conditional formats carry the same colours, so a corrected count recolours the cell
    colours = {
        rule.formula[0]: rule.dxf.fill.bgColor.rgb[-6:]
        for cf in ws.conditional_formatting
        for rule in cf.rules
    }
    assert colours == {
        '$A9="A"': fills["rank_a"],
        '$A9="B"': fills["rank_b"],
        '$A9="C"': fills["rank_c"],
    }
    assert ws["F9"].number_format == "0%"


def test_ranking_row_heights_and_widths(rich_book: Any) -> None:
    wb, _body, _model = rich_book
    ws = wb["Platform Ranking"]
    assert ws.row_dimensions[8].height == 45 and ws.row_dimensions[1].height == 18
    for row in (3, 4, 5, 6):
        assert (ws.row_dimensions[row].height or 0) > 0  # real heights for the merged notes
    assert ws["B9"].alignment.wrap_text in (False, None)


def test_sap_is_case_sensitive_and_the_patterns_follow_the_reference(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    by = {e.name: e for e in cfg.platforms}
    assert by["SAP"].case_sensitive
    assert by["SAP"].names("Runs SAP S/4HANA and Concur")
    assert not by["SAP"].names("a sap of maple, concur with you, sap")
    for text in ("experience cloud", "commerce cloud", "field service lightning", "salesforce cpq"):
        assert by["Salesforce"].names(text), text
    assert not by["Salesforce"].names("life sciences cloud")
    assert by["ServiceNow"].names("service now platform") and by["ServiceNow"].names("ServiceNow")
    assert not by["ServiceNow"].names("service nowhere")
    for text in ("eTMF", "CTMS", "veeva vault", "Vault CRM"):
        assert by["Veeva"].names(text), text
    assert by["Epic"].names("Epic Clarity reports") and by["Epic"].names("Epic Systems")
    assert not by["Epic"].names("an epic journey")
    assert by["Jira and Confluence"].names("Confluence")
    assert not by["Jira and Confluence"].names("confluence kafka")
    assert by["Jira and Confluence"].names("jira and confluence kafka")  # jira still matches
    # SAS is matched without regard to case, as the reference's build script does
    assert by["SAS"].names("SAS") and by["SAS"].names("sas macros")


def test_ranking_sorts_by_postings_with_the_dictionary_order_for_ties(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    postings = (
        *(_posting(f"w{i}", "Workday.") for i in range(2)),
        *(_posting(f"s{i}", "Salesforce.") for i in range(2)),
        *(_posting(f"t{i}", "Tableau.") for i in range(2)),
    )
    model = _model(cfg, grading, _content(postings))
    assert [r.platform for r in model.ranking] == ["Salesforce", "Tableau", "Workday"]


def test_a_stack_pull_missing_prints_not_pulled(cfg: ExportConfig, grading: GradingConfig) -> None:
    postings = tuple(_posting(f"p{i}", "ServiceNow.") for i in range(2))
    model = _model(cfg, grading, _content(postings))
    assert model.ranking[0].in_stack == cfg.rules.labels.not_pulled


@office
def test_ranking_formulas_recalculate_to_the_code_values(rich_book: Any, cfg: ExportConfig) -> None:
    _wb, body, model = rich_book
    result = recalc_check(body, expected_tech(model, cfg), LibreOffice())
    assert result.status == "passed", result.details


# ---- the Confirmed stack ------------------------------------------------------------------------


def test_stack_lane_words_follow_the_reference(cfg: ExportConfig) -> None:
    want = {
        "Salesforce": "Salesforce",
        "Salesforce Marketing Cloud": "Salesforce",
        "MuleSoft": "Salesforce (MuleSoft)",
        "ServiceNow GRC": "ServiceNow",
        "SAP SuccessFactors": "SAP",
        "Tableau": "Tableau and analytics",
        "Veeva Vault": "Veeva",
        "Workday Adaptive Planning": "Outside the lanes (Workday)",
        "Python": "",
        "Microsoft Excel": "",
    }
    for name, lane in want.items():
        assert stack_lane(cfg, name) == lane, name
    # several words: the last in the configured order wins (the reference's loop)
    assert stack_lane(cfg, "Salesforce MuleSoft connector") == "Salesforce (MuleSoft)"
    assert set(cfg.rules.workbooks.tech.stack_lanes) <= {e.name for e in cfg.platforms}


def test_stack_rows_keep_the_lookup_order_and_the_feeds_counts(
    rich_book: Any, cfg: ExportConfig
) -> None:
    wb, _body, _model = rich_book
    ws = wb[cfg.rules.workbooks.tech.sheet]
    rows = [r for r in ws.iter_rows(min_row=LIST_HEADER + 1, values_only=True)]
    assert rows == [
        ("Sales", "CRM", "Salesforce", 30, "Salesforce"),
        ("Finance", "ERP", "SAP", 9, "SAP"),
    ]
    assert ws.auto_filter.ref == f"A{LIST_HEADER}:E{LIST_HEADER + 2}"


def test_stack_row_keys_are_read_back_from_the_file(rich_book: Any, cfg: ExportConfig) -> None:
    _wb, body, model = rich_book
    keys = xlsx_stack_keys(body, cfg.rules.workbooks.tech.sheet, LIST_HEADER)
    assert keys == (
        "Sales | CRM | Salesforce",
        "Finance | ERP | SAP",
    )
    assert len(keys) == len(model.stack)


def test_no_stack_prints_the_none_found_line(cfg: ExportConfig, grading: GradingConfig) -> None:
    model = _model(cfg, grading, _content())
    wb, body = _workbook(cfg, model)
    ws = wb[cfg.rules.workbooks.tech.sheet]
    assert ws["A5"].value == cfg.rules.none_found.text and ws.auto_filter.ref == "A4:E4"
    assert xlsx_stack_keys(body, cfg.rules.workbooks.tech.sheet, LIST_HEADER) == ()


# ---- the Platform teams -------------------------------------------------------------------------


def test_teams_rows_print_when_teams_are_stored(rich_book: Any, cfg: ExportConfig) -> None:
    wb, _body, _model = rich_book
    ws = wb["Platform teams"]
    first = [c.value for c in ws[5]]
    assert first == [
        "Synthetic Commercial Systems",
        "Salesforce, SAP",
        40,
        "2025-01-02",
        "2026-08-30",
        "Salesforce (30); Microsoft Excel (12)",
        "United States (40); Ohio (12)",
        "Sales (30)",
    ]
    assert all(c.alignment.wrap_text for c in ws[5])
    assert (ws.row_dimensions[5].height or 0) > 0
    second = [c.value for c in ws[6]]
    assert second[0] == "Synthetic Finance" and second[3] in (None, "")
    assert ws.auto_filter.ref is None and ws.max_row == 6


def test_no_teams_prints_none_recorded_and_the_plain_title(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    model = _model(cfg, grading, _content())
    wb, _body = _workbook(cfg, model)
    ws = wb["Platform teams"]
    assert ws["A1"].value == "Synthetic Alpha Ltd: teams whose public records name a lane platform"
    assert ws["A5"].value == cfg.rules.none_found.teams
    assert model.teams_note == cfg.rules.none_found.teams


# ---- the whole build ----------------------------------------------------------------------


@pytest.mark.asyncio
async def test_a_full_build_verifies_the_stack_rows_in_the_produced_file(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = account()
    builder = ExportBuilder(FakeReader(content), MemoryStore(), InProcessRenderer(), cfg, grading)
    out = builder.assemble(content)
    assert not out.errors, out.errors
    tech = out.mapped["tech"]
    assert len(tech.row_ids) == len(content.stack) == 6
    assert all(re.search(r" \| ", k) for k in tech.refs)
    # the detections stay in the Evidence file's part H
    assert out.mapped["evidence"].model.detections
