"""The People List workbook against the owner's reference shape (reference set of 2026-10-20,
measured 2026-10-29; ``docs/parity/reference_manifest.json`` and ``docs/parity/reference_spec.md``):
sheet order, headers, widths, freeze pane, header height, the Priority fills, the profile link
texts, the Summary layout (title, six notes, Level table at row 11, Lane table at row 24) with its
formulas citing the label cells, the Found by strings and the former-employers cell. Every person
here is synthetic; no reference person or account name appears in this file."""

from __future__ import annotations

import io
import json
import re
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from exports.builder import ExportBuilder
from exports.config import ExportConfig, load_export_config
from exports.content import FoundVia, PeopleCounts, PersonRow, SourceRow
from exports.excel.people import expected_people
from exports.readers.common import MapContext
from exports.readers.people import map_people
from tests.exports.corpus import AS_OF, account
from tests.fakes.exports import MemoryStore

pytestmark = pytest.mark.asyncio
Make = Callable[..., ExportBuilder]

MANIFEST = Path(__file__).resolve().parents[2] / "docs/parity/reference_manifest.json"
PREFIX = "Synthetic Account"
CFG_LABELS = load_export_config().rules.labels


def _manifest() -> dict[str, Any]:
    """The People List part of the reference manifest, ``{prefix}`` replaced."""
    raw = json.loads(MANIFEST.read_text())["People_List.xlsx"]
    return json.loads(json.dumps(raw).replace("{prefix}", PREFIX))


def _person(
    src: list[SourceRow],
    host: str,
    i: int,
    title: str,
    level: str,
    role_lane: str,
    priority: str,
    **over: Any,
) -> PersonRow:
    row = SourceRow(
        code=f"S-{len(src) + 1:04d}",
        url=f"https://{host}/people/p-{i}",
        published=AS_OF,
        kind="secondary",
    )
    src.append(row)
    values: dict[str, Any] = {
        "id": f"pk-{i:04d}",
        "name": f"Synthetic Person {i:03d}",
        "title": title,
        "level": level,
        "lanes": "",
        "role_lane": role_lane,
        "priority": priority,
        "profile_url": row.url,
        "source_code": row.code,
        "origin": "title_search",
        "found_via": (FoundVia(origin="title_search", pass_no=1),),
    }
    values.update(over)
    return PersonRow(**values)


def _content(people_for: Callable[[list[SourceRow], str], list[PersonRow]]) -> Any:
    base = account(PREFIX, "synthetic-account.example", n_people=0)
    src = list(base.sources)
    people = people_for(src, base.domain or "")
    return base.model_copy(
        update={
            "people": tuple(people),
            "sources": tuple(src),
            "people_counts": PeopleCounts(
                summary=(("in scope", "", len(people)),),
                sentences=(f"{len(people)} people in scope",),
                in_scope=len(people),
            ),
        }
    )


def _spread(src: list[SourceRow], host: str) -> list[PersonRow]:
    """One person for each of several levels, lanes and priorities (all synthetic)."""
    return [
        _person(src, host, 1, "Chief Information Officer", "c_level", "technology_it", "A"),
        _person(src, host, 2, "Vice President, Shared Services", "vp", "servicenow_shared", "A"),
        _person(src, host, 3, "Senior Director, Quality", "senior_director", "business_other", "C"),
        _person(src, host, 4, "Director, Master Data", "director", "sap_erp", "B"),
        _person(
            src,
            host,
            5,
            "Associate Director, CRM",
            "associate_director",
            "salesforce_commercial",
            "A",
        ),
        _person(src, host, 6, "Head of Analytics", "head", "analytics_data_ai", "B"),
        _person(
            src, host, 7, "Senior Manager, Veeva", "senior_manager", "veeva_life_sciences", "A"
        ),
        _person(src, host, 8, "Manager, Integration", "manager", "salesforce_integration", "B"),
        _person(src, host, 9, "SAP Specialist", "other", "sap_erp", "A"),
    ]


async def _build(make_builder: Make, content: Any) -> bytes:
    store = MemoryStore()
    result = await make_builder(content, store).build_all(content.tenant_id, content.run_id)
    assert result.status == "complete", result.errors
    kinds = {d.kind: d for d in result.deliverables}
    return store.objects[kinds["people"].object_key]


# ---- the manifest ----------------------------------------------------------------------------


async def test_sheets_headers_and_widths_equal_the_reference_manifest(make_builder: Make) -> None:
    manifest = _manifest()
    wb = load_workbook(io.BytesIO(await _build(make_builder, _content(_spread))))
    assert wb.sheetnames == [s["name"] for s in sorted(manifest, key=lambda s: s["order"])]
    for sheet in manifest:
        ws = wb[sheet["name"]]
        row = [c.value for c in ws[sheet["header_row"]]]
        headers = sheet["headers"]
        assert row[: len(headers)] == headers, sheet["name"]
        for column, width in sheet["widths"].items():
            assert ws.column_dimensions[column].width == width, (sheet["name"], column)
        for merged in sheet["merged"]:
            assert merged in {str(r) for r in ws.merged_cells.ranges}, merged
        assert ws.freeze_panes == sheet["freeze"], sheet["name"]
    people = wb["People"]
    assert people.row_dimensions[1].height == 32
    # the product's own two columns follow the thirteen of the reference, nothing else
    assert [c.value for c in people[1]][13:] == ["Ref", "Source"]


async def test_people_sheet_look_fills_and_links(make_builder: Make, cfg: ExportConfig) -> None:
    wb = load_workbook(io.BytesIO(await _build(make_builder, _content(_spread))))
    ws = wb["People"]
    fills = cfg.theme.excel.fills
    rows = list(ws.iter_rows(min_row=2))
    assert len(rows) == 9
    for row in rows:
        cell = row[6]
        assert cell.value in ("A", "B", "C")
        assert cell.fill.fgColor.rgb.endswith(fills[f"priority_{cell.value.lower()}"])
        assert cell.alignment.horizontal == "center"
        # the other cells carry no priority fill
        assert row[5].fill.fill_type is None
        link = row[9]
        assert link.value == "Open profile"
        assert str(link.hyperlink.target).startswith("https://")
        assert link.font.color.rgb.endswith("1155CC") and link.font.underline == "single"
        assert row[10].value in (None, "")  # no second profile here
    assert {r[6].fill.fgColor.rgb[-6:] for r in rows} == {"DDEFE0", "FFF1CC", "EEEEEE"}
    header = ws["A1"]
    assert header.font.b and header.fill.fgColor.rgb.endswith("1F3A5F")
    assert ws["A1"].font.name == "Arial" and ws["B2"].font.name == "Arial"


async def test_rows_are_sorted_by_level_priority_lane_and_name(make_builder: Make) -> None:
    people = _spread

    # stored in reverse order; two people share a level and priority and differ by lane or name
    def many(src: list[SourceRow], host: str) -> list[PersonRow]:
        out = people(src, host)
        out.append(
            _person(
                src, host, 10, "Director, CRM Platforms", "director", "salesforce_commercial", "B"
            )
        )
        out.append(_person(src, host, 11, "Director, Cloud", "director", "technology_it", "B"))
        out.append(_person(src, host, 12, "Director, Cloud", "director", "technology_it", "A"))
        out.reverse()
        return out

    wb = load_workbook(io.BytesIO(await _build(make_builder, _content(many))))
    rows = [[c.value for c in r] for r in wb["People"].iter_rows(min_row=2)]
    keys = [(r[3], r[6], r[5], r[1]) for r in rows]
    level_order = [
        "C-level, EVP and SVP",
        "VP",
        "Senior Director",
        "Director",
        "Associate Director",
        "Head",
        "Senior Manager",
        "Manager",
        "Lead, owner, specialist or developer",
    ]
    lane_order = [
        "Salesforce and commercial",
        "Salesforce (MuleSoft and integration)",
        "ServiceNow and shared services",
        "SAP and ERP",
        "Veeva and life sciences systems",
        "Analytics, data and AI",
        "Technology, digital and IT leadership",
        "Business leadership and other",
    ]
    assert keys == sorted(
        keys, key=lambda k: (level_order.index(k[0]), k[1], lane_order.index(k[2]), k[3])
    )
    assert [r[0] for r in rows] == list(range(1, len(rows) + 1))
    directors = [r for r in rows if r[3] == "Director"]
    assert [(r[6], r[5]) for r in directors] == [
        ("A", "Technology, digital and IT leadership"),
        ("B", "Salesforce and commercial"),
        ("B", "SAP and ERP"),
        ("B", "Technology, digital and IT leadership"),
    ]


# ---- the Summary sheet -------------------------------------------------------------------------


async def test_summary_layout_formulas_and_totals(make_builder: Make, cfg: ExportConfig) -> None:
    content = _content(_spread)
    wb = load_workbook(io.BytesIO(await _build(make_builder, content)))
    s = wb["Summary"]
    assert s["A1"].value == f"{PREFIX} people list, deduplicated"
    # six notes in A3..A8, merged A:F, 58 points tall
    for row in range(3, 9):
        assert s.cell(row, 1).value, row
        assert f"A{row}:F{row}" in {str(r) for r in s.merged_cells.ranges}
        assert s.row_dimensions[row].height == 58
    assert s["A8"].value.startswith("Entity labels: ")
    assert "kept only when the title names a platform or technology" in s["A6"].value
    assert [s.cell(11, c).value for c in (1, 2, 3)] == ["Level", "People", "Priority A"]
    labels = list(cfg.rules.labels.people_level_order)
    assert [s.cell(r, 1).value for r in range(12, 21)] == labels
    assert [s.cell(21, 1).value, s.cell(22, 1).value] == ["Total", "Meets seniority rule"]
    # the formulas cite the label cell, as the reference's do
    n = 10  # nine people, the data rows end at row 10
    assert s["B12"].value == f"=COUNTIF(People!$D$2:$D${n},A12)"
    assert s["C12"].value == f'=COUNTIFS(People!$D$2:$D${n},A12,People!$G$2:$G${n},"A")'
    assert s["B21"].value == "=SUM(B12:B20)" and s["C21"].value == "=SUM(C12:C20)"
    assert s["B22"].value == f'=COUNTIF(People!$E$2:$E${n},"Yes")'
    assert s["A21"].font.b and s["B21"].font.b and s["A33"].font.b
    assert [s.cell(24, c).value for c in range(1, 6)] == [
        "Lane, senior people only",
        "A",
        "B",
        "C",
        "Total",
    ]
    lanes = [s.cell(r, 1).value for r in range(25, 33)]
    assert lanes == list(cfg.rules.labels.people_lane.values())
    assert len(lanes) == 8 and s["A33"].value == "Total"
    assert s["B25"].value == (
        f'=COUNTIFS(People!$F$2:$F${n},$A25,People!$G$2:$G${n},"A",People!$E$2:$E${n},"Yes")'
    )
    assert s["E25"].value == "=SUM(B25:D25)" and s["B33"].value == "=SUM(B25:B32)"
    assert s["A34"].value is None  # no count sentence after the Lane table
    assert s.column_dimensions["A"].width == 44 and s.column_dimensions["B"].width == 12


async def test_summary_values_equal_the_planned_counts(
    make_builder: Make, cfg: ExportConfig, grading: Any
) -> None:
    content = _content(_spread)
    ctx = MapContext(content, cfg, grading)
    model = map_people(ctx).model
    assert model is not None and not ctx.violations
    assert model.senior_count == 7  # the nine people less the two below the rule
    assert [x.total for x in model.lanes_table] == [1, 0, 1, 1, 1, 1, 1, 1]  # senior people only
    assert sum(x.people for x in model.levels) == 9
    assert len(model.paragraphs) == 6
    expected = expected_people(model, cfg)
    assert expected[("Summary", "B22")] == 7
    assert expected[("Summary", "B21")] == 9


async def test_a_run_with_one_entity_and_the_judgment_note_are_account_neutral(
    make_builder: Make, cfg: ExportConfig, grading: Any
) -> None:
    content = _content(_spread).model_copy(
        update={
            "entity_lock": {
                "subsidiaries": [{"name": "Synthetic Subsidiary Ltd", "domain": "sub.example"}]
            }
        }
    )
    ctx = MapContext(content, cfg, grading)
    model = map_people(ctx).model
    assert model is not None
    entity = model.paragraphs[-1]
    assert "Synthetic Subsidiary Ltd" in entity and PREFIX in entity
    bare = map_people(MapContext(_content(_spread), cfg, grading)).model
    assert bare is not None and "Synthetic Subsidiary" not in bare.paragraphs[-1]
    assert bare.paragraphs[-1].startswith("Entity labels: people of a business")


# ---- Found by, former employers, one lane ------------------------------------------------------


async def test_found_by_prints_the_reference_strings(cfg: ExportConfig, grading: Any) -> None:
    def ways(src: list[SourceRow], host: str) -> list[PersonRow]:
        return [
            _person(src, host, 1, "Director, X", "director", "business_other", "B"),
            _person(
                src,
                host,
                2,
                "Senior Manager, X",
                "senior_manager",
                "business_other",
                "B",
                found_via=(FoundVia(origin="title_search", pass_no=2),),
            ),
            _person(
                src,
                host,
                3,
                "Director, Y",
                "director",
                "business_other",
                "B",
                origin="association",
                found_via=(
                    FoundVia(origin="association", lane="tableau_analytics"),
                    FoundVia(origin="association", lane="servicenow"),
                ),
            ),
            _person(
                src,
                host,
                4,
                "Director, Z",
                "director",
                "business_other",
                "B",
                origin="mention",
                found_via=(
                    FoundVia(
                        origin="mention",
                        mention_url="https://synthetic-account.example/about/leadership",
                    ),
                ),
            ),
            _person(
                src,
                host,
                5,
                "Director, W",
                "director",
                "business_other",
                "B",
                origin="mention",
                found_via=(
                    FoundVia(origin="mention", mention_url="https://synthetic-news.example/story"),
                ),
            ),
            _person(
                src,
                host,
                6,
                "Director, V",
                "director",
                "business_other",
                "B",
                found_via=(
                    FoundVia(origin="title_search", pass_no=1),
                    FoundVia(origin="title_search", pass_no=2),
                    FoundVia(origin="association", lane="veeva"),
                ),
            ),
            _person(src, host, 7, "Director, U", "director", "business_other", "B", found_via=()),
        ]

    ctx = MapContext(_content(ways), cfg, grading)
    model = map_people(ctx).model
    assert model is not None and not ctx.violations
    by_name = {r.name: r.found_by for r in model.rows}
    assert by_name["Synthetic Person 001"] == "title search source 1 (leadership and platform)"
    assert by_name["Synthetic Person 002"] == "title search source 2 (manager)"
    assert by_name["Synthetic Person 003"] == "technology association search (Tableau, ServiceNow)"
    assert by_name["Synthetic Person 004"] == "Leadership page; profile found by search"
    assert by_name["Synthetic Person 005"] == "Research page; profile found by search"
    assert by_name["Synthetic Person 006"] == (
        "title search source 1 (leadership and platform); title search source 2 (manager);"
        " technology association search (Veeva)"
    )
    assert by_name["Synthetic Person 007"] == "title search"  # a record without its pass


async def test_former_employers_cell_is_empty_without_a_profile_and_prints_dated_firms(
    cfg: ExportConfig, grading: Any
) -> None:
    def people(src: list[SourceRow], host: str) -> list[PersonRow]:
        return [
            _person(src, host, 1, "Director, X", "director", "business_other", "B"),
            _person(
                src,
                host,
                2,
                "Director, Y",
                "director",
                "business_other",
                "B",
                former_employers=("Alpha (2015 to 2019)", "Beta (2020 to date)"),
            ),
            _person(
                src, host, 3, "Director, Z", "director", "business_other", "B", former_employers=()
            ),
        ]

    model = map_people(MapContext(_content(people), cfg, grading)).model
    assert model is not None
    cells = {r.name: r.former_employers for r in model.rows}
    assert cells["Synthetic Person 001"] == ""  # nothing read: no claim, and not "Not pulled"
    assert cells["Synthetic Person 002"] == "Alpha (2015 to 2019); Beta (2020 to date)"
    assert cells["Synthetic Person 003"] == ""
    assert "Not pulled" not in " ".join(cells.values())


async def test_one_lane_per_person_and_a_stored_record_without_one_falls_back(
    cfg: ExportConfig, grading: Any
) -> None:
    def people(src: list[SourceRow], host: str) -> list[PersonRow]:
        return [
            _person(
                src, host, 1, "Director, X", "director", "sap_erp", "B", lanes="salesforce,sap"
            ),
            # a record stored before the role lanes existed: the first platform lane decides
            _person(
                src, host, 2, "Director, Y", "director", "", "B", lanes="tableau_analytics,sap"
            ),
            _person(src, host, 3, "Director, Z", "director", "", "B", lanes=""),
        ]

    model = map_people(MapContext(_content(people), cfg, grading)).model
    assert model is not None
    lane = {r.name: r.lanes for r in model.rows}
    assert lane["Synthetic Person 001"] == "SAP and ERP"
    assert lane["Synthetic Person 002"] == "Analytics, data and AI"
    assert lane["Synthetic Person 003"] == "Business leadership and other"
    assert all("Also in lane" not in r.note for r in model.rows)


async def test_search_and_listing_counts_are_registered(cfg: ExportConfig, grading: Any) -> None:
    ctx = MapContext(_content(_spread), cfg, grading)
    assert map_people(ctx).model is not None
    keys = {n.key for n in ctx.registry.items()}
    assert {"people.listed.title_search.pass1", "people.searches.title.pass1.records"} <= keys


@pytest.mark.requires_postgres
async def test_the_file_built_from_the_real_people_pipeline_has_the_reference_columns() -> None:
    """Corpus -> people pipeline -> Postgres -> reader -> workbook: one lane, a level label and the
    reference Found by on every row, people below the rule listed with No."""
    from tests.exports.test_db import builder, seed_world

    w, payloads = await seed_world()
    result = await builder(w, payloads).build_all(str(w.tenant), str(w.run))
    assert result.status == "complete", result.errors
    kinds = {d.kind: d for d in result.deliverables}
    wb = load_workbook(io.BytesIO(payloads.objects[kinds["people"].object_key]))
    rows = [[c.value for c in r] for r in wb["People"].iter_rows(min_row=2)]
    assert rows
    lanes = set(CFG_LABELS.people_lane.values())
    levels = set(CFG_LABELS.people_level_order)
    for row in rows:
        assert row[3] in levels and row[4] in ("Yes", "No") and row[5] in lanes
        assert row[6] in ("A", "B", "C") and row[9] == "Open profile"
        assert re.fullmatch(
            r"(title search source [12] \((leadership and platform|manager)\)"
            r"|technology association search \([^)]+\)"
            r"|Leadership page; profile found by search|Research page; profile found by search)"
            r"(; .+)?",
            row[8],
        ), row[8]
        assert row[7] in (None, "")  # no profile pull in production: nothing claimed
    assert {r[4] for r in rows} == {"Yes", "No"}  # the corpus has people below the rule
    below = [r for r in rows if r[4] == "No"]
    assert all(r[3] in ("Manager", "Lead, owner, specialist or developer") for r in below)
