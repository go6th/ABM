"""The owner's output specification for the three workbooks (master prompt v2, 2026-10-24,
sections 7.3, 9.3 to 9.5 and 12): the People List labels and Summary tables, the Job Postings
columns and matrix, the Platform Ranking, the masking of out-of-scope platforms in every file,
and the third-party scan over every file. Every formula cell is registered with its value."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from datetime import date
from pathlib import Path
from typing import Any

import pytest
import yaml
from openpyxl import load_workbook

from account_intelligence.config import CONFIG_DIR
from account_intelligence.grading.config import GradingConfig
from exports.builder import ExportBuilder
from exports.config import ExportConfig, load_export_config
from exports.content import (
    FoundVia,
    PeopleCounts,
    PersonRow,
    PostingRow,
    SourceRow,
    StackRow,
)
from exports.errors import ConfigError
from exports.excel.jobs import expected_jobs
from exports.excel.people import expected_people
from exports.excel.tech import expected_tech
from exports.models import BuildResult
from exports.readers.common import MapContext
from exports.readers.jobs import map_jobs
from exports.readers.people import map_people
from exports.readers.tech import map_tech
from jobs.window import months_before
from tests.exports.corpus import AS_OF, account
from tests.fakes.exports import MemoryStore

pytestmark = pytest.mark.asyncio
Make = Callable[..., ExportBuilder]


async def _build(make_builder: Make, content: Any) -> tuple[BuildResult, dict[str, bytes]]:
    store = MemoryStore()
    result = await make_builder(content, store).build_all(content.tenant_id, content.run_id)
    files = {}
    for d in result.deliverables:
        files[d.kind] = store.objects[d.object_key]
    return result, files


def _texts(data: bytes) -> str:
    """Every text member of a package, nested packages included (what a reader could see)."""
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in z.namelist():
            body = z.read(name)
            if name.endswith((".docx", ".xlsx")):
                out.append(_texts(body))
            else:
                out.append(body.decode("utf-8", "ignore"))
    return "\n".join(out)


def _formula_cells(data: bytes) -> dict[tuple[str, str], str]:
    wb = load_workbook(io.BytesIO(data))
    return {
        (ws.title, c.coordinate): str(c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.data_type == "f"
    }


def _posting(src: Any, key: str, text: str, **over: Any) -> PostingRow:
    s = src.add(f"jobs/{key}", kind="primary")
    values: dict[str, Any] = {
        "posting_key": key,
        "title": f"Synthetic role {key}",
        "url": s.url,
        "status": "open",
        "posting_class": "build",
        "platforms": ("Salesforce",),
        "platform_sentences": (text,),
        "description": text,
        "posted_at": date(2026, 6, 1),
        "first_seen": date(2026, 6, 2),
        "full_text_read": True,
    }
    values.update(over)
    return PostingRow(**values)


class _Registry:
    """Adds sources to an existing content's registry, continuing its codes."""

    def __init__(self, content: Any) -> None:
        self.rows = list(content.sources)
        self.host = content.domain

    def add(self, path: str, kind: str = "primary", host: str | None = None) -> SourceRow:
        row = SourceRow(
            code=f"S-{len(self.rows) + 1:04d}",
            url=f"https://{host or self.host}/{path}",
            published=AS_OF,
            kind=kind,
        )
        self.rows.append(row)
        return row


# ---- 4. masking ------------------------------------------------------------------------------


async def test_out_of_scope_platforms_are_masked_in_every_file(
    make_builder: Make, cfg: ExportConfig
) -> None:
    c = account()
    src = _Registry(c)
    never = cfg.scope.never_report
    postings = (
        *c.postings,
        _posting(src, "mask-1", f"Administers Salesforce and {never[0]} daily."),
    )
    det = c.detections[0].model_copy(update={"technology": never[2], "id": "mask-det"})
    # the organisation's stack also returns an out-of-scope platform: the row stays, its name
    # is masked
    hidden = StackRow(
        id="mask-stack",
        business_function="Data & Analytics",
        category="Cloud Data Warehouse",
        technology=never[2],
        postings=9,
    )
    content = c.model_copy(
        update={
            "postings": postings,
            "detections": (*c.detections, det),
            "stack": (*c.stack, hidden),
            "sources": tuple(src.rows),
        }
    )
    result, files = await _build(make_builder, content)
    assert result.status == "complete", result.errors
    assert result.manifest is not None
    # the description, the platform sentence, the detection technology and the stack technology
    # at least; the evidence file repeats the sentence
    assert result.manifest.masked_terms >= 4
    for kind, body in files.items():
        text = _texts(body).casefold()
        for name in never:
            assert name.casefold() not in text, (kind, name)
    # the mask itself is on the sheets that carried the names
    assert cfg.scope.mask in _texts(files["jobs"]) and cfg.scope.mask in _texts(files["tech"])
    # the masked stack row is still a row of the Confirmed stack, with its postings count
    stack = load_workbook(io.BytesIO(files["tech"]))[cfg.rules.workbooks.tech.sheet]
    masked = [r for r in stack.iter_rows(min_row=5, values_only=True) if r[2] == cfg.scope.mask]
    assert masked == [("Data & Analytics", "Cloud Data Warehouse", cfg.scope.mask, 9, None)]
    note = stack["A2"].value
    assert "masked" in note and "kept as returned" not in note


def test_never_report_names_are_never_platforms_or_lane_labels(cfg: ExportConfig) -> None:
    names = {n.casefold() for n in cfg.scope.never_report}
    for e in cfg.platforms:
        assert e.name.casefold() not in names
        for n in cfg.scope.never_report:
            assert not e.names(n)
    for lane_id in cfg.scope.lanes:
        label = cfg.lane_label(lane_id) or ""
        assert not any(n in label.casefold() for n in names), label
    for label in cfg.rules.labels.people_lane.values():  # the People List's own lane wording
        assert not any(n in label.casefold() for n in names), label


def test_platform_dictionary_refuses_a_bad_pattern_or_a_masked_name(tmp_path: Path) -> None:
    for name in ("export_theme.yaml", "export_rules.yaml", "lanes.yaml", "output_spec.yaml"):
        (tmp_path / name).write_bytes((CONFIG_DIR / name).read_bytes())
    doc = yaml.safe_load((CONFIG_DIR / "platform_dictionary.yaml").read_text())
    bad = {**doc, "platforms": [{**doc["platforms"][0], "pattern": "("}]}
    (tmp_path / "platform_dictionary.yaml").write_text(yaml.safe_dump(bad))
    with pytest.raises(ConfigError, match="does not compile"):
        load_export_config(tmp_path)
    scope = yaml.safe_load((CONFIG_DIR / "skill" / "scope.yaml").read_text())
    masked = scope["never_report"]["value"][0]
    bad = {**doc, "platforms": [{**doc["platforms"][0], "name": masked}]}
    (tmp_path / "platform_dictionary.yaml").write_text(yaml.safe_dump(bad))
    with pytest.raises(ConfigError, match="never ranked"):
        load_export_config(tmp_path)
    for e in load_export_config().platforms:
        assert e.regex.pattern == e.pattern  # every pattern compiles


# ---- 1. People List ----------------------------------------------------------------------------


def test_people_labels_seniority_lane_and_found_by(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    c = account(n_people=0)
    src = _Registry(c)
    # (level, role lane, platform lanes, priority, found via, level label, meets the rule)
    cases = (
        (
            "c level",
            "technology_it",
            "salesforce, data ai",
            "A",
            (FoundVia(origin="title_search", pass_no=1),),
            "C-level, EVP and SVP",
            True,
        ),
        (
            "vp",
            "",
            "",
            "B",
            (FoundVia(origin="association", lane="tableau_analytics"),),
            "VP",
            True,
        ),
        (
            "senior manager",
            "",
            "data_ai",
            "B",
            (FoundVia(origin="mention", mention_url="https://x.example/about/leadership"),),
            "Senior Manager",
            True,
        ),
        (
            "manager",
            "veeva_life_sciences",
            "veeva",
            "C",
            (FoundVia(origin="title_search", pass_no=2),),
            "Manager",
            False,
        ),
        ("analyst", "", "", "C", (), "Lead, owner, specialist or developer", False),
    )
    people = []
    for i, (level, role_lane, lanes, prio, via, _label, _senior) in enumerate(cases):
        s = src.add(f"people/p-{i}", kind="secondary")
        people.append(
            PersonRow(
                id=f"pk-{i}",
                name=f"Synthetic Person {i}",
                title=f"Synthetic title {i}",
                level=level,
                lanes=lanes,
                role_lane=role_lane,
                priority=prio,
                profile_url=s.url,
                source_code=s.code,
                origin=via[0].origin if via else "title_search",
                found_via=via,
                former_employers=("Synthetic Integrator (2015 to 2019)",) if i == 0 else None,
                second_profile_url=f"https://{c.domain}/people/alt-{i}" if i == 1 else None,
                location="Synthetic City, Synthetic Country" if i == 0 else None,
            )
        )
    # stored in reverse: the list orders them by level, priority, lane and name
    people.reverse()
    content = c.model_copy(update={"people": tuple(people), "sources": tuple(src.rows)})
    ctx = MapContext(content, cfg, grading)
    model = map_people(ctx).model
    assert model is not None and not ctx.violations
    labels = cfg.rules.labels
    for row, (_, _lane, _plat, _, _via, label, senior) in zip(model.rows, cases, strict=True):
        assert row.level == label and row.meets_seniority is senior
    assert [r.found_by for r in model.rows] == [
        "title search source 1 (leadership and platform)",
        "technology association search (Tableau)",
        "Leadership page; profile found by search",
        "title search source 2 (manager)",
        "title search",
    ]
    assert [r.number for r in model.rows] == [1, 2, 3, 4, 5]
    assert [r.row_id for r in model.rows] == [f"pk-{i}" for i in range(5)]
    # ONE lane per person: the stored role lane; a record without one takes the lane of its first
    # platform lane (stored before the role lanes existed), else the default lane
    assert [r.lanes for r in model.rows] == [
        labels.people_lane["technology_it"],
        labels.people_lane["business_other"],
        labels.people_lane["analytics_data_ai"],
        labels.people_lane["veeva_life_sciences"],
        labels.people_lane["business_other"],
    ]
    assert all(r.note == "" for r in model.rows)
    assert model.rows[0].location == "Synthetic City, Synthetic Country"
    assert model.rows[1].location == ""
    assert model.rows[0].former_employers == "Synthetic Integrator (2015 to 2019)"
    assert model.rows[1].former_employers == ""  # no profile read: the cell stays empty
    assert model.rows[1].second_profile_url and model.rows[0].second_profile_url is None
    assert model.senior_count == 3
    by_label = {x.label: x for x in model.levels}
    assert list(by_label) == list(labels.people_level_order)
    assert by_label["C-level, EVP and SVP"].people == 1
    assert by_label["C-level, EVP and SVP"].priority_a == 1
    assert sum(x.people for x in model.levels) == 5
    lanes_table = {x.label: x for x in model.lanes_table}
    assert list(lanes_table) == list(labels.people_lane.values()) and len(lanes_table) == 8
    assert lanes_table[labels.people_lane["technology_it"]].a == 1
    # the senior manager (analytics) and the vp (default lane) are senior; the manager is not
    assert lanes_table[labels.people_lane["analytics_data_ai"]].b == 1
    assert lanes_table[labels.people_lane["business_other"]].b == 1
    assert lanes_table[labels.people_lane["veeva_life_sciences"]].total == 0
    assert cfg.scope.seniority_rule in model.seniority_text
    assert "kept only when the title names a platform or technology" in model.seniority_text
    assert model.seniority_text.endswith("Filter Meets seniority rule to Yes for the target list.")
    # the Summary sheet's title and its six paragraphs in the reference's order
    assert model.title_line == f"{c.account_name} people list, deduplicated"
    assert len(model.paragraphs) == 6
    what = model.paragraphs[0]
    assert what.startswith(f"This is the single people list for the {c.account_name} account.")
    # the title search, one line per pass with its own count of listed people
    assert "leadership and platform titles, 1 person on this list" in what
    assert "manager titles, 1 person on this list" in what
    assert "a technology association search for the Tableau gaps (1 person on this list)" in what
    assert "people named in research (1 person on this list)" in what
    assert "Pulled by 6th agents on " in what
    assert model.paragraphs[-3] == model.seniority_text
    assert model.paragraphs[-2] == model.judgment_note
    assert model.paragraphs[-1].startswith("Entity labels: ")
    registered = {n.key for n in ctx.registry.items()}
    assert {"people.listed", "people.withheld", "people.senior", "people.level.00"} <= registered


async def test_people_summary_names_the_excluded_with_reason_and_formulas_are_registered(
    make_builder: Make, cfg: ExportConfig
) -> None:
    c = account()
    counts = c.people_counts
    assert counts is not None
    content = c.model_copy(
        update={
            "people_counts": PeopleCounts(
                summary=counts.summary,
                sentences=counts.sentences,
                in_scope=counts.in_scope,
                excluded=(
                    ("Synthetic Excluded Person", "security", "Director, Information Security"),
                ),
            )
        }
    )
    result, files = await _build(make_builder, content)
    assert result.status == "complete", result.errors
    wb = load_workbook(io.BytesIO(files["people"]))
    assert wb.sheetnames[:2] == ["Summary", "People"]  # the reference's sheet order
    people = [[c.value for c in row] for row in wb["People"].iter_rows()]
    assert people[0][:13] == [
        "No.",
        "Name",
        "Title as reported",
        "Level",
        "Meets seniority rule",
        "Lane",
        "Priority",
        "Former SI and consulting employers (firm and years)",
        "Found by",
        "LinkedIn",
        "Second profile",
        "Location",
        "Note",
    ]
    assert people[0][-1] == "Source" and wb["People"].freeze_panes == "C2"
    link = wb["People"].cell(row=2, column=10)
    assert link.value == "Open profile" and str(link.hyperlink.target).startswith("https://")
    assert "Synthetic Excluded Person" not in str(people)
    summary = [[c.value for c in row] for row in wb["Summary"].iter_rows()]
    assert summary[0][0] == f"{c.account_name} people list, deduplicated"
    text = " ".join(str(v) for row in summary for v in row if v is not None)
    assert (
        "Security people are excluded at the account owner's instruction (1 row): "
        "Synthetic Excluded Person (Director, Information Security)."
    ) in text
    assert cfg.scope.seniority_rule in text and "self reported" in text
    assert "Entity labels: " in text
    assert "Rows removed before this list: " in text and "1 row merged as duplicates" in text
    labels_col = [row[0] for row in summary]
    assert labels_col.index("Level") < labels_col.index("Lane, senior people only")
    level_at = labels_col.index("Level")
    assert labels_col[level_at + 1 : level_at + 10] == list(cfg.rules.labels.people_level_order)
    assert labels_col[level_at + 10 : level_at + 12] == ["Total", "Meets seniority rule"]
    # the merged person stays guarded: never in any file
    for body in files.values():
        assert "Synthetic Merged Person" not in _texts(body)
    mapped = make_builder(content).assemble(content).mapped["people"].model
    assert mapped is not None
    expected = expected_people(mapped, cfg)
    formulas = _formula_cells(files["people"])
    assert set(expected) == set(formulas), set(expected) ^ set(formulas)
    assert all(sheet == "Summary" for sheet, _ in expected)


# ---- 2. Job Postings ---------------------------------------------------------------------------


def test_jobs_labels_entity_and_source(cfg: ExportConfig, grading: GradingConfig) -> None:
    c = account(n_postings=0).model_copy(update={"careers_host": "jobs.alpha-careers.example"})
    src = _Registry(c)
    postings = (
        _posting(src, "a", "Builds on Salesforce.", posting_class="build"),
        _posting(
            src,
            "b",
            "Salesforce named in the title only.",
            posting_class=None,
            full_text_read=False,
        ),
        _posting(
            src,
            "c",
            "A role with no lane platform in the text.",
            posting_class=None,
            platforms=("Workday",),
            platform_sentences=("Works in Workday.",),
            status="closed",
            posted_at=None,
            origin="live_board",
        ),
        _posting(src, "d", "Salesforce admin.", posting_class="admin_owner", origin="live_board"),
    )
    s_sub = src.add("jobs/sub", host=f"careers.{c.domain}")
    s_car = src.add("jobs/car", host="jobs.alpha-careers.example")
    s_other = src.add("jobs/other", host="boards.example", kind="aggregator")
    postings += (
        postings[0].model_copy(update={"posting_key": "e", "url": s_sub.url}),
        postings[0].model_copy(update={"posting_key": "f", "url": s_car.url}),
        postings[0].model_copy(update={"posting_key": "g", "url": s_other.url}),
    )
    content = c.model_copy(update={"postings": postings, "sources": tuple(src.rows)})
    ctx = MapContext(content, cfg, grading)
    model = map_jobs(ctx).model
    assert model is not None and not ctx.violations, ctx.violations
    labels = cfg.rules.labels
    rows = {r.row_id: r for r in model.rows}
    assert rows["a"].posting_class == "Build role"
    # the job pipeline gives every exported posting one of the four classes (a reading confirms
    # Build and Admin and owner roles, the heuristic decides the rest); a posting with no stored
    # class at all is the pipeline's own "Unclassified", never a derived "title only" label
    assert rows["b"].posting_class == "Unclassified" and rows["b"].description == ""
    assert rows["c"].posting_class == "Unclassified"
    assert rows["c"].status == "Closed"  # the stored status word
    assert rows["c"].posted_or_seen == date(2026, 6, 2)  # first seen when not posted
    assert rows["a"].source_type == labels.posting_source["feed_read"]
    assert rows["b"].source_type == labels.posting_source["feed_title_only"]
    assert rows["d"].source_type == labels.posting_source["live_board_read"]
    for key in ("a", "e", "f"):
        assert rows[key].entity == c.account_name, key
    assert rows["g"].entity == "boards.example" + labels.entity_to_confirm_suffix
    # the matrix is the seven reference platforms, the account's own entity only (row g is on
    # another website, so it is listed on the Jobs sheet and left out of the counts)
    assert [m.platform for m in model.matrix] == list(cfg.rules.workbooks.jobs.matrix_platforms)
    matrix = {m.platform: m for m in model.matrix}
    assert matrix["Salesforce"].rows == 5 and matrix["Workday"].rows == 1
    assert model.class_labels == ("Build role", "Admin and owner role", "Working tool", "User role")
    build_index = model.class_labels.index("Build role")
    assert matrix["Salesforce"].by_class[build_index] == 3
    assert matrix["Workday"].closed == 1 and matrix["Workday"].open == 0
    assert matrix["Salesforce"].closed == 0 and matrix["Salesforce"].open == 5
    assert {n.key for n in ctx.registry.items()} >= {"jobs.matrix.00", "jobs.matrix.00.00"}


def test_jobs_reference_columns(cfg: ExportConfig, grading: GradingConfig) -> None:
    """The reference file's columns: the requisition the posting prints, the stored location,
    the sentence the reading quoted (read postings only) and the full description apart."""
    c = account(n_postings=0)
    src = _Registry(c)
    text = "Job ID\n:   R-557322\n\nCategory\n: IT\n\nYou administer Salesforce for the sales team."
    postings = (
        _posting(
            src,
            "a",
            text,
            location="Heidelberg, Germany",
            class_sentence="You administer Salesforce for the sales team.",
        ),
        _posting(
            src,
            "b",
            "Requisition ID 1234 uses Salesforce.",
            full_text_read=False,
            class_sentence="uses Salesforce.",
        ),
        _posting(src, "c", "No job number printed here; works in Salesforce."),
    )
    content = c.model_copy(update={"postings": postings, "sources": tuple(src.rows)})
    model = map_jobs(MapContext(content, cfg, grading)).model
    assert model is not None
    rows = {r.row_id: r for r in model.rows}
    assert rows["a"].requisition == "R-557322" and rows["a"].location == "Heidelberg, Germany"
    assert rows["a"].read_text == "You administer Salesforce for the sales team."
    assert rows["a"].description.startswith("Job ID")
    # not read in full: no quoted sentence and no description, the printed ID still shown
    assert rows["b"].requisition == "1234" and rows["b"].read_text == ""
    assert rows["b"].description == ""
    assert rows["c"].requisition == "" and rows["c"].location == ""
    # the Summary sheet's title line and paragraphs
    assert model.summary_title == (
        f"{c.account_name} platform job postings, last 12 months"
        " (first seen 30 September 2025 to 30 September 2026)"
    )
    joined = " ".join(model.summary_notes)
    assert len(model.summary_notes) == 5  # the reference Summary's five paragraphs
    assert "6th agents" in joined and "entity to confirm" in joined
    # the class paragraph names the four classes of the matrix, the platform paragraph the rules
    for key in cfg.rules.workbooks.jobs.matrix_classes:
        assert cfg.rules.labels.posting_class[key] in model.summary_notes[2], key
    for word in (
        "S/4HANA",
        "Ariba",
        "Concur",
        "SuccessFactors",
        "SFDC",
        "eTMF",
        "CTMS",
        "PromoMats",
    ):
        assert word in model.summary_notes[4], word
    headers = {col.key: col.header for col in cfg.rules.workbooks.jobs.columns}
    for key in ("requisition", "read_text", "platform_sentence", "description", "archive_url"):
        assert headers[key] in joined, key
    for key in ("new_flag", "ref", "source_code"):
        assert headers[key] in joined, key
    assert c.account_name in model.matrix_heading and "positions only" in model.matrix_heading


async def test_jobs_sheet_and_matrix_formulas_are_registered(
    make_builder: Make, cfg: ExportConfig
) -> None:
    content = account()
    result, files = await _build(make_builder, content)
    assert result.status == "complete", result.errors
    wb = load_workbook(io.BytesIO(files["jobs"]))
    book = cfg.rules.workbooks.jobs
    # the reference file's order: Summary first, then Jobs; Counts and Long text after
    order = [book.summary_sheet, book.sheet, book.counts_sheet, book.long_text_sheet]
    assert wb.sheetnames == order
    headers = [c.value for c in wb["Jobs"][1]]
    assert headers == [c.header for c in book.columns]
    assert headers[:3] == ["Requisition", "Title", "Platforms named"]
    mapped = make_builder(content).assemble(content).mapped["jobs"].model
    assert mapped is not None
    summary = wb[book.summary_sheet]
    assert summary["A1"].value == mapped.summary_title
    assert summary["A1"].value.startswith(f"{content.account_name} platform job postings")
    assert [summary.cell(row=3 + i, column=1).value for i in range(len(mapped.summary_notes))] == (
        list(mapped.summary_notes)
    )
    texts = cfg.rules.sheet_texts.jobs
    header_row = 3 + len(mapped.summary_notes) + 2  # two blank rows, as the reference (row 10)
    assert header_row == 10
    matrix_header = [c.value for c in summary[header_row]][:8]
    assert matrix_header == [
        texts.matrix_heading.format(prefix=mapped.prefix),
        texts.total,
        texts.closed,
        texts.open,
        *mapped.class_labels,
    ]
    expected = expected_jobs(mapped, cfg)
    formulas = _formula_cells(files["jobs"])
    assert set(expected) == set(formulas), set(expected) ^ set(formulas)
    assert {sheet for sheet, _ in expected} == {book.summary_sheet, book.counts_sheet}


# ---- 3. Platform Ranking -----------------------------------------------------------------------


def _ranking_content() -> Any:
    c = account(n_postings=0, n_detections=0)
    src = _Registry(c)
    postings = []
    # Salesforce named 6 times (A), ServiceNow 3 times (B), SAP twice (C), Workday once, and
    # one old posting outside the window naming Tableau twice over
    for i in range(6):
        postings.append(_posting(src, f"sf-{i}", "Salesforce admin work.", full_text_read=i < 3))
    for i in range(3):
        postings.append(_posting(src, f"sn-{i}", "ServiceNow ITSM work.", status="closed"))
    for i in range(2):
        postings.append(_posting(src, f"sap-{i}", "SAP S/4HANA work.", full_text_read=False))
    postings.append(_posting(src, "wd", "Workday payroll.", posted_at=None, first_seen=None))
    postings.append(_posting(src, "old", "Tableau Tableau.", posted_at=date(2024, 1, 1)))
    return c.model_copy(update={"postings": tuple(postings), "sources": tuple(src.rows)})


def test_platform_ranking_counts_rank_and_window(cfg: ExportConfig, grading: GradingConfig) -> None:
    content = _ranking_content()
    ctx = MapContext(content, cfg, grading)
    model = map_tech(ctx).model
    assert model is not None and not ctx.violations
    assert months_before(AS_OF, 12) == date(2025, 9, 30)
    assert model.postings_in_window == 12 and model.base == 12
    ranked = [(r.platform, r.count, r.rank) for r in model.ranking]
    assert ranked == [("Salesforce", 6, "A"), ("ServiceNow", 3, "B"), ("SAP", 2, "C")]
    assert [r.platform for r in model.seen_once] == ["Workday"]
    assert model.rank_counts == (("A", 1), ("B", 1), ("C", 1))
    sf, sn, sap = model.ranking
    # last seen: the feed's last seen, else the first seen (the posted date is not a last seen)
    assert sf.share == 0.5 and sf.open_now == 6 and sf.last_seen == "2026-06-02"
    assert sf.read_checked == "Yes" and sap.read_checked == "No"
    assert sf.in_stack == cfg.rules.labels.not_pulled  # the run has no stack
    assert sn.lane == cfg.rules.labels.tech_lane["servicenow"] == "ServiceNow"
    assert sf.lane == "Salesforce"  # the Tech workbook's own lane words, not the people list's
    assert model.seen_once[0].lane == cfg.rules.labels.lane_outside
    assert "12 months" in model.base_line and model.title.startswith(content.account_name)
    registered = {n.key: n.value for n in ctx.registry.items()}
    assert registered["tech.ranking.base"] == 12 and registered["tech.ranking.rank.A"] == 1


def test_in_confirmed_stack_follows_the_stack_items(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = _ranking_content()
    item = StackRow(
        id="st-1",
        business_function="Support",
        category="Contact Center Technologies",
        technology="Salesforce Service Cloud",
        postings=10,
    )
    content = content.model_copy(update={"stack": (item,)})
    model = map_tech(MapContext(content, cfg, grading)).model
    assert model is not None
    by = {r.platform: r for r in model.ranking}
    assert by["Salesforce"].in_stack == "Yes" and by["ServiceNow"].in_stack == "No"
    # the technology detections no longer decide it
    assert model.rows == ()


async def test_ranking_sheet_is_first_and_its_formulas_are_registered(
    make_builder: Make, cfg: ExportConfig
) -> None:
    content = _ranking_content()
    result, files = await _build(make_builder, content)
    assert result.status == "complete", result.errors
    wb = load_workbook(io.BytesIO(files["tech"]))
    book = cfg.rules.workbooks.tech
    assert wb.sheetnames == ["Platform Ranking", "Confirmed stack", "Platform teams"]
    assert book.sheet == "Confirmed stack"
    ws = wb["Platform Ranking"]
    # the reference layout: title, base line with the base in the share column, four rule
    # paragraphs, a gap, the header on row 8 and the first platform on row 9
    title = f"{content.account_name}: platforms seen in job postings, ranked A, B, C"
    assert ws["A1"].value == title
    assert ws["A2"].value.startswith("Postings in the last 12 months") and ws["F2"].value == 12
    assert [c.value for c in ws[8]] == [c.header for c in book.ranking_columns]
    assert ws.freeze_panes == "A9" and ws.auto_filter.ref == "A8:J11"
    assert ws["A9"].value.startswith("=IF(") and ws["F9"].value == "=IF($F$2=0,0,E9/$F$2)"
    assert ws["B9"].value == "Salesforce" and ws["B11"].value == "SAP"
    assert ws["B14"].value == '=COUNTIF($A$9:$A$11,"A")' and ws["A14"].value == "Rank A"
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    for name in ("rank_rule", "count_rule", "signal_rule", "research_by"):
        assert getattr(cfg.rules.sheet_texts.tech, name) in text, name
    assert "Seen once (not ranked)" in text and "Workday" in text
    mapped = make_builder(content).assemble(content).mapped["tech"].model
    assert mapped is not None
    expected = expected_tech(mapped, cfg)
    formulas = _formula_cells(files["tech"])
    assert set(expected) == set(formulas), set(expected) ^ set(formulas)
    assert expected[("Platform Ranking", "A9")] == "A"
    assert expected[("Platform Ranking", "F9")] == 0.5
    # the seen-once block: platform, category and lane only, under its heading
    assert ws["A18"].value == "Seen once (not ranked)" and ws["B19"].value == "Workday"
    assert ws["E19"].value is None
    stack = wb[book.sheet]
    assert stack["A1"].value.startswith(content.account_name)
    assert [c.value for c in stack[4]] == [c.header for c in book.columns]
    teams = wb["Platform teams"]
    assert teams["A2"].value == cfg.rules.sheet_texts.tech.teams_note
    assert [c.value for c in teams[4]] == [c.header for c in book.teams_columns]
    assert teams["A5"].value == cfg.rules.none_found.teams and teams.max_row == 5


# ---- 5. third-party scan -----------------------------------------------------------------------


async def test_third_party_scan_groups_hosts_and_withholds_the_cost_word(
    make_builder: Make, cfg: ExportConfig
) -> None:
    c = account()
    src = _Registry(c)
    s_agg = src.add("board/1", host="boards.example", kind="aggregator")
    s_other = src.add("news/1", host="press.example", kind="secondary")
    postings = (
        *c.postings,
        _posting(src, "agg", "Salesforce role on a board.").model_copy(update={"url": s_agg.url}),
        _posting(src, "oth", "Salesforce role in the press.").model_copy(
            update={"url": s_other.url}
        ),
    )
    content = c.model_copy(update={"postings": postings, "sources": tuple(src.rows)})
    result, _files = await _build(make_builder, content)
    assert result.status == "complete", result.errors
    groups = result.manifest.third_party_domains
    assert groups["data_tools"] == ()
    assert c.domain in groups["account_and_press"]
    assert groups["aggregators"] == ("boards.example",)
    assert "press.example" in groups["other"]
    assert "jobs.alpha" not in str(groups)
    # the word "credits" is never printed in a client file: the text that carries it is withheld
    # where it is printed and listed by rule; the set is delivered (it used to fail the whole
    # build, and the log called the rule tool_name)
    first = c.postings[0].model_copy(update={"description": "Earn credits here."})
    result, files = await _build(
        make_builder, c.model_copy(update={"postings": (first, *c.postings[1:])})
    )
    assert result.status == "complete", result.errors
    assert not result.withheld_files
    assert any(":cost_word=" in t for t in result.manifest.texts_withheld)
    assert "credits" not in "".join(_texts(b) for b in files.values()).lower()


async def test_a_data_tool_host_is_withheld_not_fatal(
    make_builder: Make, cfg: ExportConfig
) -> None:
    from exports.config import provider_hosts

    host = next(h for h in provider_hosts() if "parallel" in h)
    c = account()
    src = _Registry(c)
    s = src.add("jobs/x", host=host)
    postings = (
        *c.postings,
        _posting(src, "tool", "Salesforce role.").model_copy(update={"url": s.url}),
    )
    result, _ = await _build(
        make_builder, c.model_copy(update={"postings": postings, "sources": tuple(src.rows)})
    )
    # a posting on a data tool's own host is withheld and listed, never printed or linked (the
    # third-party scan used to fail the whole build on it)
    assert result.status == "complete", result.errors
    assert any(b.endswith(":tool_link") for b in result.manifest.withheld_blocks)
