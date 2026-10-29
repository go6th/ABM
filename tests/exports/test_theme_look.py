"""The workbook look the shared ``Book`` helpers produce from the theme (the owner's reference
workbook set): header, body, link cells, named fills, column overrides, heights, bold totals, and
a stylesheet that does not depend on the content. Every expected value is read from the theme;
no colour or face is typed here."""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from exports.config import Column, ExportConfig
from exports.excel.common import Book, Formula, prefix_criterion, sheet_ref
from exports.fingerprint import xlsx_fingerprint

AS_OF = date(2026, 1, 1)
COLUMNS = (
    Column(key="ref", header="Ref", kind="ref"),
    Column(key="name", header="Name", kind="text_short", width=26),
    Column(key="priority", header="Priority", kind="label", width=9, align="center", wrap=False),
    Column(key="count", header="Postings", kind="count"),
    Column(key="url", header="LinkedIn", kind="url", link_text="Open profile"),
)
HEADER_ROW = 3


def _build(cfg: ExportConfig, n_rows: int, *, with_fills: bool = True) -> bytes:
    """A small workbook: a title, a bordered Summary table, a data table with a link column, a
    bold total row, a filled Priority cell and a new-row highlight."""
    book = Book(cfg)
    ws = book.sheet("Rows")
    book.title_line(ws, 1, "Synthetic title")
    book.header(ws, COLUMNS, row=HEADER_ROW, freeze_columns=2, height_pt=32)
    first = HEADER_ROW + 1
    for i in range(n_rows):
        values = {
            "ref": f"P-{i + 1:04d}",
            "name": "Synthetic Person " * (i + 1),
            "priority": "ABC"[i % 3],
            "count": i,
            "url": f"https://profile.example/{i}",
        }
        book.body(ws, COLUMNS, values, first + i, values["ref"])
        if with_fills:
            book.fill(ws, first + i, 3, f"priority_{values['priority'].lower()}")
    book.finish_table(ws, COLUMNS, n_rows, header_row=HEADER_ROW)
    total = first + n_rows
    book.body(ws, COLUMNS, {"ref": "Total", "name": ""}, total, "Total", bold=True)
    book.formula(Formula("Rows", total, 4, f"=SUM(D{first}:D{total - 1})", 0), bold=True)
    if with_fills:
        book.fill_row(ws, total, COLUMNS, "new_row")
    book.link(ws, total + 1, 5, "Open archived copy", "https://archive.example/copy")
    summary = book.sheet("Summary")
    book.header(summary, COLUMNS[:2], border=True, freeze=False)
    book.paragraph(summary, HEADER_ROW, "A note paragraph. " * 20, 2)
    book.paragraph(summary, HEADER_ROW + 1, "A fixed-height paragraph.", 2, height_pt=46)
    return book.bytes("Synthetic", AS_OF)


@pytest.fixture(scope="module")
def sheets(cfg: ExportConfig) -> tuple[Worksheet, Worksheet]:
    wb = load_workbook(io.BytesIO(_build(cfg, 3)))
    return wb["Rows"], wb["Summary"]


def _rgb(value: str | None) -> str:
    """openpyxl reads colours as ARGB; the theme holds RGB."""
    return (value or "")[-6:]


def _border(cell: Cell) -> tuple[str | None, str]:
    side = cell.border.left
    return (side.style if side else None, _rgb(side.color.rgb) if side and side.color else "")


def test_header_row_is_the_reference_header(
    cfg: ExportConfig, sheets: tuple[Worksheet, Worksheet]
) -> None:
    ws, _ = sheets
    t = cfg.theme.excel
    cell = ws.cell(row=HEADER_ROW, column=1)
    assert (cell.font.name, cell.font.sz, cell.font.b) == (
        t.fonts.header.name,
        t.fonts.header.size_pt,
        True,
    )
    assert _rgb(cell.font.color.rgb) == t.fonts.header.color
    assert _rgb(cell.fill.fgColor.rgb) == t.header_fill
    assert cell.alignment.vertical == t.header_vertical and cell.alignment.wrap_text
    assert _border(cell) == (None, "")  # a data sheet's header has no border
    assert ws.row_dimensions[HEADER_ROW].height == 32
    assert ws.freeze_panes == f"C{HEADER_ROW + 1}"
    assert ws.auto_filter.ref == f"A{HEADER_ROW}:E{HEADER_ROW + 3}"


def test_summary_table_header_is_bordered_when_asked(
    cfg: ExportConfig, sheets: tuple[Worksheet, Worksheet]
) -> None:
    _, summary = sheets
    cell = summary.cell(row=1, column=1)
    assert _rgb(cell.fill.fgColor.rgb) == cfg.theme.excel.header_fill
    assert _border(cell) == ("thin", cfg.theme.excel.border_color)


def test_body_cells_carry_the_body_font_border_and_overrides(
    cfg: ExportConfig, sheets: tuple[Worksheet, Worksheet]
) -> None:
    ws, _ = sheets
    t = cfg.theme.excel
    name = ws.cell(row=HEADER_ROW + 1, column=2)
    assert (name.font.name, name.font.sz, _rgb(name.font.color.rgb)) == (
        t.fonts.body.name,
        t.fonts.body.size_pt,
        t.fonts.body.color,
    )
    assert _border(name) == ("thin", t.border_color)
    assert name.alignment.vertical == t.body_vertical
    assert name.alignment.wrap_text is True  # text_short wraps
    assert name.data_type == "s"
    priority = ws.cell(row=HEADER_ROW + 1, column=3)
    assert priority.alignment.horizontal == "center" and not priority.alignment.wrap_text
    count = ws.cell(row=HEADER_ROW + 1, column=4)
    assert count.number_format == t.column_kinds["count"].number_format


def test_column_widths_follow_the_column_then_its_kind(
    cfg: ExportConfig, sheets: tuple[Worksheet, Worksheet]
) -> None:
    ws, _ = sheets
    kinds = cfg.theme.excel.column_kinds
    assert ws.column_dimensions["A"].width == kinds["ref"].width  # a kind width under 13 applies
    assert ws.column_dimensions["B"].width == 26  # the column's own width
    assert ws.column_dimensions["C"].width == 9
    assert ws.column_dimensions["E"].width == kinds["url"].width


def test_link_cells_show_the_label_in_the_link_font(
    cfg: ExportConfig, sheets: tuple[Worksheet, Worksheet]
) -> None:
    ws, _ = sheets
    link = cfg.theme.excel.fonts.link
    cell = ws.cell(row=HEADER_ROW + 1, column=5)
    assert cell.value == "Open profile" and cell.data_type == "s"
    assert cell.hyperlink is not None and cell.hyperlink.target == "https://profile.example/0"
    assert (cell.font.name, cell.font.sz, _rgb(cell.font.color.rgb)) == (
        link.name,
        link.size_pt,
        link.color,
    )
    assert cell.font.u == "single"
    assert _border(cell) == ("thin", cfg.theme.excel.border_color)
    archived = ws.cell(row=HEADER_ROW + 1 + 3 + 1, column=5)
    assert archived.value == "Open archived copy"
    assert archived.hyperlink.target == "https://archive.example/copy"
    assert _rgb(archived.font.color.rgb) == link.color and archived.font.u == "single"


def test_named_fills_apply_by_name_and_keep_the_cell_look(
    cfg: ExportConfig, sheets: tuple[Worksheet, Worksheet]
) -> None:
    ws, _ = sheets
    fills = cfg.theme.excel.fills
    for i, letter in enumerate("abc"):
        cell = ws.cell(row=HEADER_ROW + 1 + i, column=3)
        assert _rgb(cell.fill.fgColor.rgb) == fills[f"priority_{letter}"]
        assert cell.alignment.horizontal == "center"  # the column's alignment survives the fill
    total = HEADER_ROW + 1 + 3
    for j in range(1, 6):
        cell = ws.cell(row=total, column=j)
        assert _rgb(cell.fill.fgColor.rgb) == fills["new_row"], j
    assert ws.cell(row=total, column=1).font.b is True  # bold survives the fill
    assert ws.cell(row=total, column=4).font.b is True and ws.cell(
        row=total, column=4
    ).value.startswith("=SUM(")


def test_an_unknown_fill_is_refused(cfg: ExportConfig) -> None:
    book = Book(cfg)
    ws = book.sheet("Rows")
    with pytest.raises(ValueError, match="not a theme fill"):
        book.fill(ws, 1, 1, "no_such_fill")


def test_title_and_paragraph_rows(cfg: ExportConfig, sheets: tuple[Worksheet, Worksheet]) -> None:
    ws, summary = sheets
    t = cfg.theme.excel
    title = ws.cell(row=1, column=1)
    assert (title.font.name, title.font.sz, title.font.b) == (
        t.fonts.title.name,
        t.fonts.title.size_pt,
        True,
    )
    assert _rgb(title.font.color.rgb) == t.fonts.title.color
    assert ws.row_dimensions[1].height == t.title_height_pt
    estimated = summary.cell(row=HEADER_ROW, column=1)
    assert estimated.alignment.wrap_text and estimated.font.name == t.fonts.body.name
    assert summary.row_dimensions[HEADER_ROW].height > t.fonts.body.size_pt  # several lines
    assert summary.row_dimensions[HEADER_ROW + 1].height == 46
    assert "A1:B3" in [str(r) for r in summary.merged_cells.ranges] or any(
        str(r).startswith("A3") for r in summary.merged_cells.ranges
    )


def test_the_stylesheet_does_not_depend_on_the_content(cfg: ExportConfig) -> None:
    small = xlsx_fingerprint(_build(cfg, 1, with_fills=False))
    large = xlsx_fingerprint(_build(cfg, 9))
    assert small == large
    assert _build(cfg, 1) != _build(cfg, 9)


def test_link_text_belongs_to_a_url_column() -> None:
    with pytest.raises(ValueError, match="link_text"):
        Column(key="name", header="Name", kind="label", link_text="Open")


def test_formula_helpers() -> None:
    assert sheet_ref("Jobs") == "Jobs"
    assert sheet_ref("Platform Ranking") == "'Platform Ranking'"
    assert sheet_ref("A1") == "'A1'"
    assert sheet_ref("O'Brien") == "'O''Brien'"
    assert prefix_criterion("Synthetic Corp") == '"Synthetic Corp*"'
    assert prefix_criterion('A*B?C~ "x"') == '"A~*B~?C~~ ""x""*"'
