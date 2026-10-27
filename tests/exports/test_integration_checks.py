"""Integration checks (spec 9.7): Word to PDF with the row-presence check, workbook recalculation
with zero errors and totals equal to the counts, the skill's ``scan_client_output.py`` over the
unzipped package, and the Internal Only file kept out. When LibreOffice or poppler is missing the
checks report ``not_run`` with a reason; they are never reported as passed."""

from __future__ import annotations

import io
import shutil
import subprocess
import sys
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import Workbook

from account_intelligence.grading.config import GradingConfig
from exports.builder import ExportBuilder
from exports.config import ExportConfig
from exports.verify import (
    LibreOffice,
    docx_row_keys,
    keys_once_in_order,
    pdf_row_presence,
    recalc_check,
)
from tests.exports.corpus import account, broker, rich
from tests.fakes.exports import FakeReader, InProcessRenderer, MemoryStore

ROOT = Path(__file__).resolve().parents[2]
SCANNER = ROOT / ".claude/skills/account-intelligence-platform/scripts/scan_client_output.py"
HAVE_OFFICE = shutil.which("soffice") is not None and shutil.which("pdftotext") is not None
office = pytest.mark.skipif(not HAVE_OFFICE, reason="not_run: LibreOffice or poppler missing")


def test_keys_once_in_order() -> None:
    text = "L-0001 a L-0002 b L-0010 c"
    assert keys_once_in_order(text, ["L-0001", "L-0002", "L-0010"]) == ()
    assert keys_once_in_order(text, ["L-0002", "L-0001"]) == ("L-0001: out of order",)
    assert keys_once_in_order(text, ["L-0003"]) == ("L-0003: found 0 times",)
    assert keys_once_in_order("L-0001 L-0001", ["L-0001"]) == ("L-0001: found 2 times",)
    assert keys_once_in_order("XL-0001 L-00012", ["L-0001"]) == ("L-0001: found 0 times",)


def test_missing_converters_are_not_run_never_passed() -> None:
    absent = LibreOffice(binary=None)
    absent.binary = None
    assert pdf_row_presence(b"", ["L-0001"], absent).status == "not_run"
    result = recalc_check(b"", {}, absent)
    assert (result.status, result.reason) == ("not_run", "no spreadsheet engine installed")


@pytest.mark.asyncio
@office
async def test_builder_with_render_checks_passes_on_a_rich_account(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = rich()
    lo = LibreOffice()
    store = MemoryStore()
    builder = ExportBuilder(
        FakeReader(content), store, InProcessRenderer(), cfg, grading, pdf=lo, recalc=lo
    )
    out = builder.assemble(content)
    assert not out.errors, out.errors
    statuses = {(c.name, c.status) for c in out.checks}
    assert statuses == {
        ("pdf_row_presence", "passed"),
        ("workbook_recalc", "passed"),
        ("docx_rowcheck", "passed"),
        ("doc_checks", "passed"),
        ("third_party_scan", "passed"),
    }
    # two documents (4 checks each), three workbooks (recalculation and the third-party scan),
    # the workspace zip (the third-party scan)
    assert len(out.checks) == 15
    # the lost-row check really looks at the PDF: every summary and evidence key is there
    for kind in ("summary", "evidence"):
        _, body = out.files[kind]
        keys = docx_row_keys(body)
        assert keys and pdf_row_presence(body, keys, lo).status == "passed"


@office
def test_recalculation_finds_formula_errors_and_wrong_totals() -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Counts"
    ws["A1"] = 5
    ws["B1"] = "=1/0"
    ws["B2"] = "=A1*2"
    buf = io.BytesIO()
    wb.save(buf)
    result = recalc_check(buf.getvalue(), {("Counts", "B2"): 11}, LibreOffice())
    assert result.status == "failed"
    assert any("#DIV/0!" in d for d in result.details)
    assert any("B2: 10 != 11" in d for d in result.details)


@office
def test_pdf_check_finds_a_lost_row(cfg: ExportConfig, grading: GradingConfig) -> None:
    content = account()
    builder = ExportBuilder(FakeReader(content), MemoryStore(), InProcessRenderer(), cfg, grading)
    out = builder.assemble(content)
    _, body = out.files["summary"]
    keys = list(docx_row_keys(body)) + ["L-9999"]
    result = pdf_row_presence(body, keys, LibreOffice())
    assert result.status == "failed" and result.details == ("L-9999: found 0 times",)


@pytest.mark.asyncio
async def test_skill_scanner_finds_nothing_in_the_unzipped_package(
    make_builder: Callable[..., ExportBuilder], tmp_path: Path
) -> None:
    content = rich()
    store = MemoryStore()
    result = await make_builder(content, store).build_all(content.tenant_id, content.run_id)
    assert result.status == "complete", result.errors
    zipped = next(d for d in result.deliverables if d.kind == "workspace")
    target = tmp_path / "package"
    with zipfile.ZipFile(io.BytesIO(store.objects[zipped.object_key])) as z:
        z.extractall(target)
    (tmp_path / "terms.txt").write_text("\n".join(content.forbidden_terms) + "\n")
    done = subprocess.run(
        [sys.executable, str(SCANNER), "--extra", str(tmp_path / "terms.txt"), str(target)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert done.returncode == 0, done.stdout[-2000:]
    assert "0 hit(s)" in done.stdout
    # the scanner is not a no-op: a planted broker name is found
    (target / "planted.md").write_text(f"pulled from {broker('p_signals')}")
    done = subprocess.run(
        [sys.executable, str(SCANNER), str(target)], capture_output=True, text=True, check=False
    )
    assert done.returncode == 1


@pytest.mark.asyncio
async def test_the_internal_only_file_is_never_built(
    make_builder: Callable[..., ExportBuilder], cfg: ExportConfig
) -> None:
    content = account()
    store = MemoryStore()
    result = await make_builder(content, store).build_all(content.tenant_id, content.run_id)
    assert "operator_only" in cfg.filenames  # it exists in the output spec ...
    assert {d.kind for d in result.deliverables} == {
        "summary",
        "evidence",
        "people",
        "jobs",
        "tech",
        "workspace",
    }  # ... never built here
    zipped = next(d for d in result.deliverables if d.kind == "workspace")
    with zipfile.ZipFile(io.BytesIO(store.objects[zipped.object_key])) as z:
        assert not any("only" in n.lower() for n in z.namelist())
