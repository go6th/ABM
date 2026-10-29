"""The hand-written oracle (spec 9.1): a small synthetic account whose every expected output was
worked out by hand (``oracle_expected.json``) and is compared with the models, the manifest and
the produced files."""

from __future__ import annotations

import io
import json
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
import pytest_asyncio
from openpyxl import load_workbook

from exports.builder import ExportBuilder
from exports.content import SnapshotContent
from exports.models import BuildResult
from exports.verify import docx_text
from tests.exports.corpus import tracker
from tests.fakes.exports import MemoryStore

pytestmark = pytest.mark.asyncio
FIX = Path(__file__).resolve().parents[1] / "fixtures/synthetic/exports"


def oracle_content() -> SnapshotContent:
    """The oracle account with the tracker's client view added."""
    data = json.loads((FIX / "oracle_account.json").read_text())
    data.pop("_label")
    lines, tally = tracker()
    return SnapshotContent.model_validate({**data, "tracker": lines, "tracker_tally": tally})


@pytest.fixture(scope="module")
def expected() -> dict[str, Any]:
    return json.loads((FIX / "oracle_expected.json").read_text())


@pytest_asyncio.fixture
async def built(
    make_builder: Callable[..., ExportBuilder],
) -> tuple[ExportBuilder, BuildResult, MemoryStore]:
    store = MemoryStore()
    content = oracle_content()
    builder = make_builder(content, store)
    result = await builder.build_all(content.tenant_id, content.run_id)
    assert result.status == "complete", result.errors
    return builder, result, store


def _file(result: BuildResult, store: MemoryStore, kind: str) -> bytes:
    d = next(x for x in result.deliverables if x.kind == kind)
    return store.objects[d.object_key]


async def test_file_names_and_row_counts(built: Any, expected: dict[str, Any]) -> None:
    _, result, _ = built
    assert {d.kind: d.filename for d in result.deliverables} == expected["filenames"]
    assert result.manifest.rows_per_file == expected["rows_per_file"]
    assert result.manifest.withheld_sentences == expected["withheld_sentences"]


async def test_summary_model(built: Any, expected: dict[str, Any]) -> None:
    builder, _, _ = built
    out = builder.assemble(oracle_content())
    summary = out.mapped["summary"]
    model = summary.model
    lanes = [
        {
            "lane": t.lane,
            "rows": [
                {
                    "ref": r.ref,
                    "row_id": r.row_id,
                    "grade_label": r.grade_label,
                    "use": r.use,
                    "codes": list(r.codes),
                    "question": r.question,
                }
                for r in t.rows
            ],
        }
        for t in model.lanes
    ]
    assert lanes == expected["lanes"]
    assert [{"row_id": a.row_id, "question": a.question} for a in model.absences] == expected[
        "absences"
    ]
    assert [
        {"ref": p.ref, "row_id": p.row_id, "source_code": p.source_code} for p in model.key_people
    ] == expected["key_people"]
    assert [
        {"ref": f.ref, "person": f.person, "source_code": f.source_code, "withheld": f.withheld}
        for f in model.linkedin_findings
    ] == expected["linkedin_findings"]
    assert list(summary.refs[:6]) == expected["summary_ref_prefix"]
    assert len(model.tracker) == expected["tracker_rows"]
    # the excluded person's finding never appears, not even by its link
    dumped = json.dumps(model.model_dump(mode="json"))
    assert "pk-cal" not in dumped and "/in/cal/" not in dumped and "Never shown" not in dumped


async def test_workbook_models(built: Any, expected: dict[str, Any]) -> None:
    builder, _, _ = built
    out = builder.assemble(oracle_content())
    people = out.mapped["people"].model
    assert [
        {"ref": r.ref, "row_id": r.row_id, "source_code": r.source_code} for r in people.rows
    ] == expected["people"]
    jobs = out.mapped["jobs"].model
    assert [
        {
            "ref": r.ref,
            "row_id": r.row_id,
            "status": r.status,
            "posting_class": r.posting_class,
            "source_code": r.source_code,
            "description": r.description,
        }
        for r in jobs.rows
    ] == expected["jobs"]
    tech = out.mapped["tech"].model
    got = [
        {
            "ref": r.ref,
            "row_id": r.row_id,
            "postings": r.postings,
            "source_code": r.source_code,
        }
        for r in tech.rows
    ]
    want = [{k: v for k, v in t.items() if k != "first_seen"} for t in expected["tech"]]
    assert got == want
    assert tech.rows[1].first_seen == expected["tech"][1]["first_seen"]
    # the Confirmed stack rows, in the lookup's order, with the lane the platform word gives
    assert [
        {
            "row_id": r.row_id,
            "business_function": r.business_function,
            "category": r.category,
            "technology": r.technology,
            "postings": r.postings,
            "lane": r.lane,
        }
        for r in tech.stack
    ] == expected["tech_stack"]


async def test_evidence(built: Any, expected: dict[str, Any]) -> None:
    builder, result, store = built
    out = builder.assemble(oracle_content())
    ev = out.mapped["evidence"]
    model = ev.model
    assert [s.code for s in model.sources] == expected["evidence_codes"]
    assert list(ev.refs) == expected["evidence_refs"]
    assert {
        s.code: [[u.grade_label, u.sentence] for u in s.uses] for s in model.sources
    } == expected["evidence_uses"]
    kinds = {s.code: s.kind_label for s in model.sources}
    assert {k: kinds[k] for k in expected["evidence_kinds"]} == expected["evidence_kinds"]
    assert [j.row_id for j in model.job_descriptions] == expected["job_descriptions_read"]
    assert list(result.manifest.source_codes_cited) == expected["evidence_codes"]
    text = docx_text(_file(result, store, "evidence"))
    for code in expected["evidence_codes"]:
        assert code in text
    assert "S-0009" not in text  # registered but never cited: not listed


async def test_numbers(built: Any, expected: dict[str, Any]) -> None:
    _, result, _ = built
    registered = {n.key: n.value for n in result.manifest.numbers}
    for key, value in expected["numbers"].items():
        assert registered[key] == value, key
    for n in result.manifest.numbers:
        assert n.base.strip()  # every number names its base


async def test_produced_workbook_cells(built: Any, expected: dict[str, Any]) -> None:
    _, result, store = built
    wb = load_workbook(io.BytesIO(_file(result, store, "jobs")))
    ws = wb["Jobs"]
    headers = [c.value for c in ws[1]]
    assert headers[0] == "Requisition"
    col = headers.index("Full description")
    ref_col = headers.index("Ref")
    # rows are sorted by First seen, newest first: compare every row by its Ref
    cells = {
        ws.cell(row=r, column=ref_col + 1).value: ws.cell(row=r, column=col + 1).value or ""
        for r in range(2, ws.max_row + 1)
    }
    assert cells == {j["ref"]: j["description"] for j in expected["jobs"]}
    assert [ws.cell(row=r, column=ref_col + 1).value for r in range(2, ws.max_row + 1)] == [
        j["ref"] for j in expected["jobs"]
    ]
    counts = wb["Counts"]
    assert counts["B2"].value.startswith("=COUNTA(")
    people = load_workbook(io.BytesIO(_file(result, store, "people")))
    cells = [[c.value for c in row] for row in people["Summary"].iter_rows()]
    # the reference layout: a title line, then paragraphs; a person excluded by title is named
    # with the reason on the Summary sheet and never listed on the People sheet
    assert str(cells[0][0]).endswith("people list, deduplicated")
    summary_text = " ".join(str(v) for row in cells for v in row if v is not None)
    # the reference layout: six notes (rows 3 to 8), the Level table from row 11, the Lane table
    # from row 24 and nothing after its Total row (the count sentences are not on this sheet)
    assert all(cells[i][0] for i in range(2, 8)) and cells[8][0] is None
    assert [cells[10][0], cells[10][1], cells[10][2]] == ["Level", "People", "Priority A"]
    assert cells[20][0] == "Total" and cells[21][0] == "Meets seniority rule"
    assert cells[23][0] == "Lane, senior people only" and cells[32][0] == "Total"
    assert len(cells) == 33
    assert "Entity labels: " in summary_text
    names = " ".join(str(c.value) for row in people["People"].iter_rows() for c in row)
    assert "Cal Excluded" not in names
