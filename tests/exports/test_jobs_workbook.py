"""The Job Postings workbook against the owner's reference (docs/parity/reference_spec.md and
reference_manifest.json, measured 2026-10-29): sheets, headers, widths, freeze, link cells, the
archived link, the new-row fill, sort order, platform order, the sentence cell, the status words,
the entity rules and the matrix formulas. Synthetic data only; the reference account's name is
never written here (the manifest's ``{prefix}`` stands for the account's entity prefix)."""

from __future__ import annotations

import io
import json
import shutil
from collections.abc import Callable
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from account_intelligence.grading.config import GradingConfig
from exports.builder import ExportBuilder
from exports.config import ExportConfig
from exports.content import PostingRow, SourceRow
from exports.excel.jobs import expected_jobs
from exports.models import BuildResult
from exports.readers.common import MapContext
from exports.readers.jobs import (
    entity_of,
    entity_prefix,
    entity_rules,
    map_jobs,
    ordered_platforms,
    sentence_cell,
)
from exports.verify import LibreOffice, recalc_check
from tests.exports.corpus import AS_OF, account
from tests.fakes.exports import MemoryStore

ROOT = Path(__file__).resolve().parents[2]
MANIFEST = ROOT / "docs/parity/reference_manifest.json"
Make = Callable[..., ExportBuilder]
office = pytest.mark.skipif(shutil.which("soffice") is None, reason="not_run: LibreOffice missing")


async def _build(make_builder: Make, content: Any) -> tuple[BuildResult, dict[str, bytes]]:
    store = MemoryStore()
    result = await make_builder(content, store).build_all(content.tenant_id, content.run_id)
    return result, {d.kind: store.objects[d.object_key] for d in result.deliverables}


class _Sources:
    """Adds sources to an account's registry, continuing its codes."""

    def __init__(self, content: Any) -> None:
        self.rows = list(content.sources)
        self.host = content.domain

    def add(self, path: str, host: str | None = None) -> SourceRow:
        row = SourceRow(
            code=f"S-{len(self.rows) + 1:04d}",
            url=f"https://{host or self.host}/{path}",
            published=AS_OF,
            kind="primary",
        )
        self.rows.append(row)
        return row


def _post(src: _Sources, key: str, text: str, **over: Any) -> PostingRow:
    values: dict[str, Any] = {
        "posting_key": key,
        "title": f"Synthetic role {key}",
        "url": src.add(f"jobs/{key}").url,
        "status": "open",
        "posting_class": "user",
        "platforms": ("Salesforce",),
        "platform_sentences": (text,),
        "platform_sentence_pairs": (("Salesforce", text),),
        "description": text,
        "posted_at": date(2026, 6, 1),
        "first_seen": date(2026, 6, 2),
        "full_text_read": True,
    }
    values.update(over)
    return PostingRow(**values)


def _content(postings_of: Callable[[_Sources], tuple[PostingRow, ...]], **update: Any) -> Any:
    c = account(n_postings=0)
    src = _Sources(c)
    postings = postings_of(src)
    return c.model_copy(update={"postings": postings, "sources": tuple(src.rows), **update})


def _model(content: Any, cfg: ExportConfig, grading: GradingConfig) -> Any:
    ctx = MapContext(content, cfg, grading)
    model = map_jobs(ctx).model
    assert model is not None and not ctx.violations, ctx.violations
    return model, ctx


# ---- the reference manifest --------------------------------------------------------------------


@pytest.mark.asyncio
async def test_the_jobs_workbook_has_the_reference_sheets_headers_and_widths(
    make_builder: Make, cfg: ExportConfig
) -> None:
    manifest = json.loads(MANIFEST.read_text())["Job_Postings.xlsx"]
    content = account()
    result, files = await _build(make_builder, content)
    assert result.status == "complete", result.errors
    wb = load_workbook(io.BytesIO(files["jobs"]))
    # the reference sheets first, in its order; the product's Counts and Long text after them
    assert wb.sheetnames[: len(manifest)] == [s["name"] for s in manifest]
    assert wb.sheetnames[len(manifest) :] == ["Counts", "Long text"]
    summary_ref, jobs_ref = manifest
    # Jobs: the reference's 14 headers in order, then the owner's two columns, then the product's
    headers = [c.value for c in wb["Jobs"][1]]
    assert headers[: len(jobs_ref["headers"])] == jobs_ref["headers"]
    assert headers[len(jobs_ref["headers"]) :] == [
        "Archived link",
        "New since last research",
        "Ref",
        "Source code",
    ]
    jobs = wb["Jobs"]
    for letter, width in jobs_ref["widths"].items():
        assert jobs.column_dimensions[letter].width == pytest.approx(width), letter
    assert jobs.column_dimensions["G"].width == pytest.approx(jobs_ref["widths"]["F"])
    # archived link 16, new flag 14 (the owner's two columns)
    assert jobs.column_dimensions["O"].width == pytest.approx(16)
    assert jobs.column_dimensions["P"].width == pytest.approx(14)
    # the reference freezes below the header and right of Title (C2; the manifest's C9 is the
    # pane the file was last scrolled to, not a freeze)
    assert jobs.freeze_panes == "C2"
    assert jobs.auto_filter.ref == f"A1:{get_column_letter(len(headers))}{jobs.max_row}"
    # Summary: the matrix header with the account's entity prefix, merged notes A:H, widths
    prefix = content.account_name
    wanted = [h.replace("{prefix}", prefix) for h in summary_ref["headers"]]
    summary = wb["Summary"]
    assert [c.value for c in summary[summary_ref["header_row"]]][: len(wanted)] == wanted
    assert {str(r) for r in summary.merged_cells.ranges} == set(summary_ref["merged"])
    for letter, width in summary_ref["widths"].items():
        assert summary.column_dimensions[letter].width == pytest.approx(width), letter
    assert [c.value for c in summary["A"]][10:17] == [  # rows 11 to 17: the seven platforms
        "Salesforce",
        "MuleSoft",
        "ServiceNow",
        "SAP",
        "Tableau",
        "Veeva",
        "Workday",
    ]


# ---- link cells, the archived copy and the new-row fill ----------------------------------------


@pytest.mark.asyncio
async def test_link_cells_carry_the_label_and_the_address(
    make_builder: Make, cfg: ExportConfig
) -> None:
    content = account(n_postings=6)
    # every other posting has a stored archived copy
    content = content.model_copy(
        update={
            "postings": tuple(
                p.model_copy(
                    update={
                        "archive_url": (
                            f"https://app.example/accounts/{content.domain}/archive"
                            f"?url=https%3A%2F%2Fjobs.example%2F{i:04d}"
                        )
                        if i % 2 == 0
                        else ""
                    }
                )
                for i, p in enumerate(content.postings)
            )
        }
    )
    _, files = await _build(make_builder, content)
    ws = load_workbook(io.BytesIO(files["jobs"]))["Jobs"]
    headers = [c.value for c in ws[1]]
    link, archive = headers.index("Link") + 1, headers.index("Archived link") + 1
    ref_col = headers.index("Ref") + 1
    by_key = {p.posting_key: p for p in content.postings}
    code_to_post = {s.code: p for p in content.postings for s in content.sources if s.url == p.url}
    code_col = headers.index("Source code") + 1
    seen = 0
    for r in range(2, ws.max_row + 1):
        post = code_to_post[ws.cell(row=r, column=code_col).value]
        cell = ws.cell(row=r, column=link)
        assert cell.value == "Open posting" and cell.hyperlink.target == post.url
        assert cell.font.u == "single" and cell.font.color.rgb.endswith("1155CC")
        arch = ws.cell(row=r, column=archive)
        if post.archive_url:
            assert arch.value == "Open archived copy" and arch.hyperlink.target == post.archive_url
        else:
            assert not arch.value and arch.hyperlink is None
        seen += 1
        assert by_key[post.posting_key] is post and ws.cell(row=r, column=ref_col).value
    assert seen == 6


@pytest.mark.asyncio
async def test_rows_new_since_the_last_research_are_flagged_and_filled(
    make_builder: Make, cfg: ExportConfig
) -> None:
    content = account(n_postings=6)
    _, files = await _build(make_builder, content)
    ws = load_workbook(io.BytesIO(files["jobs"]))["Jobs"]
    headers = [c.value for c in ws[1]]
    new_col, code_col = (
        headers.index("New since last research") + 1,
        headers.index("Source code") + 1,
    )
    code_to_post = {s.code: p for p in content.postings for s in content.sources if s.url == p.url}
    fill = cfg.theme.excel.fills["new_row"]
    flagged = 0
    for r in range(2, ws.max_row + 1):
        post = code_to_post[ws.cell(row=r, column=code_col).value]
        cells = [ws.cell(row=r, column=c) for c in range(1, len(headers) + 1)]
        colours = {c.fill.fgColor.rgb[-6:] if c.fill.fill_type else None for c in cells}
        if post.new_since_previous:
            flagged += 1
            assert ws.cell(row=r, column=new_col).value == "Yes"
            assert colours == {fill}, post.posting_key
        else:
            assert not ws.cell(row=r, column=new_col).value
            assert fill not in colours
    assert flagged == sum(1 for p in content.postings if p.new_since_previous)


def test_the_new_column_is_blank_when_there_was_no_previous_research(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = _content(
        lambda s: (_post(s, "a", "Salesforce."), _post(s, "b", "Salesforce.", archive_url=""))
    )
    model, _ = _model(content, cfg, grading)
    assert all(r.new_flag == "" and r.new_since_previous is None for r in model.rows)


# ---- order, platforms, sentences, status, read text --------------------------------------------


def test_rows_are_sorted_by_first_seen_newest_first_then_key(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = _content(
        lambda s: (
            _post(s, "k2", "Salesforce.", first_seen=date(2026, 3, 1)),
            _post(s, "k1", "Salesforce.", first_seen=date(2026, 3, 1)),
            _post(s, "new", "Salesforce.", first_seen=date(2026, 8, 9)),
            _post(s, "old", "Salesforce.", first_seen=date(2025, 12, 1)),
            _post(s, "none", "Salesforce.", first_seen=None, posted_at=None),
            # a posting only the live board holds: dated by the capture date (as_of)
            _post(s, "live", "Salesforce.", first_seen=None, posted_at=None, origin="live_board"),
        )
    )
    model, _ = _model(content, cfg, grading)
    assert [r.row_id for r in model.rows] == ["live", "new", "k1", "k2", "old", "none"]
    assert [r.ref for r in model.rows] == [f"J-{i:04d}" for i in range(1, 7)]
    by = {r.row_id: r for r in model.rows}
    assert by["live"].first_seen_text == AS_OF.isoformat()
    assert by["new"].first_seen_text == "2026-08-09" and by["none"].first_seen_text == ""


def test_platforms_named_follow_the_fixed_order_then_alphabetical() -> None:
    order = ("Salesforce", "MuleSoft", "ServiceNow", "SAP", "Tableau", "Veeva", "Workday")
    assert ordered_platforms(
        ["Workday", "Tableau", "Salesforce", "Zeta", "Alpha", "SAP"], order
    ) == (
        "Salesforce",
        "SAP",
        "Tableau",
        "Workday",
        "Alpha",
        "Zeta",
    )
    assert ordered_platforms(["Veeva", "Veeva", "mulesoft"], order) == ("mulesoft", "Veeva")


def _sentence_post(pairs: tuple[tuple[str, str], ...], flat: tuple[str, ...] = ()) -> PostingRow:
    return PostingRow(
        posting_key="k",
        title="t",
        url="https://x.example/1",
        status="open",
        posting_class="user",
        platforms=tuple(dict.fromkeys(p for p, _ in pairs)),
        platform_sentences=flat,
        platform_sentence_pairs=pairs,
    )


ORDER = ("Salesforce", "MuleSoft", "ServiceNow", "SAP", "Tableau", "Veeva", "Workday")


def test_the_sentence_cell_groups_by_platform_in_the_fixed_order() -> None:
    post = _sentence_post(
        (
            ("SAP", "Uses SAP daily."),
            ("Salesforce", "* Experience  with\n Salesforce."),
            ("Salesforce", "Salesforce is a plus."),
            ("Workday", "Workday admin."),
        )
    )
    cell = sentence_cell(post, ORDER, per_platform=4, max_chars=3000)
    assert cell == (
        "Salesforce: Experience with Salesforce. | Salesforce is a plus."
        " || SAP: Uses SAP daily. || Workday: Workday admin."
    )
    assert "\n" not in cell


def test_the_sentence_cell_caps_sentences_and_characters_and_never_breaks_a_line() -> None:
    many = tuple(("Salesforce", f"Sentence number {i} names Salesforce.") for i in range(9))
    cell = sentence_cell(_sentence_post(many), ORDER, per_platform=4, max_chars=3000)
    assert cell.count(" | ") == 3 and "number 3" in cell and "number 4" not in cell
    # the same sentence twice is one sentence; bullets and bullet glyphs are removed
    twice = _sentence_post((("SAP", "* SAP is used."), ("SAP", "SAP is used."), ("SAP", "• SAP.")))
    assert sentence_cell(twice, ORDER, per_platform=4, max_chars=3000) == "SAP: SAP is used. | SAP."
    huge = _sentence_post((("SAP", "word " * 2000),))
    out = sentence_cell(huge, ORDER, per_platform=4, max_chars=3000)
    assert len(out) <= 3000 and out.startswith("SAP: word word") and "\n" not in out


def test_the_sentence_cell_from_a_flat_list_names_the_platform_of_each_sentence() -> None:
    post = _sentence_post((), flat=("Uses SAP daily.", "Salesforce is a plus.", "Hello."))
    post = post.model_copy(update={"platforms": ("Salesforce", "SAP")})
    assert sentence_cell(post, ORDER, per_platform=4, max_chars=3000) == (
        "Salesforce: Salesforce is a plus. | Hello. || SAP: Uses SAP daily."
    )
    empty = post.model_copy(update={"platform_sentences": ()})
    assert sentence_cell(empty, ORDER, per_platform=4, max_chars=3000) == ""


def test_status_words_and_the_read_sentence_only_for_confirmed_classes(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = _content(
        lambda s: (
            _post(s, "b", "Builds on Salesforce.", posting_class="build", class_sentence="B."),
            _post(
                s, "a", "Administers Salesforce.", posting_class="admin_owner", class_sentence="A."
            ),
            _post(
                s, "w", "Works in Salesforce.", posting_class="working_tool", class_sentence="W."
            ),
            _post(
                s,
                "u",
                "Uses Salesforce.",
                posting_class="user",
                class_sentence="U.",
                status="closed",
            ),
            _post(
                s,
                "x",
                "Unread Salesforce.",
                posting_class="build",
                class_sentence="X.",
                full_text_read=False,
            ),
        )
    )
    model, _ = _model(content, cfg, grading)
    by = {r.row_id: r for r in model.rows}
    assert {k: by[k].read_text for k in by} == {"b": "B.", "a": "A.", "w": "", "u": "", "x": ""}
    assert {r.status for r in model.rows} == {"Open", "Closed"}
    assert by["u"].status == "Closed"
    assert by["b"].seniority == "" and by["b"].last_seen_text == ""


def test_last_seen_and_seniority_print_as_given(cfg: ExportConfig, grading: GradingConfig) -> None:
    content = _content(
        lambda s: (
            _post(s, "a", "Salesforce.", last_seen=date(2026, 9, 1), seniority="vice_president"),
        )
    )
    row = _model(content, cfg, grading)[0].rows[0]
    assert (row.last_seen_text, row.seniority) == ("2026-09-01", "vice_president")


# ---- entity rules ------------------------------------------------------------------------------


def _lock(*entities: dict[str, Any], **more: Any) -> dict[str, Any]:
    return {"legal_name": "Synthetic Alpha Ltd", "entities": list(entities), **more}


DIVESTED = {
    "label": "Synthetic Spin (divested business)",
    "patterns": [r"Spin Diagnostics", r"\bSD-\d+\b"],
    "counts_as_account": False,
}


def test_entity_rules_apply_first_then_the_host(cfg: ExportConfig, grading: GradingConfig) -> None:
    content = _content(
        lambda s: (
            _post(s, "own", "Salesforce for the sales team."),
            _post(s, "spin", "Spin Diagnostics needs Salesforce."),
            _post(s, "code", "Reference SD-1234 uses Salesforce."),
            _post(s, "case", "spin diagnostics, lower case, uses Salesforce."),
        ),
    )
    content = content.model_copy(update={"entity_lock": _lock(DIVESTED)})
    model, ctx = _model(content, cfg, grading)
    by = {r.row_id: r for r in model.rows}
    assert by["own"].entity == content.account_name  # the host rule (the account's domain)
    assert by["spin"].entity == "Synthetic Spin (divested business)"
    assert by["code"].entity == "Synthetic Spin (divested business)"
    assert by["case"].entity == content.account_name  # the patterns are case sensitive
    assert not [w for w in ctx.withheld_blocks if w[0] == "jobs.entities"]
    # the matrix counts the account's own positions only, as the workbook's COUNTIFS do
    sf = model.matrix[0]
    assert sf.platform == "Salesforce" and sf.rows == 2
    assert model.prefix == content.account_name


def test_a_rule_can_end_before_a_date(cfg: ExportConfig, grading: GradingConfig) -> None:
    rule = {**DIVESTED, "before": "2026-02-01", "counts_as_account": False}
    content = _content(
        lambda s: (
            _post(s, "early", "Spin Diagnostics. Salesforce.", first_seen=date(2026, 1, 15)),
            _post(s, "late", "Spin Diagnostics. Salesforce.", first_seen=date(2026, 3, 1)),
            _post(s, "undated", "Spin Diagnostics. Salesforce.", first_seen=None, posted_at=None),
        )
    ).model_copy(update={"entity_lock": _lock(rule)})
    by = {r.row_id: r for r in _model(content, cfg, grading)[0].rows}
    assert by["early"].entity == "Synthetic Spin (divested business)"
    assert by["late"].entity == content.account_name
    assert (
        by["undated"].entity == content.account_name
    )  # no date: the rule cannot be shown to apply


def test_a_subsidiary_label_with_the_prefix_counts_with_the_account(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    sub = {
        "label": "Synthetic Alpha Ltd (Beta Supply)",
        "patterns": ["Beta Supply"],
        "counts_as_account": True,
    }
    content = _content(
        lambda s: (_post(s, "own", "Salesforce."), _post(s, "sub", "Beta Supply uses Salesforce."))
    ).model_copy(update={"entity_lock": _lock(sub, DIVESTED)})
    model, _ = _model(content, cfg, grading)
    assert {r.entity for r in model.rows} == {
        "Synthetic Alpha Ltd",
        "Synthetic Alpha Ltd (Beta Supply)",
    }
    assert model.matrix[0].rows == 2


def test_an_entity_prefix_in_the_lock_is_the_short_name(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = _content(lambda s: (_post(s, "own", "Salesforce."),))
    named = content.model_copy(update={"entity_lock": {"entity_prefix": "Synthetic"}})
    assert entity_prefix(named) == "Synthetic" and entity_prefix(content) == content.account_name
    # a prefix that is not a start of the account name is ignored
    wrong = content.model_copy(update={"entity_lock": {"entity_prefix": "Other"}})
    assert entity_prefix(wrong) == content.account_name
    model, _ = _model(named, cfg, grading)
    assert model.prefix == "Synthetic" and model.matrix[0].rows == 1
    assert "Synthetic positions only" in model.matrix_heading


def test_bad_entity_rules_are_left_out_and_listed_never_guessed(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    bad = [
        {"label": "Broken", "patterns": ["(unclosed"], "counts_as_account": False},
        {"label": "", "patterns": ["x"]},
        {"label": "No patterns", "patterns": []},
        {"label": "Bad date", "patterns": ["x"], "before": "soon"},
        {"label": "Synthetic Alpha Ltd (Wrong)", "patterns": ["x"], "counts_as_account": False},
        {"label": "Elsewhere", "patterns": ["x"], "counts_as_account": True},
        "not a mapping",
        DIVESTED,
    ]
    content = _content(lambda s: (_post(s, "spin", "Spin Diagnostics. Salesforce."),)).model_copy(
        update={"entity_lock": _lock(*bad)}
    )
    rules, problems = entity_rules(content, content.account_name)
    assert [r.label for r in rules] == [DIVESTED["label"]]
    assert [why for _, why in problems] == [
        "pattern_does_not_compile",
        "label_missing",
        "patterns_missing",
        "before_not_an_iso_date",
        "counts_as_account_disagrees_with_label",
        "counts_as_account_disagrees_with_label",
        "not_a_mapping",
    ]
    model, ctx = _model(content, cfg, grading)
    assert model.rows[0].entity == DIVESTED["label"]
    listed = [w for w in ctx.withheld_blocks if w[0] == "jobs.entities"]
    assert len(listed) == 7 and listed[0][2] == "entity_rule_pattern_does_not_compile"
    # no entities at all (the usual case) is no rule and no problem
    plain = content.model_copy(update={"entity_lock": None})
    assert entity_rules(plain, plain.account_name) == ((), ())


def test_entity_of_keeps_the_host_logic_without_rules() -> None:
    c = account(n_postings=0).model_copy(update={"careers_host": "jobs.alpha-careers.example"})
    suffix = " (entity to confirm)"
    assert entity_of(f"https://careers.{c.domain}/x", c, suffix) == c.account_name
    assert entity_of("https://jobs.alpha-careers.example/x", c, suffix) == c.account_name
    assert entity_of("https://boards.example/x", c, suffix) == "boards.example" + suffix


# ---- the Summary matrix ------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_the_summary_matrix_is_the_reference_shape(
    make_builder: Make, cfg: ExportConfig
) -> None:
    content = _content(
        lambda s: (
            _post(
                s,
                "a",
                "Salesforce and SAP.",
                platforms=("SAP", "Salesforce"),
                posting_class="build",
            ),
            _post(s, "b", "Salesforce.", status="closed", posting_class="working_tool"),
            _post(s, "c", "Spin Diagnostics Salesforce.", posting_class="user"),
        ),
    ).model_copy(update={"entity_lock": _lock(DIVESTED)})
    result, files = await _build(make_builder, content)
    assert result.status == "complete", result.errors
    wb = load_workbook(io.BytesIO(files["jobs"]))
    ws = wb["Summary"]
    prefix = content.account_name
    assert ws["A1"].value.startswith(f"{prefix} platform job postings, last 12 months (first seen ")
    assert ws["A10"].value == f"Platform (named anywhere), {prefix} positions only"
    n = wb["Jobs"].max_row
    sf = f"Jobs!$C$2:$C${n}"
    assert ws["B11"].value == (f'=COUNTIFS({sf},"*Salesforce*",Jobs!$J$2:$J${n},"{prefix}*")')
    assert ws["C11"].value.endswith(f',Jobs!$E$2:$E${n},"Closed")')
    assert ws["D11"].value.endswith(f',Jobs!$E$2:$E${n},"Open")')
    assert [ws.cell(row=10, column=j).value for j in range(5, 9)] == [
        "Build role",
        "Admin and owner role",
        "Working tool",
        "User role",
    ]
    assert ws["E11"].value.endswith(f',Jobs!$D$2:$D${n},"Build role")')
    assert ws["A19"].value == "Rows on the Jobs sheet (all entities)"
    assert ws["B19"].value == f"=COUNTA(Jobs!$B$2:$B${n})"
    assert ws["A20"].value == f"Rows that are {prefix} positions"
    assert ws["B20"].value == f'=COUNTIF(Jobs!$J$2:$J${n},"{prefix}*")'
    assert ws["B21"].value == f'=COUNTA(Jobs!$B$2:$B${n})-COUNTIF(Jobs!$J$2:$J${n},"{prefix}*")'
    assert ws.freeze_panes is None and ws.auto_filter.ref is None


@pytest.mark.asyncio
@office
async def test_the_matrix_formulas_recalculate_to_the_code_values(
    make_builder: Make, cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = _content(
        lambda s: (
            _post(
                s,
                "a",
                "Salesforce and SAP.",
                platforms=("SAP", "Salesforce"),
                posting_class="build",
            ),
            _post(s, "b", "Salesforce.", status="closed", posting_class="working_tool"),
            _post(s, "c", "Spin Diagnostics Salesforce.", posting_class="user"),
            _post(s, "d", "Workday.", platforms=("Workday",), posting_class="admin_owner"),
        ),
    ).model_copy(update={"entity_lock": _lock(DIVESTED)})
    _, files = await _build(make_builder, content)
    model, _ctx = _model(content, cfg, grading)
    expected = expected_jobs(model, cfg)
    result = recalc_check(files["jobs"], expected, LibreOffice())
    assert result.status == "passed", result.details
    assert expected[("Summary", "B11")] == 2  # Salesforce, the account's own positions: a and b
    assert expected[("Summary", "C11")] == 1 and expected[("Summary", "D11")] == 1
    assert expected[("Summary", "B20")] == 3 and expected[("Summary", "B21")] == 1


@pytest.mark.asyncio
@office
async def test_the_rich_corpus_workbook_recalculates_with_the_builders_own_check(
    cfg: ExportConfig, grading: GradingConfig
) -> None:
    from exports.builder import ExportBuilder as Builder
    from tests.fakes.exports import FakeReader, InProcessRenderer

    content = account(n_postings=8)
    lo = LibreOffice()
    builder = Builder(
        FakeReader(content), MemoryStore(), InProcessRenderer(), cfg, grading, recalc=lo
    )
    out = builder.assemble(content)
    assert not out.errors, out.errors
    assert ("workbook_recalc", "passed") in {(c.name, c.status) for c in out.checks}


# ---- properties --------------------------------------------------------------------------------


@settings(max_examples=40, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(
    seen=st.lists(
        st.one_of(st.none(), st.dates(min_value=date(2025, 10, 1), max_value=date(2026, 9, 30))),
        max_size=14,
    )
)
def test_rows_are_always_newest_first_and_conserved(
    seen: list[date | None], cfg: ExportConfig, grading: GradingConfig
) -> None:
    content = _content(
        lambda s: tuple(
            _post(s, f"k{i:02d}", "Salesforce.", first_seen=d, posted_at=None)
            for i, d in enumerate(seen)
        )
    )
    model, _ = _model(content, cfg, grading)
    assert len(model.rows) == len(seen)
    texts = [r.first_seen_text for r in model.rows]
    dated = [t for t in texts if t]
    assert dated == sorted(dated, reverse=True)
    assert texts == dated + [""] * (len(texts) - len(dated))  # undated rows last
    assert [r.ref for r in model.rows] == [f"J-{i:04d}" for i in range(1, len(seen) + 1)]
