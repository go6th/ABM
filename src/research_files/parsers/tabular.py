"""Tabular parsers: CSV and TSV, XLSX, Parquet, JSON and JSON Lines (streaming).

Row numbers follow spreadsheet convention: the header is row 1, the first data row is row 2. A
unit holds at most ``rows_per_unit`` rows or ``unit_text_bytes`` of text. Cursors:

* delimited text: byte offset after the last row, the row number, encoding, delimiter, header;
* xlsx: sheet index and next row; parquet: rows done; JSON array: byte offset and item count;
  JSON Lines: byte offset and line number.

An xlsx sheet's header is its first row with a value. A cell that shows a link label ("Open
profile") reads as the address its hyperlink opens (read-only openpyxl drops hyperlinks, so they
are read from the sheet's own part, ``part_links``). The package is untrusted: only http(s)
addresses are kept, parts with a DTD or past their byte cap give no links, targets that leave the
package are not read, and the cells a hyperlink covers are bounded.

GUARANTEE(test_csv_resume_from_a_byte_offset_reads_the_same_rows)
GUARANTEE(test_only_http_links_become_cell_values)
GUARANTEE(test_a_huge_or_malformed_hyperlink_range_is_read_in_bounded_time)
GUARANTEE(test_a_part_that_declares_a_dtd_gives_no_links)
GUARANTEE(test_a_part_that_inflates_past_its_cap_gives_no_links)
"""

from __future__ import annotations

import codecs
import csv
import datetime as dt
import io
import json
import posixpath
import re
import zipfile
from collections.abc import Iterator
from typing import IO, Any
from urllib.parse import urlsplit

from research_files.checks import Finding
from research_files.parsers.base import ParseContext, Unit, detect_encoding, fatal

csv.field_size_limit(16 * 1024 * 1024)


def header_findings(table: str, header: list[str], location: dict[str, Any]) -> list[Finding]:
    """Empty, duplicate and numeric-looking header names."""
    out: list[Finding] = []
    blanks = [i + 1 for i, h in enumerate(header) if not str(h).strip()]
    if blanks:
        out.append(
            Finding(
                "warn",
                "header_empty_name",
                {**location, "columns": blanks[:20]},
                f"{len(blanks)} column(s) have no header name.",
                "Give every column a name.",
            )
        )
    seen: dict[str, int] = {}
    dups = []
    for h in header:
        key = str(h).strip().lower()
        if key:
            seen[key] = seen.get(key, 0) + 1
            if seen[key] == 2:
                dups.append(str(h))
    if dups:
        out.append(
            Finding(
                "warn",
                "header_duplicate_name",
                location,
                f"Duplicate column names: {', '.join(dups[:10])}.",
                "Rename the columns so each name is unique.",
            )
        )
    numeric = [
        h for h in header if str(h).strip().replace(".", "", 1).replace("-", "", 1).isdigit()
    ]
    if header and len(numeric) >= max(1, len(header) // 2):
        out.append(
            Finding(
                "warn",
                "header_looks_like_data",
                location,
                "The first row looks like data, not column names.",
                "Add a header row with column names.",
            )
        )
    return out


def unique_header(header: list[str]) -> list[str]:
    """Names usable as keys: blanks become column_N, duplicates get a suffix."""
    out: list[str] = []
    used: dict[str, int] = {}
    for i, raw in enumerate(header):
        name = str(raw).strip() or f"column_{i + 1}"
        if name in used:
            used[name] += 1
            name = f"{name}_{used[name]}"
        else:
            used[name] = 1
        out.append(name)
    return out


def _render(cells: list[Any]) -> str:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\n").writerow(["" if c is None else c for c in cells])
    return buf.getvalue()


class DelimitedParser:
    """CSV and TSV over a byte stream, resumable by byte offset."""

    def __init__(self, delimiter: str | None = None) -> None:
        self.delimiter = delimiter

    def units(self, stream: IO[bytes], cursor: dict[str, Any], ctx: ParseContext) -> Iterator[Unit]:
        """See the module docstring."""
        if cursor.get("done"):
            return
        table = cursor.get("table") or ctx.name
        if "offset" not in cursor:
            stream.seek(0)
            sample = stream.read(65536)
            encoding = detect_encoding(sample)
            text_sample = sample.decode(encoding, errors="replace")
            delimiter = self.delimiter or _sniff(text_sample)
            stream.seek(0)
            cursor = {"offset": 0, "row": 0, "encoding": encoding, "delimiter": delimiter}
        yield from self._rows(stream, cursor, ctx, table)

    def _rows(
        self, stream: IO[bytes], cursor: dict[str, Any], ctx: ParseContext, table: str
    ) -> Iterator[Unit]:
        encoding, delimiter = cursor["encoding"], cursor["delimiter"]
        stream.seek(int(cursor["offset"]))
        state = {"offset": int(cursor["offset"]), "nul": 0, "bad": 0}
        line_rows: list[int] = []
        decoder_name = "utf-8" if encoding == "utf-8-sig" and state["offset"] > 0 else encoding

        def lines() -> Iterator[str]:
            # LOOP(delimited lines): cap=the file's lines; exit=end of stream
            for raw in stream:
                state["offset"] += len(raw)
                text = raw.decode(decoder_name, errors="replace")
                if "\x00" in text:
                    state["nul"] += 1
                    text = text.replace("\x00", "")
                if "�" in text:
                    state["bad"] += 1
                    line_rows.append(row_no + 1)
                record_lines.append(text)
                yield text

        record_lines: list[str] = []  # the raw lines of the record being read
        reader = csv.reader(lines(), delimiter=delimiter)
        row_no = int(cursor["row"])
        header = cursor.get("header")
        header_unit: list[str] | None = None
        findings: list[Finding] = []
        if header is None:
            try:
                first = next(reader)
            except StopIteration:
                yield Unit(
                    "note",
                    {"table": table},
                    "",
                    {"done": True},
                    [
                        Finding(
                            "warn",
                            "empty_file",
                            {"table": table},
                            "The file has no rows.",
                            "Upload a file with a header and data rows.",
                        )
                    ],
                )
                return
            except csv.Error as first_error:
                yield fatal(
                    "note",
                    {"table": table, "row": 1},
                    "corrupt_file",
                    f"The first row cannot be read: {str(first_error)[:200]}.",
                    "Check that the file is a valid delimited text file.",
                )
                return
            findings += encoding_finding(encoding, {"table": table})
            row_no = 1
            raw_header = [h.lstrip("﻿") for h in first]
            findings += header_findings(table, raw_header, {"table": table, "row": 1})
            header = unique_header(raw_header)
            header_unit = header
        limit_rows = ctx.cfg.parsing.rows_per_unit
        limit_bytes = ctx.cfg.parsing.unit_text_bytes
        # LOOP(delimited units): cap=rows / rows_per_unit units; exit=end of rows, corrupt row
        while True:
            rows: list[tuple[int, list[Any]]] = []
            texts: list[str] = []
            size = 0
            start = row_no + 1
            error: csv.Error | None = None
            # LOOP(delimited unit rows): cap=rows_per_unit rows; exit=unit full, end of rows
            while len(rows) < limit_rows and size < limit_bytes:
                try:
                    cells = next(reader)
                except StopIteration:
                    break
                except csv.Error as caught:
                    error = caught
                    break
                # the lenient reader accepts a quote left open at the end of the file: an odd
                # number of quote characters in the record's raw lines means it was never closed
                unbalanced = sum(line.count('"') for line in record_lines) % 2 == 1
                record_lines.clear()
                if unbalanced and not _peek(stream):
                    error = csv.Error("the file ends inside a quoted value")
                    break
                row_no += 1
                rows.append((row_no, cells))
                line = _render(cells)
                texts.append(line)
                size += len(line)
            if state["bad"]:
                findings.append(
                    Finding(
                        "warn",
                        "encoding_problem",
                        {
                            "table": table,
                            "rows": [line_rows[0], line_rows[-1]] if line_rows else [start, row_no],
                        },
                        f"{state['bad']} line(s) had characters that could not be decoded"
                        f" as {encoding}.",
                        "Save the file as UTF-8 and upload it again.",
                    )
                )
                state["bad"] = 0
                line_rows.clear()
            if state["nul"]:
                findings.append(
                    Finding(
                        "warn",
                        "nul_bytes",
                        {"table": table, "rows": [start, row_no]},
                        f"{state['nul']} line(s) contained NUL bytes (removed for parsing).",
                        "Check whether the file was saved from a binary format.",
                    )
                )
                state["nul"] = 0
            next_cursor = {
                "offset": state["offset"],
                "row": row_no,
                "encoding": encoding,
                "delimiter": delimiter,
                "header": header,
                "table": table,
            }
            if rows:
                yield Unit(
                    "rows",
                    {"table": table, "rows": [start, row_no]},
                    "".join(texts),
                    next_cursor,
                    findings,
                    table=table,
                    header=header_unit,
                    header_row=1 if header_unit else None,
                    rows=rows,
                )
                header_unit, findings = None, []
            if error is not None:
                yield fatal(
                    "note",
                    {"table": table, "row": row_no + 1},
                    "truncated_or_corrupt",
                    f"Row {row_no + 1} cannot be read: {str(error)[:200]}. The file may be"
                    " truncated (for example, it ends inside a quoted value).",
                    "Re-export the file and upload it again.",
                )
                return
            if len(rows) < limit_rows and size < limit_bytes:
                yield Unit("note", {"table": table}, "", {"done": True}, findings)
                return


def _peek(stream: IO[bytes]) -> bytes:
    """One byte past the reader's position, without consuming it (empty at end of file)."""
    here = stream.tell()
    data = stream.read(1)
    stream.seek(here)
    return data


def encoding_finding(encoding: str, location: dict[str, Any]) -> list[Finding]:
    """A note when the text is not UTF-8 (it was decoded with the detected encoding)."""
    if encoding.replace("-sig", "").lower() in ("utf-8", "utf_8", "ascii"):
        return []
    return [
        Finding(
            "warn",
            "encoding_not_utf8",
            location,
            f"The file is not UTF-8; it was read as {encoding}. Some characters may be wrong.",
            "Save the file as UTF-8 and upload it again.",
        )
    ]


def _sniff(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample[:32768], delimiters=",;\t|").delimiter
    except csv.Error:
        return "\t" if sample.count("\t") > sample.count(",") else ","


def _cell(value: Any) -> Any:
    if isinstance(value, dt.datetime | dt.date | dt.time):
        return value.isoformat()
    if isinstance(value, bytes):
        return value[:64].hex()
    return value


# A link cell that shows a label ("Open profile", "Open posting", "View") is read as the address
# it opens: the People List's LinkedIn column is such a cell, and its label alone is no link.
_LINK_LABEL = re.compile(r"(?i)^\s*(?:open|view|see|visit|link|click)\b[\w\s]{0,30}$")
_XL_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_XL_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_CELL_REF = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d{1,7})(?::\$?([A-Za-z]{1,3})\$?(\d{1,7}))?$")
_LAST_ROW, _LAST_COL = 1_048_576, 16_384  # the largest sheet Excel can hold
# The link reader reads an uploaded (untrusted) package, so everything it reads is bounded:
MAX_LINKS = 200_000  # linked cells kept per sheet (the parser), per workbook (sheet_links)
MAX_LINK_CELLS = 1_000  # cells one hyperlink may cover; a larger range keeps its first cell only
MAX_URL_CHARS = 2_000  # a longer address is not kept
MAX_URL_BYTES = 32 * 1024 * 1024  # all kept addresses together
MAX_PART_BYTES = 64 * 1024 * 1024  # a workbook.xml or .rels part (decompressed) read for links
MAX_SHEET_BYTES = 1024 * 1024 * 1024  # a sheet part (decompressed) scanned for its hyperlinks


class _Refused(Exception):
    """A package part the link reader will not read: too long, or it declares a DTD."""


class _Capped:
    """A reader over one decompressed zip member that raises ``_Refused`` past ``cap`` bytes, so
    a member that inflates to gigabytes (a zip bomb) costs at most ``cap`` bytes of reading."""

    def __init__(self, raw: IO[bytes], cap: int) -> None:
        self.raw = raw
        self.left = cap

    def read(self, size: int = -1) -> bytes:
        """Up to ``size`` bytes (``size`` is bounded by the cap)."""
        want = self.left + 1 if size is None or size < 0 else min(size, self.left + 1)
        data = self.raw.read(want)
        self.left -= len(data)
        if self.left < 0:
            raise _Refused("part too long")
        return data


class CellLinks(dict[tuple[int, int], str]):
    """A sheet's linked cells: (row, column) -> address. ``anchors`` are the cells a hyperlink
    starts at: only there may an empty cell take the address (the other cells of a merged range
    are empty in read-only mode and stay empty)."""

    def __init__(self) -> None:
        super().__init__()
        self.anchors: set[tuple[int, int]] = set()


def safe_url(value: Any) -> str | None:
    """``value`` when it is an absolute http(s) address a reader may open, else None: never a
    ``javascript:``, ``file:``, ``mailto:`` or relative target, never one with a user name in
    front of the host (``https://trusted@other``), spaces, control or invisible characters."""
    if not isinstance(value, str):
        return None
    url = value.strip()
    if not url or len(url) > MAX_URL_CHARS or not url.isprintable() or " " in url:
        return None
    try:
        parts = urlsplit(url)
        host = parts.hostname
    except ValueError:
        return None
    if parts.scheme.lower() not in ("http", "https") or not host or "@" in parts.netloc:
        return None
    return url


def _elements(book: zipfile.ZipFile, name: str, tags: tuple[str, ...], cap: int) -> Iterator[Any]:
    """Each ``tags`` element of one package part as it ends, streamed and then dropped, so a big
    sheet's rows never pile up in memory. A part longer than ``cap`` bytes, or one that declares
    a DTD (no Office part does: it is how an entity bomb or an external entity is written),
    raises ``_Refused``. Entities are never expanded and nothing is fetched."""
    from lxml import etree

    with book.open(name) as raw:
        events = etree.iterparse(
            _Capped(raw, cap),
            events=("end",),
            tag=tags,
            resolve_entities=False,
            no_network=True,
            load_dtd=False,
            huge_tree=False,
        )
        checked = False
        # LOOP(package part elements): cap=`cap` bytes of the part; exit=end of the part
        for _, el in events:
            if not checked:
                if el.getroottree().docinfo.doctype:
                    raise _Refused("the part declares a DTD")
                checked = True
            yield el
            el.clear()
            # LOOP(drop read siblings): cap=the elements already read; exit=none left before
            while el.getprevious() is not None:
                del el.getparent()[0]


def _rels_name(part: str) -> str:
    """The relationships part of a package part (``xl/_rels/workbook.xml.rels``)."""
    folder, _, name = part.rpartition("/")
    return f"{folder}/_rels/{name}.rels" if folder else f"_rels/{name}.rels"


def _part_name(source: str, target: str) -> str | None:
    """A relationship target as the part it names, resolved as openpyxl resolves it (relative to
    the source part's folder; a leading ``/`` is the package root), or None when it climbs out of
    the package (``../../x``) or is no part name (an address)."""
    if not target or "://" in target or "\\" in target:
        return None
    if target.startswith("/"):
        path = posixpath.normpath(target.lstrip("/"))
    else:
        path = posixpath.normpath(posixpath.join(posixpath.dirname(source), target))
    if path.startswith(("../", "/")) or path in (".", ".."):
        return None
    return path


def _hyperlink_targets(book: zipfile.ZipFile, part: str, budget: list[int]) -> dict[str, str]:
    """Relationship id -> address of every external hyperlink a sheet part declares, kept only
    when it is a safe http(s) address (``safe_url``); bounded by ``MAX_LINKS`` and
    ``MAX_URL_BYTES`` (``budget`` is the bytes left, shared by the workbook)."""
    rels = _rels_name(part)
    if rels not in book.NameToInfo:
        return {}
    out: dict[str, str] = {}
    for el in _elements(book, rels, (f"{_PKG_REL}Relationship",), MAX_PART_BYTES):
        if el.get("TargetMode") != "External" or not el.get("Type", "").endswith("/hyperlink"):
            continue
        url = safe_url(el.get("Target"))
        if url is None:
            continue
        budget[0] -= len(url)
        if budget[0] < 0 or len(out) >= MAX_LINKS:
            break
        out[el.get("Id", "")] = url
    return out


def _column(letters: str) -> int:
    """A column's number from its letters (A is 1, XFD is 16384)."""
    number = 0
    for ch in letters.upper():
        number = number * 26 + ord(ch) - 64
    return number


def _span(ref: str | None) -> tuple[int, int, int, int] | None:
    """A hyperlink's ``ref`` as (first row, first column, last row, last column), or None when it
    is no cell or range of the sheet (malformed, a whole column, past the sheet's last cell)."""
    m = _CELL_REF.match((ref or "").strip().split(" ")[0])
    if m is None:
        return None
    c1, c2 = _column(m.group(1)), _column(m.group(3) or m.group(1))
    r1, r2 = int(m.group(2)), int(m.group(4) or m.group(2))
    top, bottom, left, right = min(r1, r2), max(r1, r2), min(c1, c2), max(c1, c2)
    if top < 1 or bottom > _LAST_ROW or left < 1 or right > _LAST_COL:
        return None
    return top, left, bottom, right


def part_links(
    book: zipfile.ZipFile, part: str, cells_left: list[int], url_bytes: list[int]
) -> CellLinks:
    """The linked cells of one worksheet part (see ``sheet_links``). Raises ``_Refused`` (or a
    zip or XML error) for a part it will not read; ``cells_left`` and ``url_bytes`` are the
    workbook's remaining budgets."""
    out = CellLinks()
    targets = _hyperlink_targets(book, part, url_bytes) if part in book.NameToInfo else {}
    if not targets:
        return out
    hyperlinks = f"{_XL_MAIN}hyperlinks"
    for el in _elements(
        book, part, (f"{_XL_MAIN}row", f"{_XL_MAIN}hyperlink", hyperlinks), MAX_SHEET_BYTES
    ):
        if el.tag == hyperlinks:
            break  # the sheet's hyperlinks are all read; what follows has none
        if el.tag != f"{_XL_MAIN}hyperlink":
            continue  # a row: read only to get past it (and dropped)
        url = targets.get(el.get(f"{_XL_REL}id", ""))  # an internal link (location=) has none
        span = _span(el.get("ref")) if url else None
        if url is None or span is None:
            continue
        top, left, bottom, right = span
        if (bottom - top + 1) * (right - left + 1) > MAX_LINK_CELLS:
            bottom, right = top, left  # a range this wide is no list of link cells
        out.anchors.add((top, left))
        # LOOP(hyperlink cells): cap=MAX_LINK_CELLS per hyperlink, MAX_LINKS per workbook;
        # exit=range end, budget spent
        for row in range(top, bottom + 1):
            for col in range(left, right + 1):
                if cells_left[0] <= 0:
                    return out
                if (row, col) not in out:
                    out[(row, col)] = url
                    cells_left[0] -= 1
    return out


def _workbook_part(book: zipfile.ZipFile) -> str:
    """The workbook part the package's root relationships name (``xl/workbook.xml`` almost
    always)."""
    if "_rels/.rels" in book.NameToInfo:
        for el in _elements(book, "_rels/.rels", (f"{_PKG_REL}Relationship",), MAX_PART_BYTES):
            if el.get("Type", "").endswith("/officeDocument"):
                part = _part_name("", el.get("Target", ""))
                if part is not None:
                    return part
    return "xl/workbook.xml"


def sheet_links(stream: IO[bytes]) -> dict[str, CellLinks]:
    """The external hyperlink of every linked cell, by sheet name and (row, column): openpyxl's
    read-only mode does not read them. Only safe http(s) addresses are kept (``safe_url``); a
    sheet whose relationship target leaves the package is skipped; a repeated sheet name is
    renamed as openpyxl renames it. Bounded (``MAX_LINKS`` and the byte caps above); a package
    it cannot read safely gives no links rather than an error (the values are still read).
    ``XlsxParser`` reads one sheet's links at a time (``_links_for``)."""
    from openpyxl.workbook.child import avoid_duplicate_name

    if not _seekable(stream):
        return {}
    here = stream.tell()
    out: dict[str, CellLinks] = {}
    try:
        with zipfile.ZipFile(stream) as book:
            workbook = _workbook_part(book)
            sheets = [
                (str(el.get("name", "")), el.get(f"{_XL_REL}id", ""))
                for el in _elements(book, workbook, (f"{_XL_MAIN}sheet",), MAX_PART_BYTES)
            ]
            targets: dict[str, str] = {}
            if _rels_name(workbook) in book.NameToInfo:
                for el in _elements(
                    book, _rels_name(workbook), (f"{_PKG_REL}Relationship",), MAX_PART_BYTES
                ):
                    if el.get("TargetMode") != "External":
                        targets[el.get("Id", "")] = el.get("Target", "")
            cells_left, url_bytes = [MAX_LINKS], [MAX_URL_BYTES]
            # LOOP(workbook sheets for links): cap=the workbook's sheets; exit=last sheet
            for name, rel in sheets:
                part = _part_name(workbook, targets.get(rel, ""))
                if part is None or part not in book.NameToInfo:
                    continue  # openpyxl skips such a sheet too
                key = avoid_duplicate_name(list(out), name)
                out[key] = part_links(book, part, cells_left, url_bytes)
    except Exception:  # noqa: BLE001 - zip, zlib, XML and limit errors: no links, values stay
        return {}
    finally:
        stream.seek(here)
    return out


def _seekable(stream: Any) -> bool:
    """Whether the stream can be read twice (a zip package needs it)."""
    try:
        return bool(stream.seekable())
    except (AttributeError, OSError, ValueError):
        return False


def _links_for(stream: IO[bytes], ws: Any) -> CellLinks:
    """The linked cells of one read-only worksheet, read from its own part (openpyxl's resolved
    ``_worksheet_path``, so two sheets with one name never share links). Never raises: a package
    or part it cannot read safely, or a stream it cannot rewind, gives no links. The stream's
    position is restored for openpyxl."""
    part = getattr(ws, "_worksheet_path", None)
    if not isinstance(part, str) or not _seekable(stream):
        return CellLinks()
    here = stream.tell()
    try:
        with zipfile.ZipFile(stream) as book:
            return part_links(book, part, [MAX_LINKS], [MAX_URL_BYTES])
    except Exception:  # noqa: BLE001 - zip, zlib, XML and limit errors: no links, values stay
        return CellLinks()
    finally:
        stream.seek(here)


def linked(cells: list[Any], row_no: int, links: dict[tuple[int, int], str]) -> list[Any]:
    """The row's cells with each labelled link cell read as its address (``_LINK_LABEL``); an
    empty cell takes the address only where a hyperlink starts (``CellLinks.anchors``)."""
    if not links:
        return cells
    anchors = getattr(links, "anchors", None)
    for j, value in enumerate(cells):
        url = links.get((row_no, j + 1))
        if not url:
            continue
        label = isinstance(value, str) and _LINK_LABEL.match(value) is not None
        empty_anchor = value is None and (anchors is None or (row_no, j + 1) in anchors)
        if label or empty_anchor:
            cells[j] = url
    return cells


class XlsxParser:
    """Every sheet of a workbook, row batches in read-only (streaming) mode."""

    def units(self, stream: IO[bytes], cursor: dict[str, Any], ctx: ParseContext) -> Iterator[Unit]:
        """See the module docstring."""
        if cursor.get("done"):
            return
        from openpyxl import load_workbook

        try:
            book = load_workbook(stream, read_only=True, data_only=True)
        except Exception as error:  # noqa: BLE001 - zip, zlib, XML (an entity bomb) and value errors
            yield fatal(
                "note",
                {},
                "corrupt_file",
                f"The workbook cannot be opened: {type(error).__name__}.",
                "Open and save the workbook in a spreadsheet program, then upload it again.",
            )
            return
        name = ""
        try:
            sheets = book.sheetnames
            sheet_i = int(cursor.get("sheet", 0))
            next_row = int(cursor.get("row", 1))
            limit = ctx.cfg.parsing.rows_per_unit
            # LOOP(xlsx sheets): cap=the workbook's sheets; exit=last sheet
            while sheet_i < len(sheets):
                name = sheets[sheet_i]
                ws = book[name]
                header: list[str] | None = (
                    cursor.get("header") if cursor.get("sheet") == sheet_i else None
                )
                header_unit = None
                findings: list[Finding] = []
                # read before the rows: a sheet's hyperlinks follow its rows in the part
                links = _links_for(stream, ws)
                rows_iter = ws.iter_rows(min_row=next_row, values_only=True)
                row_no = next_row - 1
                header_row = next_row
                if header is None:
                    # the header is the first row with a value: a sheet may open with empty rows
                    first = None
                    # LOOP(xlsx leading empty rows): cap=the sheet's rows; exit=a row with a value
                    for values in rows_iter:
                        row_no += 1
                        if any(v is not None and str(v).strip() for v in values):
                            first = values
                            break
                    header_row = row_no
                    if first is None:
                        yield Unit(
                            "note",
                            {"sheet": name},
                            "",
                            {"sheet": sheet_i + 1, "row": 1},
                            [
                                Finding(
                                    "info",
                                    "empty_sheet",
                                    {"sheet": name},
                                    f'The sheet "{name}" is empty.',
                                    None,
                                )
                            ],
                        )
                        sheet_i, next_row = sheet_i + 1, 1
                        continue
                    raw = [("" if c is None else str(c)) for c in first]
                    # LOOP(xlsx header trailing blanks): cap=the header's cells; exit=a named cell
                    while raw and raw[-1] == "":
                        raw.pop()
                    findings += header_findings(name, raw, {"sheet": name, "row": header_row})
                    header = unique_header(raw)
                    header_unit = header
                batch: list[tuple[int, list[Any]]] = []
                texts: list[str] = []
                start = row_no + 1
                for values in rows_iter:
                    row_no += 1
                    cells = linked([_cell(v) for v in values], row_no, links)
                    # LOOP(xlsx row trailing blanks): cap=the row's cells; exit=a value or the
                    # header's width
                    while cells and cells[-1] is None and len(cells) > len(header):
                        cells.pop()
                    if all(c is None for c in cells):
                        continue
                    batch.append((row_no, cells))
                    texts.append(_render(cells))
                    if len(batch) >= limit:
                        yield Unit(
                            "rows",
                            {"sheet": name, "rows": [start, row_no]},
                            "".join(texts),
                            {"sheet": sheet_i, "row": row_no + 1, "header": header},
                            findings,
                            table=name,
                            header=header_unit,
                            header_row=header_row if header_unit else None,
                            rows=batch,
                        )
                        batch, texts, findings, header_unit, start = [], [], [], None, row_no + 1
                if batch or header_unit or findings:
                    yield Unit(
                        "rows" if batch else "note",
                        {"sheet": name, "rows": [start, row_no]} if batch else {"sheet": name},
                        "".join(texts),
                        {"sheet": sheet_i + 1, "row": 1},
                        findings,
                        table=name,
                        header=header_unit,
                        header_row=header_row if header_unit else None,
                        rows=batch or None,
                    )
                sheet_i, next_row = sheet_i + 1, 1
                cursor = {}
        except Exception as error:  # noqa: BLE001 - a sheet part that is corrupt or hostile XML
            yield fatal(
                "note",
                {"sheet": name},
                "truncated_or_corrupt",
                f'The sheet "{name}" cannot be read: {type(error).__name__}.',
                "Open and save the workbook in a spreadsheet program, then upload it again.",
            )
            return
        finally:
            book.close()
        yield Unit("note", {}, "", {"done": True})


class ParquetParser:
    """Row batches of a Parquet file (row groups read one at a time)."""

    def units(self, stream: IO[bytes], cursor: dict[str, Any], ctx: ParseContext) -> Iterator[Unit]:
        """See the module docstring."""
        if cursor.get("done"):
            return
        try:
            import pyarrow.parquet as pq  # type: ignore[import-untyped]
        except ImportError:  # pragma: no cover - installed with the project
            yield fatal(
                "note",
                {},
                "parser_unavailable",
                "Parquet support is not installed.",
                "Ask an operator to install the parquet parser.",
            )
            return
        table = ctx.name
        try:
            pf = pq.ParquetFile(stream)
        except Exception as error:  # noqa: BLE001 - pyarrow raises many types for bad files
            yield fatal(
                "note",
                {"table": table},
                "corrupt_file",
                f"The Parquet file cannot be opened: {type(error).__name__}.",
                "Re-export the file and upload it again.",
            )
            return
        header = unique_header([f.name for f in pf.schema_arrow])
        done = int(cursor.get("rows_done", 0))
        groups, start_of = [], 0
        for g in range(pf.num_row_groups):
            n = pf.metadata.row_group(g).num_rows
            if start_of + n > done:
                groups.append((g, start_of))
            start_of += n
        first_unit = done == 0
        limit = ctx.cfg.parsing.rows_per_unit
        for g, g_start in groups:
            skip = max(0, done - g_start)
            try:
                batches = pf.iter_batches(batch_size=limit, row_groups=[g])
                offset = g_start
                for batch in batches:
                    records = batch.to_pylist()
                    if offset + len(records) <= done:
                        offset += len(records)
                        continue
                    if skip and offset < done:
                        records = records[done - offset :]
                        offset = done
                    rows = []
                    texts = []
                    for k, rec in enumerate(records):
                        cells = [_cell(rec.get(name)) for name in pf.schema_arrow.names]
                        row_no = offset + k + 2
                        rows.append((row_no, cells))
                        texts.append(_render(cells))
                    offset += len(records)
                    done = offset
                    yield Unit(
                        "rows",
                        {"table": table, "rows": [rows[0][0], rows[-1][0]]},
                        "".join(texts),
                        {"rows_done": done},
                        [],
                        table=table,
                        header=header if first_unit else None,
                        header_row=1 if first_unit else None,
                        rows=rows,
                    )
                    first_unit = False
            except Exception as error:  # noqa: BLE001
                yield fatal(
                    "note",
                    {"table": table, "row_group": g},
                    "corrupt_file",
                    f"Row group {g} cannot be read: {type(error).__name__}.",
                    "Re-export the file and upload it again.",
                )
                return
        yield Unit("note", {"table": table}, "", {"done": True})


def _flatten(record: Any) -> dict[str, Any]:
    if not isinstance(record, dict):
        return {"value": record}
    out: dict[str, Any] = {}
    for key, value in record.items():
        if isinstance(value, dict):
            for k2, v2 in value.items():
                out[f"{key}.{k2}"] = json.dumps(v2) if isinstance(v2, dict | list) else v2
        elif isinstance(value, list):
            out[str(key)] = json.dumps(value)[:2000]
        else:
            out[str(key)] = value
    return out


class _Records:
    """Turn dict records into rows over a growing header."""

    def __init__(self, header: list[str] | None) -> None:
        self.header = list(header or [])
        self.index = {h: i for i, h in enumerate(self.header)}

    def row(self, record: Any) -> list[Any]:
        flat = _flatten(record)
        for key in flat:
            if key not in self.index:
                self.index[key] = len(self.header)
                self.header.append(key)
        cells: list[Any] = [None] * len(self.header)
        for key, value in flat.items():
            cells[self.index[key]] = value
        return cells


class JsonLinesParser:
    """One JSON value per line; a bad line is a finding and the parse goes on."""

    def units(self, stream: IO[bytes], cursor: dict[str, Any], ctx: ParseContext) -> Iterator[Unit]:
        """See the module docstring."""
        if cursor.get("done"):
            return
        table = ctx.name
        offset, line_no = int(cursor.get("offset", 0)), int(cursor.get("line", 0))
        records = _Records(cursor.get("header"))
        first = offset == 0
        stream.seek(offset)
        limit, limit_bytes = ctx.cfg.parsing.rows_per_unit, ctx.cfg.parsing.unit_text_bytes
        rows: list[tuple[int, list[Any]]] = []
        texts: list[str] = []
        findings: list[Finding] = []
        bad = 0
        size = 0
        for raw in stream:
            offset += len(raw)
            line_no += 1
            text = raw.decode("utf-8", errors="replace").strip()
            if not text:
                continue
            try:
                value = json.loads(text)
            except json.JSONDecodeError as error:
                bad += 1
                if bad <= ctx.cfg.audit.max_findings_per_check:
                    findings.append(
                        Finding(
                            "error",
                            "invalid_json_line",
                            {"table": table, "line": line_no},
                            f"Line {line_no} is not valid JSON ({error.msg}).",
                            "Fix or remove the line.",
                        )
                    )
                continue
            cells = records.row(value)
            rows.append((line_no, cells))
            texts.append(text[:4000] + "\n")
            size += len(text)
            if len(rows) >= limit or size >= limit_bytes:
                yield Unit(
                    "rows",
                    {"table": table, "lines": [rows[0][0], rows[-1][0]]},
                    "".join(texts),
                    {"offset": offset, "line": line_no, "header": records.header},
                    findings,
                    table=table,
                    header=list(records.header) if first else None,
                    header_row=None,
                    rows=rows,
                )
                rows, texts, findings, size, first = [], [], [], 0, False
        if rows or findings:
            yield Unit(
                "rows" if rows else "note",
                {"table": table, "lines": [rows[0][0], rows[-1][0]]} if rows else {"table": table},
                "".join(texts),
                {"offset": offset, "line": line_no, "header": records.header},
                findings,
                table=table if rows else None,
                header=list(records.header) if first and rows else None,
                rows=rows or None,
            )
        yield Unit("note", {"table": table}, "", {"done": True})


class JsonParser:
    """A top-level JSON array is streamed item by item; any other document is read whole up to
    ``json_max_document_bytes`` (a list of objects inside it becomes the table)."""

    CHUNK = 1 << 20

    def units(self, stream: IO[bytes], cursor: dict[str, Any], ctx: ParseContext) -> Iterator[Unit]:
        """See the module docstring."""
        if cursor.get("done"):
            return
        table = ctx.name
        if "offset" not in cursor:
            stream.seek(0)
            head = stream.read(4096).lstrip(b"\xef\xbb\xbf \t\r\n")
            if not head.startswith(b"["):
                yield from self._document(stream, ctx)
                return
            stream.seek(0)
            raw = stream.read(4096)
            cursor = {"offset": raw.index(b"[") + 1, "item": 0}
        yield from self._array(stream, cursor, ctx, table)

    def _array(
        self, stream: IO[bytes], cursor: dict[str, Any], ctx: ParseContext, table: str
    ) -> Iterator[Unit]:
        decoder = json.JSONDecoder()
        records = _Records(cursor.get("header"))
        offset, item = int(cursor["offset"]), int(cursor["item"])
        first = item == 0
        stream.seek(offset)
        text_decoder = codecs.getincrementaldecoder("utf-8")(errors="replace")
        buf = ""
        eof = False
        limit, limit_bytes = ctx.cfg.parsing.rows_per_unit, ctx.cfg.parsing.unit_text_bytes
        rows: list[tuple[int, list[Any]]] = []
        texts: list[str] = []
        size = 0
        # LOOP(json array items): cap=the array's items; exit=closing bracket, corrupt item
        while True:
            stripped = buf.lstrip(" \t\r\n,")
            offset += len(buf[: len(buf) - len(stripped)].encode("utf-8"))
            buf = stripped
            if buf.startswith("]"):
                break
            try:
                value, end = decoder.raw_decode(buf)
            except json.JSONDecodeError as error:
                if not eof and len(buf) < 256 * self.CHUNK:
                    chunk = stream.read(self.CHUNK)
                    if not chunk:
                        eof = True
                        buf += text_decoder.decode(b"", final=True)
                    else:
                        buf += text_decoder.decode(chunk)
                    continue
                if rows:
                    yield self._unit(table, rows, texts, offset, item, records, first)
                yield fatal(
                    "note",
                    {"table": table, "item": item + 1, "byte": offset},
                    "truncated_or_corrupt",
                    f"Item {item + 1} (byte {offset}) is not valid JSON ({error.msg}); the file"
                    " may be truncated.",
                    "Re-export the file and upload it again.",
                )
                return
            item += 1
            consumed = buf[:end]
            offset += len(consumed.encode("utf-8"))
            buf = buf[end:]
            rows.append((item, records.row(value)))
            texts.append(consumed[:4000] + "\n")
            size += len(consumed)
            if len(rows) >= limit or size >= limit_bytes:
                yield self._unit(table, rows, texts, offset, item, records, first)
                rows, texts, size, first = [], [], 0, False
        if rows:
            yield self._unit(table, rows, texts, offset, item, records, first)
        yield Unit("note", {"table": table}, "", {"done": True})

    @staticmethod
    def _unit(
        table: str,
        rows: list[tuple[int, list[Any]]],
        texts: list[str],
        offset: int,
        item: int,
        records: _Records,
        first: bool,
    ) -> Unit:
        return Unit(
            "rows",
            {"table": table, "items": [rows[0][0], rows[-1][0]]},
            "".join(texts),
            {"offset": offset, "item": item, "header": records.header},
            [],
            table=table,
            header=list(records.header) if first else None,
            rows=list(rows),
        )

    def _document(self, stream: IO[bytes], ctx: ParseContext) -> Iterator[Unit]:
        table = ctx.name
        stream.seek(0)
        data = stream.read(ctx.cfg.parsing.json_max_document_bytes + 1)
        if len(data) > ctx.cfg.parsing.json_max_document_bytes:
            yield fatal(
                "note",
                {"table": table},
                "too_large_to_parse",
                "This JSON file is a single object larger than the parse limit; only JSON"
                " arrays and JSON Lines are read in streaming mode.",
                "Export the data as a JSON array or JSON Lines.",
            )
            return
        try:
            value = json.loads(data.decode(detect_encoding(data[:65536]), errors="replace"))
        except json.JSONDecodeError as error:
            yield fatal(
                "note",
                {"table": table, "line": error.lineno, "column": error.colno},
                "truncated_or_corrupt",
                f"The JSON is not valid at line {error.lineno}"
                f" ({error.msg}); the file may be truncated.",
                "Re-export the file and upload it again.",
            )
            return
        records_list = _largest_record_list(value)
        if records_list is not None:
            path, items = records_list
            records = _Records(None)
            rows = [(i + 1, records.row(v)) for i, v in enumerate(items)]
            for start in range(0, len(rows), ctx.cfg.parsing.rows_per_unit):
                chunk = rows[start : start + ctx.cfg.parsing.rows_per_unit]
                yield Unit(
                    "rows",
                    {"table": table, "path": path, "items": [chunk[0][0], chunk[-1][0]]},
                    "".join(
                        json.dumps(v, default=str)[:4000] + "\n"
                        for v in items[start : start + len(chunk)]
                    ),
                    {"done": False},
                    [],
                    table=f"{table}:{path}",
                    header=list(records.header) if start == 0 else None,
                    rows=chunk,
                )
        text = json.dumps(value, indent=1, default=str)
        yield Unit(
            "text", {"table": table}, text[: ctx.cfg.parsing.unit_text_bytes], {"done": True}
        )


def _largest_record_list(value: Any, path: str = "$") -> tuple[str, list[Any]] | None:
    best: tuple[str, list[Any]] | None = None
    if isinstance(value, list) and value and all(isinstance(v, dict) for v in value[:50]):
        best = (path, value)
    if isinstance(value, dict):
        for key, child in value.items():
            found = _largest_record_list(child, f"{path}.{key}")
            if found and (best is None or len(found[1]) > len(best[1])):
                best = found
    return best
