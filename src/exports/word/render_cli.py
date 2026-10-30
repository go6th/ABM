"""Interim Word renderer behind the D-26 contract: ``DocumentSpec`` JSON on stdin, docx bytes on
stdout. Exit 0 on success, 2 when the input violates the contract, 3 when rendering fails.

This is NOT the legacy Node library (not supplied); it is a neutral python-docx renderer that runs
in the same sandbox (subprocess, cleared environment, no database credentials, bounded output) so
the whole pipeline is exercised for real. When the Node library is ported it replaces this
command; nothing else changes. Every style value comes from ``spec.theme``.

What it writes, following the owner's reference files: Word ``Heading 1`` and ``Heading 2``
paragraph styles (navigation pane and outline) with a page break before a section that asks for
one; a hyperlink (``w:hyperlink`` over an external relationship) for every run or paragraph with
an ``href``; a cell fill for every body cell whose whole text is a theme fill word (grades and
tracker statuses); the label column of a table in bold; a right-aligned footer with a PAGE field
and a header part only for a white-label brand name.
"""

from __future__ import annotations

import copy
import io
import re
import sys
import weakref
import zipfile
from functools import lru_cache
from typing import Any
from xml.sax.saxutils import escape  # nosec B406 -- escapes generated XML; parses nothing.

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from pydantic import ValidationError

from exports.config import Font, WordTheme
from exports.word.contract import (
    BulletsBlock,
    Cell,
    DocumentSpec,
    HeadingBlock,
    ParagraphBlock,
    TableBlock,
    TitleBlock,
    cell_text,
)

EXIT_BAD_INPUT = 2
EXIT_RENDER = 3
_HEADING_STYLES = {1: "Heading 1", 2: "Heading 2"}


@lru_cache(maxsize=64)
def _rpr(spec: Font) -> Any:
    """The run properties of one theme font, built once (schema order: rFonts, b, i, color,
    sz, szCs, u) and copied into every run: building them per run is the renderer's hot path."""
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), spec.name)
    rpr.append(fonts)
    for flag, on in (("w:b", spec.bold), ("w:i", spec.italic)):
        if on:
            rpr.append(OxmlElement(flag))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), spec.color)
    rpr.append(color)
    half_points = str(round(spec.size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        size = OxmlElement(tag)
        size.set(qn("w:val"), half_points)
        rpr.append(size)
    if spec.underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    return rpr


def _link_font(theme: WordTheme, size_of: Font) -> Font:
    """The hyperlink font at the size of the text around it (the theme's link font is the table
    size; a body paragraph link keeps the body size)."""
    link = theme.fonts.link
    if link.size_pt == size_of.size_pt:
        return link
    return link.model_copy(update={"size_pt": size_of.size_pt})


def _text(paragraph: Paragraph, text: str, spec: Font) -> None:
    """One run in the theme font; newlines in the text become line breaks."""
    _append_run(paragraph._p, text, spec)


def _run(text: str, spec: Font) -> Any:
    run = OxmlElement("w:r")
    run.append(copy.deepcopy(_rpr(spec)))
    for i, line in enumerate(text.split("\n")):
        if i:
            run.append(OxmlElement("w:br"))
        t = OxmlElement("w:t")
        t.text = line
        t.set(qn("xml:space"), "preserve")
        run.append(t)
    return run


def _append_run(p: Any, text: str, spec: Font) -> None:
    p.append(_run(text, spec))


class _LinkIds:
    """The external hyperlink relationships of one document part: one per distinct target, made
    in constant time. python-docx's ``relate_to`` walks every relationship of the part to find a
    match and again to find a free id, so a document with thousands of links took quadratic
    time (7,000 postings: minutes). Ids are handed out in order of first use, the same ids
    ``relate_to`` would give (a part without gaps), so the file is byte for byte the same."""

    def __init__(self, part: Any) -> None:
        self.rels = part.rels
        self.ids: dict[str, str] = {}
        self.next = len(self.rels) + 1

    def rid(self, href: str) -> str:
        found = self.ids.get(href)
        if found is not None:
            return found
        # LOOP(free relationship id): cap=the part's relationships plus one; exit=an unused id
        while f"rId{self.next}" in self.rels:
            self.next += 1
        rid = f"rId{self.next}"
        self.next += 1
        self.rels.add_relationship(RT.HYPERLINK, href, rid, is_external=True)
        self.ids[href] = rid
        return rid


_LINKS: weakref.WeakKeyDictionary[Any, _LinkIds] = weakref.WeakKeyDictionary()


def _link_id(part: Any, href: str) -> str:
    """The relationship id of the external hyperlink to ``href`` on ``part``."""
    ids = _LINKS.get(part)
    if ids is None:
        ids = _LINKS[part] = _LinkIds(part)
    return ids.rid(href)


def _append_link(p: Any, text: str, href: str, spec: Font, part: Any) -> None:
    """A hyperlink run: ``w:hyperlink`` over an external relationship to ``href`` (the same
    target gets one relationship, however often it is linked)."""
    rid = _link_id(part, href)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    link.set(qn("w:history"), "1")
    link.append(_run(text, spec))
    p.append(link)


def _append_cell(p: Any, cell: Cell, spec: Font, theme: WordTheme, part: Any) -> None:
    """A cell's content: plain text, or its runs in order, links as hyperlinks."""
    if isinstance(cell, str):
        _append_run(p, cell, spec)
        return
    link = _link_font(theme, spec)
    for run in cell:
        if run.href:
            _append_link(p, run.text, run.href, link, part)
        else:
            _append_run(p, run.text, spec)


def _fresh_numbering(doc: Any) -> int:
    """A new numbering instance of the List Number style's definition that starts at 1: every
    numbered list of the document counts from 1, not on from the previous list."""
    style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
    numbering = doc.part.numbering_part.element
    abstract = next(n.abstractNumId.val for n in numbering.num_lst if n.numId == style_num)
    num = numbering.add_num(abstract)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return int(num.numId)


def _bullets(doc: Any, block: BulletsBlock, theme: WordTheme) -> None:
    """A bulleted or numbered list; a note goes under its item, indented like the item text."""
    sp = theme.spacing
    num_id = _fresh_numbering(doc) if block.numbered else None
    for i, item in enumerate(block.items):
        p = doc.add_paragraph(style="List Number" if block.numbered else "List Bullet")
        if num_id is not None:
            num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = num_id
        _text(p, item, theme.fonts.body)
        _spacing(p, 0, sp.paragraph_after_pt)
        note = block.notes[i] if block.notes else ""
        if note:
            q = doc.add_paragraph(style="List Continue")  # indented like the item text
            _text(q, note, theme.fonts.note)
            _spacing(q, 0, sp.paragraph_after_pt)


def _spacing(paragraph: Paragraph, before: float, after: float) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)


def _border(tbl: Table, theme: WordTheme) -> None:
    tbl_pr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(theme.table.border_size_eighth_pt))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), theme.table.border_color)
        borders.append(el)
    tbl_pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, twips in (
        ("top", theme.table.cell_margin_twips),
        ("left", theme.table.cell_margin_side_twips),
        ("bottom", theme.table.cell_margin_twips),
        ("right", theme.table.cell_margin_side_twips),
    ):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(twips))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tbl_pr.append(margins)


def _row_props(row_el: Any, theme: WordTheme, header: bool) -> None:
    tr_pr = OxmlElement("w:trPr")
    row_el.append(tr_pr)
    if theme.table.row_cant_split:
        tr_pr.append(OxmlElement("w:cantSplit"))
    if header and theme.table.repeat_header_row:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _shade(tc_pr: Any, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _row_template(theme: WordTheme, twips: list[int], header: bool) -> Any:
    """An empty table row: row properties, and per column a cell with its width (and the header
    fill on the header row) holding one paragraph with no spacing."""
    tr = OxmlElement("w:tr")
    _row_props(tr, theme, header)
    for width in twips:
        tc = OxmlElement("w:tc")
        tc_pr = OxmlElement("w:tcPr")
        tc_w = OxmlElement("w:tcW")
        tc_w.set(qn("w:w"), str(width))
        tc_w.set(qn("w:type"), "dxa")
        tc_pr.append(tc_w)
        if header:
            _shade(tc_pr, theme.table.header_fill)
        tc.append(tc_pr)
        p = OxmlElement("w:p")
        p_pr = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "0")
        spacing.set(qn("w:after"), "0")
        p_pr.append(spacing)
        p.append(p_pr)
        tc.append(p)
        tr.append(tc)
    return tr


def _table(doc: Any, block: TableBlock, theme: WordTheme) -> None:
    page = theme.page
    text_width = page.width_in - page.margins_in.left - page.margins_in.right
    tbl = doc.add_table(rows=0, cols=len(block.headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _border(tbl, theme)
    widths = [Inches(text_width * w / 100) for w in block.widths]
    for k, w in enumerate(widths):
        tbl.columns[k].width = w
    twips = [round(text_width * w / 100 * 1440) for w in block.widths]
    templates = (_row_template(theme, twips, header=True), _row_template(theme, twips, False))
    body = theme.fonts.table_body
    label = body.model_copy(update={"bold": True}) if theme.table.bold_label_column else body
    fills = theme.fills
    part = doc.part
    if len(block.rows) >= STREAM_ROWS:
        _table_rows_streamed(doc, tbl, block, theme, twips, label)
        return
    # rows are built from a template row (python-docx's add_row and cell accessors walk the
    # table and re-create properties per cell: the renderer's hot path on large tables)
    for r, values in enumerate((block.headers, *block.rows)):
        header = r == 0
        tr = copy.deepcopy(templates[0 if header else 1])
        for k, tc in enumerate(tr.findall(qn("w:tc"))):
            cell: Cell = values[k]
            if header:
                spec = theme.fonts.table_header
            else:
                spec = label if k == block.label_column else body
                fill = fills.get(cell_text(cell))
                if fill:
                    _shade(tc.find(qn("w:tcPr")), fill)
            _append_cell(tc.find(qn("w:p")), cell, spec, theme, part)
        tbl._tbl.append(tr)


# A table of this many body rows is written as text, not as a tree of elements: a Word table of a
# million cells held as an element tree took over a gigabyte at 7,000 postings (fix wave C). The
# rows are the same XML (``tests/exports/test_fix_wave_c.py`` compares both, canonically) and are
# put into the saved document in place of a marker row.
STREAM_ROWS = 400
_MARK = "@@ROWS-{}@@"
_BAD_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


@lru_cache(maxsize=64)
def _rpr_xml(spec: Font) -> str:
    """``_rpr(spec)`` as XML text without the namespace declaration (the document root has it)."""
    return re.sub(r' xmlns:\w+="[^"]*"', "", etree.tostring(_rpr(spec)).decode())


def _run_xml(text: str, spec: Font) -> str:
    """``_run`` as text: the run properties, then each line in a ``w:t``, a ``w:br`` between."""
    lines = _BAD_XML.sub("", text).split("\n")
    body = "<w:br/>".join(f'<w:t xml:space="preserve">{escape(line)}</w:t>' for line in lines)
    return f"<w:r>{_rpr_xml(spec)}{body}</w:r>"


def _cell_xml(cell: Cell, spec: Font, theme: WordTheme, part: Any) -> str:
    """``_append_cell`` as text: plain text, or its runs in order, links as hyperlinks."""
    if isinstance(cell, str):
        return _run_xml(cell, spec)
    link = _link_font(theme, spec)
    out = []
    for run in cell:
        if run.href:
            rid = _link_id(part, run.href)
            out.append(
                f'<w:hyperlink r:id="{rid}" w:history="1">{_run_xml(run.text, link)}</w:hyperlink>'
            )
        else:
            out.append(_run_xml(run.text, spec))
    return "".join(out)


def _row_xml(
    theme: WordTheme, twips: list[int], header: bool, cells: list[tuple[str, str | None]]
) -> str:
    """``_row_template`` plus the content of each cell, as text. ``cells`` is per column the
    content XML and the fill (None: no fill)."""
    flags = ""
    if theme.table.row_cant_split:
        flags += "<w:cantSplit/>"
    if header and theme.table.repeat_header_row:
        flags += "<w:tblHeader/>"
    tcs = []
    for width, (content, fill) in zip(twips, cells, strict=True):
        shade = ""
        fill = theme.table.header_fill if header else fill
        if fill:
            shade = f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>'
        tcs.append(
            f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{shade}</w:tcPr>'
            f'<w:p><w:pPr><w:spacing w:before="0" w:after="0"/></w:pPr>{content}</w:p></w:tc>'
        )
    return f"<w:tr><w:trPr>{flags}</w:trPr>{''.join(tcs)}</w:tr>"


def _table_rows_streamed(
    doc: Any, tbl: Table, block: TableBlock, theme: WordTheme, twips: list[int], label: Font
) -> None:
    """The header and body rows of ``tbl`` as text, spliced into the document when it is saved
    (``_splice``); the table gets one marker row now. Hyperlink relationships are made here, so
    they are in the package when it is saved."""
    body = theme.fonts.table_body
    part = doc.part
    rows = [
        _row_xml(
            theme,
            twips,
            True,
            [(_cell_xml(c, theme.fonts.table_header, theme, part), None) for c in block.headers],
        )
    ]
    for values in block.rows:
        rows.append(
            _row_xml(
                theme,
                twips,
                False,
                [
                    (
                        _cell_xml(cell, label if k == block.label_column else body, theme, part),
                        theme.fills.get(cell_text(cell)),
                    )
                    for k, cell in enumerate(values)
                ],
            )
        )
    pending: list[list[str]] = doc.__dict__.setdefault("_pending_rows", [])
    marker = OxmlElement("w:tr")
    tc = OxmlElement("w:tc")
    p = OxmlElement("w:p")
    _append_run(p, _MARK.format(len(pending)), body)
    tc.append(p)
    marker.append(tc)
    tbl._tbl.append(marker)
    pending.append(rows)


def _splice(data: bytes, pending: list[list[str]]) -> bytes:
    """``data`` (a saved docx) with each marker row replaced by the rows written as text."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            raw = src.read(info)
            if info.filename != "word/document.xml":
                dst.writestr(info, raw, compress_type=zipfile.ZIP_DEFLATED)
                continue
            text = raw.decode("utf-8")
            with dst.open(zipfile.ZipInfo(info.filename, date_time=info.date_time), "w") as w:
                cursor = 0
                for k, rows in enumerate(pending):
                    at = text.index(_MARK.format(k))
                    start = text.rfind("<w:tr>", 0, at)
                    end = text.index("</w:tr>", at) + len("</w:tr>")
                    w.write(text[cursor:start].encode("utf-8"))
                    for row in rows:
                        w.write(row.encode("utf-8"))
                    cursor = end
                w.write(text[cursor:].encode("utf-8"))
    return out.getvalue()


def _footer(doc: Any, page_label: str, theme: WordTheme) -> None:
    """The reference footer: "Page N" right-aligned, N a PAGE field; nothing else."""
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    spec = theme.fonts.footer
    _append_run(p._p, f"{page_label} ", spec)
    for kind in ("begin", "instr", "separate", "text", "end"):
        run = OxmlElement("w:r")
        run.append(copy.deepcopy(_rpr(spec)))
        if kind == "instr":
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = " PAGE "
            run.append(instr)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = "1"
            run.append(t)
        else:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run.append(fld)
        p._p.append(run)


def _rule_under(paragraph: Paragraph, theme: WordTheme) -> None:
    """The rule under the opening paragraph (the reference's orange line)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(theme.title_rule.size_eighth_pt))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), theme.title_rule.color)
    borders.append(bottom)
    p_pr.append(borders)


def _title(doc: Any, block: TitleBlock, theme: WordTheme) -> None:
    sp = theme.spacing
    p = doc.add_paragraph()
    _text(p, block.text, theme.fonts.title)
    _spacing(p, 0, sp.paragraph_after_pt)
    if block.note:
        p = doc.add_paragraph()
        _text(p, block.note, theme.fonts.note)
        _spacing(p, 0, sp.paragraph_after_pt)
    p = doc.add_paragraph()
    _text(p, block.subtitle, theme.fonts.subtitle)
    _rule_under(p, theme)
    _spacing(p, 0, sp.heading_before_pt)


def _heading(doc: Any, block: HeadingBlock, theme: WordTheme) -> None:
    """A Word heading style (outline level for the navigation pane) in the theme font."""
    sp = theme.spacing
    p = doc.add_paragraph(style=_HEADING_STYLES[block.level])
    _text(p, block.text, theme.fonts.heading1 if block.level == 1 else theme.fonts.heading2)
    _spacing(p, sp.heading_before_pt, sp.heading_after_pt)
    fmt = p.paragraph_format
    fmt.keep_with_next = True
    if block.page_break:
        fmt.page_break_before = True


def _paragraph(doc: Any, block: ParagraphBlock, theme: WordTheme) -> None:
    p = doc.add_paragraph()
    font = {"body": theme.fonts.body, "quote": theme.fonts.quote, "note": theme.fonts.note}[
        block.style
    ]
    if block.href:
        _append_link(p._p, block.text, block.href, _link_font(theme, font), doc.part)
    else:
        _text(p, block.text, font)
    _spacing(p, 0, theme.spacing.paragraph_after_pt)


def render(spec: DocumentSpec) -> bytes:
    """docx bytes for one spec (not yet normalised; the caller normalises the package)."""
    theme = spec.theme
    doc = Document()
    section = doc.sections[0]
    page = theme.page
    section.orientation = (
        WD_ORIENT.LANDSCAPE if page.orientation == "landscape" else WD_ORIENT.PORTRAIT
    )
    section.page_width, section.page_height = Inches(page.width_in), Inches(page.height_in)
    m = page.margins_in
    section.top_margin, section.bottom_margin = Inches(m.top), Inches(m.bottom)
    section.left_margin, section.right_margin = Inches(m.left), Inches(m.right)
    normal = doc.styles["Normal"]
    normal.font.name = theme.fonts.body.name
    normal.font.size = Pt(theme.fonts.body.size_pt)
    _footer(doc, spec.page_label, theme)
    if spec.header_text:
        _text(section.header.paragraphs[0], spec.header_text, theme.fonts.header)
    sp = theme.spacing
    for block in spec.blocks:
        if isinstance(block, TitleBlock):
            _title(doc, block, theme)
        elif isinstance(block, HeadingBlock):
            _heading(doc, block, theme)
        elif isinstance(block, ParagraphBlock):
            _paragraph(doc, block, theme)
        elif isinstance(block, BulletsBlock):
            _bullets(doc, block, theme)
        else:
            _table(doc, block, theme)
            _spacing(doc.add_paragraph(), 0, sp.paragraph_after_pt)
    out = io.BytesIO()
    doc.save(out)
    pending = doc.__dict__.get("_pending_rows")
    return _splice(out.getvalue(), pending) if pending else out.getvalue()


def main() -> int:
    """Read the spec from stdin, write the docx to stdout."""
    try:
        spec = DocumentSpec.model_validate_json(sys.stdin.buffer.read())
    except (ValidationError, ValueError) as error:
        sys.stderr.write(f"contract violation: {error.__class__.__name__}\n")
        return EXIT_BAD_INPUT
    try:
        data = render(spec)
    except Exception as error:  # noqa: BLE001 - any renderer failure is one exit code
        sys.stderr.write(f"render failed: {error.__class__.__name__}\n")
        return EXIT_RENDER
    sys.stdout.buffer.write(data)
    sys.stdout.buffer.flush()
    return 0


if __name__ == "__main__":  # pragma: no cover - entry point, exercised through the subprocess
    sys.exit(main())
