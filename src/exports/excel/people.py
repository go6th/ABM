"""People List workbook in the shape of the owner's reference People List (reference set of
2026-10-20, measured 2026-10-29). The Summary sheet comes first: the title, the paragraphs in the
reference's order (what the list merges and from which searches, the rows removed before it, the
people excluded by title by name, the seniority rule, the judgment note, the entity labels), the
Level table (nine levels, People and Priority A, Total, Meets seniority rule) and the Lane table
(eight lanes, senior people only, A, B, C and Total, Total row), every count a formula citing its
label cell. The People sheet follows with the reference's 13 columns and widths, the product's
Ref and Source columns to their right, the Priority cell filled A, B or C, profile links reading
"Open profile", panes frozen right of the name. Layout from ``workbooks.people``; look from the
theme; every formula cell is planned once (``plan_people``) and checked after recalculation
(B12)."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter

from exports.config import Column, ExportConfig
from exports.excel.common import Book, Formula, criterion, expected_of, sheet_ref
from exports.models import PeopleModel

# the Priority cell's theme fill by letter (``excel.fills``)
_PRIORITY_FILL = {"A": "priority_a", "B": "priority_b", "C": "priority_c"}


@dataclass(frozen=True)
class _Layout:
    """Row positions of the Summary sheet blocks (one computation for writer and checker)."""

    paragraphs_first: int
    level_header: int
    level_first: int
    level_total: int
    senior_row: int
    lane_header: int
    lane_first: int
    lane_total: int


def _layout(model: PeopleModel) -> _Layout:
    paragraphs_first = 3
    level_header = paragraphs_first + len(model.paragraphs) + 2
    level_first = level_header + 1
    level_total = level_first + len(model.levels)
    senior_row = level_total + 1
    lane_header = senior_row + 2
    lane_first = lane_header + 1
    lane_total = lane_first + len(model.lanes_table)
    return _Layout(
        paragraphs_first,
        level_header,
        level_first,
        level_total,
        senior_row,
        lane_header,
        lane_first,
        lane_total,
    )


def plan_people(model: PeopleModel, cfg: ExportConfig) -> tuple[_Layout, tuple[Formula, ...]]:
    """The Summary sheet layout and every formula cell with its code-computed value. A count
    cites the cell that holds its label (``=COUNTIF(People!$D$2:$D$N,A12)``), as the reference
    does, so a renamed label in the sheet still counts."""
    book = cfg.rules.workbooks.people
    labels = cfg.rules.labels
    keys = [c.key for c in book.columns]
    col = {k: get_column_letter(keys.index(k) + 1) for k in keys}
    last = max(2, len(model.rows) + 1)

    def span(key: str) -> str:
        return f"{sheet_ref(book.sheet)}!${col[key]}$2:${col[key]}${last}"

    lay = _layout(model)
    sheet = book.summary_sheet
    out: list[Formula] = []
    for i, level in enumerate(model.levels):
        row = lay.level_first + i
        label = f"A{row}"
        out.append(Formula(sheet, row, 2, f"=COUNTIF({span('level')},{label})", level.people))
        out.append(
            Formula(
                sheet,
                row,
                3,
                f"=COUNTIFS({span('level')},{label},{span('priority')},{criterion('A')})",
                level.priority_a,
            )
        )
    if model.levels:
        first, before = lay.level_first, lay.level_total - 1
        out.append(Formula(sheet, lay.level_total, 2, f"=SUM(B{first}:B{before})", len(model.rows)))
        out.append(
            Formula(
                sheet,
                lay.level_total,
                3,
                f"=SUM(C{first}:C{before})",
                sum(x.priority_a for x in model.levels),
            )
        )
    yes = criterion(labels.yes)
    out.append(
        Formula(
            sheet,
            lay.senior_row,
            2,
            f"=COUNTIF({span('meets_seniority')},{yes})",
            model.senior_count,
        )
    )
    for i, lane in enumerate(model.lanes_table):
        row = lay.lane_first + i
        for j, (letter, n) in enumerate((("A", lane.a), ("B", lane.b), ("C", lane.c)), start=2):
            text = (
                f"=COUNTIFS({span('lanes')},$A{row},{span('priority')},{criterion(letter)},"
                f"{span('meets_seniority')},{yes})"
            )
            out.append(Formula(sheet, row, j, text, n))
        out.append(Formula(sheet, row, 5, f"=SUM(B{row}:D{row})", lane.total))
    if model.lanes_table:
        first, before = lay.lane_first, lay.lane_total - 1
        totals = (
            sum(x.a for x in model.lanes_table),
            sum(x.b for x in model.lanes_table),
            sum(x.c for x in model.lanes_table),
            sum(x.total for x in model.lanes_table),
        )
        for j, n in enumerate(totals, start=2):
            letter = get_column_letter(j)
            text = f"=SUM({letter}{first}:{letter}{before})"
            out.append(Formula(sheet, lay.lane_total, j, text, n))
    return lay, tuple(out)


def expected_people(
    model: PeopleModel, cfg: ExportConfig
) -> dict[tuple[str, str], int | str | float]:
    """Formula cells and the values they must show (the builder's ``_expected``)."""
    return expected_of(plan_people(model, cfg)[1])


def render_people(model: PeopleModel, cfg: ExportConfig, title: str) -> tuple[bytes, int]:
    """(xlsx bytes, continuation parts)."""
    layout = cfg.rules.workbooks.people
    texts = cfg.rules.sheet_texts.people
    book = Book(cfg)
    summary = book.sheet(layout.summary_sheet)
    ws = book.sheet(layout.sheet)
    book.header(
        ws,
        layout.columns,
        freeze_columns=layout.freeze_columns,
        height_pt=layout.header_height_pt,
    )
    for i, row in enumerate(model.rows, start=2):
        values = row.model_dump()
        book.body(ws, layout.columns, values, i, row.ref)
        for j, column in enumerate(layout.columns, start=1):
            if column.key == "priority" and row.priority in _PRIORITY_FILL:
                book.fill(ws, i, j, _PRIORITY_FILL[row.priority])
    if not model.rows:
        book.note(ws, 3, cfg.rules.none_found.text)
    book.finish_table(ws, layout.columns, len(model.rows))

    lay, formulas = plan_people(model, cfg)
    level_cols = (
        Column(key="label", header=texts.level_heading, kind="text_short"),
        Column(key="people", header="People", kind="count"),
        Column(key="priority_a", header="Priority A", kind="count"),
    )
    lane_cols = (
        Column(key="label", header=texts.lane_heading, kind="text_short"),
        Column(key="a", header="A", kind="count"),
        Column(key="b", header="B", kind="count"),
        Column(key="c", header="C", kind="count"),
        Column(key="total", header="Total", kind="count"),
    )
    book.header(summary, level_cols, row=lay.level_header, freeze=False, border=True)
    for i, level in enumerate(model.levels):
        book.body(summary, level_cols, {"label": level.label}, lay.level_first + i, "")
    book.body(summary, level_cols, {"label": texts.total_label}, lay.level_total, "", bold=True)
    book.body(summary, level_cols[:1], {"label": texts.senior_label}, lay.senior_row, "")
    book.header(summary, lane_cols, row=lay.lane_header, freeze=False, border=True)
    for i, lane in enumerate(model.lanes_table):
        book.body(summary, lane_cols, {"label": lane.label}, lay.lane_first + i, "")
    book.body(summary, lane_cols, {"label": texts.total_label}, lay.lane_total, "", bold=True)
    # the label column and the count columns (the Level and Lane tables share them)
    book.set_width(summary, 1, layout.summary_label_width)
    for j in range(2, layout.summary_span + 1):
        book.set_width(summary, j, layout.summary_count_width)
    # the title and the paragraphs last, so their heights follow the final column widths
    book.title_line(summary, 1, model.title_line or title)
    for k, text in enumerate(model.paragraphs, start=lay.paragraphs_first):
        book.paragraph(
            summary, k, text, layout.summary_span, height_pt=layout.summary_paragraph_height_pt
        )
    total_rows = {lay.level_total, lay.lane_total}
    for f in formulas:
        book.formula(f, bold=f.row in total_rows)
    return book.bytes(title, model.as_of), book.counted
