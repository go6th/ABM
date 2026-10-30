"""Tech Stack and Teams workbook, in the shape of the owner's reference file set (Tech Stack and
Teams workbook): three sheets in this order.

* Platform Ranking: the title line, the base line with its number in the share column (the cell
  every share formula divides by), the rule paragraphs (merged across the table, real row
  heights), the ranked table (rank and share as formulas, master prompt v2 section 7.3; the
  rank cell bold on its theme fill and under a conditional format, so a corrected count
  recolours it), the counts per rank letter (COUNTIF formulas) and the platforms seen once.
* Confirmed stack: the title line, the note merged across the five columns, the header on
  row 4 and one row per technology of the organisation's tech stack (business function,
  category, technology, postings naming it, lane). No Ref column: the reference has none.
* Platform teams: the title line, the note, the reference's header on row 4 and one wrapped row
  per team when teams were stored, else the visible "None recorded" line.

Every formula is planned once (``plan_tech``) so the writer and the recalculation check (B12)
read the same cells.
"""

from __future__ import annotations

from collections.abc import Sequence

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from exports.config import Column, ExportConfig
from exports.excel.common import (
    LINE_SPACING,
    WRAP_SLACK,
    Book,
    Formula,
    criterion,
    expected_of,
    font,
    text_cell,
)
from exports.models import TechModel
from exports.readers.tech import RANK_LETTERS

# the title line, then the base line, then the rule paragraphs, a gap and the table header
_TITLE_ROW = 1
_BASE_ROW = 2
_NOTES_FIRST = 3
# the stack and teams sheets: title, note, a gap, then the header
LIST_HEADER = 4
# the reference's row heights: the Platform Ranking header, the Confirmed stack and Platform teams
# notes, the Platform teams header
_RANKING_HEADER_HEIGHT = 45.0
_LIST_NOTE_HEIGHT = 40.0
_TEAMS_HEADER_HEIGHT = 28.0
# Excel's tallest row
_MAX_ROW_HEIGHT = 409.0
# a rank letter's theme fill
_RANK_FILL = {"A": "rank_a", "B": "rank_b", "C": "rank_c"}


def _rows(model: TechModel) -> dict[str, int]:
    """Row positions of the Platform Ranking blocks (one computation for writer and checker)."""
    header = _NOTES_FIRST + len(model.notes) + 1
    first = header + 1
    last = header + len(model.ranking)
    counts_heading = last + 2
    counts_first = counts_heading + 1
    seen_heading = counts_first + len(RANK_LETTERS) + 1
    return {
        "header": header,
        "first": first,
        "last": max(first, last),
        "counts_heading": counts_heading,
        "counts_first": counts_first,
        "seen_heading": seen_heading,
        "seen_first": seen_heading + 1,
    }


def plan_tech(model: TechModel, cfg: ExportConfig) -> tuple[Formula, ...]:
    """The rank, share and rank-count formula cells with their code-computed values."""
    book = cfg.rules.workbooks.tech
    keys = [c.key for c in book.ranking_columns]
    col = {k: keys.index(k) + 1 for k in keys}
    letter = {k: get_column_letter(n) for k, n in col.items()}
    sheet = book.ranking_sheet
    rows = _rows(model)
    out: list[Formula] = []
    count, share = letter["count"], letter["share"]
    base = f"${share}${_BASE_ROW}"
    for i, r in enumerate(model.ranking, start=rows["first"]):
        rank = (
            f'=IF(AND({share}{i}>=0.15,{count}{i}>=5),"A",'
            f'IF(AND({share}{i}>=0.05,{count}{i}>=3),"B",IF({count}{i}>=2,"C","")))'
        )
        out.append(Formula(sheet, i, col["rank"], rank, r.rank, "label"))
        share_text = f"=IF({base}=0,0,{count}{i}/{base})"
        out.append(Formula(sheet, i, col["share"], share_text, r.share, "percent"))
    rank_span = f"${letter['rank']}${rows['first']}:${letter['rank']}${rows['last']}"
    for j, (x, n) in enumerate(model.rank_counts):
        out.append(
            Formula(sheet, rows["counts_first"] + j, 2, f"=COUNTIF({rank_span},{criterion(x)})", n)
        )
    return tuple(out)


def expected_tech(model: TechModel, cfg: ExportConfig) -> dict[tuple[str, str], int | str | float]:
    """Formula cells and the values they must show (the Platform Ranking's)."""
    return expected_of(plan_tech(model, cfg))


def _plain_style(book: Book, *, bold: bool = False, kind: str = "label") -> str:
    """The named style of a line outside the tables: the body font, no border, no wrapping (the
    base line and its number, the Counts labels)."""
    theme = book.theme
    name = f"Export plain {kind}" + (" bold" if bold else "")
    return book._named(
        name,
        font(theme.fonts.body, bold=True if bold else None),
        None,
        None,
        theme.column_kinds[kind],  # type: ignore[index]
        wrap=False,
    )


def _plain(
    book: Book, ws: Worksheet, row: int, column: int, text: str, *, bold: bool = False
) -> None:
    """A text line outside the tables in the plain style."""
    cell = ws.cell(row=row, column=column)
    text_cell(cell, text)
    cell.style = _plain_style(book, bold=bold)


def _filter(book: Book, ws: Worksheet, columns: Sequence[Column], header: int, n: int) -> None:
    if book.theme.autofilter:
        ws.auto_filter.ref = f"A{header}:{get_column_letter(len(columns))}{header + max(n, 0)}"


def _rank_formats(book: Book, ws: Worksheet, first: int, last: int) -> None:
    """The rank cells' conditional formats (the theme's rank fills): a corrected count that
    changes a rank letter changes its colour, as in the reference."""
    for letter, name in _RANK_FILL.items():
        color = book.theme.fills[name]
        ws.conditional_formatting.add(
            f"A{first}:A{last}",
            FormulaRule(  # type: ignore[no-untyped-call]
                formula=[f'$A{first}="{letter}"'],
                fill=PatternFill(fill_type="solid", start_color=color, end_color=color),
            ),
        )


def _ranking(book: Book, model: TechModel, cfg: ExportConfig) -> None:
    layout = cfg.rules.workbooks.tech
    texts = cfg.rules.sheet_texts.tech
    cols = layout.ranking_columns
    keys = [c.key for c in cols]
    ws = book.sheet(layout.ranking_sheet)
    rows = _rows(model)
    book.title_line(ws, _TITLE_ROW, model.title)
    _plain(book, ws, _BASE_ROW, 1, model.base_line)
    base_cell = ws.cell(row=_BASE_ROW, column=keys.index("share") + 1)
    base_cell.value = model.base
    base_cell.style = _plain_style(book, bold=True, kind="count")
    # the header first: it sets the column widths the paragraphs' row heights are estimated from
    book.header(ws, cols, row=rows["header"], height_pt=_RANKING_HEADER_HEIGHT)
    for k, note in enumerate(model.notes):
        book.paragraph(ws, _NOTES_FIRST + k, note, len(cols), font_role="note")
    for i, row in enumerate(model.ranking, start=rows["first"]):
        book.body(ws, cols, row.model_dump(), i, "")
    _filter(book, ws, cols, rows["header"], len(model.ranking))
    _rank_formats(book, ws, rows["first"], rows["last"])
    _plain(book, ws, rows["counts_heading"], 1, texts.counts_heading, bold=True)
    for j, (x, _n) in enumerate(model.rank_counts):
        _plain(book, ws, rows["counts_first"] + j, 1, texts.rank_label.format(rank=x))
    _plain(book, ws, rows["seen_heading"], 1, texts.seen_once_heading, bold=True)
    shown = [c for c in cols if c.key in ("platform", "category", "lane")]
    for i, r in enumerate(model.seen_once, start=rows["seen_first"]):
        values = r.model_dump()
        for c in shown:
            _plain(book, ws, i, keys.index(c.key) + 1, str(values[c.key]))
    if not model.seen_once:
        _plain(book, ws, rows["seen_first"], keys.index("platform") + 1, texts.seen_once_none)
    ranks = {(f.row, f.column): f.expected for f in plan_tech(model, cfg)}
    for f in plan_tech(model, cfg):
        if f.column == 1 and f.row >= rows["first"] and f.row <= rows["last"]:
            # the rank cell: bold on the fill of its letter
            book.formula(f, bold=True, fill=_RANK_FILL.get(str(ranks[(f.row, f.column)])))
        else:
            book.formula(f)


def _stack(book: Book, model: TechModel, cfg: ExportConfig) -> None:
    layout = cfg.rules.workbooks.tech
    ws = book.sheet(layout.sheet)
    book.title_line(ws, _TITLE_ROW, model.stack_title)
    book.header(ws, layout.columns, row=LIST_HEADER)
    book.paragraph(
        ws,
        _TITLE_ROW + 1,
        model.stack_note,
        len(layout.columns),
        font_role="note",
        height_pt=_LIST_NOTE_HEIGHT,
    )
    for i, line in enumerate(model.stack, start=LIST_HEADER + 1):
        book.body(ws, layout.columns, line.model_dump(), i, "")
    if not model.stack:
        book.note(ws, LIST_HEADER + 1, cfg.rules.none_found.text)
    _filter(book, ws, layout.columns, LIST_HEADER, len(model.stack))


def _lines(text: str, width: float) -> int:
    """Wrapped lines of ``text`` in a column ``width`` characters wide (words that do not fill a
    line are allowed for)."""
    return max(1, -(-int(len(text) * WRAP_SLACK) // max(1, int(width))))


def _teams(book: Book, model: TechModel, cfg: ExportConfig) -> None:
    layout = cfg.rules.workbooks.tech
    cols = layout.teams_columns
    ws = book.sheet(layout.teams_sheet)
    book.title_line(ws, _TITLE_ROW, model.teams_title)
    book.header(ws, cols, row=LIST_HEADER, height_pt=_TEAMS_HEADER_HEIGHT)
    book.paragraph(
        ws,
        _TITLE_ROW + 1,
        model.teams_sheet_note,
        len(cols),
        font_role="note",
        height_pt=_LIST_NOTE_HEIGHT,
    )
    size = book.theme.fonts.body.size_pt
    for i, team in enumerate(model.teams, start=LIST_HEADER + 1):
        values = team.model_dump()
        book.body(ws, cols, values, i, "")
        # Excel does not size a wrapped row itself in a file written by a library
        most = max(
            _lines(str(values[c.key]), c.width or 10) for c in cols if c.key not in ("postings",)
        )
        ws.row_dimensions[i].height = min(_MAX_ROW_HEIGHT, most * size * LINE_SPACING)
    if not model.teams:
        book.note(ws, LIST_HEADER + 1, cfg.rules.none_found.teams)


def render_tech(model: TechModel, cfg: ExportConfig, title: str) -> tuple[bytes, int]:
    """(xlsx bytes, continuation parts)."""
    book = Book(cfg)
    _ranking(book, model, cfg)
    _stack(book, model, cfg)
    _teams(book, model, cfg)
    return book.bytes(title, model.as_of), book.counted
