"""Job Postings workbook, in the order and shape of the owner's reference file (Summary, Jobs;
``docs/parity/reference_spec.md``): Summary (title line, five explanatory paragraphs merged
across A:H, the platform matrix of the seven reference platforms with Total, Closed, Open and the
four classes as COUNTIFS over the Jobs sheet restricted to the account's own entity, then the
three row-count lines), Jobs (the reference's 14 columns, then the owner's Archived link and New
since last research, then the product's Ref and Source code), Counts (the count lines as formulas
and the job pipeline's own counts) and Long text (counted continuation of descriptions over the
cell limit). Counts and Long text come after the reference sheets (owner decision pending)."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.utils import get_column_letter

from exports.config import Column, ExportConfig
from exports.excel.common import (
    Book,
    Formula,
    counts_sheet,
    criterion,
    expected_of,
    prefix_criterion,
    sheet_ref,
)
from exports.models import JobsModel

NEW_ROW_FILL = "new_row"


@dataclass(frozen=True)
class _Layout:
    """Summary sheet rows (1-based): the reference's title on row 1, the notes from row 3, two
    blank rows, the matrix header (row 10 for five notes), the platform rows, a blank row and
    the three row-count lines (rows 19 to 21)."""

    title: int
    paragraphs_first: int
    matrix_header: int
    matrix_first: int
    rows_all: int
    rows_account: int
    rows_other: int


def _layout(model: JobsModel) -> _Layout:
    first = 3
    header = first + len(model.summary_notes) + 2
    rows_all = header + len(model.matrix) + 2
    return _Layout(
        title=1,
        paragraphs_first=first,
        matrix_header=header,
        matrix_first=header + 1,
        rows_all=rows_all,
        rows_account=rows_all + 1,
        rows_other=rows_all + 2,
    )


def _wildcard(name: str) -> str:
    """A COUNTIF criterion for a cell that contains ``name`` (any case)."""
    return f'"*{name.replace("~", "~~").replace("*", "~*").replace("?", "~?")}*"'


def _prefix(model: JobsModel) -> str:
    return model.prefix or model.account


def plan_jobs(model: JobsModel, cfg: ExportConfig) -> tuple[Formula, ...]:
    """The Summary sheet's formula cells with their code-computed values: the platform matrix
    (Total, Closed, Open, then one column per class, every cell restricted to the rows whose
    Entity starts with the account prefix) and the three row counts (the count lines of the
    Counts sheet are written by ``counts_sheet`` and checked through ``formula``)."""
    book = cfg.rules.workbooks.jobs
    labels = cfg.rules.labels
    keys = [c.key for c in book.columns]
    col = {k: get_column_letter(keys.index(k) + 1) for k in keys}
    last = max(2, len(model.rows) + 1)
    sheet_name = sheet_ref(book.sheet)

    def span(key: str) -> str:
        return f"{sheet_name}!${col[key]}$2:${col[key]}${last}"

    platforms, classes, status = span("platforms"), span("posting_class"), span("status")
    entity, titles = span("entity"), span("title")
    closed_word = criterion(labels.posting_status["closed"])
    open_word = criterion(labels.posting_status["open"])
    own = prefix_criterion(_prefix(model))
    lay = _layout(model)
    sheet = book.summary_sheet
    out: list[Formula] = []
    for i, m in enumerate(model.matrix):
        row = lay.matrix_first + i
        base = f"{platforms},{_wildcard(m.platform)},{entity},{own}"
        cells = [
            (f"=COUNTIFS({base})", m.rows),
            (f"=COUNTIFS({base},{status},{closed_word})", m.closed),
            (f"=COUNTIFS({base},{status},{open_word})", m.open),
            *(
                (f"=COUNTIFS({base},{classes},{criterion(label)})", n)
                for label, n in zip(model.class_labels, m.by_class, strict=True)
            ),
        ]
        out.extend(Formula(sheet, row, 2 + j, text, n) for j, (text, n) in enumerate(cells))
    prefix = _prefix(model).casefold()
    account = sum(1 for r in model.rows if r.entity.casefold().startswith(prefix))
    mine = f"COUNTIF({entity},{own})"
    out.append(Formula(sheet, lay.rows_all, 2, f"=COUNTA({titles})", len(model.rows)))
    out.append(Formula(sheet, lay.rows_account, 2, f"={mine}", account))
    out.append(
        Formula(sheet, lay.rows_other, 2, f"=COUNTA({titles})-{mine}", len(model.rows) - account)
    )
    return tuple(out)


def expected_jobs(model: JobsModel, cfg: ExportConfig) -> dict[tuple[str, str], int | str | float]:
    """Formula cells and the values they must show: the count lines and the Summary sheet."""
    sheet = cfg.rules.workbooks.jobs.counts_sheet
    lines: dict[tuple[str, str], int | str | float] = {
        (sheet, f"B{i}"): line.value
        for i, line in enumerate(model.counts, start=2)
        if line.formula != "none"
    }
    return {**lines, **expected_of(plan_jobs(model, cfg))}


def render_jobs(model: JobsModel, cfg: ExportConfig, title: str) -> tuple[bytes, int]:
    """(xlsx bytes, continuation parts)."""
    layout = cfg.rules.workbooks.jobs
    texts = cfg.rules.sheet_texts.jobs
    book = Book(cfg, long_text_sheet=layout.long_text_sheet)
    summary = book.sheet(layout.summary_sheet)  # first, as in the reference file
    ws = book.sheet(layout.sheet)
    # freeze below the header and right of Requisition and Title (the reference's C2)
    book.header(ws, layout.columns, freeze_columns=2)
    for i, row in enumerate(model.rows, start=2):
        book.body(ws, layout.columns, row.model_dump(), i, row.ref)
        if row.new_since_previous:  # a row first seen after the previous research
            book.fill_row(ws, i, layout.columns, NEW_ROW_FILL)
    if not model.rows:
        book.note(ws, 3, cfg.rules.none_found.text)
    book.finish_table(ws, layout.columns, len(model.rows))
    keys = [c.key for c in layout.columns]
    status_col = get_column_letter(keys.index("status") + 1)
    counts_sheet(
        book,
        layout.counts_sheet,
        model.counts,
        layout.sheet,
        len(model.rows),
        status_col,
        cfg.rules.labels.posting_status,
        count_col=get_column_letter(keys.index("ref") + 1),
    )
    cws = book.wb[layout.counts_sheet]
    for k, sentence in enumerate(model.sentences, start=len(model.counts) + 3):
        book.note(cws, k, sentence)

    lay = _layout(model)
    prefix = _prefix(model)
    matrix_cols = (
        Column(
            key="platform", header=model.matrix_heading or texts.matrix_heading, kind="text_short"
        ),
        Column(key="rows", header=texts.total, kind="count"),
        Column(key="closed", header=texts.closed, kind="count"),
        Column(key="open", header=texts.open, kind="count"),
        *(
            Column(key=f"class_{j}", header=label, kind="count")
            for j, label in enumerate(model.class_labels)
        ),
    )
    book.header(summary, matrix_cols, row=lay.matrix_header, freeze=False, border=True)
    for i, m in enumerate(model.matrix):
        values = {"platform": m.platform, "rows": m.rows, "closed": m.closed, "open": m.open}
        values.update({f"class_{j}": n for j, n in enumerate(m.by_class)})
        book.body(summary, matrix_cols, values, lay.matrix_first + i, "")
    line_cols = matrix_cols[:2]
    for line_row, label in (
        (lay.rows_all, texts.rows_all),
        (lay.rows_account, texts.rows_account.format(account=model.account, prefix=prefix)),
        (lay.rows_other, texts.rows_other.format(account=model.account, prefix=prefix)),
    ):
        book.body(summary, line_cols, {"platform": label, "rows": 0}, line_row, "")
    for f in plan_jobs(model, cfg):
        book.formula(f)
    # the label column and the count columns, then the title and the notes (their heights follow
    # the final widths)
    book.set_width(summary, 1, layout.summary_label_width)
    for j in range(2, layout.summary_span + 1):
        book.set_width(summary, j, layout.summary_count_width)
    book.title_line(summary, lay.title, model.summary_title or model.account)
    for k, text in enumerate(model.summary_notes, start=lay.paragraphs_first):
        book.paragraph(summary, k, text, layout.summary_span)
    book.write_long_text()
    return book.bytes(title, model.as_of), book.counted
