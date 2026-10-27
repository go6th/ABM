"""The deep-dive tables of an account, as spreadsheet sheets: People, Signals, Job postings,
Tech stack and teams, News, From your mail and drive, Sources and Coverage (the 33-item research
tracker of the latest finished run, with each item's status and reason).

Each sheet is a fixed definition: a base query over the account's rows, its columns (each with an
allow-listed sort expression), and its filter chips (each a fixed SQL predicate with a live
count). Sorting, filtering and paging run in SQL; nothing from the request is ever put into SQL
except through these allow-lists and bound parameters. Row numbers are absolute (offset + 1).

"From your mail and drive" holds only the signed-in user's own matched items, and the export of
that sheet is labelled private.
"""

from __future__ import annotations

import json
import re
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import UTC, date, datetime
from functools import cache, lru_cache
from typing import Any
from uuid import UUID

from sqlalchemy import text

from account_intelligence.replica.answers import LANE_LABEL, day_month
from account_intelligence.replica.source_label import source_link
from account_intelligence.replica.workspace import FACT_CLASS, relative
from exports.policy import leak_hits

GRADE_LABEL = {
    "confirmed": "Confirmed",
    "probable": "Probable",
    "hypothesis": "Hypothesis",
    "lead": "Lead",
    "internal": "Internal",
}
GRADE_ORDER = (
    "CASE {} WHEN 'confirmed' THEN 0 WHEN 'probable' THEN 1 WHEN 'hypothesis' THEN 2 ELSE 3 END"
)


@dataclass(frozen=True)
class Col:
    """One column: its key in the row, its heading, its sort expression and how it is shown.
    ``link_text`` is what a link column of a file-shaped sheet shows in place of the address
    (the People List's "Open profile"); the address goes in the row's ``links``."""

    key: str
    label: str
    sort: str | None
    # text, strong, grade, priority, status, date, ago, link, num, url, code, long
    kind: str = "text"
    link_text: str | None = None


@dataclass(frozen=True)
class Filter:
    """One filter chip: a fixed predicate over the sheet's rows."""

    key: str
    label: str
    where: str


@dataclass(frozen=True)
class Sheet:
    """One deep-dive table."""

    key: str
    title: str
    file: str
    scope: str
    base: str
    cols: tuple[Col, ...]
    filters: tuple[Filter, ...]
    default_sort: str
    default_dir: str = "asc"
    private: bool = False
    footnote: str = "Only Confirmed facts should be said to a prospect."
    tiebreak: str = "s._id"
    # Rows that must never leave in a shareable workbook (internal facts from uploaded files).
    export_where: str = "true"
    # A file-shaped sheet (``file_sheet``): the client table of the export file set it shows
    # (``views.tables``), read from the delivered run's materialised rows, and the sheet or
    # table of the file it matches.
    table: str | None = None
    sheet_name: str = "Sheet 1"


SORT_HINT = "Click a column heading to sort."
_MONTH = "date_trunc('month', CAST(:now AS timestamptz))"


def _fact_rows(category: str, columns: str) -> str:
    """The public facts of a category as rows of its sheet (the research graph records what it
    finds as graded facts; see ``workspace.FACT_CLASS``), with each fact's latest evidence."""
    return (
        f"SELECT f.id AS _id, {columns} FROM facts f LEFT JOIN LATERAL (SELECT url, evidence_date"
        " FROM fact_evidence x WHERE x.fact_id = f.id ORDER BY evidence_date DESC NULLS LAST"
        f" LIMIT 1) e ON true WHERE {FACT_CLASS[category]}"
    )


SHEETS: dict[str, Sheet] = {
    "people": Sheet(
        key="people",
        title="People",
        file="People list.xlsx",
        scope=f"Director and above, global, security excluded. {SORT_HINT}",
        base=(
            "SELECT p.id AS _id, coalesce(p.name, '[Name]') AS name, p.title,"
            " coalesce(p.classification_json->>'area', '') AS area, coalesce(p.priority, '') AS"
            " priority,"
            " coalesce(p.classification_json->>'grade', '') AS grade, p.created_at AS seen,"
            " p.profile_url AS link, p.profile_url AS profile,"
            " (p.run_id IS DISTINCT FROM (SELECT r.id FROM account_runs r WHERE r.account_id ="
            " p.account_id"
            "   AND r.status IN ('COMPLETED', 'PARTIAL') ORDER BY r.finished_at LIMIT 1)) AS"
            " later_run"
            " FROM people p WHERE p.account_id = :a AND NOT p.excluded AND p.merged_into IS NULL"
            " AND p.erased_at IS NULL AND NOT coalesce(p.internal_only, false)"
            " UNION ALL "
            + _fact_rows(
                "people",
                "coalesce(f.value_text, f.value_json->>'text') AS name,"
                " coalesce(nullif(f.value_json->>'title', ''),"
                "   replace(split_part(f.attribute_key, '.', 3), '_', ' ')) AS title,"
                " coalesce(f.value_json->>'lane', '') AS area,"
                " coalesce(f.value_json->>'priority', '') AS priority, f.product_grade AS grade,"
                " f.created_at AS seen, coalesce(e.url, f.value_json->>'link') AS link,"
                # the person's own profile: an imported People List row's LinkedIn cell, a
                # profile page the research cited (never the page that merely named them), or
                # the profile the people search found for the same name on this account
                " coalesce(nullif(f.value_json->>'link', ''), nullif(f.value_json->>'profile_url',"
                "   ''), CASE WHEN e.url ~* '^https?://([a-z]+[.])?linkedin[.]com/(in|pub)/'"
                "   THEN e.url END,"
                "   (SELECT p2.profile_url FROM people p2 WHERE p2.account_id = f.account_id"
                "     AND p2.profile_url IS NOT NULL AND NOT p2.excluded AND p2.merged_into IS NULL"
                "     AND p2.erased_at IS NULL AND lower(p2.name) = lower(trim(split_part("
                "       coalesce(f.value_text, f.value_json->>'text', ''), ',', 1)))"
                "     ORDER BY p2.created_at LIMIT 1)) AS profile,"
                " false AS later_run",
            )
        ),
        cols=(
            Col("name", "Name", "s.name", "strong"),
            Col("title", "Title as reported", "s.title"),
            Col("area", "Area", "s.area"),
            Col("priority", "Priority", "s.priority", "priority"),
            Col("profile", "LinkedIn", None, "url", link_text="Open profile"),
            Col("grade", "Fact grade", GRADE_ORDER.format("s.grade"), "grade"),
            Col("seen", "Last seen", "s.seen", "ago"),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("priority_a", "Priority A", "s.priority = 'A'"),
            Filter("technology", "Technology", "s.area = 'Technology'"),
            Filter("commercial", "Commercial", "s.area = 'Commercial'"),
            Filter("changed", "Changed this month", f"s.later_run AND s.seen >= {_MONTH}"),
        ),
        default_sort="priority",
        tiebreak=GRADE_ORDER.format("s.grade") + ", s.title, s._id",
    ),
    "signals": Sheet(
        key="signals",
        title="Signals",
        file="Signals.xlsx",
        scope=f"Hiring, leadership, technology, news and speaking signals. {SORT_HINT}",
        base=(
            "SELECT s.id AS _id, s.title, initcap(s.signal_type) AS type, s.grade, s.detected_at"
            " AS detected,"
            " (SELECT e.url FROM fact_evidence e WHERE e.fact_id = s.fact_id LIMIT 1) AS link,"
            " (s.detected_at > coalesce((SELECT r.finished_at FROM account_runs r WHERE"
            " r.account_id = s.account_id"
            "   AND r.status IN ('COMPLETED', 'PARTIAL') ORDER BY r.finished_at DESC OFFSET 1"
            " LIMIT 1),"
            "   'infinity'::timestamptz)) AS is_new,"
            " coalesce((SELECT f.visibility = 'private' FROM facts f WHERE f.id = s.fact_id),"
            " false) AS _private"
            " FROM signal_events s WHERE s.account_id = :a"
        ),
        cols=(
            Col("title", "Signal", "s.title", "strong"),
            Col("type", "Type", "s.type"),
            Col("grade", "Fact grade", GRADE_ORDER.format("s.grade"), "grade"),
            Col("detected", "Detected", "s.detected", "date"),
            Col("link", "Source", None, "link"),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("new", "New since last refresh", "s.is_new"),
            Filter("hiring", "Hiring", "s.type = 'Hiring'"),
            Filter("leadership", "Leadership", "s.type = 'Leadership'"),
            Filter("news", "News", "s.type = 'News'"),
        ),
        default_sort="detected",
        default_dir="desc",
        export_where="NOT s._private",
    ),
    "jobs": Sheet(
        key="jobs",
        title="Job postings",
        file="Job postings.xlsx",
        scope=f"Open roles from the account's careers pages and job feeds. {SORT_HINT}",
        base=(
            "SELECT j.id AS _id, j.title, array_to_string(j.lanes, ', ') AS lanes,"
            " coalesce(j.location, '') AS location,"
            " coalesce(j.posted_at, j.first_seen) AS posted, initcap(j.status) AS status, j.url AS"
            " link"
            " FROM job_postings j WHERE j.account_id = :a"
            " UNION ALL "
            + _fact_rows(
                "jobs",
                "coalesce(f.value_text, f.value_json->>'text') AS title,"
                " coalesce(f.value_json->>'lanes', '') AS lanes,"
                " coalesce(f.value_json->>'location', '') AS location,"
                " coalesce(CASE WHEN f.value_json->>'posted' ~ '^[0-9]{4}-[0-9]{2}-[0-9]{2}'"
                "   THEN CAST(left(f.value_json->>'posted', 10) AS timestamptz) END,"
                "   CAST(e.evidence_date AS timestamptz), f.created_at) AS posted,"
                " coalesce(nullif(f.value_json->>'status', ''), 'Found') AS status,"
                " coalesce(e.url, f.value_json->>'url') AS link",
            )
        ),
        cols=(
            Col("title", "Role", "s.title", "strong"),
            Col("lanes", "Lane", "s.lanes"),
            Col("location", "Location", "s.location"),
            Col("posted", "Posted", "s.posted", "date"),
            Col("status", "Status", "s.status"),
            Col("link", "Posting", None, "link"),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("month", "Posted this month", f"s.posted >= {_MONTH}"),
            Filter("salesforce", "Salesforce", "s.lanes ILIKE '%salesforce%'"),
            Filter("data", "Data and AI", "s.lanes ILIKE '%data%'"),
            Filter("open", "Open", "s.status = 'Open'"),
        ),
        default_sort="posted",
        default_dir="desc",
    ),
    "technology": Sheet(
        key="technology",
        title="Tech stack and teams",
        file="Tech stack and teams.xlsx",
        scope=f"Technologies detected, with how many sources show each. {SORT_HINT}",
        base=(
            "SELECT t.id AS _id, t.technology, t.source_count AS sources,"
            " CASE t.cross_check_result WHEN 'confirmed' THEN 'Confirmed' WHEN 'contradicted' THEN"
            " 'Contradicted'"
            "   WHEN 'not_found' THEN 'Not cross-checked' ELSE '' END AS check,"
            " coalesce(t.first_seen_at, '') AS first_seen, coalesce(t.last_seen_at, '') AS"
            " last_seen,"
            " (SELECT s.url FROM sources s WHERE s.account_id = t.account_id AND s.source_code ="
            " t.source_ref) AS link"
            " FROM tech_detections t WHERE t.account_id = :a"
            " UNION ALL "
            + _fact_rows(
                "tech",
                "coalesce(f.value_text, f.value_json->>'text') AS technology,"
                " coalesce(CASE WHEN f.value_json->>'sources' ~ '^[0-9]+$'"
                "   THEN CAST(f.value_json->>'sources' AS int) END,"
                "   (SELECT count(*) FROM fact_evidence x WHERE x.fact_id = f.id)) AS sources,"
                " coalesce(f.value_json->>'rank', '') AS check,"
                " coalesce(CAST(e.evidence_date AS text), '') AS first_seen,"
                " coalesce(f.value_json->>'last_seen', CAST(e.evidence_date AS text), '')"
                " AS last_seen, e.url AS link",
            )
        ),
        cols=(
            Col("technology", "Technology", "s.technology", "strong"),
            Col("sources", "Sources", "s.sources", "num"),
            Col("check", "Cross check", "s.check"),
            Col("first_seen", "First seen", "s.first_seen"),
            Col("last_seen", "Last seen", "s.last_seen"),
            Col("link", "Source", None, "link"),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("confirmed", "Cross-checked", "s.check = 'Confirmed'"),
            Filter("multi", "Two or more sources", "s.sources >= 2"),
        ),
        default_sort="technology",
    ),
    "news": Sheet(
        key="news",
        title="News",
        file="News.xlsx",
        scope=f"Announcements and coverage about the account. {SORT_HINT}",
        base=(
            "SELECT f.id AS _id, coalesce(f.value_text, f.value_json->>'text') AS headline,"
            " f.product_grade AS grade,"
            " e.evidence_date AS published, e.url AS link"
            " FROM facts f LEFT JOIN LATERAL (SELECT url, evidence_date FROM fact_evidence x WHERE"
            " x.fact_id = f.id"
            "   ORDER BY evidence_date DESC NULLS LAST LIMIT 1) e ON true"
            f" WHERE {FACT_CLASS['news']}"
        ),
        cols=(
            Col("headline", "Headline", "s.headline", "strong"),
            Col("grade", "Fact grade", GRADE_ORDER.format("s.grade"), "grade"),
            Col("published", "Published", "s.published", "date"),
            Col("link", "Source", None, "link"),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("month", "This month", f"s.published >= {_MONTH}"),
            Filter("confirmed", "Confirmed", "s.grade = 'confirmed'"),
        ),
        default_sort="published",
        default_dir="desc",
    ),
    "mail-and-drive": Sheet(
        key="mail-and-drive",
        title="From your mail and drive",
        file="From your mail and drive.xlsx",
        scope="Your own threads and files that mention this account. Shown only to you.",
        base=(
            "SELECT i.id AS _id, coalesce(i.title, '') AS title, CASE i.item_kind WHEN 'file' THEN"
            " 'File'"
            "   ELSE 'Thread' END AS kind, coalesce(i.location, '') AS location, i.item_date AS"
            " dated,"
            " coalesce(i.snippet, '') AS snippet"
            " FROM connector_items i WHERE i.account_id = :a AND i.owner_user_id = :u AND i.state"
            " = 'matched'"
        ),
        cols=(
            Col("title", "Title", "s.title", "strong"),
            Col("kind", "Kind", "s.kind"),
            Col("location", "From", "s.location"),
            Col("dated", "Date", "s.dated", "date"),
            Col("snippet", "Summary", None),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("threads", "Mail threads", "s.kind = 'Thread'"),
            Filter("files", "Files", "s.kind = 'File'"),
        ),
        default_sort="dated",
        default_dir="desc",
        private=True,
        footnote=(
            "From your mail and drive. Private to you, and never included in the research files"
            " you share."
        ),
    ),
    "sources": Sheet(
        key="sources",
        title="Sources",
        file="Sources.xlsx",
        scope=f"Every page the research read for this account. {SORT_HINT}",
        # graded sources, plus pages a run read that were not graded yet (a run that stopped
        # before grading still shows what it read; audit 2026-10-15)
        base=(
            "SELECT s0.id::text AS _id, s0.source_code AS code, coalesce(s0.title, '') AS title,"
            " coalesce(s0.publisher, '') AS publisher, s0.published_at AS published, s0.url AS"
            " link, false AS ungraded,"
            " EXISTS (SELECT 1 FROM fact_evidence e WHERE e.source_id = s0.id) AS cited"
            " FROM sources s0 WHERE s0.account_id = :a"
            " UNION ALL"
            " SELECT DISTINCT ON (t.response->'args'->>'url') t.id::text AS _id, 'Read' AS code,"
            " '' AS title, '' AS publisher, NULL::timestamptz AS published,"
            " t.response->'args'->>'url' AS link, true AS ungraded, false AS cited"
            " FROM tool_calls t JOIN account_runs r ON r.id = t.run_id"
            " WHERE r.account_id = :a"
            " AND t.tool IN ('p_search:read_page', 'p_scrape:scrape_job_page')"
            " AND t.response->'args'->>'url' IS NOT NULL"
            " AND NOT EXISTS (SELECT 1 FROM sources x WHERE x.account_id = :a"
            "   AND x.url = t.response->'args'->>'url')"
            # the client's own uploaded files that facts cite (documents.deliverables)
            " UNION ALL"
            " SELECT d.id::text AS _id, 'File' AS code, d.name AS title,"
            " 'Uploaded by you' AS publisher, d.created_at AS published, NULL::text AS link,"
            " false AS ungraded, true AS cited FROM documents d WHERE d.account_id = :a"
            " AND EXISTS (SELECT 1 FROM fact_evidence e JOIN facts f ON f.id = e.fact_id"
            "   WHERE e.document_id = d.id AND f.account_id = :a AND f.status = 'current')"
        ),
        cols=(
            Col("code", "Code", "s.code"),
            Col("title", "Title", "s.title", "strong"),
            Col("publisher", "Publisher", "s.publisher"),
            Col("published", "Published", "s.published", "date"),
            Col("link", "Link", None, "link"),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("cited", "Cited in findings", "s.cited"),
            Filter("ungraded", "Read, not graded yet", "s.ungraded"),
            Filter("month", "Published this month", f"s.published >= {_MONTH}"),
        ),
        default_sort="code",
    ),
    "searches": Sheet(
        key="searches",
        title="Web searches",
        file="Web searches.xlsx",
        scope=(
            "Every web search the research made for this account: what it searched for, when,"
            f" how many pages it returned and the sites they came from. {SORT_HINT}"
        ),
        base=(
            "SELECT t.id AS _id,"
            " coalesce((SELECT string_agg(q, ' | ') FROM jsonb_array_elements_text("
            "   CASE WHEN jsonb_typeof(t.response->'args'->'queries') = 'array'"
            "   THEN t.response->'args'->'queries' ELSE '[]'::jsonb END) q),"
            "   t.response->'args'->>'objective', '') AS query,"
            " t.created_at AS searched,"
            " coalesce((t.response->>'record_count')::int, 0) AS results,"
            " coalesce((SELECT string_agg(DISTINCT regexp_replace(p->>'url',"
            "   '^https?://(www[.])?([^/:]+).*$', '\\2'), ', ') FROM jsonb_array_elements("
            "   CASE WHEN jsonb_typeof(t.response->'pages') = 'array'"
            "   THEN t.response->'pages' ELSE '[]'::jsonb END) p), '') AS sites,"
            " (t.cost_source = 'cache') AS reused"
            " FROM tool_calls t JOIN account_runs r ON r.id = t.run_id"
            " WHERE r.account_id = :a AND t.tool = 'p_search:web_search'"
        ),
        cols=(
            Col("query", "Searched for", "s.query", "strong"),
            Col("searched", "When", "s.searched", "date"),
            Col("results", "Pages returned", "s.results", "num"),
            Col("sites", "Sites", "s.sites"),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("month", "This month", f"s.searched >= {_MONTH}"),
            Filter("reused", "Reused from earlier research", "s.reused"),
        ),
        default_sort="searched",
        default_dir="desc",
        footnote=(
            "Sites are listed for searches from 15 October 2026 on; earlier searches show counts."
        ),
    ),
    "coverage": Sheet(
        key="coverage",
        title="Coverage",
        file="Coverage tracker.xlsx",
        scope=(
            "The research tracker of the latest finished run: every item, whether it was covered"
            f" and why not. {SORT_HINT}"
        ),
        base=(
            "SELECT c.row_key AS _id,"
            " coalesce(CAST(:tracker AS jsonb)->'label'->>c.row_key, c.row_key) AS item,"
            " initcap(c.priority) AS priority_label,"
            " CASE c.priority WHEN 'high' THEN 0 WHEN 'medium' THEN 1 ELSE 2 END AS prio,"
            " CASE c.status WHEN 'done' THEN 'Done' WHEN 'partial' THEN 'Partial'"
            "   WHEN 'not_done' THEN 'Not done' WHEN 'blocked' THEN 'Blocked' WHEN 'none' THEN"
            " 'None found'"
            "   ELSE 'Not checked' END AS status,"
            " CASE c.status WHEN 'blocked' THEN 0 WHEN 'not_done' THEN 1 WHEN 'partial' THEN 2"
            "   WHEN 'none' THEN 3 WHEN 'done' THEN 4 ELSE 5 END AS status_order,"
            " coalesce(CAST(:tracker AS jsonb)->'reason'->>c.reason_code, '') AS reason,"
            " coalesce(CAST(CAST(:tracker AS jsonb)->'order'->>c.row_key AS int), 999) AS ord,"
            " c.last_checked_at AS checked"
            " FROM coverage_rows c WHERE c.run_id = (SELECT r.id FROM account_runs r"
            "   WHERE r.account_id = :a AND r.status IN ('COMPLETED', 'PARTIAL', 'BUDGET_EXCEEDED')"
            "   AND EXISTS (SELECT 1 FROM coverage_rows x WHERE x.run_id = r.id)"
            "   ORDER BY r.finished_at DESC NULLS LAST LIMIT 1)"
        ),
        cols=(
            Col("item", "Research item", "s.ord", "strong"),
            Col("priority_label", "Priority", "s.prio"),
            Col("status", "Status", "s.status_order", "status"),
            Col("reason", "Reason", None),
            Col("checked", "Checked", "s.checked", "date"),
        ),
        filters=(
            Filter("all", "All", "true"),
            Filter("done", "Done", "s.status = 'Done'"),
            Filter("partial", "Partial", "s.status = 'Partial'"),
            Filter("open", "Not done or blocked", "s.status IN ('Not done', 'Blocked')"),
            Filter("high", "High priority", "s.prio = 0"),
        ),
        default_sort="item",
        tiebreak="s.ord, s._id",
        footnote=(
            "Each item ends Done, Partial, Not done, Blocked or None found. A source that failed"
            " shows Blocked with its reason; nothing is left out silently."
        ),
    ),
}

MAX_PAGE = 500
MAX_EXPORT = 5000

# ---- file-shaped sheets ------------------------------------------------------------------------
# Once an account has delivered files, its People, Job postings, Tech stack, Platform picture and
# Coverage sheets show the client tables of that file set (``views.tables``: the columns, headers
# and order of the files, from the export configuration; ``views.rowsources``: the cells as the
# files print them), read from the delivered run's materialised rows (``views.materialize``).
# Before any file is delivered (a research still running, an account known only from uploaded
# files) the live sheets above stand in.

# sheet key -> client table (views.tables.TableKind)
FILE_TABLES: dict[str, str] = {
    "people": "people",
    "jobs": "jobs",
    "technology": "tech",
    "facts": "facts",
    "coverage": "coverage",
}
# the order the file prints (each table's rows are materialised in the file's order)
FILE_ORDER = Col("_file", "File order", "s._pos")
_KEY = re.compile(r"^[a-z][a-z0-9_]{0,39}$")
_RESERVED = frozenset({"_id", "_pos", "_cells", "link"})
# how a few columns are shown on screen (a pill, bold); the cell text stays the file's
_SHOWN: dict[tuple[str, str], str] = {
    ("people", "name"): "strong",
    ("people", "priority"): "priority",
    ("jobs", "title"): "strong",
    ("tech", "technology"): "strong",
    ("facts", "item"): "strong",
    ("facts", "grade"): "grade",
    ("coverage", "row"): "strong",
    ("coverage", "status"): "status",
}
# how an export column kind is shown when no entry above says otherwise
_KIND_SHOWN = {"count": "num", "percent": "num", "url": "url", "code": "code", "text_long": "long"}


def _lit(value: str) -> str:
    """A configured word as an SQL string literal (configuration only, never request input)."""
    return "'" + value.replace("'", "''") + "'"


def _heading(text_value: str) -> str:
    """A document heading without its section number ("5. Platform picture by lane")."""
    return re.sub(r"^\d+\.\s*", "", text_value)


@cache
def file_sheet(key: str) -> Sheet:
    """The file-shaped sheet of one sheet key, built from the export configuration (the same
    source the files and ``views.tables`` read, so a column added to a file appears here)."""
    from exports.config import load_export_config
    from views.tables import specs

    table = FILE_TABLES[key]
    spec = specs()[table]  # type: ignore[index]
    rules = load_export_config().rules
    live = SHEETS.get(key)
    cols: list[Col] = []
    for c in spec.exported_columns:
        if not _KEY.match(c.key) or c.key in _RESERVED:
            raise ValueError(f"unusable column key {c.key!r} in table {table}")
        shown = _SHOWN.get((table, c.key), _KIND_SHOWN.get(c.kind, "text"))
        sortable = c.kind not in ("text_long", "url")
        link_text = (
            rules.sheet_texts.people.link_text if (table == "people" and c.kind == "url") else None
        )
        cols.append(
            Col(c.key, c.header, f"s._cells->'{c.key}'" if sortable else None, shown, link_text)
        )
    selected = ", ".join(f"coalesce(v.cells->>'{c.key}', '') AS \"{c.key}\"" for c in cols)
    base = (
        "SELECT v.position AS _id, v.position AS _pos, v.cells AS _cells, v.source_url AS link,"
        f" {selected} FROM view_rows v WHERE v.run_id = CAST(:run AS uuid)"
        f" AND v.table_kind = {_lit(table)}"
    )

    def chip(k: str, label: str, column: str, value: str) -> Filter:
        return Filter(k, label, f's."{column}" = {_lit(value)}')

    chips: list[Filter] = [Filter("all", "All", "true")]
    labels = rules.labels
    if table == "people":
        header = {c.key: c.header for c in spec.exported_columns}
        chips += [
            chip("priority_a", "Priority A", "priority", "A"),
            chip("senior", header["meets_seniority"], "meets_seniority", labels.yes),
        ]
    elif table == "jobs":
        chips += [
            chip(f"status_{k}", word, "status", word) for k, word in labels.posting_status.items()
        ]
    elif table == "facts":
        st = rules.statement_text
        chips += [
            chip("may_state", st.may_state, "use", st.may_state),
            chip("ask", st.ask, "use", st.ask),
        ]
    elif table == "coverage":
        chips += [
            chip(f"status_{i}", word, "status", word)
            for i, word in enumerate(rules.documents.labels.remains, start=1)
        ]
    # ``file`` names the sheet's own export (Open in Excel); ``sheet_name`` the file's sheet, or
    # the Account Summary table, it matches
    if table == "facts":
        title = sheet_name = _heading(rules.headings.summary.lane_tables)
        file_name = f"{title}.xlsx"
    elif table == "coverage":
        assert live is not None
        title, file_name = live.title, live.file
        sheet_name = _heading(rules.headings.summary.coverage_tracker)
    else:
        assert live is not None
        title, file_name, sheet_name = live.title, live.file, spec.sheet or "Sheet 1"
    return Sheet(
        key=key,
        title=title,
        file=file_name,
        scope=f"The table of your delivered files, row for row. {SORT_HINT}",
        base=base,
        cols=tuple(cols),
        filters=tuple(chips),
        default_sort=FILE_ORDER.key,
        tiebreak="s._pos",
        table=table,
        sheet_name=sheet_name,
    )


async def delivered_run(session: Any, account_id: UUID) -> UUID | None:
    """The run of the account's newest delivered file (the file set the sheets match)."""
    found = await session.scalar(
        text(
            "SELECT d.run_id FROM deliverables d WHERE d.account_id = :a AND d.run_id IS NOT NULL"
            " ORDER BY d.built_at DESC LIMIT 1"
        ),
        {"a": account_id},
    )
    return UUID(str(found)) if found else None


# the delivered file each client table is printed in (deliverables.kind)
_FILE_KIND = {
    "people": "people",
    "jobs": "jobs",
    "tech": "tech",
    "facts": "summary",
    "coverage": "summary",
}


async def is_materialised(session: Any, run: UUID) -> bool:
    """Whether the run's client tables were built (``views.materialize``), in any state."""
    return bool(
        await session.scalar(
            text("SELECT true FROM view_materializations WHERE run_id = :r"), {"r": run}
        )
    )


@dataclass(frozen=True)
class _FilePrefetch:
    """What ``_file_table`` reads about a delivered run, fetched once for every file-shaped sheet
    (a screen asking for all five tables paid five times as many queries)."""

    run: UUID
    ready: bool
    manifests: Mapping[str, Any]  # deliverables.kind -> the newest version's manifest
    first: Mapping[str, Any]  # table kind -> the first materialised row's cells
    counts: Mapping[str, int]  # table kind -> materialised rows


async def _prefetch(session: Any, run: UUID) -> _FilePrefetch:
    ready = await session.scalar(
        text("SELECT status = 'ready' FROM view_materializations WHERE run_id = :r"), {"r": run}
    )
    manifests: dict[str, Any] = {}
    first: dict[str, Any] = {}
    counts: dict[str, int] = {}
    if ready:
        for r in (
            await session.execute(
                text(
                    "SELECT DISTINCT ON (kind) kind, manifest_json FROM deliverables"
                    " WHERE run_id = :r ORDER BY kind, version DESC"
                ),
                {"r": run},
            )
        ).all():
            m = r.manifest_json
            manifests[r.kind] = json.loads(m) if isinstance(m, str) else m
        for r in (
            await session.execute(
                text(
                    "SELECT DISTINCT ON (table_kind) table_kind, cells FROM view_rows"
                    " WHERE run_id = :r ORDER BY table_kind, position"
                ),
                {"r": run},
            )
        ).all():
            first[r.table_kind] = r.cells
        counts = {
            r.table_kind: int(r.n)
            for r in (
                await session.execute(
                    text(
                        "SELECT table_kind, count(*) AS n FROM view_rows WHERE run_id = :r"
                        " GROUP BY table_kind"
                    ),
                    {"r": run},
                )
            ).all()
        }
    return _FilePrefetch(run, bool(ready), manifests, first, counts)


def _judge(key: str, pre: _FilePrefetch) -> tuple[UUID, int] | None:
    """(delivered run, rows) of the file-shaped sheet ``key`` from what ``_prefetch`` read."""
    table = FILE_TABLES.get(key)
    if table is None or not pre.ready:
        return None
    manifest = pre.manifests.get(_FILE_KIND[table])
    printed = manifest.get("rows") if isinstance(manifest, dict) else None
    if not isinstance(printed, int):
        return None
    from views.tables import specs

    want = {c.key for c in specs()[table].columns}  # type: ignore[index]
    first = pre.first.get(table)
    if first is not None and set(dict(first)) != want:
        return None
    count = pre.counts.get(table, 0)
    if _FILE_KIND[table] != "summary" and count != printed:
        return None
    return pre.run, count


async def _file_table(session: Any, account_id: UUID, key: str) -> tuple[UUID, int] | None:
    """(delivered run, rows) of the file-shaped sheet ``key``, or None when the live sheet
    stands in. The rows are shown only when they are the delivered file's: the run's
    materialisation is ready and was built with this configuration's columns (one from before a
    column change is passed over rather than shown with blank columns), the file was made by the
    export builder (its stored manifest counts its rows), and a workbook holds exactly the rows
    materialised."""
    table = FILE_TABLES.get(key)
    run = await delivered_run(session, account_id) if table else None
    if table is None or run is None:
        return None
    return _judge(key, await _prefetch(session, run))


async def file_run(session: Any, account_id: UUID, key: str) -> UUID | None:
    """The delivered run whose materialised rows the sheet ``key`` shows (``_file_table``)."""
    found = await _file_table(session, account_id, key)
    return found[0] if found else None


async def file_counts(session: Any, account_id: UUID, run: UUID | None = None) -> dict[str, int]:
    """The rows of each file-shaped sheet with a delivered run, by sheet key (``run``: the
    account's delivered run when the caller already read it)."""
    run = run if run is not None else await delivered_run(session, account_id)
    if run is None:
        return {}
    pre = await _prefetch(session, run)
    out: dict[str, int] = {}
    for key in FILE_TABLES:
        found = _judge(key, pre)
        if found is not None:
            out[key] = found[1]
    return out


class SheetNotFound(LookupError):
    """No such sheet."""


class BadSheetQuery(ValueError):
    """An unknown sort column, filter or direction."""


@dataclass(frozen=True)
class Query:
    """A validated sheet request. ``run_id`` is the delivered run a file-shaped sheet reads."""

    sheet: Sheet
    filter: Filter
    sort: Col
    desc: bool
    offset: int
    limit: int
    run_id: UUID | None = None


def parse(
    sheet_key: str,
    filter_key: str | None,
    sort_key: str | None,
    direction: str | None,
    offset: int = 0,
    limit: int = 200,
    *,
    run_id: UUID | None = None,
) -> Query:
    """Validate a request against the sheet's allow-lists. With ``run_id`` (the account's
    delivered run, ``file_run``) a sheet the files print is the file-shaped one; the Platform
    picture has no live form, so it is always file-shaped (and empty before a delivery)."""
    if sheet_key in FILE_TABLES and (run_id is not None or sheet_key not in SHEETS):
        sheet: Sheet | None = file_sheet(sheet_key)
    else:
        sheet = SHEETS.get(sheet_key)
    if sheet is None:
        raise SheetNotFound(sheet_key)
    filt = next((f for f in sheet.filters if f.key == (filter_key or "all")), None)
    cols = (*sheet.cols, FILE_ORDER) if sheet.table else sheet.cols
    col = next((c for c in cols if c.key == (sort_key or sheet.default_sort) and c.sort), None)
    if filt is None or col is None or (direction or sheet.default_dir) not in ("asc", "desc"):
        raise BadSheetQuery("unknown filter, sort or direction")
    return Query(
        sheet,
        filt,
        col,
        (direction or sheet.default_dir) == "desc",
        max(0, offset),
        max(1, min(limit, MAX_PAGE)),
        run_id if sheet.table else None,
    )


def params(q: Query, account_id: UUID, user_id: str, now: datetime) -> dict[str, Any]:
    """The bound values of the sheet's SQL."""
    out: dict[str, Any] = {"a": account_id, "u": user_id, "now": now}
    if q.sheet.key == "coverage" and q.sheet.table is None:
        out["tracker"] = tracker_labels()
    if q.sheet.table is not None:
        out["run"] = str(q.run_id) if q.run_id else None
    return out


def row_out(q: Query, n: int, r: Any, now: datetime) -> dict[str, Any]:
    """One row as the sheet shows it: its number, cells, source link and, on a file-shaped
    sheet, the address of each link column (the cell shows the file's link text)."""
    if q.sheet.table is None:
        out_cells: dict[str, str] = {}
        out_links: dict[str, str] = {}
        for c in q.sheet.cols:
            if c.kind != "url":
                out_cells[c.key] = display(c, r.get(c.key), now)
                continue
            # a profile link: the cell reads "Open profile" and opens the page itself; a data
            # provider's own page is never shown (``source_label``)
            found = source_link(r.get(c.key))
            href = found.href if found else None
            out_cells[c.key] = (c.link_text or "Open") if href else ""
            if href:
                out_links[c.key] = href
        return {"n": n, "cells": out_cells, "links": out_links, **link_fields(r.get("link"))}
    from views.render_safe import safe_link

    cells: dict[str, str] = {}
    links: dict[str, str] = {}
    for c in q.sheet.cols:
        value = "" if r.get(c.key) is None else str(r.get(c.key))
        href = safe_link(value) if c.kind == "url" else None
        if href:
            links[c.key] = href
            value = c.link_text or value
        cells[c.key] = value
    return {"n": n, "cells": cells, "links": links, **link_fields(r.get("link"))}


def link_fields(raw: Any) -> dict[str, str | None]:
    """A row's source link as the client sees it (``source_label``): ``link`` opens the original
    page, ``link_label`` reads "6th / LinkedIn"; a data provider's own page has no link."""
    found = source_link(raw)
    return {"link": found.href if found else None, "link_label": found.label if found else None}


def display(col: Col, value: Any, now: datetime) -> str:
    """The cell as the sheet shows it (and as the export writes it)."""
    if value is None:
        return ""
    if col.kind == "grade":
        return GRADE_LABEL.get(str(value), str(value).capitalize())
    if col.kind == "ago" and isinstance(value, datetime):
        return relative(value, now)
    if col.kind == "date" and isinstance(value, datetime | date):
        return f"{day_month(value)} {value.year}"
    if isinstance(value, list):
        return ", ".join(str(LANE_LABEL.get(v, v)) for v in value)
    return str(value)


@lru_cache(maxsize=1)
def tracker_labels() -> str:
    """The tracker's item labels, reason wording and checklist order, as the JSON the Coverage
    sheet binds (the labels live in configuration, not in the database)."""
    from account_intelligence.tracker.config import load_checklist, load_tracker_config

    rows = load_checklist().rows
    return json.dumps(
        {
            "label": {r.key: r.label for r in rows},
            "reason": load_tracker_config().reason_codes,
            "order": {r.key: i for i, r in enumerate(sorted(rows, key=lambda r: r.pass_no))},
        }
    )


async def read(
    session: Any, q: Query, account_id: UUID, user_id: str, now: datetime, *, export: bool = False
) -> dict[str, Any]:
    """One page of the sheet (or the whole filtered sheet for an export) and the chip counts."""
    bound = params(q, account_id, user_id, now)
    counts_sql = ", ".join(
        f"count(*) FILTER (WHERE {f.where}) AS c{i}" for i, f in enumerate(q.sheet.filters)
    )
    counts_query = f"SELECT {counts_sql} FROM ({q.sheet.base}) s"  # nosec B608 -- allow-listed sheet SQL; values are bound.
    counts = (await session.execute(text(counts_query), bound)).one()
    order = f"{q.sort.sort} {'DESC' if q.desc else 'ASC'} NULLS LAST, {q.sheet.tiebreak}"
    limit = MAX_EXPORT if export else q.limit
    offset = 0 if export else q.offset
    where = f"({q.filter.where}) AND ({q.sheet.export_where})" if export else q.filter.where
    page_query = f"SELECT * FROM ({q.sheet.base}) s WHERE {where} ORDER BY {order}"  # nosec B608 -- allow-listed sheet SQL; values are bound.
    rows = (
        (
            await session.execute(
                text(page_query + " OFFSET :o LIMIT :l"),
                {**bound, "o": offset, "l": limit},
            )
        )
        .mappings()
        .all()
    )
    total = int(counts[[f.key for f in q.sheet.filters].index(q.filter.key)])
    return {
        "key": q.sheet.key,
        "title": q.sheet.title,
        "file": q.sheet.file,
        "scope": q.sheet.scope,
        "footnote": q.sheet.footnote,
        "private": q.sheet.private,
        "sheet_name": q.sheet.sheet_name,
        "from_file": q.sheet.table is not None,
        "columns": [
            {"key": c.key, "label": c.label, "sortable": c.sort is not None, "kind": c.kind}
            for c in q.sheet.cols
        ],
        "filters": [
            {"key": f.key, "label": f.label, "count": int(counts[i])}
            for i, f in enumerate(q.sheet.filters)
        ],
        "filter": q.filter.key,
        "sort": q.sort.key,
        "dir": "desc" if q.desc else "asc",
        "total": total,
        "offset": offset,
        "rows": [row_out(q, offset + i + 1, r, now) for i, r in enumerate(rows)],
    }


class ExportBlocked(Exception):
    """The finished workbook still failed a client-file check; nothing is sent."""

    def __init__(self, checks: tuple[str, ...]) -> None:
        super().__init__("export blocked: " + ", ".join(checks))
        self.checks = checks


WITHHELD = "Withheld for review"
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_DASH = "\u2012\u2013\u2014\u2015"
_SPACED_DASH = re.compile(f"(?<=\\S)\\s+(?:[{_DASH}]|-{{1,2}})\\s+(?=\\S)")
_ANY_DASH = re.compile(f"[{_DASH}]")
# The pre-delivery gate's rule for client files (documents.gate): dash characters, or a hyphen
# used as punctuation between words.
_GATE_DASHES = re.compile(f"[{_DASH}]|\\S\\s+-{{1,2}}\\s+\\S")


def safe_cell(value: Any, *, shareable: bool) -> tuple[Any, bool]:
    """One cell made safe for a client file. Returns (value, withheld).

    Dashes used as punctuation become commas and dash characters become hyphens (the format
    rule); in a shareable workbook tool names are scrubbed, provider domains and remaining
    guarded terms are withheld; text that a spreadsheet would run as a formula (it starts with
    ``=``, ``+``, ``-``, ``@``, a tab or a return) gets a leading apostrophe."""
    if not isinstance(value, str):
        return value, False
    text_value = _ANY_DASH.sub("-", _SPACED_DASH.sub(", ", value))
    if shareable:
        from agents.naming import third_party_domain_hits, tool_name_hits
        from agents.skill_context import strip_tool_names

        if third_party_domain_hits(text_value):
            return WITHHELD, True
        text_value = strip_tool_names(text_value)
        if tool_name_hits(text_value) or leak_hits(text_value):
            return WITHHELD, True
    if text_value.startswith(_FORMULA_START):
        text_value = "'" + text_value
    return text_value, False


def workbook(sheet: dict[str, Any], account_name: str) -> bytes:
    """The filtered sheet as a real Excel workbook: one sheet, a frozen header, links as text.

    It passes the same checks as every other client file before it leaves: no dash used as
    punctuation, no data tool name, credit count or internal reference (those cells are withheld
    and counted on the About sheet), and no cell a spreadsheet would run as a formula; every
    text cell is stored as text. The finished file is scanned once more as a whole
    (``exports.verify.scan_package``) and refused with ``ExportBlocked`` if anything is left.
    The private sheet (the user's own mail and drive) is labelled private on both sheets and
    is never scanned as shareable: it is exported only to the user it belongs to."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from exports.excel.common import text_cell

    shareable = not sheet["private"]
    withheld = 0

    def put(ws: Any, row: int, values: list[Any]) -> None:
        nonlocal withheld
        for col, raw in enumerate(values, start=1):
            value, held = safe_cell(raw, shareable=shareable)
            withheld += held
            cell = ws.cell(row=row, column=col)
            if isinstance(value, str):
                text_cell(cell, value)
            else:
                cell.value = value

    book = Workbook()
    ws = book.create_sheet("Private to you" if not shareable else "Sheet 1", 0)
    book.remove(book.worksheets[1])
    cols = sheet["columns"]
    put(ws, 1, ["#"] + [c["label"] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    link_font = Font(color="0563C1", underline="single")
    for i, row in enumerate(sheet["rows"], start=2):
        # a link cell reads "6th / LinkedIn" and opens the original page (source_label)
        put(
            ws,
            i,
            [row["n"]]
            + [
                row.get("link_label") or "" if c["kind"] == "link" else row["cells"].get(c["key"])
                for c in cols
            ],
        )
        for col_no, c in enumerate(cols, start=2):
            # a file-shaped sheet's link column shows the file's text and opens its address
            href = (
                row.get("link")
                if c["kind"] == "link"
                else (row.get("links") or {}).get(c["key"])
                if c["kind"] == "url"
                else None
            )
            if href:
                cell = ws.cell(row=i, column=col_no)
                cell.hyperlink = href
                cell.font = link_font
    ws.freeze_panes = "B2"
    for i, c in enumerate(cols, start=2):
        width = max(
            [len(c["label"])]
            + [
                len(str(r["cells"].get(c["key"]) or r.get("link_label") or ""))
                for r in sheet["rows"]
            ]
        )
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    ws.column_dimensions["A"].width = 6
    meta = book.create_sheet("About")
    about: list[list[Any]] = [
        ["Account", account_name],
        ["Table", sheet["title"]],
        ["Filter", next(f["label"] for f in sheet["filters"] if f["key"] == sheet["filter"])],
        ["Rows", len(sheet["rows"])],
        ["Exported", datetime.now(UTC).strftime("%d %b %Y %H:%M UTC")],
    ]
    if withheld:
        about.append(["Withheld", f"{withheld} cells held back for review"])
    if not shareable:
        about.append(["Note", "Private to you. Do not share."])
    for r, line in enumerate(about, start=1):
        put(meta, r, line)
    book.properties.creator = "6th agents"
    book.properties.title = "Private to you" if not shareable else sheet["title"]
    out = io.BytesIO()
    book.save(out)
    data = out.getvalue()
    if shareable:
        check_finished(data)
    return data


def check_finished(data: bytes) -> None:
    """The last check of a shareable workbook as a whole: no client term, no dash used as
    punctuation (``ExportBlocked``)."""
    from agents.naming import third_party_domain_hits, tool_name_hits
    from documents.gate import client_text
    from exports.verify import scan_package

    body = client_text(data)
    failed = tuple(f"client_terms:{h.split(': ')[-1]}" for h in scan_package(data, ()))
    if third_party_domain_hits(body):
        failed += ("provider_host",)
    if tool_name_hits(body):
        failed += ("tool_name",)
    if _GATE_DASHES.search(body):
        failed += ("client_dashes",)
    if failed:
        raise ExportBlocked(failed)
