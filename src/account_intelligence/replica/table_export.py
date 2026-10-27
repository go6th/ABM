"""Whole tables out (owner decision 2026-10-14, Phase 4B section 6): every row of a deep-dive
table under the chosen filter and sort, as CSV or as an Excel workbook, never cut at a page or an
export cap.

* Rows are read page by page (``PAGE`` at a time), so a table of any length is never held whole.
* Every cell passes the client-file rules on its way out (``sheets.safe_cell``): a dash used as
  punctuation becomes a comma, tool names are scrubbed, provider domains and remaining guarded
  terms are withheld in a shareable table (counted on the About sheet, and in the CSV's last
  line), and text a spreadsheet would run as a formula gets a leading apostrophe. The private
  table ("From your mail and drive") is the user's own and says so.
* A shareable workbook is checked once more as a whole before it leaves
  (``sheets.check_finished``); anything left blocks the export.
* Excel holds 1,048,576 rows per sheet: a longer table continues on "Sheet 2", "Sheet 3", each
  with the header row. The workbook is written in streaming mode to a temporary file.

GUARANTEE(test_a_table_exports_whole_past_the_page_and_excel_limits)
"""

from __future__ import annotations

import asyncio
import csv
import io
from collections.abc import AsyncIterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID

from sqlalchemy import text

from account_intelligence.replica import sheets as SH

PAGE = 2000
EXCEL_ROWS = 1_048_576  # rows per worksheet, the header row included


async def rows(
    session: Any, q: SH.Query, account_id: UUID, user_id: str, now: datetime, page: int = PAGE
) -> AsyncIterator[dict[str, Any]]:
    """Every row of the filtered, sorted table, as the sheet shows it (``n``, cells, link)."""
    params = SH.params(q, account_id, user_id, now)
    order = f"{q.sort.sort} {'DESC' if q.desc else 'ASC'} NULLS LAST, {q.sheet.tiebreak}"
    where = f"({q.filter.where}) AND ({q.sheet.export_where})"
    sql = f"SELECT * FROM ({q.sheet.base}) s WHERE {where} ORDER BY {order} OFFSET :o LIMIT :l"  # nosec B608 -- allow-listed sheet SQL; values are bound.
    n = 0
    # LOOP(table export pages): cap=the table's rows / page pages; exit=a short page
    while True:
        got = (await session.execute(text(sql), {**params, "o": n, "l": page})).mappings().all()
        for r in got:
            n += 1
            yield SH.row_out(q, n, r, now)
        if len(got) < page:
            return


def _values(q: SH.Query, row: dict[str, Any], *, excel: bool = False) -> list[Any]:
    """The row's cells in column order. A link cell is the page's address in a CSV; in Excel it
    reads "6th / LinkedIn" and opens the page (``_link_cell``). A file-shaped sheet's link
    column is its address in a CSV and the file's link text in Excel."""
    links = row.get("links") or {}
    return [row["n"]] + [
        (row.get("link_label") or "" if excel else row["link"])
        if c.kind == "link"
        else (row["cells"][c.key] if excel else links.get(c.key, row["cells"][c.key]))
        if c.kind == "url"
        else row["cells"][c.key]
        for c in q.sheet.cols
    ]


def _link_cell(ws: Any, label: Any, href: str | None) -> Any:
    """A write-only cell that shows ``label`` and opens ``href``."""
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font

    cell = WriteOnlyCell(ws, value=label)
    cell.data_type = "s"
    if href:
        cell.hyperlink = href
        cell.font = Font(color="0563C1", underline="single")
    return cell


def header(q: SH.Query) -> list[str]:
    """The header row."""
    shareable = not q.sheet.private
    return [
        SH.safe_cell(value, shareable=shareable)[0]
        for value in ("#", *(c.label for c in q.sheet.cols))
    ]


async def csv_chunks(q: SH.Query, source: AsyncIterator[dict[str, Any]]) -> AsyncIterator[bytes]:
    """The table as CSV (UTF-8 with a byte order mark, so Excel reads it), a few rows a chunk."""
    shareable = not q.sheet.private
    buffer = io.StringIO()
    out = csv.writer(buffer, lineterminator="\r\n")
    out.writerow(header(q))
    withheld = 0
    yield ("﻿" + buffer.getvalue()).encode("utf-8")
    buffer.seek(0)
    buffer.truncate()
    count = 0
    async for row in source:
        cells = []
        for raw in _values(q, row):
            value, held = SH.safe_cell(raw, shareable=shareable)
            withheld += held
            cells.append("" if value is None else value)
        out.writerow(cells)
        count += 1
        if count % 500 == 0:
            yield buffer.getvalue().encode("utf-8")
            buffer.seek(0)
            buffer.truncate()
    if withheld:
        out.writerow([f"{withheld} cells withheld for review"])
    if q.sheet.private:
        out.writerow(["Private to you. Do not share."])
    yield buffer.getvalue().encode("utf-8")


async def write_xlsx(
    path: Path, q: SH.Query, account_name: str, source: AsyncIterator[dict[str, Any]]
) -> int:
    """The table as an Excel workbook at ``path`` (streaming); the rows written."""
    from openpyxl import Workbook

    shareable = not q.sheet.private
    book = Workbook(write_only=True)
    book.properties.creator = "6th agents"
    book.properties.title = "Private to you" if not shareable else q.sheet.title
    sheet_no, written, in_sheet, withheld = 0, 0, EXCEL_ROWS, 0
    ws: Any = None
    async for row in source:
        if in_sheet >= EXCEL_ROWS:
            sheet_no += 1
            name = "Private to you" if not shareable else "Sheet"
            ws = book.create_sheet(f"{name} {sheet_no}" if sheet_no > 1 or shareable else name)
            ws.append(header(q))
            in_sheet = 1
        values: list[Any] = []
        cols = (None, *q.sheet.cols)
        links = row.get("links") or {}
        for raw, col in zip(_values(q, row, excel=True), cols, strict=True):
            value, held = SH.safe_cell(raw, shareable=shareable)
            withheld += held
            if col is not None and col.kind == "link":
                values.append(_link_cell(ws, value, row["link"]))
            elif col is not None and col.kind == "url" and links.get(col.key):
                values.append(_link_cell(ws, value, links[col.key]))
            else:
                values.append(value)
        ws.append(values)
        in_sheet += 1
        written += 1
    if ws is None:
        ws = book.create_sheet("Sheet 1")
        ws.append(header(q))
    about = book.create_sheet("About")
    lines: list[list[Any]] = [
        ["Account", account_name],
        ["Table", q.sheet.title],
        ["Filter", q.filter.label],
        ["Rows", written],
        ["Exported", datetime.now(UTC).strftime("%d %b %Y %H:%M UTC")],
    ]
    if withheld:
        lines.append(["Withheld", f"{withheld} cells held back for review"])
    if not shareable:
        lines.append(["Note", "Private to you. Do not share."])
    for line in lines:
        about.append([SH.safe_cell(v, shareable=shareable)[0] for v in line])
    # saving zips every sheet and the check reads the whole file: both run off the event loop, so
    # a large table never stalls the worker's other requests (or its liveness probe)
    await asyncio.to_thread(book.save, str(path))
    if shareable:
        # the whole file once more (ExportBlocked)
        await asyncio.to_thread(lambda: SH.check_finished(path.read_bytes()))
    return written
