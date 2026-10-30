"""The delivered workbooks' own sheets, as tabs of the deep-dive tables (owner 2026-10-27).

The owner's reference workspace (its file set) has several sheets per workbook: People List
(Summary, People), Job Postings (Summary, Jobs, Counts, Long text), Tech Stack and Teams
(Platform Ranking, Stack detections, Platform teams). The web tables showed only each file's
main sheet, so the Summary matrix, the Platform Ranking and the teams were not on screen. Here
the delivered workbook itself is read back (the bytes the client downloads), sheet by sheet, in
the file's order, and each sheet becomes a tab:

* a sheet is a list of blocks in row order: text lines (the title, notes and paragraphs) and
  tables (a header row the builder styled ``Export header``, then its rows down to the next
  blank row); a sheet from elsewhere with no styled header takes its first row of two or more
  cells as the header;
* a formula cell shows the value the formula computes (``workbook_formulas``), or the value the
  file cached when the function is one the evaluator does not have (a re-saved file); a cell's
  link opens its address and shows the file's text ("Open profile", "Open posting"); a cell's
  solid fill is carried as a theme token (``Cell.fill``: ``priority_a``, ``rank_b``, ``new_row``,
  the hex values from ``config/export_theme.yaml`` ``excel.fills``) for the screen's classes;
* the main sheet (``workbooks.<kind>.sheet``, else the sheet with the largest table) is marked,
  so the screen keeps its sortable, filterable table there and shows the other sheets as they
  are. Once a workbook is delivered its main table is what the screen shows, by the file's own
  header (``builder_main`` lets the header gain columns, ``main_table``), never the live
  research table with other columns.

Reading is bounded: the stored file (``MAX_BYTES``), its unpacked sheets (``MAX_XML_BYTES``,
checked before the workbook is opened), its cells (``MAX_CELLS``), the rows of one table
(``MAX_TABLE_ROWS``) and the workbooks read at once (``READS_AT_ONCE``, in a worker thread). A
read book is cached by tenant, object and SHA-256, after the object is checked to be the
tenant's, and two requests for the same uncached book read it once; a tab switch does not read
the workbook again. A tab's response is capped (``SCREEN_ROWS`` per table, ``SCREEN_CELLS`` and
``SCREEN_CHARS`` per tab). The cells passed every client-file check when the file was built; the
API scrubs the response once more like every sheet, and an export of the file's table passes
the export checks (``table_export``).

GUARANTEE(test_the_web_tabs_are_the_delivered_workbooks_sheets)
GUARANTEE(test_the_book_cache_is_the_tenants_and_reads_a_book_once)
GUARANTEE(test_an_oversized_workbook_is_refused_before_it_is_opened)
GUARANTEE(test_the_file_table_sorts_blanks_last_and_pages_past_the_end)
GUARANTEE(test_a_tab_response_is_capped)
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
import re
import threading
import zipfile
from collections import OrderedDict
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache
from typing import Any
from uuid import UUID

from sqlalchemy import text

from account_intelligence.replica.workbook_formulas import Evaluator

MAX_BYTES = 80 * 1024 * 1024
# the unpacked parts of the package (sheets, shared text, styles): openpyxl holds every cell of
# a workbook it opens, so this, not the stored size, bounds the memory a read takes
MAX_XML_BYTES = 64 * 1024 * 1024
MAX_CELLS = 2_000_000
MAX_TABLE_ROWS = 20_000
CACHE_BOOKS = 16
CACHE_COUNTS = 4096
READS_AT_ONCE = 2
HEADER_STYLE = "Export header"
# sheet key -> the deliverable that prints it (deliverables.kind) and its workbook settings
KINDS = {"people": "people", "jobs": "jobs", "technology": "tech"}
_SLUG = re.compile(r"[^a-z0-9]+")


@dataclass(frozen=True)
class Cell:
    """One cell as the file shows it, the address its link opens (None: no link) and the theme
    token of its solid fill (None: no fill, or a fill the theme does not name)."""

    text: str
    href: str | None = None
    fill: str | None = None


@dataclass(frozen=True)
class Block:
    """A text line (``kind="text"``) or a table (``kind="table"``) of one sheet."""

    kind: str
    text: str = ""
    header: tuple[str, ...] = ()
    rows: tuple[tuple[Cell, ...], ...] = ()
    cut: int = 0  # rows past MAX_TABLE_ROWS left out (the file holds them)


@dataclass(frozen=True)
class Tab:
    """One sheet of the workbook."""

    key: str
    name: str
    blocks: tuple[Block, ...]
    main: bool = False

    @property
    def table(self) -> Block | None:
        """The sheet's largest table (the main sheet's data table)."""
        tables = [b for b in self.blocks if b.kind == "table"]
        return max(tables, key=lambda b: len(b.rows)) if tables else None


@dataclass(frozen=True)
class Book:
    """A workbook read back: its tabs in the file's order."""

    tabs: tuple[Tab, ...]

    def tab(self, key: str | None) -> Tab | None:
        """The tab with this key; the main tab when ``key`` is None."""
        if key is None:
            return next((t for t in self.tabs if t.main), None)
        return next((t for t in self.tabs if t.key == key), None)


@dataclass(frozen=True)
class Delivered:
    """The delivered file a sheet's tabs come from."""

    id: str
    filename: str
    sha256: str
    object_key: str
    # the data rows the builder counted into the file's manifest (None: not a builder's file)
    rows: int | None = None
    size: int | None = None  # the stored bytes (deliverables.size_bytes)


class TooLarge(ValueError):
    """A workbook past the read bounds (``MAX_XML_BYTES``, ``MAX_CELLS``)."""


def tab_key(name: str) -> str:
    """``Platform Ranking`` -> ``platform-ranking`` (unique within a workbook: see ``read``)."""
    return _SLUG.sub("-", name.lower()).strip("-")[:40].strip("-") or "sheet"


def _text(cell: Any, value: Any) -> str:
    """A value as the file shows it: dates as ISO dates, a percent format as a percentage,
    whole numbers without a decimal point, booleans as Yes or No."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int | float):
        fmt = str(getattr(cell, "number_format", "") or "")
        if "%" in fmt:
            decimals = len(fmt.split(".", 1)[1].split("%", 1)[0]) if "." in fmt else 0
            return f"{value * 100:.{decimals}f}%"
        if float(value).is_integer():
            return str(int(value))
        return f"{value:.4f}".rstrip("0").rstrip(".")
    return str(value)


def _href(cell: Any) -> str | None:
    from views.render_safe import safe_link

    link = getattr(cell, "hyperlink", None)
    target = getattr(link, "target", None) if link is not None else None
    return safe_link(target) if isinstance(target, str) else None


# The fills the screen knows by name: the theme (``excel.fills``) gives each its hex value, so no
# colour value lives here. A fill the theme does not name is dropped.
FILL_TOKENS = ("priority_a", "priority_b", "priority_c", "rank_a", "rank_b", "rank_c", "new_row")


@lru_cache(maxsize=1)
def fill_tokens() -> dict[str, str]:
    """RGB hex (six upper-case digits) -> fill token, from the export theme's ``excel.fills``
    (``config/export_theme.yaml``, validated by ``exports.config``); empty when the theme names
    none or cannot be loaded."""
    from exports.config import load_export_config
    from exports.errors import ConfigError

    try:
        fills = load_export_config().theme.excel.fills
    except ConfigError:
        return {}
    out: dict[str, str] = {}
    for token in FILL_TOKENS:
        value = fills.get(token)
        if isinstance(value, str) and re.fullmatch(r"[0-9A-Fa-f]{6,8}", value.strip()):
            out[value.strip().upper()[-6:]] = token
    return out


def _fill(cell: Any) -> str | None:
    """The theme token of a cell's solid fill, else None."""
    fill = getattr(cell, "fill", None)
    if fill is None or getattr(fill, "fill_type", None) != "solid":
        return None
    rgb = getattr(getattr(fill, "start_color", None), "rgb", None)
    if not isinstance(rgb, str) or len(rgb) < 6:
        return None
    return fill_tokens().get(rgb.upper()[-6:])


def read(data: bytes, main_sheet: str | None) -> Book:
    """Every sheet of the workbook as a tab (see the module docstring); ``TooLarge`` past the
    read bounds."""
    from openpyxl import load_workbook

    _check_package(data)
    wb = load_workbook(io.BytesIO(data), data_only=False)
    grid: dict[tuple[str, int, int], Any] = {}
    formulas: list[tuple[str, int, int]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    grid[(ws.title, c.row, c.column)] = c.value
                    if isinstance(c.value, str) and c.value.startswith("="):
                        formulas.append((ws.title, c.row, c.column))
        if len(grid) > MAX_CELLS:
            raise TooLarge(f"more than {MAX_CELLS} cells")
    ev = Evaluator(grid, [ws.title for ws in wb.worksheets])
    values = _Values(ev, _cached_values(data, formulas) if formulas else {})
    tabs: list[Tab] = []
    seen: set[str] = set()
    for ws in wb.worksheets:
        key = tab_key(ws.title)
        n = 2
        # LOOP(tab key suffix): cap=the workbook's sheets; exit=an unused key
        while key in seen:
            key, n = f"{tab_key(ws.title)}-{n}", n + 1
        seen.add(key)
        tabs.append(Tab(key, ws.title, _blocks(ws, values), main=ws.title == main_sheet))
    if tabs and not any(t.main for t in tabs):
        # the configured sheet name is not in the file: the sheet with the largest table is the
        # data sheet (a Summary of a few rows is never the main table)
        sized = [(len(t.table.rows) if t.table is not None else -1, -i) for i, t in enumerate(tabs)]
        i = -max(sized)[1]
        tabs[i] = Tab(tabs[i].key, tabs[i].name, tabs[i].blocks, main=True)
    return Book(tuple(tabs))


def _cached_values(data: bytes, formulas: list[tuple[str, int, int]]) -> dict[Any, Any]:
    """The values the file cached for its formula cells (a workbook re-saved by Excel or
    LibreOffice carries them; the builder's own files do not), read from the sheets' own XML
    (``<c><f>..</f><v>..</v></c>``) of the sheets that hold formulas: a second workbook load
    would parse every shared string again (13 MB on the reference Jobs file). An error value
    (``#DIV/0!``) is not a value; a text result is inline, never a shared string."""
    from lxml import etree

    wanted: dict[str, set[tuple[int, int]]] = {}
    for title, r, c in formulas:
        wanted.setdefault(title, set()).add((r, c))
    out: dict[Any, Any] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            paths = _sheet_paths(package)
            for title, cells in wanted.items():
                path = paths.get(title)
                if path is None or path not in package.namelist():
                    continue
                with package.open(path) as part:
                    _read_cached(part, title, cells, out)
    except (zipfile.BadZipFile, KeyError, ValueError, etree.LxmlError):
        return {}
    return out


_SS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_REF = re.compile(r"^([A-Z]{1,3})(\d+)$")


def _sheet_paths(package: zipfile.ZipFile) -> dict[str, str]:
    """Sheet title -> the package path of its XML (workbook.xml and its relationships)."""
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)
    rels = etree.fromstring(package.read("xl/_rels/workbook.xml.rels"), parser=parser)
    targets = {
        str(r.get("Id")): str(r.get("Target") or "")
        for r in rels.iter(f"{{{_PKG_REL}}}Relationship")
    }
    book = etree.fromstring(package.read("xl/workbook.xml"), parser=parser)
    out: dict[str, str] = {}
    for sheet in book.iter(f"{{{_SS}}}sheet"):
        target = targets.get(str(sheet.get(f"{{{_REL}}}id")), "")
        if not target:
            continue
        path = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
        out[str(sheet.get("name"))] = path
    return out


def _read_cached(part: Any, title: str, cells: set[tuple[int, int]], out: dict[Any, Any]) -> None:
    """The cached values of the wanted formula cells of one sheet part, streamed."""
    from lxml import etree
    from openpyxl.utils.cell import column_index_from_string

    for _, c in etree.iterparse(
        part, events=("end",), tag=f"{{{_SS}}}c", resolve_entities=False, no_network=True,
        huge_tree=True,
    ):  # fmt: skip
        ref = _REF.match(str(c.get("r") or ""))
        if ref is not None:
            key = (int(ref.group(2)), column_index_from_string(ref.group(1)))
            if key in cells and c.find(f"{{{_SS}}}f") is not None:
                v = c.find(f"{{{_SS}}}v")
                kind = c.get("t")
                if v is not None and v.text is not None and kind not in ("e", "s"):
                    text = str(v.text)
                    if kind == "str" or kind == "inlineStr":
                        out[(title, *key)] = text
                    elif kind == "b":
                        out[(title, *key)] = text.strip() == "1"
                    else:
                        try:
                            number = float(text)
                        except ValueError:
                            continue
                        out[(title, *key)] = int(number) if number.is_integer() else number
        c.clear()


class _Values:
    """A cell's value for the screen: the evaluator's, else the file's cached value."""

    def __init__(self, ev: Evaluator, cached: dict[Any, Any]) -> None:
        self.ev = ev
        self.cached = cached

    def value(self, sheet: str, row: int, col: int) -> Any:
        got = self.ev.value(sheet, row, col)
        if got is None and self.cached:
            return self.cached.get((sheet, row, col))
        return got


def _check_package(data: bytes) -> None:
    """``TooLarge`` when the package unpacks past ``MAX_XML_BYTES`` (read from the zip's
    directory, before anything is unpacked)."""
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        unpacked = sum(i.file_size for i in package.infolist())
    if unpacked > MAX_XML_BYTES:
        raise TooLarge(f"unpacks to {unpacked} bytes")


def _blocks(ws: Any, ev: Any) -> tuple[Block, ...]:
    """A sheet's text lines and tables, in row order. A blank row closes a table before the next
    row is judged, so an unstyled header right after a gap starts its own table (the reference
    People Summary's lane table; audit 2026-10-29, F7)."""
    rows: dict[int, list[tuple[int, Any]]] = {}
    # the filled cells of a row, empty ones too: a row highlight covers the blanks of the row
    filled: dict[int, dict[int, Any]] = {}
    styled = False
    for row in ws.iter_rows():
        cells = [(c.column, c) for c in row if c.value is not None]
        if cells:
            rows[row[0].row] = cells
            styled = styled or any(_is_header(c) for _, c in cells)
            blank = {c.column: c for c in row if c.value is None and _fill(c)}
            if blank:
                filled[row[0].row] = blank
    # without the builder's header style, a header is a row at least 60% as wide as the widest
    widest = max((len(c) for c in rows.values()), default=0)
    unstyled_min = max(2, -(-widest * 3 // 5))
    out: list[Block] = []
    table: dict[str, Any] | None = None
    last_row = 0
    for r in sorted(rows):
        cells = rows[r]
        gap = r > last_row + 1
        last_row = r
        if table is not None and gap:
            out.append(_close(table))
            table = None
        header = (
            any(_is_header(c) for _, c in cells)
            if styled
            else (table is None and len(cells) >= unstyled_min)
        )
        if table is not None and header:
            out.append(_close(table))
            table = None
        if header:
            width = max(col for col, _ in cells)
            names = {col: _text(c, ev.value(ws.title, r, col)) for col, c in cells}
            table = {"header": tuple(names.get(j, "") for j in range(1, width + 1)), "rows": []}
            continue
        if table is not None:
            width = len(table["header"])
            if len(table["rows"]) >= MAX_TABLE_ROWS:
                table["cut"] = table.get("cut", 0) + 1
                continue
            by_col = dict(cells)
            blanks = filled.get(r, {})
            table["rows"].append(
                tuple(
                    Cell(
                        _text(by_col.get(j), ev.value(ws.title, r, j)),
                        _href(by_col.get(j)),
                        _fill(by_col.get(j) or blanks.get(j)),
                    )
                    for j in range(1, width + 1)
                )
            )
            continue
        parts = [_text(c, ev.value(ws.title, r, col)) for col, c in cells]
        parts = [p for p in parts if p]
        if parts:
            # a label and its count ("Rank A: 3"; a label that ends in a colon keeps one), else
            # the cells side by side
            two = len(parts) == 2
            line = f"{parts[0].rstrip(':')}: {parts[1]}" if two else " · ".join(parts)
            out.append(Block("text", text=line))
    if table is not None:
        out.append(_close(table))
    return tuple(out)


def _is_header(cell: Any) -> bool:
    try:
        return bool(cell.style == HEADER_STYLE)
    except (AttributeError, IndexError, KeyError):
        return False


def _close(table: dict[str, Any]) -> Block:
    return Block(
        "table", header=table["header"], rows=tuple(table["rows"]), cut=table.get("cut", 0)
    )


# (tenant, object key, SHA-256, main sheet) -> the book read from it. The cache is used from
# the event loop and from the reading threads, so it is guarded by a thread lock (never an
# asyncio primitive, which belongs to one event loop).
CacheKey = tuple[str, str, str, str]
_CACHE: OrderedDict[CacheKey, Book] = OrderedDict()
_CACHE_LOCK = threading.Lock()
_READING: dict[CacheKey, threading.Lock] = {}
_READS = threading.BoundedSemaphore(READS_AT_ONCE)


def _cached(key: CacheKey) -> Book | None:
    with _CACHE_LOCK:
        book = _CACHE.get(key)
        if book is not None:
            _CACHE.move_to_end(key)
        return book


def _keep(key: CacheKey, book: Book) -> None:
    with _CACHE_LOCK:
        _CACHE[key] = book
        _CACHE.move_to_end(key)
        # LOOP(book cache trim): cap=one entry over CACHE_BOOKS; exit=within the cap
        while len(_CACHE) > CACHE_BOOKS:
            _CACHE.popitem(last=False)


def _read_once(key: CacheKey, data: bytes, main_sheet: str, keep: bool) -> Book:
    """``read`` in a worker thread, once per key: a second request for the same book waits for
    the first and takes its result; at most ``READS_AT_ONCE`` workbooks are open at a time."""
    with _CACHE_LOCK:
        gate = _READING.setdefault(key, threading.Lock())
    try:
        with gate:
            book = _cached(key)
            if book is None:
                with _READS:
                    book = read(data, main_sheet)
                if keep:
                    _keep(key, book)
            return book
    finally:
        with _CACHE_LOCK:
            if _READING.get(key) is gate and not gate.locked():
                del _READING[key]


async def book_of(payloads: Any, tenant_id: UUID, found: Delivered, main_sheet: str) -> Book | None:
    """The delivered workbook read back, or None when it cannot be read: not the tenant's,
    missing, too large, or not a workbook. Never raises.

    The object is checked to be the tenant's before anything else, the cache included, and the
    cache is keyed by tenant and object as well as by content, so a book read for one tenant is
    never served to another. A file whose bytes do not match its recorded SHA-256 is shown but
    not cached."""
    from obs.logging import logger
    from tenancy.storage import PathRefused, require_owned

    try:
        require_owned(tenant_id, found.object_key)
    except PathRefused:
        return None
    key: CacheKey = (str(tenant_id), found.object_key, found.sha256, main_sheet)
    hit = _cached(key)
    if hit is not None:
        return hit
    if found.size is not None and found.size > MAX_BYTES:
        return None
    try:
        data = await payloads.get(found.object_key)
        if not isinstance(data, bytes | bytearray) or len(data) > MAX_BYTES:
            return None
        body = bytes(data)
        same = hashlib.sha256(body).hexdigest() == found.sha256
        if not same:
            logger(__name__).warning("file_tabs.sha_mismatch", extra={"file": found.id})
        return await asyncio.to_thread(_read_once, key, body, main_sheet, same)
    except Exception as error:  # noqa: BLE001 - the screen keeps its table without the tabs
        logger(__name__).warning("file_tabs.unreadable", extra={"error": type(error).__name__})
        return None


async def delivered_file(session: Any, account_id: UUID, sheet_key: str) -> Delivered | None:
    """The newest delivered workbook of the account's delivered run that prints ``sheet_key``."""
    from account_intelligence.replica.sheets import delivered_run

    kind = KINDS.get(sheet_key)
    run = await delivered_run(session, account_id) if kind else None
    if run is None:
        return None
    row = (
        await session.execute(
            text(
                "SELECT id, filename, sha256, object_key, size_bytes, manifest_json"
                " FROM deliverables"
                " WHERE account_id = :a AND run_id = :r AND kind = :k"
                " ORDER BY version DESC LIMIT 1"
            ),
            {"a": account_id, "r": run, "k": kind},
        )
    ).first()
    # a file built under another exclusion list than the tenant's current one is not read (F1)
    if row is not None and _built_under(row.manifest_json) != await _current_hash(session):
        return None
    return _delivered(row)


async def _current_hash(session: Any) -> str:
    from views.materialize import current_exclusion_hash

    return await current_exclusion_hash(session)


def _built_under(manifest: Any) -> str:
    """The exclusion set a stored file was built under (``none`` for a file built before the set
    was recorded)."""
    if isinstance(manifest, str):
        try:
            manifest = json.loads(manifest)
        except ValueError:
            manifest = None
    built = manifest.get("exclusion_hash") if isinstance(manifest, dict) else None
    return str(built or "none")


def _delivered(row: Any) -> Delivered | None:
    """The ``Delivered`` of a deliverables row (None: no row, or not a workbook)."""
    if row is None or not str(row.filename).endswith(".xlsx"):
        return None
    manifest = row.manifest_json
    if isinstance(manifest, str):
        try:
            manifest = json.loads(manifest)
        except ValueError:
            manifest = None  # an unreadable manifest: not a builder's file
    rows = manifest.get("rows") if isinstance(manifest, dict) else None
    return Delivered(
        str(row.id),
        str(row.filename),
        str(row.sha256),
        str(row.object_key),
        rows if isinstance(rows, int) and not isinstance(rows, bool) else None,
        int(row.size_bytes) if row.size_bytes is not None else None,
    )


async def delivered_files(
    session: Any, account_id: UUID, sheet_keys: Sequence[str], run: UUID | None = None
) -> dict[str, Delivered]:
    """``delivered_file`` of several sheets in one query (and one ``delivered_run`` read, or none
    when the caller has it): the newest workbook of the delivered run for each sheet with one."""
    from account_intelligence.replica.sheets import delivered_run

    kinds = {k: KINDS[k] for k in sheet_keys if k in KINDS}
    if not kinds:
        return {}
    run = run if run is not None else await delivered_run(session, account_id)
    if run is None:
        return {}
    rows = (
        await session.execute(
            text(
                "SELECT DISTINCT ON (kind) kind, id, filename, sha256, object_key, size_bytes,"
                " manifest_json FROM deliverables"
                " WHERE account_id = :a AND run_id = :r AND kind = ANY(:k)"
                " ORDER BY kind, version DESC"
            ),
            {"a": account_id, "r": run, "k": sorted(set(kinds.values()))},
        )
    ).all()
    current = await _current_hash(session)
    fresh = [r for r in rows if _built_under(r.manifest_json) == current]
    by_kind = {r.kind: _delivered(r) for r in fresh}
    return {k: d for k, kind in kinds.items() if (d := by_kind.get(kind)) is not None}


# (tenant, object key, SHA-256, sheet, rows the manifest counts, header signature) -> the rows
# of the file's own main table when the file is this builder's (``builder_main``), else None.
# Only a number, so it outlives the book itself in the book cache: the workspace screen reads it
# on every open, and a book evicted (or a new instance) is parsed again only for a tab.
_COUNTS: OrderedDict[tuple[str, str, str, str, int | None, int], int | None] = OrderedDict()


def _count_cached(key: tuple[str, str, str, str, int | None, int]) -> tuple[bool, int | None]:
    with _CACHE_LOCK:
        if key in _COUNTS:
            _COUNTS.move_to_end(key)
            return True, _COUNTS[key]
    return False, None


def _count_keep(key: tuple[str, str, str, str, int | None, int], value: int | None) -> None:
    with _CACHE_LOCK:
        _COUNTS[key] = value
        _COUNTS.move_to_end(key)
        # LOOP(count cache trim): cap=one entry over CACHE_COUNTS; exit=within the cap
        while len(_COUNTS) > CACHE_COUNTS:
            _COUNTS.popitem(last=False)


async def main_rows(payloads: Any, tenant_id: UUID, found: Delivered, sheet_key: str) -> int | None:
    """The rows of the delivered file's own main table when it stands in for the stored table
    (``builder_main``, else ``main_table``), else None. The answer is remembered by file content,
    so a file is opened for its count once per process. A file whose manifest states no row
    count is not this builder's: None."""
    from exports.config import load_export_config

    if found.rows is None:
        return None

    columns = getattr(load_export_config().rules.workbooks, KINDS[sheet_key]).columns
    signature = hash(tuple(c.header for c in columns))
    key = (str(tenant_id), found.object_key, found.sha256, sheet_key, found.rows, signature)
    hit, value = _count_cached(key)
    if hit:
        return value
    book = await book_of(payloads, tenant_id, found, main_sheet_of(sheet_key))
    if book is None:
        return None  # not remembered: the file may be readable on the next request
    main = builder_main(book, found, sheet_key) or main_table(book)
    table = main.table if main is not None else None
    value = len(table.rows) if table is not None else None
    _count_keep(key, value)
    return value


def builder_main(book: Book, found: Delivered, sheet_key: str) -> Tab | None:
    """The main tab when the file is this builder's own: its manifest counts exactly the main
    table's rows and the table's header agrees with the configured headers (``headers_agree``:
    the file may carry columns added since, or lack columns the configuration gained since it
    was delivered), else None."""
    from exports.config import load_export_config

    main = book.tab(None)
    table = main.table if main is not None else None
    if table is None or found.rows is None or len(table.rows) != found.rows:
        return None
    columns = getattr(load_export_config().rules.workbooks, KINDS[sheet_key]).columns
    return main if headers_agree(table.header, [c.header for c in columns]) else None


def headers_agree(file_header: Sequence[str], configured: Sequence[str]) -> bool:
    """Whether a file's header names are the configured columns by name: the names both have
    come in the configured order, and one list holds every name of the other (the file gained
    columns, or the configuration did after the file was delivered). An empty overlap, or two
    lists that each lack names of the other, is another table."""
    names = [h for h in file_header if h]
    shared = [h for h in configured if h in names]
    if not shared:
        return False
    if [h for h in names if h in configured] != shared:
        return False  # the shared names come in another order
    return set(configured) <= set(names) or set(names) <= set(configured)


def main_table(book: Book) -> Tab | None:
    """The main tab whenever it holds a table: what a delivered workbook shows on screen when the
    stored table cannot stand for it (never the live research table, owner 2026-10-29)."""
    main = book.tab(None)
    return main if main is not None and main.table is not None else None


def main_sheet_of(sheet_key: str) -> str:
    """The workbook sheet the deep-dive table shows (``workbooks.<kind>.sheet``)."""
    from exports.config import load_export_config

    books = load_export_config().rules.workbooks
    return str(getattr(books, KINDS[sheet_key]).sheet)


# ---- responses -------------------------------------------------------------------------------

SCREEN_ROWS = 2000  # rows of one table a tab shows (the file holds the rest)
SCREEN_CELLS = 200_000  # cells of every table of one tab
SCREEN_CHARS = 4_000_000  # characters of every cell and line of one tab
SCREEN_LINES = 1000  # text lines of one tab


def tab_refs(book: Book) -> list[dict[str, Any]]:
    """The tab bar: every sheet in the file's order."""
    return [{"key": t.key, "label": t.name, "main": t.main} for t in book.tabs]


def _cell_out(c: Cell) -> dict[str, Any]:
    out: dict[str, Any] = {"text": c.text, "href": c.href}
    if c.fill:
        out["fill"] = c.fill
    return out


def scrub_blocks(blocks: list[dict[str, Any]], extra_terms: Sequence[str] = ()) -> int:
    """The client-output guard on a tab's blocks, withholding per cell and per line (as
    ``sheets.scrub_rows`` does for a page): the tab always loads. ``extra_terms``: the other
    accounts' names and the excluded companies, which no cell may carry. Returns the cells
    withheld."""
    from account_intelligence.replica.sheets import WITHHELD
    from agents.naming import client_hits, scrub_client_text, third_party_domain_hits

    def guard(value: str) -> tuple[str, bool]:
        cleaned, withheld = scrub_client_text(value)
        if not withheld and extra_terms and client_hits(cleaned, extra_terms):
            withheld = True
        return cleaned, withheld

    held = 0
    for b in blocks:
        if b.get("kind") == "text":
            cleaned, withheld = guard(str(b.get("text") or ""))
            b["text"] = WITHHELD if withheld else cleaned
            held += withheld
            continue
        header = [str(h) for h in b.get("header") or []]
        for j, h in enumerate(header):
            cleaned, withheld = guard(h)
            header[j] = WITHHELD if withheld else cleaned
            held += withheld
        b["header"] = header
        for row in b.get("rows") or []:
            for cell in row:
                cleaned, withheld = guard(str(cell.get("text") or ""))
                cell["text"] = WITHHELD if withheld else cleaned
                held += withheld
                href = cell.get("href")
                if isinstance(href, str) and third_party_domain_hits(href):
                    cell["href"] = None
                    held += 1
    return held


def blocks_out(tab: Tab) -> list[dict[str, Any]]:
    """A tab's blocks as the API returns them, capped (``SCREEN_ROWS`` per table, then
    ``SCREEN_CELLS``, ``SCREEN_CHARS`` and ``SCREEN_LINES`` for the whole tab): the rows left out
    are counted in each table's ``cut``, so the screen says the file holds more."""
    out: list[dict[str, Any]] = []
    cells, chars, lines = 0, 0, 0
    for b in tab.blocks:
        if b.kind == "text":
            if lines < SCREEN_LINES and chars + len(b.text) <= SCREEN_CHARS:
                out.append({"kind": "text", "text": b.text})
                lines += 1
                chars += len(b.text)
            continue
        shown: list[list[dict[str, Any]]] = []
        for r in b.rows[:SCREEN_ROWS]:
            size = sum(len(c.text) + len(c.href or "") for c in r)
            if cells + len(r) > SCREEN_CELLS or chars + size > SCREEN_CHARS:
                break
            shown.append([_cell_out(c) for c in r])
            cells += len(r)
            chars += size
        out.append(
            {
                "kind": "table",
                "header": list(b.header),
                "rows": shown,
                "cut": b.cut + len(b.rows) - len(shown),
            }
        )
    return out


# The main tab read from the file, as a sortable sheet: what the screen shows when the delivered
# run's stored table cannot be used (it was stored under an earlier column set, or not stored):
# the file's columns and cells, in the file's order or sorted by one column, never the live
# research table with other columns.
_NUMERIC = re.compile(r"^-?\d+(?:\.\d+)?%?$")


def _columns(table: Block) -> list[Any]:
    """The table's columns as ``sheets.Col``: a link column (not sortable, like a file-shaped
    sheet's), a number column, a long text column or a text column."""
    from account_intelligence.replica.sheets import Col

    cols = []
    for j, label in enumerate(table.header):
        values = [r[j] for r in table.rows if j < len(r)]
        texts = [c.text for c in values if c.text]
        if any(c.href for c in values):
            kind = "url"
        elif texts and all(_NUMERIC.match(t.strip()) for t in texts):
            kind = "num"
        elif texts and sum(len(t) for t in texts) / len(texts) > 120:
            kind = "long"
        else:
            kind = "text"
        sort = None if kind == "url" else f"c{j + 1}"
        cols.append(Col(f"c{j + 1}", label or f"Column {j + 1}", sort, kind))
    return cols


def grid_sheet(tab: Tab, live: Any) -> Any:
    """The main tab as a ``sheets.Sheet`` (file order, one filter), for the screen and exports."""
    from account_intelligence.replica.sheets import FILE_ORDER, SORT_HINT, Filter, Sheet

    table = tab.table or Block("table")
    return Sheet(
        key=live.key,
        title=live.title,
        file=live.file,
        scope=f"The table of your delivered files, row for row. {SORT_HINT}",
        base="",
        cols=tuple(_columns(table)),
        filters=(Filter("all", "All", "true"), *(f for f, _, _ in _chips(tab, live.key))),
        default_sort=FILE_ORDER.key,
        table="file",
        sheet_name=tab.name,
    )


def grid_query(
    tab: Tab,
    live: Any,
    sort: str | None,
    direction: str | None,
    offset: int,
    limit: int,
    filter_key: str | None = None,
) -> Any:
    """A validated request over the main tab (``sheets.BadSheetQuery`` for an unknown or
    unsortable column, a direction other than asc or desc, or a filter the file's table does not
    have: it has All and the stored table's chips whose column the file carries); the page is
    clamped like ``sheets.parse`` clamps it."""
    from account_intelligence.replica.sheets import FILE_ORDER, MAX_PAGE, BadSheetQuery, Query

    sheet = grid_sheet(tab, live)
    cols = {c.key: c for c in (*sheet.cols, FILE_ORDER)}
    col = cols.get(sort or FILE_ORDER.key)
    if col is None or col.sort is None or (direction or "asc") not in ("asc", "desc"):
        raise BadSheetQuery("unknown sort or direction")
    filt = next((f for f in sheet.filters if f.key == (filter_key or "all")), None)
    if filt is None:
        raise BadSheetQuery("unknown filter")
    return Query(
        sheet,
        filt,
        col,
        direction == "desc",
        max(0, offset),
        max(1, min(limit, MAX_PAGE)),
    )


def _sort_value(text_value: str) -> tuple[int, Any]:
    """Numbers (a percentage as its number) before text, text without case."""
    t = text_value.strip()
    if _NUMERIC.match(t):
        return (0, float(t.rstrip("%")))
    return (1, t.casefold())


@lru_cache(maxsize=8)
def _chip_headers(kind: str) -> tuple[tuple[str, str, str, str], ...]:
    """The stored table's chips beside All as ``(key, label, header, value)``, the column given
    by its header in the file (``sheets.chip_specs``). Read once per process: the export
    configuration is parsed from YAML on every ``load_export_config`` call."""
    from account_intelligence.replica.sheets import chip_specs
    from exports.config import load_export_config
    from views.tables import specs

    columns = specs()[kind].exported_columns  # type: ignore[index]
    header_of = {c.key: c.header for c in columns}
    return tuple(
        (key, label, header_of[column], value)
        for key, label, column, value in chip_specs(kind, columns, load_export_config().rules)
        if column in header_of
    )


def _chips(tab: Tab, sheet_key: str) -> list[tuple[Any, int, str]]:
    """The filter chips of the main tab beside All, as ``(filter, column index, value)``: the
    stored table's own chips (Priority A, Meets seniority rule, the posting statuses) matched to
    the file's columns by header, so a chip filters the file's table the way it filters the
    stored one. A chip whose column the file lacks is left out."""
    from account_intelligence.replica.sheets import Filter

    kind = KINDS.get(sheet_key)
    table = tab.table
    if kind is None or table is None:
        return []
    return [
        (Filter(key, label, "true"), table.header.index(name), value)
        for key, label, name, value in _chip_headers(kind)
        if name in table.header
    ]


def _kept(tab: Tab, q: Any) -> list[tuple[Cell, ...]]:
    """The main tab's rows under the request's filter (All: every row)."""
    table = tab.table or Block("table")
    if q.filter.key == "all":
        return list(table.rows)
    chip = next((c for c in _chips(tab, q.sheet.key) if c[0].key == q.filter.key), None)
    if chip is None:
        return []
    _, j, value = chip
    return [r for r in table.rows if j < len(r) and r[j].text.strip() == value]


def _order(tab: Tab, q: Any) -> list[tuple[Cell, ...]]:
    """The main tab's rows in the request's order. A sort is stable (equal cells keep the file's
    order) and puts blank cells last in either direction, as the stored tables do (NULLS LAST)."""
    rows = _kept(tab, q)
    if q.sort.key == "_file":
        return rows[::-1] if q.desc else rows
    j = int(q.sort.key[1:]) - 1

    def cell(r: tuple[Cell, ...]) -> str:
        return r[j].text if j < len(r) else ""

    filled = [r for r in rows if cell(r).strip()]
    blank = [r for r in rows if not cell(r).strip()]
    filled.sort(key=lambda r: _sort_value(cell(r)), reverse=q.desc)
    return filled + blank


def _row(width: int, n: int, r: tuple[Cell, ...]) -> dict[str, Any]:
    cells = {f"c{j + 1}": (r[j].text if j < len(r) else "") for j in range(width)}
    links = {f"c{j + 1}": r[j].href for j in range(min(width, len(r))) if r[j].href}
    fills = {f"c{j + 1}": r[j].fill for j in range(min(width, len(r))) if r[j].fill}
    out: dict[str, Any] = {"n": n, "cells": cells, "links": links, "link": None, "link_label": None}
    if fills:
        out["fills"] = fills
    return out


def grid_cell(tab: Tab, q: Any, n: int, col_key: str) -> dict[str, Any] | None:
    """One whole cell of the main tab in the request's order: row ``n`` (absolute, 1-based) and
    column ``col_key`` (``c1``...), or None past the table."""
    ordered = _order(tab, q)
    if n < 1 or n > len(ordered) or not re.fullmatch(r"c\d{1,3}", col_key):
        return None
    j = int(col_key[1:]) - 1
    r = ordered[n - 1]
    if j < 0 or j >= len((tab.table or Block("table")).header):
        return None
    cell = r[j] if j < len(r) else Cell("")
    return {"n": n, "col": col_key, "text": cell.text, "href": cell.href}


def grid_rows(tab: Tab, q: Any) -> list[dict[str, Any]]:
    """Every row of the main tab in the request's order, as ``sheets.row_out`` shapes them."""
    width = len((tab.table or Block("table")).header)
    return [_row(width, n, r) for n, r in enumerate(_order(tab, q), start=1)]


def _filter_counts(tab: Tab, q: Any) -> list[dict[str, Any]]:
    """The chips with the rows each holds over the whole table (not the filtered page)."""
    table = tab.table or Block("table")
    out = [{"key": "all", "label": "All", "count": len(table.rows)}]
    for f, j, value in _chips(tab, q.sheet.key):
        n = sum(1 for r in table.rows if j < len(r) and r[j].text.strip() == value)
        out.append({"key": f.key, "label": f.label, "count": n})
    return out


def grid_page(tab: Tab, q: Any) -> dict[str, Any]:
    """One page of the main tab in the shape ``sheets.read`` returns (an offset past the end
    is an empty page with the table's total); only the page's rows are shaped."""
    ordered = _order(tab, q)
    width = len((tab.table or Block("table")).header)
    page = [
        _row(width, q.offset + i + 1, r)
        for i, r in enumerate(ordered[q.offset : q.offset + q.limit])
    ]
    return {
        "key": q.sheet.key,
        "title": q.sheet.title,
        "file": q.sheet.file,
        "scope": q.sheet.scope,
        "footnote": q.sheet.footnote,
        "private": False,
        "sheet_name": q.sheet.sheet_name,
        "from_file": True,
        "columns": [
            {"key": c.key, "label": c.label, "sortable": c.sort is not None, "kind": c.kind}
            for c in q.sheet.cols
        ],
        "filters": _filter_counts(tab, q),
        "filter": q.filter.key,
        "sort": q.sort.key,
        "dir": "desc" if q.desc else "asc",
        "total": len(ordered),
        "offset": q.offset,
        "rows": page,
    }


async def iter_rows(tab: Tab, q: Any) -> Any:
    """The main tab's rows for an export (``table_export``), every row in the request's order."""
    width = len((tab.table or Block("table")).header)
    for n, r in enumerate(_order(tab, q), start=1):
        yield _row(width, n, r)
