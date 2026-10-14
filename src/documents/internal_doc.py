"""The Internal Only Word file of a file-mode run (two-output rule, rule 5).

Built by its own code path into its own storage location (``tenants/<t>/internal/<run>/``),
never into the client file set, the shared folder or the workspace package. The cover and the
header of every page carry "Internal Only" and the label "Internal: from your uploaded
material". Every finding names its grade and its citation (document and location); public
findings of a linked public research run, when there is one, sit in their own section labelled
"Public research".

GUARANTEE(test_file_run_builds_a_labelled_internal_file)
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from typing import Any

from documents.gate import INTERNAL_LABEL

GRADE_TEXT = {"internal": "Internal", "confirmed": "Confirmed (authoritative document)"}
CATEGORY_TEXT = {
    "platforms": "Platforms and systems",
    "incumbent_partners": "Incumbent partners and vendors",
    "programs": "Programmes and initiatives",
    "budget_cycles": "Budget cycles and dates",
    "objections": "Objections and risks",
    "decisions": "Decisions and stages",
    "next_steps": "Next steps",
    "people": "People and roles",
}


@dataclass(frozen=True)
class FileFinding:
    """One file-derived fact as the internal file shows it."""

    category: str
    value: str
    grade: str
    citation: str
    quote: str


@dataclass(frozen=True)
class TrackerLine:
    """One tracker row as the internal file shows it."""

    label: str
    status: str
    reason: str


def build(
    *,
    account: str,
    as_of: date,
    documents: Sequence[tuple[str, int]],
    findings: Sequence[FileFinding],
    tracker: Sequence[TrackerLine],
    public: Sequence[tuple[str, str, str]] = (),
    label: str = INTERNAL_LABEL,
) -> bytes:
    """The docx bytes (see the module docstring). ``public``: (finding, grade, link). ``label``
    is the file's label (a company research run says it has no uploaded material)."""
    from docx import Document

    doc = Document()
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = f"Internal Only. {label}"
        footer = section.footer.paragraphs[0]
        footer.text = f"Internal Only. {account}. Not for the shared folder or the client."
    doc.add_heading("Internal Only", level=0)
    doc.add_paragraph(label)
    doc.add_paragraph(f"Account: {account}")
    source = f"{len(documents)} uploaded documents" if documents else "this research run"
    doc.add_paragraph(f"Prepared on {as_of.isoformat()} from {source}.")
    doc.add_heading("Documents read", level=1)
    for name, passages in documents:
        doc.add_paragraph(f"{name}: {passages} passages", style="List Bullet")
    doc.add_heading("From your files", level=1)
    if not findings:
        doc.add_paragraph("The uploaded documents state nothing in the researched areas.")
    by_cat: dict[str, list[FileFinding]] = {}
    for f in findings:
        by_cat.setdefault(f.category, []).append(f)
    for cat in CATEGORY_TEXT:
        items = by_cat.get(cat)
        if not items:
            continue
        doc.add_heading(CATEGORY_TEXT[cat], level=2)
        for f in items:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f.value).bold = True
            p.add_run(f". Grade: {GRADE_TEXT.get(f.grade, f.grade.title())}. Source: {f.citation}.")
            doc.add_paragraph(f"“{f.quote}”", style="Intense Quote")
    if public:
        doc.add_heading("Public research", level=1)
        for value, grade, link in public:
            doc.add_paragraph(f"{value}. Grade: {grade.title()}. Link: {link}", style="List Bullet")
    doc.add_heading("Research tracker", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    head = table.rows[0].cells
    head[0].text, head[1].text, head[2].text = "Research area", "Status", "Reason"
    for line in tracker:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = line.label, line.status, line.reason
    props = doc.core_properties
    props.title = f"Internal Only: {account}"
    props.author = "6th agents"
    props.comments = label
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def as_dict(f: FileFinding) -> dict[str, Any]:
    """JSON form for the API's internal view."""
    return {
        "category": f.category,
        "value": f.value,
        "grade": f.grade,
        "citation": f.citation,
        "quote": f.quote,
    }
