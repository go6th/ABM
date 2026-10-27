"""Workbook checks after recalculation.

Each workbook is re-saved by LibreOffice (formulas computed, values cached; the delivered file
is untouched) and then: zero ``#REF!``, ``#DIV/0!``, ``#VALUE!`` or ``#NAME?`` (and the other
error values) anywhere; every formula cell of the delivered file is still a formula in the
recalculated one; every formula has a cached value; and the Summary values equal what
``parity.counts`` computes from the bundle rows. The Summary blocks are found by their labels,
not by fixed cells: the Jobs Summary rows are the lanes of ``meta.lanes`` under a header row
that starts with Total, Closed, Open; the row totals start with "Rows"; the People Summary
rows are the level labels, "Total", "Meets seniority rule" and the lane labels; the Platform
Ranking rows follow the header "Rank, Platform" and the Counts block reads "Rank A", "Rank B",
"Rank C". A block whose label is not found is a finding, so a renderer cannot drop a block
unnoticed.

The recalculated Summary values are returned for the manifest, keyed as
``parity.counts.summary_values`` keys them.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Mapping
from dataclasses import asdict
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from parity import counts
from parity.gate.files import (
    Bundle,
    Office,
    excel_files,
    is_jobs_workbook,
    is_people_workbook,
    is_tech_workbook,
)
from parity.gate.report import CheckResult, Finding, result, skipped

NAME = "workbook"
_NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/package/2006/relationships",
}
_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
_SHARE_TOLERANCE = 1e-9


def _fold(value: object) -> str:
    return str(value if value is not None else "").strip().casefold()


def _formula_cells(path: Path) -> list[tuple[str, str]]:
    wb = load_workbook(path)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    out.append((ws.title, cell.coordinate))
    return out


def cached_formula_cells(xlsx: bytes) -> set[tuple[str, str]]:
    """``(sheet, coordinate)`` of every formula cell that carries a cached value (``<v>``) in
    the workbook's own XML. Read from the XML because openpyxl gives None both for a missing
    value and for a formula whose result is the empty string (a blank rank)."""
    out: set[tuple[str, str]] = set()
    with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
        rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"), parser=_PARSER)
        targets = {
            str(r.get("Id")): str(r.get("Target")) for r in rels.findall("p:Relationship", _NS)
        }
        book = etree.fromstring(z.read("xl/workbook.xml"), parser=_PARSER)
        for sheet in book.findall("m:sheets/m:sheet", _NS):
            name = str(sheet.get("name"))
            target = targets.get(str(sheet.get(f"{{{_NS['r']}}}id")), "")
            path = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            if path not in z.namelist():
                continue
            root = etree.fromstring(z.read(path), parser=_PARSER)
            for c in root.iter(f"{{{_NS['m']}}}c"):
                if c.find("m:f", _NS) is not None and c.find("m:v", _NS) is not None:
                    out.add((name, str(c.get("r"))))
    return out


def _rows(ws: Worksheet) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _find_row(rows: list[list[Any]], label: str, start: int = 0) -> int:
    for i in range(start, len(rows)):
        if rows[i] and _fold(rows[i][0]) == _fold(label):
            return i
    return -1


def _find_row_prefix(rows: list[list[Any]], prefix: str, start: int = 0) -> int:
    for i in range(start, len(rows)):
        if rows[i] and _fold(rows[i][0]).startswith(_fold(prefix)):
            return i
    return -1


def _cell(rows: list[list[Any]], r: int, c: int) -> Any:
    row = rows[r] if 0 <= r < len(rows) else []
    return row[c] if c < len(row) else None


def _compare(
    got: Any, expected: int | float | str, key: str, file: str, findings: list[Finding]
) -> None:
    if isinstance(expected, float):
        ok = isinstance(got, int | float) and abs(float(got) - expected) <= _SHARE_TOLERANCE
    elif isinstance(expected, str):
        ok = _fold(got) == _fold(expected)
    else:
        ok = isinstance(got, int | float) and not isinstance(got, bool) and int(got) == expected
    if not ok:
        findings.append(Finding(NAME, f"{key}: workbook {got!r}, computed {expected!r}", file))


def jobs_summary_values(
    ws: Worksheet, bundle: Bundle, file: str, findings: list[Finding]
) -> dict[str, Any]:
    """Read and compare the Jobs Summary blocks; the values read are returned."""
    rows = _rows(ws)
    values: dict[str, Any] = {}
    header = next(
        (
            i
            for i, r in enumerate(rows)
            if len(r) > 3 and [_fold(x) for x in r[1:4]] == ["total", "closed", "open"]
        ),
        -1,
    )
    if header < 0:
        findings.append(
            Finding(NAME, "Jobs Summary header row (Total, Closed, Open) not found", file)
        )
        return values
    expected = counts.jobs_summary(bundle)
    for lane, jc in expected.items():
        r = _find_row(rows, lane, header + 1)
        if r < 0:
            findings.append(Finding(NAME, f"Jobs Summary row for lane {lane} not found", file))
            continue
        for c, (key, want) in enumerate(asdict(jc).items(), start=1):
            got = _cell(rows, r, c)
            values[f"jobs.{lane}.{key}"] = got
            _compare(got, want, f"jobs.{lane}.{key}", file, findings)
    totals = counts.jobs_row_totals(bundle)
    labels = {
        "jobs.rows.all": ("rows on the", totals.all_rows),
        "jobs.rows.account": (
            f"rows that are {_fold(bundle['meta']['account_short'])}",
            totals.account_rows,
        ),
    }
    first_rows = -1
    for key, (prefix, want) in labels.items():
        r = _find_row_prefix(rows, prefix, header + 1)
        if r < 0:
            findings.append(Finding(NAME, f"Jobs Summary row total {key} not found", file))
            continue
        if key == "jobs.rows.account":
            first_rows = r
        values[key] = _cell(rows, r, 1)
        _compare(_cell(rows, r, 1), want, key, file, findings)
    if first_rows >= 0:
        r = _find_row_prefix(rows, "rows that are", first_rows + 1)
        if r >= 0:
            values["jobs.rows.other"] = _cell(rows, r, 1)
            _compare(_cell(rows, r, 1), totals.other_rows, "jobs.rows.other", file, findings)
        else:
            findings.append(Finding(NAME, "Jobs Summary row total jobs.rows.other not found", file))
    return values


def people_summary_values(
    ws: Worksheet, bundle: Bundle, file: str, findings: list[Finding]
) -> dict[str, Any]:
    """Read and compare the People Summary blocks; the values read are returned."""
    rows = _rows(ws)
    values: dict[str, Any] = {}
    last = -1
    for label, lc in counts.people_level_counts(bundle).items():
        r = _find_row(rows, label)
        if r < 0:
            findings.append(Finding(NAME, f"People Summary level row {label} not found", file))
            continue
        last = max(last, r)
        values[f"people.level.{label}.people"] = _cell(rows, r, 1)
        values[f"people.level.{label}.priority_a"] = _cell(rows, r, 2)
        _compare(_cell(rows, r, 1), lc.people, f"people.level.{label}.people", file, findings)
        _compare(
            _cell(rows, r, 2), lc.priority_a, f"people.level.{label}.priority_a", file, findings
        )
    level_total = counts.people_level_totals(bundle)
    r = _find_row(rows, "Total", last + 1) if last >= 0 else -1
    if r < 0:
        findings.append(Finding(NAME, "People Summary level Total row not found", file))
    else:
        values["people.level.total.people"] = _cell(rows, r, 1)
        values["people.level.total.priority_a"] = _cell(rows, r, 2)
        _compare(_cell(rows, r, 1), level_total.people, "people.level.total.people", file, findings)
        _compare(
            _cell(rows, r, 2),
            level_total.priority_a,
            "people.level.total.priority_a",
            file,
            findings,
        )
    r = _find_row(rows, "Meets seniority rule")
    if r < 0:
        findings.append(Finding(NAME, "People Summary Meets seniority rule row not found", file))
    else:
        values["people.senior"] = _cell(rows, r, 1)
        _compare(_cell(rows, r, 1), counts.meets_seniority(bundle), "people.senior", file, findings)
    lane_last = -1
    for lane, lane_count in counts.people_lane_counts(bundle).items():
        r = _find_row(rows, lane, max(last, 0))
        if r < 0:
            findings.append(Finding(NAME, f"People Summary lane row {lane} not found", file))
            continue
        lane_last = max(lane_last, r)
        for c, (key, want) in enumerate(asdict(lane_count).items(), start=1):
            values[f"people.lane.{lane}.{key}"] = _cell(rows, r, c)
            _compare(_cell(rows, r, c), want, f"people.lane.{lane}.{key}", file, findings)
    if lane_last >= 0:
        r = _find_row(rows, "Total", lane_last + 1)
        if r < 0:
            findings.append(Finding(NAME, "People Summary lane Total row not found", file))
        else:
            for c, (key, want) in enumerate(
                asdict(counts.people_lane_totals(bundle)).items(), start=1
            ):
                values[f"people.lane.total.{key}"] = _cell(rows, r, c)
                _compare(_cell(rows, r, c), want, f"people.lane.total.{key}", file, findings)
    return values


def ranking_values(
    ws: Worksheet, bundle: Bundle, file: str, findings: list[Finding]
) -> dict[str, Any]:
    """Read and compare the Platform Ranking sheet; the values read are returned."""
    rows = _rows(ws)
    values: dict[str, Any] = {}
    header = next(
        (
            i
            for i, r in enumerate(rows)
            if len(r) > 1 and _fold(r[0]) == "rank" and _fold(r[1]) == "platform"
        ),
        -1,
    )
    if header < 0:
        findings.append(
            Finding(NAME, "Platform Ranking header row (Rank, Platform) not found", file)
        )
        return values
    expected = counts.platform_ranking(bundle)
    by_platform: dict[str, tuple[int, list[Any]]] = {}
    i = header + 1
    # LOOP(ranking rows): cap=rows of the sheet; exit=first blank Platform cell
    while i < len(rows) and rows[i] and rows[i][1] not in (None, ""):
        by_platform[_fold(rows[i][1])] = (i, rows[i])
        i += 1
    for row in expected:
        found = by_platform.get(_fold(row.platform))
        if found is None:
            findings.append(
                Finding(NAME, f"Platform Ranking row for {row.platform} not found", file)
            )
            continue
        _, cells = found
        got_rank, got_count, got_share = (
            _cell([cells], 0, 0),
            _cell([cells], 0, 4),
            _cell([cells], 0, 5),
        )
        values[f"ranking.{row.platform}.count"] = got_count
        values[f"ranking.{row.platform}.share"] = got_share
        values[f"ranking.{row.platform}.rank"] = got_rank if got_rank is not None else ""
        _compare(got_count, row.count, f"ranking.{row.platform}.count", file, findings)
        _compare(got_share, row.share, f"ranking.{row.platform}.share", file, findings)
        _compare(
            got_rank if got_rank is not None else "",
            row.rank,
            f"ranking.{row.platform}.rank",
            file,
            findings,
        )
    for letter, want in counts.rank_counts(bundle).items():
        r = _find_row(rows, f"Rank {letter}", header + 1)
        if r < 0:
            findings.append(
                Finding(NAME, f"Platform Ranking Counts row Rank {letter} not found", file)
            )
            continue
        values[f"ranking.count.{letter}"] = _cell(rows, r, 1)
        _compare(_cell(rows, r, 1), want, f"ranking.count.{letter}", file, findings)
    return values


def check_workbook(
    path: Path, bundle: Bundle, office: Office
) -> tuple[list[Finding], dict[str, Any] | None]:
    """One workbook: findings and the recalculated Summary values (None when not run)."""
    findings: list[Finding] = []
    original_formulas = _formula_cells(path)
    done = office.recalculated(path.read_bytes())
    if done is None:
        return [Finding(NAME, "recalculation did not run", path.name)], None
    values_wb = load_workbook(io.BytesIO(done), data_only=True)
    formulas_wb = load_workbook(io.BytesIO(done))
    errors = 0
    for ws in values_wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() in ERRORS:
                    errors += 1
    if errors:
        findings.append(Finding(NAME, f"formula errors: {errors}", path.name))
    cached = cached_formula_cells(done)
    lost = missing = 0
    for sheet, coord in original_formulas:
        still = formulas_wb[sheet][coord].value
        if not (isinstance(still, str) and still.startswith("=")):
            lost += 1
        if (sheet, coord) not in cached:
            missing += 1
    if lost:
        findings.append(Finding(NAME, f"formulas no longer formulas: {lost}", path.name))
    if missing:
        findings.append(Finding(NAME, f"formulas without a cached value: {missing}", path.name))
    values: dict[str, Any] = {"formulas": len(original_formulas)}
    if is_jobs_workbook(path):
        if "Summary" in values_wb.sheetnames:
            values.update(jobs_summary_values(values_wb["Summary"], bundle, path.name, findings))
        else:
            findings.append(Finding(NAME, "no Summary sheet", path.name))
    elif is_people_workbook(path):
        if "Summary" in values_wb.sheetnames:
            values.update(people_summary_values(values_wb["Summary"], bundle, path.name, findings))
        else:
            findings.append(Finding(NAME, "no Summary sheet", path.name))
    elif is_tech_workbook(path):
        if "Platform Ranking" in values_wb.sheetnames:
            values.update(
                ranking_values(values_wb["Platform Ranking"], bundle, path.name, findings)
            )
        else:
            findings.append(Finding(NAME, "no Platform Ranking sheet", path.name))
    return findings, values


def run_workbook_checks(
    files: list[Path], bundle: Bundle, office: Office | None
) -> tuple[CheckResult, dict[str, Mapping[str, Any]]]:
    """The check over every workbook, and the recalculated Summary values per file."""
    if office is None:
        return skipped(NAME, "no spreadsheet engine"), {}
    findings: list[Finding] = []
    summary: dict[str, Mapping[str, Any]] = {}
    for path in excel_files(files):
        found, values = check_workbook(path, bundle, office)
        findings += found
        if values is not None:
            summary[path.name] = values
    return result(NAME, findings, {"files": list(summary)}), summary
