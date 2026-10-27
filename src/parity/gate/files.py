"""What the gate reads: the bundle, the rendered files and their text.

The bundle is a plain mapping (``parity.bundle`` is used to load it when that module exists).
Client files are the Word and Excel files of the output folder whose names are not refused
(``refused_name``: the internal-only file and any name carrying a guarded person's name, which
the repository holds only as word digests). Text comes straight from the Office XML: every
``w:t`` run of a Word file (body, headers, footers, footnotes), every cell of a workbook.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from collections.abc import Iterator, Mapping
from pathlib import Path
from typing import Any, Protocol

from lxml import etree
from openpyxl import load_workbook

Bundle = Mapping[str, Any]


class Office(Protocol):
    """The converter the checks need (``exports.verify.LibreOffice`` fits): a Word file's PDF
    text, and a workbook re-saved with its formulas computed. Either is None when the tool
    could not run."""

    def pdf_text(self, docx: bytes) -> str | None:
        """The text of the Word file rendered to PDF, or None without a converter."""
        ...

    def recalculated(self, xlsx: bytes) -> bytes | None:
        """The workbook re-saved with cached values, or None without a converter."""
        ...


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)
INTERNAL_ONLY_MARKER = "Internal_Only"
_WORDS = re.compile(r"[^\W\d_]+", re.UNICODE)

STANDING_SUFFIXES = (
    "_Account_Summary.docx",
    "_Sources_and_Evidence.docx",
    "_People_List.xlsx",
    "_Job_Postings.xlsx",
    "_Tech_Stack_and_Teams.xlsx",
)
OPTIONAL_SUFFIXES = ("_Lane_Highlights.docx",)


def load_bundle(path: Path) -> dict[str, Any]:
    """The bundle at ``path`` as a mapping: validated through ``parity.bundle`` when that module
    is importable (``load_bundle`` then ``dump_bundle``), else the JSON as it is."""
    try:
        import importlib

        bundle_module: Any = importlib.import_module("parity.bundle")
    except ImportError:
        bundle_module = None
    if bundle_module is not None and callable(getattr(bundle_module, "load_bundle", None)):
        loaded = bundle_module.load_bundle(path)
        if isinstance(loaded, Mapping):
            return dict(loaded)
        dump = getattr(bundle_module, "dump_bundle", None)
        data = dump(loaded) if callable(dump) else getattr(loaded, "model_dump", dict)()
        if isinstance(data, Mapping):
            return dict(data)
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("bundle is not a JSON object")
    return data


def prefix_of(bundle: Bundle) -> str:
    """The file name prefix (``meta.account_short``)."""
    return str(bundle["meta"]["account_short"])


def refused_name(name: str) -> bool:
    """Whether a file name may never be in a zip or a share: it carries the internal-only
    marker or a guarded person's name (``agents.naming.digest_hits`` on its words)."""
    base = Path(name).name
    if INTERNAL_ONLY_MARKER.casefold() in base.casefold():
        return True
    from agents.naming import digest_hits

    return bool(digest_hits(" ".join(_WORDS.findall(base))))


def standing_names(prefix: str) -> tuple[str, ...]:
    """The standing deliverable names for a prefix."""
    return tuple(f"{prefix}{s}" for s in STANDING_SUFFIXES)


def optional_names(prefix: str) -> tuple[str, ...]:
    """Deliverables that are in the zip when present."""
    return tuple(f"{prefix}{s}" for s in OPTIONAL_SUFFIXES)


def office_files(out_dir: Path) -> list[Path]:
    """Every Word and Excel file of the output folder, sorted."""
    return sorted(p for p in out_dir.iterdir() if p.suffix.lower() in (".docx", ".xlsx"))


def client_files(out_dir: Path) -> list[Path]:
    """The Office files bound for the client (refused names left out)."""
    return [p for p in office_files(out_dir) if not refused_name(p.name)]


def word_files(paths: list[Path]) -> list[Path]:
    """The Word files among ``paths``."""
    return [p for p in paths if p.suffix.lower() == ".docx"]


def excel_files(paths: list[Path]) -> list[Path]:
    """The Excel files among ``paths``."""
    return [p for p in paths if p.suffix.lower() == ".xlsx"]


def is_people_workbook(path: Path) -> bool:
    """Whether the file is the People List workbook."""
    return path.name.endswith("_People_List.xlsx")


def is_jobs_workbook(path: Path) -> bool:
    """Whether the file is the Job Postings workbook."""
    return path.name.endswith("_Job_Postings.xlsx")


def is_tech_workbook(path: Path) -> bool:
    """Whether the file is the Tech Stack and Teams workbook."""
    return path.name.endswith("_Tech_Stack_and_Teams.xlsx")


def docx_paragraphs(data: bytes) -> list[str]:
    """Every paragraph of a Word file (body, headers, footers, footnotes), one string each,
    table cells included."""
    out: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in sorted(z.namelist()):
            if name.startswith("word/") and name.endswith(".xml") and "theme" not in name:
                root = etree.fromstring(z.read(name), parser=_PARSER)
                for p in root.iter(f"{W}p"):
                    out.append("".join(t.text or "" for t in p.iter(f"{W}t")))
    return out


def docx_text(data: bytes) -> str:
    """The paragraphs of a Word file joined by newlines."""
    return "\n".join(docx_paragraphs(data))


def docx_table_first_cells(data: bytes) -> list[list[str]]:
    """Per table (document order), the first cell's text of every row."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = etree.fromstring(z.read("word/document.xml"), parser=_PARSER)
    tables: list[list[str]] = []
    for tbl in root.iter(f"{W}tbl"):
        rows = []
        for tr in tbl.findall(f"{W}tr"):
            tc = tr.find(f"{W}tc")
            if tc is None:
                continue
            rows.append("".join(t.text or "" for t in tc.iter(f"{W}t")))
        tables.append(rows)
    return tables


def office_xml_members(data: bytes) -> Iterator[tuple[str, str]]:
    """Every XML and relationship member of an Office package as text (for scans that must
    also see hyperlink targets and document properties)."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for name in sorted(z.namelist()):
            if name.endswith((".xml", ".rels")):
                yield name, z.read(name).decode("utf-8", "replace")


def xlsx_cells(path: Path, data_only: bool = False) -> Iterator[tuple[str, str, Any]]:
    """``(sheet, coordinate, value)`` of every non-empty cell of a workbook."""
    wb = load_workbook(path, data_only=data_only)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.value != "":
                    yield ws.title, cell.coordinate, cell.value


def xlsx_strings(path: Path) -> Iterator[tuple[str, str, str]]:
    """``(sheet, coordinate, text)`` of every text cell that is not a formula."""
    for sheet, coord, value in xlsx_cells(path):
        if isinstance(value, str) and not value.startswith("="):
            yield sheet, coord, value


def text_of(path: Path) -> str:
    """The visible text of a Word or Excel file."""
    if path.suffix.lower() == ".docx":
        return docx_text(path.read_bytes())
    return "\n".join(text for _, _, text in xlsx_strings(path))
