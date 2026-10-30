"""Data warnings: things a reviewer should see that never fail the gate.

* duplicate names in ``people_rows`` with different profile URLs (two people, or one person
  twice with two links);
* duplicate names in ``people_excluded.security``;
* columns that are entirely empty on a workbook sheet (below the header row);
* postings in ``jobs_rows`` without a requisition ID (a count).

Warnings carry counts, column names and sheet names, never a person's name.
"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from parity.gate.files import Bundle, excel_files
from parity.gate.report import Finding

NAME = "data_warnings"


def duplicate_profiles(bundle: Bundle) -> int:
    """How many names appear more than once with more than one distinct profile URL."""
    urls: dict[str, set[str]] = defaultdict(set)
    for row in bundle.get("people_rows") or ():
        name = str(row.get("name") or "").strip().casefold()
        if name:
            urls[name].add(str(row.get("profile_url") or ""))
    return sum(1 for seen in urls.values() if len(seen) > 1)


def duplicate_exclusions(bundle: Bundle) -> int:
    """How many names repeat in the security exclusion list."""
    seen: dict[str, int] = defaultdict(int)
    for entry in (bundle.get("people_excluded") or {}).get("security") or ():
        seen[str(entry).split(" (")[0].strip().casefold()] += 1
    return sum(1 for n in seen.values() if n > 1)


def empty_columns(path: Path) -> dict[str, list[str]]:
    """Per sheet, the header names of columns with no value below the header."""
    out: dict[str, list[str]] = {}
    wb = load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        header = rows[0]
        width = max(len(r) for r in rows)
        empties = []
        for i in range(width):
            name = header[i] if i < len(header) else None
            if name in (None, ""):
                continue
            if all(i >= len(r) or r[i] in (None, "") for r in rows[1:]):
                empties.append(str(name))
        if empties:
            out[ws.title] = empties
    return out


def postings_without_requisition(bundle: Bundle) -> int:
    """How many postings have no requisition ID."""
    return sum(
        1 for r in bundle.get("jobs_rows") or () if not str(r.get("requisition") or "").strip()
    )


def run_warnings(files: list[Path], bundle: Bundle) -> list[Finding]:
    """Every data warning."""
    out: list[Finding] = []
    n = duplicate_profiles(bundle)
    if n:
        out.append(Finding(NAME, f"duplicate names with different profile URLs: {n}", None, "warn"))
    n = duplicate_exclusions(bundle)
    if n:
        out.append(
            Finding(NAME, f"duplicate names in the security exclusion list: {n}", None, "warn")
        )
    for path in excel_files(files):
        for sheet, columns in empty_columns(path).items():
            out.append(
                Finding(
                    NAME, f"empty columns on sheet {sheet}: {', '.join(columns)}", path.name, "warn"
                )
            )
    n = postings_without_requisition(bundle)
    if n:
        out.append(Finding(NAME, f"postings without a requisition ID: {n}", None, "warn"))
    return out
