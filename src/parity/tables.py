"""The tables of the deliverable files, declared once: id, file, sheet or heading, columns (key,
header, kind, link label, chip fills) and the rows each table prints from a bundle, in the
file's own order and with the file's own cell text.

The workbook renderers (``parity.excel``) take their header rows from here; the web app's
``web/lib/parity/tableSpecs.ts`` is generated from here (``scripts/parity/gen_table_specs.py``),
so the screens, the files and the API's table pages share one definition. ``table_rows`` mirrors
each renderer cell for cell (the sort, the cleaner, the link label, the truncation), which is
what the Playwright parity test asserts against the rendered files; ``dump_file_tables`` reads
the tables back out of a rendered file for that comparison.

Kinds: ``text`` (a short cell), ``long`` (wrapped text), ``num`` (right aligned), ``link`` (the
cell text is the link's label, ``href`` the address), ``priority`` and ``rank`` (a chip in the
workbook's fill colour), ``grade`` (a grade word chip), ``status`` (Open, Closed and the tracker
statuses), ``id`` (an S or J identifier).
"""

from __future__ import annotations

import argparse
import io
import json
import sys
import zipfile
from collections.abc import Callable, Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from parity import counts
from parity.rules import LEVELS, sort_tracker_p2
from parity.text import clean, clean_cell, clean_posting

Bundle = Mapping[str, Any]
Cell = dict[str, Any]
Row = list[Cell]

PRIORITY_FILLS = {"A": "DDEFE0", "B": "FFF1CC", "C": "EEEEEE"}
RANK_FILLS = {"A": "C6EFCE", "B": "FFEB9C", "C": "E7E6E6"}
PRIORITY_ORDER = {"A": 0, "B": 1, "C": 2}
PAGE_LIMIT_DEFAULT = 100
PAGE_LIMIT_MAX = 500
FILES = ("summary", "evidence", "people", "jobs", "stack", "highlights", "internal")
FILE_SUFFIX = {
    "summary": "_Account_Summary.docx",
    "evidence": "_Sources_and_Evidence.docx",
    "highlights": "_Lane_Highlights.docx",
    "internal": "_Internal_Only.docx",
    "people": "_People_List.xlsx",
    "jobs": "_Job_Postings.xlsx",
    "stack": "_Tech_Stack_and_Teams.xlsx",
}
FILE_LABELS = {
    "summary": "Account Summary",
    "evidence": "Sources and Evidence",
    "people": "People List",
    "jobs": "Job Postings",
    "stack": "Tech Stack and Teams",
    "highlights": "Lane Highlights",
    "internal": "Internal Only",
}


@dataclass(frozen=True)
class Column:
    """One column of a table."""

    key: str
    header: str
    kind: str = "text"
    link_label: str | None = None
    fills: Mapping[str, str] | None = None


@dataclass(frozen=True)
class TableSpec:
    """One table of a file."""

    id: str
    file: str
    title: str
    columns: tuple[Column, ...]
    sheet: str | None = None
    paginated: bool = False
    rows: Callable[[Bundle], list[Row]] = field(default=lambda _b: [], repr=False, compare=False)

    @property
    def headers(self) -> list[str]:
        """The header row as the file prints it."""
        return [c.header for c in self.columns]

    def header_row(self, bundle: Bundle) -> list[str]:
        """The header row of this bundle (one heading carries the pull date)."""
        return [h.replace("{pull_day_month}", pull_day_month(bundle)) for h in self.headers]


def cell(text: object, href: str | None = None) -> Cell:
    """A cell: its text and, for a link, its address."""
    out: Cell = {"text": "" if text is None else str(text)}
    if href:
        out["href"] = href
    return out


def _row(*texts: object) -> Row:
    return [cell(t) for t in texts]


def _cleaned(rows: Sequence[Sequence[Any]]) -> list[Row]:
    return [[cell(clean(x)) for x in r] for r in rows]


def pull_day_month(bundle: Bundle) -> str:
    """The short pull date without its year, as the Summary's hiring table heads it."""
    meta = bundle["meta"]
    short = str(meta.get("pull_date_short") or meta.get("pull_date") or "")
    parts = short.split()
    if len(parts) >= 3 and parts[-1].isdigit() and len(parts[-1]) == 4:
        return " ".join(parts[:-1])
    return short


def host_of(url: str) -> str:
    """The host of a URL without a leading ``www.`` (the Word renderer's ``hostOf``)."""
    host = urlsplit(url).hostname or ""
    return host[4:] if host.startswith("www.") else host


def percent_text(share: float) -> str:
    """A share as the workbook's ``0%`` format prints it (half up)."""
    value = Decimal(str(share)) * 100
    return f"{int(value.quantize(Decimal('1'), rounding=ROUND_HALF_UP))}%"


def number_text(value: object) -> str:
    """A workbook number cell as text: a whole number without a decimal point."""
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# ---------------------------------------------------------------- workbook rows
def people_sorted(bundle: Bundle) -> list[Mapping[str, Any]]:
    """The People sheet's order: level, priority, lane, name."""
    return sorted(
        bundle.get("people_rows") or (),
        key=lambda p: (
            int(p["level"]),
            PRIORITY_ORDER.get(str(p["priority"]), 3),
            str(p["lane"]),
            str(p["name"]),
        ),
    )


def _people_rows(bundle: Bundle) -> list[Row]:
    out: list[Row] = []
    for i, p in enumerate(people_sorted(bundle), 1):
        url = p.get("profile_url") or None
        out.append(
            [
                cell(i),
                cell(p["name"]),
                cell(p["title"]),
                cell(LEVELS[int(p["level"]) - 1]),
                cell("Yes" if p["senior"] else "No"),
                cell(p["lane"]),
                cell(p["priority"]),
                cell(p.get("former_si") or ""),
                cell(p["found_by"]),
                cell("Open profile" if url else "", url),
                cell(p.get("second_profile_url") or ""),
                cell(p.get("location") or ""),
                cell(p.get("note") or ""),
            ]
        )
    return out


def jobs_sorted(bundle: Bundle) -> list[Mapping[str, Any]]:
    """The Jobs sheet's order: first seen, newest first (a stable sort)."""
    return sorted(bundle.get("jobs_rows") or (), key=lambda r: str(r["first_seen"]), reverse=True)


def _jobs_rows(bundle: Bundle) -> list[Row]:
    out: list[Row] = []
    for r in jobs_sorted(bundle):
        out.append(
            [
                cell(r.get("requisition") or ""),
                cell(clean_posting(r["title"])),
                cell(", ".join(r.get("platforms") or ())),
                cell(r["cls"]),
                cell(r["status"]),
                cell(r["first_seen"]),
                cell(r["last_seen"]),
                cell(clean_posting(r.get("location") or "")),
                cell(r.get("seniority") or ""),
                cell(r["entity"]),
                cell(r.get("what") or ""),
                cell(clean_posting(r.get("platform_sentences") or "")[:3000]),
                cell(clean_posting(r.get("description") or "")[:32000]),
                cell("Open posting", r["url"]),
            ]
        )
    return out


def _ranking_rows(bundle: Bundle) -> list[Row]:
    rows = (bundle.get("ranking") or {}).get("rows") or ()
    ranked = counts.platform_ranking(bundle)
    out: list[Row] = []
    for rec, rank in zip(rows, ranked, strict=True):
        out.append(
            [
                cell(rank.rank),
                cell(clean_cell(rec["platform"])),
                cell(rec["category"]),
                cell(rec["lane"]),
                cell(rec["count"]),
                cell(percent_text(rank.share)),
                cell(rec["open"]),
                cell(rec["last_seen"]),
                cell(rec["in_stack"]),
                cell(rec["checked"]),
            ]
        )
    return out


def _stack_rows(bundle: Bundle) -> list[Row]:
    return [
        _row(
            r["business_function"],
            r["category"],
            clean_cell(r["technology"]),
            number_text(r.get("postings")),
            r["lane"],
        )
        for r in bundle.get("stack_rows") or ()
    ]


def _team_rows(bundle: Bundle) -> list[Row]:
    return [
        _row(
            clean_cell(r["team"]),
            clean_cell(r["lane_technologies"]),
            number_text(r.get("postings")),
            r["first_activity"],
            r["last_activity"],
            clean_cell(r["technologies"]),
            clean_cell(r["locations"]),
            clean_cell(r["job_functions"]),
        )
        for r in bundle.get("team_rows") or ()
    ]


# ---------------------------------------------------------------- Word rows
def jobs_count_rows(bundle: Bundle) -> list[Row]:
    """The hiring table rows: lane, total, open and the four classes (``jobsCountRows``)."""
    return [
        _row(lane, v["total"], v["open"], v["build"], v["admin"], v["tool"], v["user"])
        for lane, v in (bundle["computed"].get("jobs_counts") or {}).items()
    ]


def _legend_rows(_bundle: Bundle) -> list[Row]:
    return [
        _row("Confirmed", "Read in a primary source during this run", "May be stated"),
        _row(
            "Probable",
            "Inferred from one or two sources, reasoning given in the row",
            "Ask, do not state",
        ),
        _row("Hypothesis", "Our synthesis, no source states it", "Ask, do not state"),
        _row("Lead", "Seen in a secondary source, not verified", "Do not use until checked"),
    ]


def _prose(bundle: Bundle, *path: str) -> Any:
    node: Any = bundle["prose"]
    for key in path:
        node = (node or {}).get(key)
    return node or []


def _key_people_rows(bundle: Bundle) -> list[Row]:
    out: list[Row] = []
    for r in _prose(bundle, "key_people"):
        url = r[4] if len(r) > 4 and r[4] else ""
        out.append(
            [
                cell(clean(r[0])),
                cell(clean(r[1])),
                cell(clean(r[2])),
                cell(r[3] if len(r) > 3 else ""),
                cell("Profile", url) if url else cell("No link found"),
                cell(clean(r[5] if len(r) > 5 else "")),
            ]
        )
    return out


def _sources_rows(bundle: Bundle) -> list[Row]:
    return [
        [cell(sid), cell(clean(s["name"]), s["url"]), cell(host_of(s["url"]))]
        for sid, s in (bundle.get("sources") or {}).items()
    ]


def _register_rows(bundle: Bundle) -> list[Row]:
    return [
        [
            cell(sid),
            cell(clean(s["name"]), s["url"]),
            cell(clean(s["type"])),
            cell(clean(s["date"])),
            cell(clean(s["how"])),
            cell(clean(s["supports"])),
        ]
        for sid, s in (bundle.get("sources") or {}).items()
    ]


def _jobs_read_rows(bundle: Bundle) -> list[Row]:
    return [
        [
            cell(jid),
            cell(clean(j["name"]), j["url"]),
            cell(clean(j["plats"])),
            cell(j["cls"]),
            cell(clean(j["what"])),
            cell(clean(j["ent"])),
        ]
        for jid, j in (bundle.get("jobs_read") or {}).items()
    ]


def part_c_lanes(bundle: Bundle) -> list[str]:
    """The Part C lane order (``meta.part_c_lanes`` when given, else ``meta.lanes``)."""
    meta = bundle["meta"]
    return [str(x) for x in (meta.get("part_c_lanes") or meta.get("lanes") or ())]


def part_c_rows(bundle: Bundle, lane: str) -> list[Row]:
    """Part C: the postings read in full that name ``lane`` (the renderer's filter)."""
    out: list[Row] = []
    for jid, j in (bundle.get("jobs_read") or {}).items():
        if lane not in str(j["plats"]).split(", "):
            continue
        what = str(j["what"])
        reasoning = ": ".join(what.split(": ")[1:]) or what
        out.append(
            [
                cell(jid),
                cell(clean(j["name"]), j["url"]),
                cell(j["cls"]),
                cell(j["status"]),
                cell(clean(reasoning)),
            ]
        )
    return out


def _live_rows(bundle: Bundle) -> list[Row]:
    out: list[Row] = []
    for r in _prose(bundle, "evidence", "live_rows"):
        out.append(
            [
                cell(clean(r[0]), r[1] or None),
                cell(clean(r[2])),
                cell(clean(r[3])),
                cell(r[4]),
                cell(r[5]),
            ]
        )
    return out


def _detection_rows(bundle: Bundle) -> list[Row]:
    return [[cell(clean(str(x))) for x in r] for r in _prose(bundle, "evidence", "detections")]


def _entity_rows(bundle: Bundle) -> list[Row]:
    return [_row(clean(k), str(v)) for k, v in (bundle["computed"].get("ent_counts") or {}).items()]


def _tracker_p2_rows(bundle: Bundle) -> list[Row]:
    return _cleaned(sort_tracker_p2(_prose(bundle, "tracker_p2")))


_TRACKER = (
    Column("row", "Source type"),
    Column("status", "Status", "status"),
    Column("priority", "Priority"),
    Column("needs", "Needs"),
    Column("done", "What was done (proof)", "long"),
    Column("remains", "What remains", "long"),
)
_SIGNAL = (
    Column("finding", "Finding", "long"),
    Column("grade", "Grade", "grade"),
    Column("lanes", "Lanes"),
    Column("date", "Date"),
    Column("sources", "Sources", "id"),
)
_HIRING_SUMMARY = (
    Column("platform", "Platform"),
    Column("total", "Postings naming it", "num"),
    Column("open", "Open on {pull_day_month}", "num"),
    Column("build", "Build roles", "num"),
    Column("admin", "Admin and owner roles", "num"),
    Column("tool", "Working tool", "num"),
    Column("user", "User role", "num"),
)
_HIRING_EVIDENCE = (
    Column("platform", "Platform"),
    Column("total", "Postings", "num"),
    Column("open", "Open", "num"),
    Column("build", "Build", "num"),
    Column("admin", "Admin and owner", "num"),
    Column("tool", "Working tool", "num"),
    Column("user", "User role", "num"),
)
PART_C_COLUMNS = (
    Column("id", "ID", "id"),
    Column("posting", "Posting", "link"),
    Column("cls", "Class"),
    Column("status", "Status", "status"),
    Column("reasoning", "Classification reasoning", "long"),
)


def _signals(tier: str, key: str, title: str) -> TableSpec:
    return TableSpec(
        f"summary_signals_{key}",
        "summary",
        title,
        _SIGNAL,
        rows=lambda b, t=tier: _cleaned(_prose(b, "signals", t)),  # type: ignore[misc]
    )


TABLES: tuple[TableSpec, ...] = (
    TableSpec(
        "people",
        "people",
        "People",
        (
            Column("n", "No.", "num"),
            Column("name", "Name"),
            Column("title", "Title as reported", "long"),
            Column("level", "Level"),
            Column("senior", "Meets seniority rule"),
            Column("lane", "Lane"),
            Column("priority", "Priority", "priority", fills=PRIORITY_FILLS),
            Column("former_si", "Former SI and consulting employers (firm and years)", "long"),
            Column("found_by", "Found by", "long"),
            Column("profile", "LinkedIn", "link", "Open profile"),
            Column("second_profile", "Second profile"),
            Column("location", "Location"),
            Column("note", "Note", "long"),
        ),
        sheet="People",
        paginated=True,
        rows=_people_rows,
    ),
    TableSpec(
        "jobs",
        "jobs",
        "Jobs",
        (
            Column("requisition", "Requisition"),
            Column("title", "Title", "long"),
            Column("platforms", "Platforms named"),
            Column("cls", "Class"),
            Column("status", "Status", "status"),
            Column("first_seen", "First seen"),
            Column("last_seen", "Last seen"),
            Column("location", "Location", "long"),
            Column("seniority", "Seniority"),
            Column("entity", "Entity"),
            Column("what", "What the posting says (read)", "long"),
            Column("platform_sentences", "Platform sentences", "long"),
            Column("description", "Full description", "long"),
            Column("link", "Link", "link", "Open posting"),
        ),
        sheet="Jobs",
        paginated=True,
        rows=_jobs_rows,
    ),
    TableSpec(
        "ranking",
        "stack",
        "Platform Ranking",
        (
            Column("rank", "Rank", "rank", fills=RANK_FILLS),
            Column("platform", "Platform"),
            Column("category", "Category"),
            Column("lane", "Lane"),
            Column("count", "Postings naming it", "num"),
            Column("share", "Share of base", "num"),
            Column("open", "Open now", "num"),
            Column("last_seen", "Last seen"),
            Column("in_stack", "In confirmed stack"),
            Column("checked", "Matches checked by reading"),
        ),
        sheet="Platform Ranking",
        rows=_ranking_rows,
    ),
    TableSpec(
        "stack",
        "stack",
        "Confirmed stack",
        (
            Column("business_function", "Business function"),
            Column("category", "Category"),
            Column("technology", "Technology"),
            Column("postings", "Postings naming it", "num"),
            Column("lane", "Lane"),
        ),
        sheet="Confirmed stack",
        rows=_stack_rows,
    ),
    TableSpec(
        "teams",
        "stack",
        "Platform teams",
        (
            Column("team", "Team"),
            Column("lane_technologies", "Lane technologies named"),
            Column("postings", "Postings count", "num"),
            Column("first_activity", "First activity"),
            Column("last_activity", "Last activity"),
            Column("technologies", "Team technologies (name, postings)", "long"),
            Column("locations", "Team locations (place, postings)", "long"),
            Column("job_functions", "Job functions", "long"),
        ),
        sheet="Platform teams",
        rows=_team_rows,
    ),
    TableSpec(
        "summary_legend",
        "summary",
        "How to read this file",
        (
            Column("grade", "Grade", "grade"),
            Column("meaning", "Meaning", "long"),
            Column("use", "Use in outreach"),
        ),
        rows=_legend_rows,
    ),
    TableSpec(
        "summary_news",
        "summary",
        "2. Recent news (newest first)",
        (
            Column("date", "Date"),
            Column("what", "What happened", "long"),
            Column("read", "Read depth"),
            Column("source", "Source", "id"),
            Column("why", "Why it matters", "long"),
        ),
        rows=lambda b: _cleaned(_prose(b, "news")),
    ),
    _signals("Tier 1", "tier1", "Tier 1, budget proximate"),
    _signals("Tier 2", "tier2", "Tier 2, one step removed"),
    _signals("Tier 3", "tier3", "Tier 3, context and openers"),
    _signals("Context and history", "context", "Context and history (older than the shelf life)"),
    TableSpec(
        "summary_platform",
        "summary",
        "5. Platform picture by lane",
        (
            Column("item", "Item"),
            Column("found", "What we found", "long"),
            Column("grade", "Grade", "grade"),
            Column("reasoning", "Reasoning and sources", "long"),
        ),
        rows=lambda b: _cleaned(_prose(b, "platform")),
    ),
    TableSpec(
        "summary_hiring",
        "summary",
        "Platform hiring over 12 months",
        _HIRING_SUMMARY,
        rows=jobs_count_rows,
    ),
    TableSpec(
        "summary_key_people",
        "summary",
        "Key people (priority A and the executives the lanes report to)",
        (
            Column("name", "Name"),
            Column("title", "Title as reported", "long"),
            Column("lane", "Lane"),
            Column("priority", "Priority", "priority", fills=PRIORITY_FILLS),
            Column("link", "Link", "link", "Profile"),
            Column("note", "Note (from the profile, the leadership page or a release)", "long"),
        ),
        rows=_key_people_rows,
    ),
    TableSpec(
        "summary_profile_findings",
        "summary",
        "Findings from profiles",
        (
            Column("topic", "Topic"),
            Column("shows", "What the profiles show", "long"),
            Column("grade", "Grade", "grade"),
            Column("interpretation", "Our interpretation", "long"),
        ),
        rows=lambda b: _cleaned(_prose(b, "profile_findings")),
    ),
    TableSpec(
        "summary_status_changes",
        "summary",
        "People whose status changed",
        (
            Column("person", "Person"),
            Column("was", "Was", "long"),
            Column("now", "Now", "long"),
            Column("source", "Source", "id"),
        ),
        rows=lambda b: _cleaned(_prose(b, "status_changes")),
    ),
    TableSpec(
        "summary_tracker_p1",
        "summary",
        "Pass 1, summary homework",
        _TRACKER,
        rows=lambda b: _cleaned(_prose(b, "tracker_p1")),
    ),
    TableSpec(
        "summary_tracker_p2",
        "summary",
        "Pass 2, detailed homework",
        _TRACKER,
        rows=_tracker_p2_rows,
    ),
    TableSpec(
        "summary_sources",
        "summary",
        "10. Source links",
        (Column("id", "ID", "id"), Column("source", "Source", "link"), Column("domain", "Domain")),
        rows=_sources_rows,
    ),
    TableSpec(
        "evidence_sources",
        "evidence",
        "Part A. Source register",
        (
            Column("id", "ID", "id"),
            Column("source", "Source", "link"),
            Column("type", "Type"),
            Column("date", "Date"),
            Column("how", "How retrieved", "long"),
            Column("supports", "What it supports", "long"),
        ),
        rows=_register_rows,
    ),
    TableSpec(
        "evidence_jobs_read",
        "evidence",
        "Part B. Job posting extracts for postings read in full",
        (
            Column("id", "ID", "id"),
            Column("posting", "Posting", "link"),
            Column("platforms", "Platforms"),
            Column("cls", "Class"),
            Column("what", "What the posting says", "long"),
            Column("entity", "Entity"),
        ),
        rows=_jobs_read_rows,
    ),
    TableSpec(
        "evidence_live",
        "evidence",
        "Part E. Job board captures",
        (
            Column("posting", "Posting", "link"),
            Column("platforms", "Platforms named"),
            Column("entity", "Entity"),
            Column("read", "Read in full"),
            Column("cls", "Class"),
        ),
        rows=_live_rows,
    ),
    TableSpec(
        "evidence_leads",
        "evidence",
        "Part F. Leads not verified, and items checked and dropped",
        (
            Column("lead", "Lead", "long"),
            Column("queries", "Queries and checks tried", "long"),
            Column("status", "Status", "status"),
            Column("source", "Source", "id"),
        ),
        rows=lambda b: _cleaned(_prose(b, "evidence", "leads")),
    ),
    TableSpec(
        "evidence_detections",
        "evidence",
        "Technology detections for lane platforms",
        (
            Column("technology", "Technology"),
            Column("first_seen", "First seen"),
            Column("last_seen", "Last seen"),
            Column("score", "Score", "num"),
            Column("postings", "Seen on postings", "num"),
            Column("pages", "Seen on pages", "num"),
        ),
        rows=_detection_rows,
    ),
    TableSpec(
        "evidence_hiring",
        "evidence",
        "Postings naming each lane platform",
        _HIRING_EVIDENCE,
        rows=jobs_count_rows,
    ),
    TableSpec(
        "evidence_entities",
        "evidence",
        "Entity split of the 12 month postings",
        (Column("entity", "Entity"), Column("postings", "Postings", "num")),
        rows=_entity_rows,
    ),
    TableSpec(
        "evidence_news_events",
        "evidence",
        "News feed events",
        (
            Column("date", "Date"),
            Column("category", "Category"),
            Column("headline", "Headline as the feed wrote it", "long"),
            Column("domain", "Most relevant source domain"),
        ),
        rows=lambda b: _cleaned(_prose(b, "evidence", "news_events")),
    ),
)
BY_ID: dict[str, TableSpec] = {t.id: t for t in TABLES}
PART_C_PREFIX = "evidence_part_c:"


def spec_of(table_id: str, bundle: Bundle | None = None) -> TableSpec | None:
    """The spec of a table id; a Part C lane table (``evidence_part_c:<lane>``) is built for
    the bundle's lanes."""
    if table_id in BY_ID:
        return BY_ID[table_id]
    if bundle is not None and table_id.startswith(PART_C_PREFIX):
        lane = table_id[len(PART_C_PREFIX) :]
        if lane in part_c_lanes(bundle):
            return TableSpec(
                table_id,
                "evidence",
                f"Part C: {lane}",
                PART_C_COLUMNS,
                rows=lambda b, lane=lane: part_c_rows(b, lane),  # type: ignore[misc]
            )
    return None


def tables_of(bundle: Bundle, file: str) -> list[TableSpec]:
    """The tables of one file in document order (Part C lane tables only when they have rows,
    as the Word file prints a sentence instead of an empty one)."""
    out: list[TableSpec] = []
    for spec in TABLES:
        if spec.file != file:
            continue
        out.append(spec)
        if spec.id == "evidence_jobs_read":
            for lane in part_c_lanes(bundle):
                lane_spec = spec_of(PART_C_PREFIX + lane, bundle)
                if lane_spec is not None and lane_spec.rows(bundle):
                    out.append(lane_spec)
    return out


def table_rows(bundle: Bundle, table_id: str) -> list[Row]:
    """Every row of the table in file order (empty for an unknown id)."""
    spec = spec_of(table_id, bundle)
    return spec.rows(bundle) if spec is not None else []


def file_name(bundle: Bundle, file: str) -> str:
    """The file's name for this bundle's prefix."""
    return f"{bundle['meta']['account_short']}{FILE_SUFFIX[file]}"


def spec_json(spec: TableSpec, bundle: Bundle | None = None) -> dict[str, Any]:
    """The spec as plain data (what the generated TypeScript and the API carry)."""
    return {
        "id": spec.id,
        "file": spec.file,
        "sheet": spec.sheet,
        "title": spec.title,
        "paginated": spec.paginated,
        "columns": [
            {
                "key": c.key,
                "header": c.header
                if bundle is None
                else c.header.replace("{pull_day_month}", pull_day_month(bundle)),
                "kind": c.kind,
                "link_label": c.link_label,
                "fills": dict(c.fills) if c.fills else None,
            }
            for c in spec.columns
        ],
    }


# ---------------------------------------------------------------- reading the files back
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_SHEET_HEADER_ROW = {
    "People": 1,
    "Jobs": 1,
    "Confirmed stack": 4,
    "Platform teams": 4,
    "Platform Ranking": 8,
}


def _docx_tables(data: bytes) -> Iterator[dict[str, Any]]:
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = etree.fromstring(z.read("word/document.xml"), parser=parser)
    for tbl in root.iter(f"{_W}tbl"):
        rows = []
        for tr in tbl.findall(f"{_W}tr"):
            rows.append(
                ["".join(t.text or "" for t in tc.iter(f"{_W}t")) for tc in tr.findall(f"{_W}tc")]
            )
        if rows:
            yield {"headers": rows[0], "rows": rows[1:]}


def _xlsx_tables(path: Path) -> Iterator[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        start = _SHEET_HEADER_ROW.get(ws.title)
        if start is None:
            continue
        rows_out: list[list[str]] = []
        headers = [number_text(c.value) for c in ws[start]]
        # LOOP(trailing blank headers): cap=len(headers) cells; exit=last header non-empty
        while headers and headers[-1] == "":
            headers.pop()
        for row in ws.iter_rows(min_row=start + 1, max_col=len(headers)):
            cells = list(row)
            if ws.title == "Platform Ranking" and (cells[1].value in (None, "")):
                break
            values = []
            for c in cells:
                if c.number_format == "0%" and isinstance(c.value, (int, float)):
                    values.append(percent_text(float(c.value)))
                else:
                    values.append(number_text(c.value))
            if all(v == "" for v in values):
                if ws.title == "Platform Ranking":
                    break
                continue
            rows_out.append(values)
        yield {"sheet": ws.title, "headers": headers, "rows": rows_out}


def dump_file_tables(path: Path) -> list[dict[str, Any]]:
    """The tables of a rendered file as text: Word tables in document order (first row the
    header), workbook data sheets from their header row down."""
    if path.suffix.lower() == ".docx":
        return list(_docx_tables(path.read_bytes()))
    return list(_xlsx_tables(path))


def main(argv: Sequence[str] | None = None) -> int:
    """``python -m parity.tables dump <file>``: the file's tables as JSON; ``specs``: the specs."""
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("dump")
    d.add_argument("file", type=Path)
    sub.add_parser("specs")
    args = ap.parse_args(argv)
    if args.cmd == "dump":
        json.dump(dump_file_tables(args.file), sys.stdout, ensure_ascii=False)
    else:
        json.dump([spec_json(t) for t in TABLES], sys.stdout, indent=2, ensure_ascii=False)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
