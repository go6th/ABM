"""The People List workbook: a Summary sheet (notes, the level table and the lane table as
formulas over the People sheet) and the People sheet, rendered from a bundle exactly as the
reference wrote it (same strings, formulas, styles, widths, freeze panes and filter)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from parity.bundle import AccountBundle
from parity.excel import (
    FONT,
    body_font,
    link_font,
    navy_fill,
    style_header,
    style_table,
    thin_border,
    title_font,
)
from parity.rules import LEVELS, PEOPLE_LANES
from parity.tables import BY_ID

# the header row of the People sheet (parity.tables, the one definition the web shares)
HEADER = BY_ID["people"].headers
WIDTHS = [6, 26, 56, 22, 11, 30, 9, 30, 30, 13, 13, 24, 50]
PRIORITY_FILLS = {"A": "DDEFE0", "B": "FFF1CC", "C": "EEEEEE"}
SECURITY_LINE = (
    "Security, IT compliance and IT audit people are excluded at the account owner's "
    "instruction ({count} rows): {names}."
)
_ORDER = {"A": 0, "B": 1, "C": 2}


def security_line(bundle: AccountBundle) -> str:
    """The Summary note that lists the excluded security, IT compliance and IT audit titles
    with their count (entries, as the reference counts rows; a name with two titles counts
    twice and the gate's data warning names it)."""
    entries = bundle.people_excluded.security
    template = bundle.prose.workbook_notes.people_security_line_template or SECURITY_LINE
    return template.format(count=len(entries), names="; ".join(entries))


def summary_notes(bundle: AccountBundle) -> list[str]:
    """The Summary notes with the security line inserted at its index."""
    notes = list(bundle.prose.workbook_notes.people)
    at = bundle.prose.workbook_notes.people_security_line_at
    notes.insert(2 if at is None else min(at, len(notes)), security_line(bundle))
    return notes


def render_people(bundle: AccountBundle, path: Path) -> Path:
    """Write ``{prefix}_People_List.xlsx`` to ``path`` and return it."""
    rows = sorted(bundle.people_rows, key=lambda p: (p.level, _ORDER[p.priority], p.lane, p.name))
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "People"
    ws.append(HEADER)
    for i, p in enumerate(rows, 1):
        ws.append(
            [
                i,
                p.name,
                p.title,
                LEVELS[p.level - 1],
                "Yes" if p.senior else "No",
                p.lane,
                p.priority,
                p.former_si or "",
                p.found_by,
                "Open profile" if p.profile_url else "",
                p.second_profile_url or "",
                p.location or "",
                p.note or "",
            ]
        )
        if p.profile_url:
            c = ws.cell(row=i + 1, column=10)
            c.hyperlink = p.profile_url
            c.font = link_font()
    style_header(ws[1])
    # the formula ranges end at the last row, or at the first empty row when the sheet has none
    last = len(rows) + 1
    # (LibreOffice normalises a one-cell range to the cell, which the workbook check reads as a
    # changed formula, so the range always spans at least two rows)
    end = max(last, 3)
    for cells in ws.iter_rows(min_row=2, max_row=last):
        for c in cells:
            if not c.hyperlink:
                c.font = body_font()
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = thin_border()
        cells[6].fill = PatternFill("solid", fgColor=PRIORITY_FILLS[str(cells[6].value)])
        cells[6].alignment = Alignment(horizontal="center", vertical="top")
    for col, w in zip("ABCDEFGHIJKLM", WIDTHS, strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{end}"
    ws.row_dimensions[1].height = 32

    s = wb.create_sheet("Summary", 0)
    t = "People"
    s["A1"] = f"{bundle.meta.account_short} people list, deduplicated"
    s["A1"].font = title_font()
    notes = summary_notes(bundle)
    for k, n in enumerate(notes):
        c = s.cell(row=3 + k, column=1, value=n)
        c.font = body_font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        s.merge_cells(start_row=3 + k, start_column=1, end_row=3 + k, end_column=6)
        s.row_dimensions[3 + k].height = 58
    r0 = 3 + len(notes) + 2
    s.cell(row=r0, column=1, value="Level")
    s.cell(row=r0, column=2, value="People")
    s.cell(row=r0, column=3, value="Priority A")
    for k, label in enumerate(LEVELS):
        r = r0 + 1 + k
        s.cell(row=r, column=1, value=label)
        s.cell(row=r, column=2, value=f"=COUNTIF({t}!$D$2:$D${end},A{r})")
        s.cell(
            row=r,
            column=3,
            value=f'=COUNTIFS({t}!$D$2:$D${end},A{r},{t}!$G$2:$G${end},"A")',
        )
    rt = r0 + 1 + len(LEVELS)
    s.cell(row=rt, column=1, value="Total")
    s.cell(row=rt, column=2, value=f"=SUM(B{r0 + 1}:B{rt - 1})")
    s.cell(row=rt, column=3, value=f"=SUM(C{r0 + 1}:C{rt - 1})")
    s.cell(row=rt + 1, column=1, value="Meets seniority rule")
    s.cell(row=rt + 1, column=2, value=f'=COUNTIF({t}!$E$2:$E${end},"Yes")')
    r2 = rt + 3
    s.cell(row=r2, column=1, value="Lane, senior people only")
    for j, h in enumerate(["A", "B", "C", "Total"]):
        s.cell(row=r2, column=2 + j, value=h)
    lanes = bundle.meta.people_lanes or list(PEOPLE_LANES)
    for k, lane in enumerate(lanes):
        r = r2 + 1 + k
        s.cell(row=r, column=1, value=lane)
        for j, pv in enumerate(["A", "B", "C"]):
            s.cell(
                row=r,
                column=2 + j,
                value=(
                    f"=COUNTIFS({t}!$F$2:$F${end},$A{r},{t}!$G$2:$G${end},"
                    f'"{pv}",{t}!$E$2:$E${end},"Yes")'
                ),
            )
        s.cell(row=r, column=5, value=f"=SUM(B{r}:D{r})")
    r3 = r2 + 1 + len(lanes)
    s.cell(row=r3, column=1, value="Total")
    for j in range(4):
        col = get_column_letter(2 + j)
        s.cell(row=r3, column=2 + j, value=f"=SUM({col}{r2 + 1}:{col}{r3 - 1})")
    style_table(s, r0, r3, {rt, r3})
    for c in list(s[r0]) + list(s[r2]):
        if c.value is not None:
            c.fill = navy_fill()
            c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    s.column_dimensions["A"].width = 44
    for col in "BCDEF":
        s.column_dimensions[col].width = 12
    wb.save(path)
    return path
