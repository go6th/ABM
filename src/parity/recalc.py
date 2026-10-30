"""Recalculate rendered workbooks with LibreOffice headless so cached formula values exist.

openpyxl writes formulas without values; a reader that does not calculate shows blanks. The
reference workbooks were written by openpyxl and re-saved by LibreOffice (their ``app.xml``
names it), which is also what the parity checker compares against, so the same re-save is the
last step of every render: one ``soffice --convert-to xlsx`` call for all files, a private
profile, a cleared environment and a timeout. ``check_workbook`` then asserts that no formula
produced an error value, that every formula still is one and that every formula has a cached
value.
"""

from __future__ import annotations

import shutil
import subprocess  # nosec B404 - fixed argv, no shell
import tempfile
from collections.abc import Sequence
from pathlib import Path

from openpyxl import load_workbook

ERROR_VALUES = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
DEFAULT_BINARY = "/usr/bin/soffice"


class RecalcError(RuntimeError):
    """LibreOffice did not produce a recalculated workbook."""


def soffice_binary() -> str | None:
    """The LibreOffice binary, or None when it is not installed."""
    if Path(DEFAULT_BINARY).exists():
        return DEFAULT_BINARY
    return shutil.which("soffice")


def recalculate(sources: Sequence[Path], out_dir: Path, timeout_s: float = 600.0) -> list[Path]:
    """Re-save every workbook in ``sources`` into ``out_dir`` (same file names) through
    LibreOffice, returning the written paths in the order given."""
    binary = soffice_binary()
    if binary is None:
        raise RecalcError("LibreOffice (soffice) is not installed")
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        cmd = [
            binary,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--norestore",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
            *[str(p) for p in sources],
        ]
        try:
            subprocess.run(  # nosec B603 - fixed argv, no shell
                cmd,
                check=True,
                capture_output=True,
                timeout=timeout_s,
                env={"HOME": tmp, "TMPDIR": tmp, "LC_ALL": "C.UTF-8", "PATH": "/usr/bin:/bin"},
            )
        except (subprocess.SubprocessError, OSError) as exc:
            raise RecalcError(f"LibreOffice conversion failed: {exc}") from exc
    written = [out_dir / p.name for p in sources]
    missing = [p for p in written if not p.exists()]
    if missing:
        raise RecalcError(f"LibreOffice wrote no file for {missing}")
    return written


def formula_cells(path: Path) -> dict[tuple[str, str], str]:
    """Every formula of a workbook, keyed by (sheet, coordinate)."""
    wb = load_workbook(path)
    out: dict[tuple[str, str], str] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out[(ws.title, c.coordinate)] = c.value
    return out


def check_workbook(path: Path, source: Path | None = None) -> list[str]:
    """Problems with a recalculated workbook: error values, formulas without cached values,
    and (when ``source`` is given) formulas that are no longer formulas or changed text."""
    problems: list[str] = []
    formulas = formula_cells(path)
    values = load_workbook(path, data_only=True)
    for (sheet, coord), formula in formulas.items():
        value = values[sheet][coord].value
        if value is None:
            problems.append(f"{sheet}!{coord}: no cached value for {formula}")
        elif isinstance(value, str) and value in ERROR_VALUES:
            problems.append(f"{sheet}!{coord}: {value} from {formula}")
    for ws in values.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value in ERROR_VALUES:
                    key = (ws.title, c.coordinate)
                    if key not in formulas:
                        problems.append(f"{ws.title}!{c.coordinate}: {c.value}")
    if source is not None:
        before = formula_cells(source)
        for key, formula in before.items():
            if key not in formulas:
                problems.append(f"{key[0]}!{key[1]}: formula lost ({formula})")
            elif formulas[key] != formula:
                problems.append(f"{key[0]}!{key[1]}: formula changed to {formulas[key]}")
    return problems
