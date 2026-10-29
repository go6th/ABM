"""Export parity (VW1): every exported row equals the API row, cell for cell, in the same order.

The check reads the produced bytes back (it never trusts the model that rendered them):

* workbooks: the table's sheet, header to column key through the table spec; a long text split
  onto the Long text sheet is rejoined (the parts join back exactly, ``exports.policy``);
* Word documents: every table whose header row is the table's headers (the Account Summary's
  Platform picture by lane, its tracker table), rows in document order, paragraphs joined;
* CSV: the file's rows, with the one leading apostrophe that neutralises a formula removed.

Values are compared after the documented formatting (``normalise``): empty for None, ISO dates,
integers as digits, whitespace collapsed. Missing, extra, reordered or different rows are
problems; any problem blocks delivery of the package (``check_package``).

GUARANTEE(test_exported_rows_equal_api_rows_and_a_mutation_is_caught)
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from lxml import etree
from openpyxl import load_workbook

from views.tables import TableSpec

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)
_POINTER = re.compile(
    r" \[continued: part 1 of (\d+); parts 2 to \d+ are on the .+ sheet under "
    r"[A-Z]{1,3}-\d{4,}\]$"
)
_REF = re.compile(r"^[A-Z]{1,3}-\d{4,}$")
_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def normalise(value: Any) -> str:
    """The documented comparison form of one cell."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


@dataclass(frozen=True)
class ParityResult:
    """One table's verdict."""

    table: str
    compared: int
    problems: tuple[str, ...] = field(default=())

    @property
    def ok(self) -> bool:
        """True when nothing differs."""
        return not self.problems


def _long_parts(wb: Any, sheet: str) -> dict[tuple[str, str], dict[int, str]]:
    if sheet not in wb.sheetnames:
        return {}
    out: dict[tuple[str, str], dict[int, str]] = {}
    for row in wb[sheet].iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            ref, column, part, _of, text = (list(row) + [None] * 5)[:5]
            out.setdefault((str(ref), str(column)), {})[int(part)] = str(text or "")
    return out


def _header_at(rows: Sequence[Sequence[Any]], headers: Sequence[str]) -> int | None:
    """The index of the sheet's header row: the first row that holds every one of the table's
    headers (a sheet may open with a title line and a note above its table, as the Stack
    detections sheet does); else the first row, as before."""
    want = set(headers)
    for i, row in enumerate(rows):
        if want <= {normalise(v) for v in row}:
            return i
    return 0 if rows else None


def _cell_value(cell: Any, kind: str) -> Any:
    """A cell's value; a link column's value is the address it opens (the People sheet shows
    "Open profile" and links to the profile)."""
    link = getattr(cell, "hyperlink", None)
    if kind == "url" and link is not None and getattr(link, "target", None):
        return link.target
    return cell.value


def xlsx_rows(data: bytes, spec: TableSpec, long_sheet: str = "Long text") -> list[dict[str, str]]:
    """The data rows of the table's sheet, as {column key: text}, long text rejoined."""
    # not read-only: a read-only sheet does not carry the cells' hyperlinks
    wb = load_workbook(io.BytesIO(data), data_only=True)
    try:
        if spec.sheet is None or spec.sheet not in wb.sheetnames:
            raise ValueError(f"sheet {spec.sheet!r} is missing")
        kinds = {c.header: c.kind for c in spec.exported_columns}
        grid = list(wb[spec.sheet].iter_rows())
        long = _long_parts(wb, long_sheet)
    finally:
        wb.close()
    head = _header_at(
        [[c.value for c in row] for row in grid], [c.header for c in spec.exported_columns]
    )
    if head is None:
        return []
    header_row = [c.value for c in grid[head]]
    rows = [
        [
            _cell_value(c, kinds.get(str(h), "")) if h is not None else c.value
            for c, h in zip(row, header_row + [None] * (len(row) - len(header_row)), strict=False)
        ]
        for row in grid[head + 1 :]
    ]
    header_to_key = {c.header: c.key for c in spec.exported_columns}
    keys = [header_to_key.get(str(h)) if h is not None else None for h in header_row]
    # the reference is the column keyed ``ref`` (the People sheet leads with a running number and
    # carries its reference further right); a sheet with no such column keeps the first one
    ref_at = keys.index("ref") if "ref" in keys else 0
    out: list[dict[str, str]] = []
    for row in rows:
        if "ref" in keys:
            ref = row[ref_at] if row and len(row) > ref_at else None
            if not isinstance(ref, str) or not _REF.match(ref):
                continue
        else:
            # a sheet with no reference column (the Confirmed stack sheet has none, as in the
            # owner's reference file): a data row has two or more filled cells (a "none found"
            # note under the header fills one)
            ref = ""
            if sum(1 for v in row if v not in (None, "")) < 2:
                continue
        record: dict[str, str] = {}
        for key, header, value in zip(keys, header_row, row, strict=False):
            if key is None:
                continue
            text = value if not isinstance(value, str) else value
            if isinstance(text, str):
                m = _POINTER.search(text)
                if m:
                    parts = long.get((ref, str(header)), {})
                    text = text[: m.start()] + "".join(
                        parts.get(k, "") for k in range(2, int(m.group(1)) + 1)
                    )
            record[key] = normalise(text)
        out.append(record)
    return out


def _cell_text(tc: Any) -> str:
    paragraphs = ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in tc.iter(f"{W}p")]
    return "\n".join(paragraphs)


def docx_rows(data: bytes, spec: TableSpec) -> list[dict[str, str]]:
    """Rows of every table whose header row is the table's headers, in document order."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = etree.fromstring(z.read("word/document.xml"), parser=_PARSER)
    headers = [c.header for c in spec.exported_columns]
    keys = [c.key for c in spec.exported_columns]
    out: list[dict[str, str]] = []
    for tbl in root.iter(f"{W}tbl"):
        trs = [tr for tr in tbl.iter(f"{W}tr")]
        if not trs:
            continue
        head = [normalise(_cell_text(tc)) for tc in trs[0].iter(f"{W}tc")]
        if head != headers:
            continue
        for tr in trs[1:]:
            cells = [_cell_text(tc) for tc in tr.iter(f"{W}tc")]
            if not cells or not _REF.match(cells[0].strip()):
                continue
            out.append({k: normalise(v) for k, v in zip(keys, cells, strict=False)})
    return out


def csv_rows(data: bytes, spec: TableSpec) -> list[dict[str, str]]:
    """The CSV's data rows, the formula-neutralising apostrophe removed."""
    reader = csv.reader(io.StringIO(data.decode("utf-8")))
    rows = list(reader)
    if not rows:
        return []
    header_to_key = {c.header: c.key for c in spec.exported_columns}
    keys = [header_to_key.get(h) for h in rows[0]]
    out = []
    for row in rows[1:]:
        record = {}
        for key, value in zip(keys, row, strict=False):
            if key is None:
                continue
            if value.startswith("'") and value[1:].lstrip(" ").startswith(_FORMULA):
                value = value[1:]
            record[key] = normalise(value)
        out.append(record)
    return out


def compare(
    table: str,
    expected: Sequence[Mapping[str, Any]],
    found: Sequence[Mapping[str, str]],
    keys: Sequence[str],
) -> ParityResult:
    """Row by row, in order: the API rows (``expected``) against the file's rows (``found``)."""
    problems: list[str] = []
    if len(expected) != len(found):
        problems.append(f"{table}: {len(expected)} API rows, {len(found)} exported rows")
    for i, (want, got) in enumerate(zip(expected, found, strict=False), start=1):
        for key in keys:
            a, b = normalise(want.get(key)), normalise(got.get(key))
            if a != b:
                problems.append(f"{table} row {i} ({normalise(want.get('ref'))}): {key} differs")
    return ParityResult(table, min(len(expected), len(found)), tuple(problems[:200]))


def check(
    spec: TableSpec, api_rows: Sequence[Mapping[str, Any]], file_kind: str, data: bytes
) -> ParityResult:
    """One table against one produced file (``xlsx``, ``docx`` or ``csv``)."""
    try:
        if file_kind == "xlsx":
            found = xlsx_rows(data, spec)
        elif file_kind == "docx":
            found = docx_rows(data, spec)
        elif file_kind == "csv":
            found = csv_rows(data, spec)
        else:
            return ParityResult(spec.kind, 0, (f"{spec.kind}: unknown file kind {file_kind}",))
    except (ValueError, KeyError, zipfile.BadZipFile, etree.XMLSyntaxError) as e:
        return ParityResult(spec.kind, 0, (f"{spec.kind}: file unreadable ({type(e).__name__})",))
    return compare(spec.kind, api_rows, found, [c.key for c in spec.exported_columns])


def check_package(
    specs: Mapping[str, TableSpec],
    api_rows: Mapping[str, Sequence[Mapping[str, Any]]],
    files: Mapping[str, tuple[str, bytes]],
) -> dict[str, ParityResult]:
    """Every table against the file that prints it. ``files`` maps a file kind (``people``,
    ``jobs``, ``tech``, ``summary``) to (extension, bytes). A table whose file is absent is a
    problem (a package must carry every table)."""
    out: dict[str, ParityResult] = {}
    for kind, spec in specs.items():
        produced = files.get(spec.file_kind)
        if produced is None:
            out[kind] = ParityResult(kind, 0, (f"{kind}: file {spec.file_kind} missing",))
            continue
        ext, data = produced
        out[kind] = check(spec, api_rows.get(kind, ()), ext, data)
    return out
